#!/usr/bin/env python3
"""
Сбор новостей Cbonds по эмитенту и дате букбилдинга через API.

Что делает:
1. Читает Excel-файл bonds_ext_enriched_rusbonds_orientir.xlsx.
2. Берёт только те выпуски, где столбец `pb_tizer` пустой.
3. Для каждого уникального сочетания (Эмитент, Дата букбилдинга) получает `emitent_id`
   через POST https://cbonds.ru/api/suggest/companies/.
4. Вызывает POST https://cbonds.ru/api/news/ с фильтрами по `emitent_id` и дате.
5. Забирает все страницы результата (`page=1,2,3...`) через API.
6. Сохраняет новости в SQLite-базу, используя текст статьи из поля `text`.
7. Привязывает найденные новости ко всем выпускам из Excel, относящимся к этой группе.

Важно:
- Для /api/news/ Cbonds может требовать авторизованную сессию.
- Проще всего положить строку Cookie из браузера в файл `cbonds_cookie.txt`
  рядом со скриптом или передать её через переменную окружения `CBONDS_COOKIE`.
- Если Cookie не указаны, скрипт всё равно попробует отработать, но API может вернуть
  пустой результат или ошибку авторизации.

Установка зависимостей:
    pip install requests openpyxl

Запуск:
    python cbonds_news_scrape.py
"""

from __future__ import annotations

import json
import os
import re
import sqlite3
import time
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

import openpyxl
import requests

# ──────────────────────────────────────────────────────────────────────────────
# Настройки
# ──────────────────────────────────────────────────────────────────────────────

INPUT_FILE = "bonds_ext_enriched_rusbonds_orientir.xlsx"
CACHE_FILE = "cbonds_news_cache.json"
DB_FILE = "cbonds_news.sqlite3"
COOKIE_FILE = "cbonds_cookie.txt"
DEBUG_DIR = "debug_api_news"

DELAY = 0.7
TIMEOUT = 25
MAX_RETRIES = 3
NEWS_PAGE_LIMIT = 30
MAX_API_PAGES = 100

SITE_BASE_URL = "https://cbonds.ru"
SUGGEST_COMPANIES_URL = f"{SITE_BASE_URL}/api/suggest/companies/"
NEWS_API_URL = f"{SITE_BASE_URL}/api/news/"
NEWS_SEARCH_BASE_URL = f"{SITE_BASE_URL}/news/"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/147.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-GB,en;q=0.9,ru-RU;q=0.8,ru;q=0.7,en-US;q=0.6",
}

REQUIRED_HEADERS = ["Бумага", "ISIN", "Эмитент", "Дата букбилдинга", "pb_tizer"]


# ──────────────────────────────────────────────────────────────────────────────
# Модели
# ──────────────────────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class IssueRow:
    excel_row: int
    paper: str | None
    isin: str | None
    emitent: str
    bookbuilding_date: str  # YYYY-MM-DD


@dataclass(frozen=True)
class EmitentMatch:
    emitent_id: int
    ttl: str
    status_id: str | None


# ──────────────────────────────────────────────────────────────────────────────
# Вспомогательные функции
# ──────────────────────────────────────────────────────────────────────────────

def normalize_text(value: str | None) -> str:
    if not value:
        return ""
    value = str(value).strip().casefold().replace("ё", "е")
    value = re.sub(r"\s+", " ", value)
    return value


def ensure_date_string(value) -> str | None:
    if value is None or str(value).strip() == "":
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass

    raise ValueError(f"Не удалось распознать дату: {value!r}")


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def group_key(emitent_id: int, date_str: str) -> str:
    return f"{emitent_id}||{date_str}"


def maybe_sleep() -> None:
    time.sleep(DELAY)


def sanitize_filename(value: str) -> str:
    value = re.sub(r"[^\w\-.]+", "_", value, flags=re.UNICODE)
    return value[:180].strip("._") or "debug"


def load_cookie_header() -> str | None:
    env_cookie = os.getenv("CBONDS_COOKIE", "").strip()
    if env_cookie:
        return env_cookie

    cookie_path = Path(COOKIE_FILE)
    if cookie_path.exists():
        text = cookie_path.read_text(encoding="utf-8").strip()
        return text or None

    return None


def dump_debug_json(prefix: str, payload: Any) -> str:
    Path(DEBUG_DIR).mkdir(parents=True, exist_ok=True)
    file_name = sanitize_filename(prefix) + ".json"
    file_path = Path(DEBUG_DIR) / file_name
    file_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(file_path)


# ──────────────────────────────────────────────────────────────────────────────
# Кэш
# ──────────────────────────────────────────────────────────────────────────────

def load_cache() -> dict:
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {"emitents": {}, "search_results": {}}


def save_cache(cache: dict) -> None:
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


# ──────────────────────────────────────────────────────────────────────────────
# HTTP
# ──────────────────────────────────────────────────────────────────────────────

session = requests.Session()
session.headers.update(HEADERS)

_cookie_header = load_cookie_header()
if _cookie_header:
    session.headers["Cookie"] = _cookie_header


def request_with_retry(method: str, url: str, **kwargs) -> requests.Response:
    last_error = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = session.request(method, url, timeout=TIMEOUT, **kwargs)
            resp.raise_for_status()
            return resp
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            if attempt == MAX_RETRIES:
                raise
            sleep_for = DELAY * attempt
            print(f"    [retry {attempt}/{MAX_RETRIES - 1}] {method} {url} -> {exc}")
            time.sleep(sleep_for)
    raise RuntimeError(f"Request failed: {method} {url}: {last_error}")


def parse_json_response(resp: requests.Response) -> dict:
    try:
        return resp.json()
    except Exception:  # noqa: BLE001
        text = resp.text.lstrip("\ufeff\n\r\t ")
        return json.loads(text)


# ──────────────────────────────────────────────────────────────────────────────
# Excel
# ──────────────────────────────────────────────────────────────────────────────

def find_target_sheet(wb: openpyxl.Workbook):
    normalized_required = {normalize_text(x) for x in REQUIRED_HEADERS}

    for ws in wb.worksheets:
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        normalized_headers = {normalize_text(h) for h in headers if h is not None}
        if normalized_required.issubset(normalized_headers):
            return ws

    available = ", ".join(wb.sheetnames)
    raise RuntimeError(
        f"Не найден лист с колонками {REQUIRED_HEADERS}. Доступные листы: {available}"
    )


def read_target_rows(input_file: str) -> list[IssueRow]:
    wb = openpyxl.load_workbook(input_file, data_only=True)
    ws = find_target_sheet(wb)
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    header_to_idx = {str(h).strip(): i for i, h in enumerate(headers)}

    paper_idx = header_to_idx["Бумага"]
    isin_idx = header_to_idx["ISIN"]
    emitent_idx = header_to_idx["Эмитент"]
    date_idx = header_to_idx["Дата букбилдинга"]
    pb_tizer_idx = header_to_idx["pb_tizer"]

    result: list[IssueRow] = []

    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        pb_tizer = row[pb_tizer_idx]
        emitent = row[emitent_idx]
        bookbuilding_value = row[date_idx]

        if not is_blank(pb_tizer):
            continue
        if is_blank(emitent) or is_blank(bookbuilding_value):
            continue

        date_str = ensure_date_string(bookbuilding_value)
        if not date_str:
            continue

        result.append(
            IssueRow(
                excel_row=excel_row,
                paper=None if is_blank(row[paper_idx]) else str(row[paper_idx]).strip(),
                isin=None if is_blank(row[isin_idx]) else str(row[isin_idx]).strip(),
                emitent=str(emitent).strip(),
                bookbuilding_date=date_str,
            )
        )

    return result


# ──────────────────────────────────────────────────────────────────────────────
# SQLite
# ──────────────────────────────────────────────────────────────────────────────

def init_db(db_path: str) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")

    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS searches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            emitent_name TEXT NOT NULL,
            emitent_id INTEGER,
            bookbuilding_date TEXT NOT NULL,
            search_url TEXT,
            total_items INTEGER,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(emitent_name, bookbuilding_date)
        );

        CREATE TABLE IF NOT EXISTS articles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            article_id INTEGER UNIQUE,
            url TEXT NOT NULL UNIQUE,
            title TEXT,
            date_label TEXT,
            date_time TEXT,
            time_label TEXT,
            introtext TEXT,
            content_text TEXT,
            source_label TEXT,
            language TEXT,
            type_label TEXT,
            importance INTEGER,
            raw_json TEXT,
            fetched_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS issue_articles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            excel_row INTEGER NOT NULL,
            paper TEXT,
            isin TEXT,
            emitent_name TEXT NOT NULL,
            emitent_id INTEGER,
            bookbuilding_date TEXT NOT NULL,
            article_id INTEGER NOT NULL,
            article_url TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(excel_row, article_id),
            FOREIGN KEY(article_url) REFERENCES articles(url)
        );
        """
    )

    return conn


def upsert_search(
    conn: sqlite3.Connection,
    emitent_name: str,
    emitent_id: int | None,
    bookbuilding_date: str,
    search_url: str,
    total_items: int | None,
) -> None:
    conn.execute(
        """
        INSERT INTO searches (emitent_name, emitent_id, bookbuilding_date, search_url, total_items)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(emitent_name, bookbuilding_date) DO UPDATE SET
            emitent_id = excluded.emitent_id,
            search_url = excluded.search_url,
            total_items = excluded.total_items
        """,
        (emitent_name, emitent_id, bookbuilding_date, search_url, total_items),
    )
    conn.commit()


def save_article(conn: sqlite3.Connection, item: dict) -> int:
    article_id = int(item["id"])
    url = str(item.get("cb_link") or f"{SITE_BASE_URL}/news/{article_id}/")
    raw_json = json.dumps(item, ensure_ascii=False)

    conn.execute(
        """
        INSERT INTO articles (
            article_id, url, title, date_label, date_time, time_label, introtext,
            content_text, source_label, language, type_label, importance, raw_json
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(article_id) DO UPDATE SET
            url = excluded.url,
            title = COALESCE(excluded.title, articles.title),
            date_label = COALESCE(excluded.date_label, articles.date_label),
            date_time = COALESCE(excluded.date_time, articles.date_time),
            time_label = COALESCE(excluded.time_label, articles.time_label),
            introtext = COALESCE(excluded.introtext, articles.introtext),
            content_text = COALESCE(excluded.content_text, articles.content_text),
            source_label = COALESCE(excluded.source_label, articles.source_label),
            language = COALESCE(excluded.language, articles.language),
            type_label = COALESCE(excluded.type_label, articles.type_label),
            importance = COALESCE(excluded.importance, articles.importance),
            raw_json = COALESCE(excluded.raw_json, articles.raw_json)
        """,
        (
            article_id,
            url,
            item.get("caption"),
            item.get("date"),
            item.get("date_time"),
            item.get("time"),
            item.get("introtext"),
            item.get("text"),
            item.get("source"),
            item.get("language"),
            item.get("type"),
            None if item.get("imp.numeric") is None else int(item.get("imp.numeric")),
            raw_json,
        ),
    )
    conn.commit()
    return article_id


def link_issue_to_article(
    conn: sqlite3.Connection,
    issue: IssueRow,
    emitent_id: int | None,
    article_id: int,
    article_url: str,
) -> None:
    conn.execute(
        """
        INSERT OR IGNORE INTO issue_articles (
            excel_row, paper, isin, emitent_name, emitent_id, bookbuilding_date, article_id, article_url
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            issue.excel_row,
            issue.paper,
            issue.isin,
            issue.emitent,
            emitent_id,
            issue.bookbuilding_date,
            article_id,
            article_url,
        ),
    )
    conn.commit()


# ──────────────────────────────────────────────────────────────────────────────
# API: suggest companies
# ──────────────────────────────────────────────────────────────────────────────

def choose_best_emitent_match(query: str, items: list[dict]) -> EmitentMatch | None:
    if not items:
        return None

    q_norm = normalize_text(query)

    exact_matches = []
    startswith_matches = []

    for item in items:
        ttl = str(item.get("ttl") or "").strip()
        ttl_norm = normalize_text(ttl)

        if ttl_norm == q_norm:
            exact_matches.append(item)
        elif q_norm and ttl_norm.startswith(q_norm):
            startswith_matches.append(item)

    chosen = (
        exact_matches[0]
        if exact_matches
        else startswith_matches[0]
        if startswith_matches
        else items[0]
    )

    try:
        emitent_id = int(chosen["id"])
    except Exception:  # noqa: BLE001
        return None

    return EmitentMatch(
        emitent_id=emitent_id,
        ttl=str(chosen.get("ttl") or query).strip(),
        status_id=None if chosen.get("emitent_statuses_id") is None else str(chosen.get("emitent_statuses_id")),
    )


def get_emitent_id(emitent_name: str, cache: dict) -> EmitentMatch | None:
    cached = cache.setdefault("emitents", {}).get(emitent_name)
    if cached:
        return EmitentMatch(
            emitent_id=int(cached["emitent_id"]),
            ttl=str(cached.get("ttl") or emitent_name),
            status_id=None if cached.get("status_id") is None else str(cached.get("status_id")),
        )

    payload = {
        "term": quote(emitent_name, safe=""),
        "emitent_stop_statuses_ids": [],
    }

    resp = request_with_retry(
        "POST",
        SUGGEST_COMPANIES_URL,
        json=payload,
        headers={
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": SITE_BASE_URL,
            "Referer": f"{SITE_BASE_URL}/",
        },
    )
    data = parse_json_response(resp)
    items = data.get("items", [])
    match = choose_best_emitent_match(emitent_name, items)

    if match:
        cache.setdefault("emitents", {})[emitent_name] = {
            "emitent_id": match.emitent_id,
            "ttl": match.ttl,
            "status_id": match.status_id,
        }
        save_cache(cache)

    return match


# ──────────────────────────────────────────────────────────────────────────────
# API: news
# ──────────────────────────────────────────────────────────────────────────────

def build_search_url(emitent_id: int, emitent_name: str, date_str: str, page: int = 1) -> str:
    query = "&".join(
        [
            f"emitent_id[]={quote(str(emitent_id), safe='')}",
            f"emitent_id_name[]={quote(emitent_name, safe='')}",
            f"date_min={quote(date_str, safe='')}",
            f"date_max={quote(date_str, safe='')}",
            f"page={page}",
        ]
    )
    return f"{NEWS_SEARCH_BASE_URL}?{query}"


def build_news_api_payload(emitent_id: int, date_str: str, page: int) -> dict:
    return {
        "filters": [
            {"field": "show_only_loans", "operator": "eq", "value": 0},
            {"field": "emitent_id", "operator": "in", "value": [str(emitent_id)]},
            {"field": "date", "operator": "ge", "value": date_str},
            {"field": "date", "operator": "le", "value": date_str},
        ],
        "lang": "rus",
        "quantity": {
            "offset": (page - 1) * NEWS_PAGE_LIMIT,
            "limit": NEWS_PAGE_LIMIT,
            "page": page,
        },
        "sorting": [],
    }


def fetch_news_items_for_group(emitent_name: str, emitent_id: int, date_str: str, cache: dict) -> list[dict]:
    key = group_key(emitent_id, date_str)
    cached_items = cache.setdefault("search_results", {}).get(key)
    if cached_items is not None:
        return cached_items

    all_items: list[dict] = []
    total_expected: int | None = None

    for page in range(1, MAX_API_PAGES + 1):
        search_url = build_search_url(emitent_id, emitent_name, date_str, page=page)
        payload = build_news_api_payload(emitent_id, date_str, page)

        resp = request_with_retry(
            "POST",
            NEWS_API_URL,
            json=payload,
            headers={
                "Accept": "application/json, text/plain, */*",
                "Content-Type": "application/json",
                "Origin": SITE_BASE_URL,
                "Referer": search_url,
            },
        )

        try:
            data = parse_json_response(resp)
        except Exception as exc:  # noqa: BLE001
            debug_path = dump_debug_json(f"failed_parse_{emitent_id}_{date_str}_page{page}", {
                "error": str(exc),
                "status_code": resp.status_code,
                "headers": dict(resp.headers),
                "text": resp.text,
                "request_payload": payload,
                "search_url": search_url,
            })
            raise RuntimeError(f"Не удалось распарсить JSON API news. Debug: {debug_path}") from exc

        error_block = data.get("error") or {}
        err_no = error_block.get("err_no", 0)
        err_str = str(error_block.get("err_str") or "").strip()
        if err_no or err_str:
            debug_path = dump_debug_json(f"api_error_{emitent_id}_{date_str}_page{page}", data)
            raise RuntimeError(f"API /api/news/ вернул ошибку err_no={err_no} err_str={err_str!r}. Debug: {debug_path}")

        response_block = data.get("response") or {}
        items = response_block.get("items") or []

        if total_expected is None:
            total_expected = response_block.get("total")
            if total_expected is not None:
                total_expected = int(total_expected)

        if not items:
            break

        all_items.extend(items)
        maybe_sleep()

        if total_expected is not None and len(all_items) >= total_expected:
            break

        if len(items) < NEWS_PAGE_LIMIT:
            break

    cache.setdefault("search_results", {})[key] = all_items
    save_cache(cache)
    return all_items


# ──────────────────────────────────────────────────────────────────────────────
# Основной сценарий
# ──────────────────────────────────────────────────────────────────────────────

def main() -> None:
    base_dir = Path(__file__).resolve().parent
    os.chdir(base_dir)

    if not Path(INPUT_FILE).exists():
        raise FileNotFoundError(f"Не найден входной файл: {base_dir / INPUT_FILE}")

    print(f"Cookie: {'загружены' if 'Cookie' in session.headers else 'не заданы'}")

    cache = load_cache()
    conn = init_db(DB_FILE)

    target_rows = read_target_rows(INPUT_FILE)
    print(f"Строк к обработке (pb_tizer пустой): {len(target_rows)}")

    grouped: dict[tuple[str, str], list[IssueRow]] = defaultdict(list)
    for issue in target_rows:
        grouped[(issue.emitent, issue.bookbuilding_date)].append(issue)

    total_groups = len(grouped)
    print(f"Уникальных групп (Эмитент + Дата букбилдинга): {total_groups}\n")

    for idx, ((emitent_name, bookbuilding_date), issues) in enumerate(sorted(grouped.items()), start=1):
        print(f"[{idx}/{total_groups}] {emitent_name} | {bookbuilding_date} | выпусков: {len(issues)}")

        try:
            match = get_emitent_id(emitent_name, cache)
            if not match:
                print("    [skip] emitent_id не найден")
                continue

            print(f"    emitent_id={match.emitent_id} | ttl={match.ttl}")
            maybe_sleep()

            search_url = build_search_url(match.emitent_id, match.ttl, bookbuilding_date, page=1)
            news_items = fetch_news_items_for_group(match.ttl, match.emitent_id, bookbuilding_date, cache)
            upsert_search(conn, emitent_name, match.emitent_id, bookbuilding_date, search_url, len(news_items))

            print(f"    найдено новостей: {len(news_items)}")
            if not news_items:
                continue

            for item in news_items:
                article_id = save_article(conn, item)
                article_url = str(item.get("cb_link") or f"{SITE_BASE_URL}/news/{article_id}/")
                for issue in issues:
                    link_issue_to_article(conn, issue, match.emitent_id, article_id, article_url)

        except Exception as exc:  # noqa: BLE001
            print(f"    [error] {exc}")

    conn.close()
    save_cache(cache)
    print(f"\nГотово. Результаты сохранены в {base_dir / DB_FILE}")


if __name__ == "__main__":
    main()
