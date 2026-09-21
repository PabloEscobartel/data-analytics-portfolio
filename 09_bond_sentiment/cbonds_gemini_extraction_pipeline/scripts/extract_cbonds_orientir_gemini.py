#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sqlite3
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from statistics import median
from typing import TYPE_CHECKING, Any, Dict, Iterable, List, Optional

import openpyxl

try:
    from google import genai
    from google.genai import types
except ImportError:  # pragma: no cover
    genai = None
    types = None

if TYPE_CHECKING:  # pragma: no cover
    from google.genai import Client as GeminiClient
else:  # pragma: no cover
    GeminiClient = Any

# -----------------------------------------------------------------------------
# Defaults
# -----------------------------------------------------------------------------

DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite-preview")
DEFAULT_MAX_GROUPS_PER_BATCH = 8
DEFAULT_MAX_BATCH_CHARS = 22000
DEFAULT_MAX_REQUESTS = 490
DEFAULT_DELAY = 5.0
DEFAULT_MAX_RETRIES = 5
DEFAULT_MAX_OUTPUT_TOKENS = 8192
DEFAULT_SAVE_EVERY = 1
DEFAULT_WRITE_EVERY = 1

REQUIRED_HEADERS = ["Бумага", "ISIN", "Эмитент", "Дата букбилдинга", "pb_tizer"]
OPTIONAL_FLOATING_HEADERS = ["Floating_rate", "floating_rate"]

OUTPUT_COLUMNS = [
    "gemini_orient_type",
    "gemini_orient_value",
    "gemini_orient_finality",
    "gemini_coupon_guide",
    "gemini_spread_guide",
    "gemini_source_title",
    "gemini_source_url",
    "gemini_source_article_ref",
    "gemini_value_text",
    "gemini_evidence_text",
    "gemini_reason_short",
    "gemini_confidence",
    "gemini_model_name",
    "gemini_model_version",
    "gemini_status",
]

SYSTEM_PROMPT = """Ты аналитик российского первичного долгового рынка.

Тебе передают несколько ГРУПП новостей Cbonds.
Каждая группа соответствует одному сочетанию (эмитент, дата букбилдинга),
и внутри группы может быть один или несколько выпусков облигаций.

Твоя задача:
для КАЖДОГО выпуска отдельно определить ОРИЕНТИР по купону или спреду (марже),
который использовался в маркетинге книги заявок ДО размещения.

КРИТИЧЕСКИ ВАЖНО:
НУЖНО ИСКАТЬ ТОЛЬКО ОРИЕНТИР.
НЕ НУЖНО возвращать фактическую ставку купона, установленную по итогам размещения,
закрытия книги, торгов или начала обращения, если в тексте нет явного указания,
что это именно ориентир.

================================
ЧТО СЧИТАЕТСЯ ЦЕЛЕВЫМ ЗНАЧЕНИЕМ
================================

Подходит только то значение, которое явно описано как ориентир / guidance / маркетируемый уровень.

Подходящие формулировки:
- "ориентир ставки 1 купона"
- "ориентир 1 купона"
- "ориентир купона"
- "финальный ориентир ставки 1 купона"
- "финальный ориентир 1 купона"
- "финальный ориентир купона"
- "ориентир спреда"
- "финальный ориентир спреда"
- "ориентир маржи"
- "финальный ориентир маржи"
- "не выше 17.5% годовых"
- "не выше 375 б.п."
- "диапазон ... %"
- "книга заявок открыта ... ориентир ..."
- "финальный ориентир установлен на уровне ..."

================================
ЧТО НЕ ПОДХОДИТ
================================

НЕ подходит фактическая ставка купона / итог размещения / результат размещения,
если в фрагменте нет явного признака, что это именно ориентир.

НЕЛЬЗЯ брать как ответ:
- "ставка 1 купона установлена на уровне 20%"
- "ставка купона установлена на уровне 20%"
- "купон установлен на уровне 20%"
- "по итогам букбилдинга ставка составила 20%"
- "размещение прошло со ставкой 20%"
- "эмитент разместил облигации со ставкой 20%"
- "ставка определена в размере 20%"
- "значение спреда для купонов установлено ..."
- любые аналогичные формулировки ФАКТИЧЕСКОГО результата,
  если рядом нет явного слова "ориентир" или явной конструкции типа
  "не выше / не ниже / диапазон" в контексте книги заявок

ОСОБОЕ ПРАВИЛО:
Фраза "ставка 1 купона установлена ..." сама по себе НЕ является ориентиром.
Фраза "значение спреда установлено ..." сама по себе НЕ является ориентиром.
Их можно использовать только если в том же фрагменте явно сказано,
что это "ориентир" или "финальный ориентир".

================================
ПРИОРИТЕТ ВЫБОРА
================================

1) Если есть ЯВНО финальный ориентир, нужно брать его.

Признаки финального ориентира:
- "финальный ориентир"
- "финальный ориентир установлен"
- "ориентир установлен"
- "финальный ориентир спреда установлен"
- "финальный ориентир ставки 1 купона"

2) Если финального ориентира нет, нужно взять просто ориентир.

3) Если ориентиров несколько и финального нет, бери самый поздний / последний
по времени релевантный ориентир.

================================
КАК ОПРЕДЕЛЯТЬ ТИП
================================

- Для floating_rate = Да предпочитай spread / margin, если они явно указаны как ориентир.
- Для floating_rate = Нет предпочитай coupon, если он явно указан как ориентир.
- Но если для конкретного выпуска явно указан только другой тип, можно вернуть его.

================================
ПРАВИЛА ПО ВЫПУСКАМ
================================

- Возвращай ровно один объект на каждый выпуск из списка ISSUES.
- Если по выпуску нельзя уверенно определить ИМЕННО ОРИЕНТИР,
  верни orient_type="none" и orient_value=0.
- Если одна и та же новость явно относится к нескольким перечисленным выпускам
  и значение у них одно и то же, можно повторить его для всех.
- Если в группе несколько выпусков, а значение явно указано только для одного,
  не распространяй его на остальные.

================================
ФОРМАТ ПОЛЕЙ
================================

- orient_type:
  - "coupon"
  - "spread"
  - "none"

- orient_value:
  - для coupon — число в процентах, например 17 или 16.5
  - для spread — число в базисных пунктах, например 350
  - для none — 0

- finality:
  - "final" — если это финальный ориентир
  - "guide" — если это просто ориентир, но не финализирован
  - "none" — если ничего не найдено

- value_text:
  короткий точный фрагмент ТОЛЬКО с самим значением,
  например "17% годовых" или "350 б.п."

- evidence_text:
  короткий точный фрагмент 1–2 предложения из новости,
  где явно видно, что это именно ориентир, а не фактическая ставка

- reason_short:
  очень короткая причина, максимум 25 слов

================================
КОНТРПРИМЕРЫ
================================

НЕПРАВИЛЬНО:
Текст: "Ставка 1 купона установлена на уровне 20% годовых."
Ответ: orient_type="none", orient_value=0, finality="none"
Потому что здесь нет явного указания, что 20% — это ориентир.

ПРАВИЛЬНО:
Текст: "Финальный ориентир ставки 1 купона установлен на уровне 17% годовых."
Ответ: orient_type="coupon", orient_value=17, finality="final"

ПРАВИЛЬНО:
Текст: "Новый ориентир спреда к КС установлен на уровне не выше 375 б.п."
Ответ: orient_type="spread", orient_value=375, finality="guide"

================================
ФОРМАТ ОТВЕТА
================================

Возвращай только JSON-массив.
Никакого markdown.
Никакого текста вне JSON.
"""

RESULTS_JSON_SCHEMA: Dict[str, Any] = {
    "type": "array",
    "items": {
        "type": "object",
        "properties": {
            "excel_row": {"type": "integer", "minimum": 2},
            "orient_type": {"type": "string", "enum": ["coupon", "spread", "none"]},
            "orient_value": {"type": "number", "minimum": 0},
            "finality": {"type": "string", "enum": ["final", "guide", "none"]},
            "value_text": {"type": "string"},
            "evidence_text": {"type": "string"},
            "source_article_ref": {"type": "string"},
            "source_url": {"type": "string"},
            "source_title": {"type": "string"},
            "confidence": {"type": "number", "minimum": 0, "maximum": 1},
            "reason_short": {"type": "string"},
        },
        "required": [
            "excel_row",
            "orient_type",
            "orient_value",
            "finality",
            "value_text",
            "evidence_text",
            "source_article_ref",
            "source_url",
            "source_title",
            "confidence",
            "reason_short",
        ],
        "additionalProperties": False,
    },
}

LAST_REQUEST_TS = 0.0


# -----------------------------------------------------------------------------
# Data models
# -----------------------------------------------------------------------------

@dataclass
class IssueInfo:
    excel_row: int
    paper: str
    isin: str
    emitent: str
    bookbuilding_date: str
    floating_rate: str
    pb_tizer: str


@dataclass
class ArticleInfo:
    ref: str
    url: str
    title: str
    date_time: str
    text: str
    sort_ts: float


@dataclass
class GroupInfo:
    group_uid: str
    emitent: str
    bookbuilding_date: str
    issues: List[IssueInfo] = field(default_factory=list)
    articles: List[ArticleInfo] = field(default_factory=list)
    estimated_chars: int = 0


@dataclass
class BatchGroup:
    group_uid: str
    payload_text: str
    estimated_chars: int


@dataclass
class CallResult:
    parsed: Any
    raw_text: str
    response_model_version: Optional[str]
    usage_metadata: Any


# -----------------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------------


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = re.sub(r"\s+", " ", text)
    return text.strip()



def normalize_key_text(value: Any) -> str:
    return normalize_text(value).casefold().replace("ё", "е")



def safe_float_01(value: Any, default: float = 0.5) -> float:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return default
    if math.isnan(x):
        return default
    return max(0.0, min(1.0, x))



def safe_nonneg_number(value: Any, default: float = 0.0) -> float:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return default
    if math.isnan(x):
        return default
    return max(0.0, x)



def truncate_words(text: str, max_words: int) -> str:
    words = normalize_text(text).split()
    if len(words) <= max_words:
        return " ".join(words)
    return " ".join(words[:max_words])



def parse_any_date(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    text = normalize_text(value)
    if not text:
        return None

    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None



def parse_datetime_to_ts(text: str) -> float:
    text = normalize_text(text)
    if not text:
        return 0.0
    for fmt in (
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%dT%H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt).timestamp()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).timestamp()
    except Exception:
        return 0.0



def group_uid(emitent: str, bookbuilding_date: str) -> str:
    return f"{emitent}||{bookbuilding_date}"



def load_json(path: Path) -> Dict[str, Any]:
    if path.exists():
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}



def save_json(path: Path, payload: Dict[str, Any]) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)



def extract_json_array(raw_text: str) -> str:
    raw_text = (raw_text or "").strip()
    if raw_text.startswith("```"):
        raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text, flags=re.IGNORECASE)
        raw_text = re.sub(r"\s*```$", "", raw_text).strip()
    start = raw_text.find("[")
    end = raw_text.rfind("]")
    if start != -1 and end != -1 and end > start:
        return raw_text[start : end + 1]
    raise ValueError("JSON array not found in model output")



def extract_retry_delay_seconds(exc: Exception) -> Optional[float]:
    text = str(exc)
    patterns = [
        r"retry in\s+([0-9]+(?:\.[0-9]+)?)s",
        r"retry_delay\s*\{\s*seconds:\s*([0-9]+)",
        r"Retry in\s+([0-9]+(?:\.[0-9]+)?)s",
    ]
    for pattern in patterns:
        m = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if m:
            try:
                return float(m.group(1))
            except ValueError:
                return None
    return None



def sleep_to_respect_min_interval(min_interval: float) -> None:
    global LAST_REQUEST_TS
    now = time.time()
    elapsed = now - LAST_REQUEST_TS
    if LAST_REQUEST_TS and elapsed < min_interval:
        time.sleep(min_interval - elapsed)



def mark_request_sent() -> None:
    global LAST_REQUEST_TS
    LAST_REQUEST_TS = time.time()



def usage_counts(usage_metadata: Any) -> tuple[Optional[int], Optional[int], Optional[int]]:
    if usage_metadata is None:
        return None, None, None
    prompt_tokens = getattr(usage_metadata, "prompt_token_count", None)
    output_tokens = getattr(usage_metadata, "candidates_token_count", None)
    total_tokens = getattr(usage_metadata, "total_token_count", None)
    return prompt_tokens, output_tokens, total_tokens


# -----------------------------------------------------------------------------
# Excel
# -----------------------------------------------------------------------------


def strip_workbook_images(wb: openpyxl.Workbook) -> None:
    for ws in wb.worksheets:
        if hasattr(ws, "_images"):
            ws._images = []



def find_target_sheet(wb: openpyxl.Workbook):
    required = {normalize_key_text(x) for x in REQUIRED_HEADERS}
    for ws in wb.worksheets:
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        normalized_headers = {normalize_key_text(h) for h in headers if h is not None}
        if required.issubset(normalized_headers):
            return ws
    raise RuntimeError(f"Не найден лист с колонками: {REQUIRED_HEADERS}")



def build_header_index(ws) -> Dict[str, int]:
    return {normalize_text(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}



def find_optional_column(headers: Dict[str, int], variants: Iterable[str]) -> Optional[int]:
    normalized_to_idx = {normalize_key_text(k): v for k, v in headers.items()}
    for name in variants:
        idx = normalized_to_idx.get(normalize_key_text(name))
        if idx is not None:
            return idx
    return None



def load_excel_issues(input_path: Path, only_blank_pb_tizer: bool) -> tuple[openpyxl.Workbook, Any, Dict[int, IssueInfo]]:
    wb = openpyxl.load_workbook(input_path)
    strip_workbook_images(wb)
    ws = find_target_sheet(wb)
    headers = build_header_index(ws)
    floating_col = find_optional_column(headers, OPTIONAL_FLOATING_HEADERS)

    issues: Dict[int, IssueInfo] = {}
    for row_idx in range(2, ws.max_row + 1):
        paper = normalize_text(ws.cell(row_idx, headers["Бумага"]).value)
        isin = normalize_text(ws.cell(row_idx, headers["ISIN"]).value)
        emitent = normalize_text(ws.cell(row_idx, headers["Эмитент"]).value)
        bookbuilding_date = parse_any_date(ws.cell(row_idx, headers["Дата букбилдинга"]).value)
        pb_tizer = normalize_text(ws.cell(row_idx, headers["pb_tizer"]).value)
        floating_rate = normalize_text(ws.cell(row_idx, floating_col).value) if floating_col else ""

        if only_blank_pb_tizer and pb_tizer:
            continue
        if not emitent or not bookbuilding_date:
            continue

        issues[row_idx] = IssueInfo(
            excel_row=row_idx,
            paper=paper,
            isin=isin,
            emitent=emitent,
            bookbuilding_date=bookbuilding_date,
            floating_rate=floating_rate,
            pb_tizer=pb_tizer,
        )

    return wb, ws, issues



def ensure_output_columns(ws) -> Dict[str, int]:
    headers = build_header_index(ws)
    for col_name in OUTPUT_COLUMNS:
        if col_name not in headers:
            col_idx = ws.max_column + 1
            ws.cell(1, col_idx, col_name)
            headers[col_name] = col_idx
    return headers



def apply_cache_to_workbook(ws, headers: Dict[str, int], cache: Dict[str, Any]) -> None:
    for group_payload in cache.values():
        if not isinstance(group_payload, dict):
            continue
        results = group_payload.get("results") or []
        status = normalize_text(group_payload.get("status")) or "unknown"
        model_name = normalize_text(group_payload.get("model_name"))
        model_version = normalize_text(group_payload.get("model_version"))

        for item in results:
            try:
                row_idx = int(item.get("excel_row"))
            except (TypeError, ValueError):
                continue
            if row_idx < 2 or row_idx > ws.max_row:
                continue

            orient_type = normalize_text(item.get("orient_type")).lower() or "none"
            orient_value = safe_nonneg_number(item.get("orient_value"), 0.0)
            finality = normalize_text(item.get("finality")).lower() or "none"
            value_text = normalize_text(item.get("value_text"))
            evidence_text = normalize_text(item.get("evidence_text"))
            source_ref = normalize_text(item.get("source_article_ref"))
            source_title = normalize_text(item.get("source_title"))
            source_url = normalize_text(item.get("source_url"))
            reason_short = normalize_text(item.get("reason_short"))
            confidence = safe_float_01(item.get("confidence"), 0.0)

            coupon_value = orient_value if orient_type == "coupon" and orient_value > 0 else None
            spread_value = orient_value if orient_type == "spread" and orient_value > 0 else None

            ws.cell(row_idx, headers["gemini_orient_type"], orient_type)
            ws.cell(row_idx, headers["gemini_orient_value"], orient_value if orient_value > 0 else None)
            ws.cell(row_idx, headers["gemini_orient_finality"], finality)
            ws.cell(row_idx, headers["gemini_coupon_guide"], coupon_value)
            ws.cell(row_idx, headers["gemini_spread_guide"], spread_value)
            ws.cell(row_idx, headers["gemini_source_title"], source_title)
            ws.cell(row_idx, headers["gemini_source_url"], source_url)
            ws.cell(row_idx, headers["gemini_source_article_ref"], source_ref)
            ws.cell(row_idx, headers["gemini_value_text"], value_text)
            ws.cell(row_idx, headers["gemini_evidence_text"], evidence_text)
            ws.cell(row_idx, headers["gemini_reason_short"], reason_short)
            ws.cell(row_idx, headers["gemini_confidence"], confidence)
            ws.cell(row_idx, headers["gemini_model_name"], model_name)
            ws.cell(row_idx, headers["gemini_model_version"], model_version)
            ws.cell(row_idx, headers["gemini_status"], status)



def save_partial_workbook(input_wb: openpyxl.Workbook, ws, cache: Dict[str, Any], output_path: Path) -> None:
    headers = ensure_output_columns(ws)
    apply_cache_to_workbook(ws, headers, cache)
    try:
        input_wb.save(output_path)
    except ValueError as exc:
        if "closed file" not in str(exc).lower():
            raise
        strip_workbook_images(input_wb)
        input_wb.save(output_path)


# -----------------------------------------------------------------------------
# SQLite
# -----------------------------------------------------------------------------


def table_columns(conn: sqlite3.Connection, table_name: str) -> List[str]:
    rows = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    return [str(row[1]) for row in rows]



def choose_existing(columns: Iterable[str], candidates: List[str]) -> Optional[str]:
    colset = set(columns)
    for candidate in candidates:
        if candidate in colset:
            return candidate
    return None



def estimate_group_chars(group: GroupInfo) -> int:
    issues_chars = sum(len(issue.paper) + len(issue.isin) + len(issue.emitent) + 80 for issue in group.issues)
    articles_chars = sum(len(a.title) + len(a.date_time) + len(a.url) + len(a.text) + 120 for a in group.articles)
    return issues_chars + articles_chars + 300



def load_groups_from_db(db_path: Path, excel_issues: Dict[int, IssueInfo]) -> List[GroupInfo]:
    if not db_path.exists():
        raise FileNotFoundError(f"DB не найдена: {db_path}")

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    issue_cols = table_columns(conn, "issue_articles")
    article_cols = table_columns(conn, "articles")

    issue_link_col = choose_existing(issue_cols, ["article_url", "cb_link", "url"])
    if not issue_link_col:
        raise RuntimeError("Не удалось найти колонку ссылки на статью в issue_articles")

    article_url_col = choose_existing(article_cols, ["cb_link", "url"])
    text_col = choose_existing(article_cols, ["text", "content_text"])
    title_col = choose_existing(article_cols, ["caption", "title"])
    date_time_col = choose_existing(article_cols, ["date_time", "published_label", "date"])
    article_id_col = choose_existing(article_cols, ["news_id", "id.numeric", "id"])

    if not article_url_col or not text_col:
        raise RuntimeError("Не удалось найти ключевые колонки в articles (ссылка и текст)")

    ia_rows = conn.execute(
        f"SELECT excel_row, {issue_link_col} AS article_url FROM issue_articles"
    ).fetchall()

    article_parts = [
        f"{article_url_col} AS article_url",
        f"{text_col} AS article_text",
        f"{title_col} AS article_title" if title_col else "'' AS article_title",
        f"{date_time_col} AS article_date_time" if date_time_col else "'' AS article_date_time",
        f"{article_id_col} AS article_id" if article_id_col else "NULL AS article_id",
    ]
    article_rows = conn.execute(f"SELECT {', '.join(article_parts)} FROM articles").fetchall()
    conn.close()

    article_map: Dict[str, Dict[str, Any]] = {}
    for row in article_rows:
        url = normalize_text(row["article_url"])
        if not url:
            continue
        article_map[url] = {
            "article_id": row["article_id"],
            "title": normalize_text(row["article_title"]),
            "date_time": normalize_text(row["article_date_time"]),
            "text": normalize_text(row["article_text"]),
        }

    groups: Dict[str, GroupInfo] = {}
    group_issue_seen: Dict[str, set[int]] = {}
    group_article_seen: Dict[str, set[str]] = {}

    for row in ia_rows:
        try:
            excel_row = int(row["excel_row"])
        except (TypeError, ValueError):
            continue
        issue = excel_issues.get(excel_row)
        if issue is None:
            continue

        g_uid = group_uid(issue.emitent, issue.bookbuilding_date)
        group = groups.setdefault(
            g_uid,
            GroupInfo(group_uid=g_uid, emitent=issue.emitent, bookbuilding_date=issue.bookbuilding_date),
        )
        issue_seen = group_issue_seen.setdefault(g_uid, set())
        article_seen = group_article_seen.setdefault(g_uid, set())

        if excel_row not in issue_seen:
            group.issues.append(issue)
            issue_seen.add(excel_row)

        article_url = normalize_text(row["article_url"])
        article_payload = article_map.get(article_url)
        if not article_url or not article_payload or article_url in article_seen:
            continue

        title = article_payload["title"]
        text = article_payload["text"]
        date_time = article_payload["date_time"]
        article_id = article_payload.get("article_id")
        ref = f"A{len(article_seen) + 1}"
        if article_id not in (None, ""):
            ref = f"A{len(article_seen) + 1}|{article_id}"

        group.articles.append(
            ArticleInfo(
                ref=ref,
                url=article_url,
                title=title,
                date_time=date_time,
                text=text,
                sort_ts=parse_datetime_to_ts(date_time),
            )
        )
        article_seen.add(article_url)

    result: List[GroupInfo] = []
    for group in groups.values():
        group.issues.sort(key=lambda x: x.excel_row)
        group.articles.sort(key=lambda x: (x.sort_ts, x.ref))
        group.estimated_chars = estimate_group_chars(group)
        if group.issues and group.articles:
            result.append(group)

    result.sort(key=lambda g: (g.bookbuilding_date, g.emitent))
    return result


# -----------------------------------------------------------------------------
# Prompt / batching
# -----------------------------------------------------------------------------


def build_group_payload_text(group: GroupInfo) -> str:
    issue_lines: List[str] = []
    for issue in group.issues:
        issue_lines.append(
            " | ".join(
                [
                    f"excel_row={issue.excel_row}",
                    f"paper={issue.paper or 'NA'}",
                    f"isin={issue.isin or 'NA'}",
                    f"floating_rate={issue.floating_rate or 'NA'}",
                ]
            )
        )

    article_blocks: List[str] = []
    for article in group.articles:
        article_blocks.append(
            "\n".join(
                [
                    f"[{article.ref}]",
                    f"date_time: {article.date_time or 'NA'}",
                    f"title: {article.title or 'NA'}",
                    f"url: {article.url}",
                    "text:",
                    article.text or "NA",
                ]
            )
        )

    return "\n\n".join(
        [
            "<GROUP>",
            f"group_uid: {group.group_uid}",
            f"emitent: {group.emitent}",
            f"bookbuilding_date: {group.bookbuilding_date}",
            "ISSUES:",
            *issue_lines,
            "ARTICLES:",
            *article_blocks,
            "</GROUP>",
        ]
    )



def prepare_pending_groups(groups: List[GroupInfo], cache: Dict[str, Any]) -> List[BatchGroup]:
    pending: List[BatchGroup] = []
    for group in groups:
        if group.group_uid in cache:
            continue
        payload_text = build_group_payload_text(group)
        pending.append(
            BatchGroup(
                group_uid=group.group_uid,
                payload_text=payload_text,
                estimated_chars=len(payload_text) + 300,
            )
        )
    return pending



def make_batches(groups: List[BatchGroup], max_groups_per_batch: int, max_batch_chars: int) -> List[List[BatchGroup]]:
    batches: List[List[BatchGroup]] = []
    current: List[BatchGroup] = []
    current_chars = 0

    for item in groups:
        if current and (
            len(current) >= max_groups_per_batch
            or current_chars + item.estimated_chars > max_batch_chars
        ):
            batches.append(current)
            current = []
            current_chars = 0
        current.append(item)
        current_chars += item.estimated_chars

    if current:
        batches.append(current)
    return batches



def build_user_prompt(batch: List[BatchGroup]) -> str:
    header = (
        "Ниже несколько групп новостей. "
        "Для каждого выпуска (excel_row) верни ровно один JSON-объект. "
        "Ищи только ориентир, а не фактическую ставку купона по итогам размещения."
    )
    body = "\n\n".join(item.payload_text for item in batch)
    return f"{header}\n\n{body}"


# -----------------------------------------------------------------------------
# Gemini call
# -----------------------------------------------------------------------------


def parse_model_items(raw_payload: Any) -> List[Dict[str, Any]]:
    if isinstance(raw_payload, list):
        parsed = raw_payload
    elif isinstance(raw_payload, str):
        parsed = json.loads(extract_json_array(raw_payload))
    else:
        raise ValueError("Unsupported model payload type")
    if not isinstance(parsed, list):
        raise ValueError("Model output is not a list")
    return parsed



def normalize_model_item(item: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    if not isinstance(item, dict):
        return None

    try:
        excel_row = int(item.get("excel_row"))
    except (TypeError, ValueError):
        return None

    orient_type = normalize_text(item.get("orient_type")).lower()
    if orient_type not in {"coupon", "spread", "none"}:
        orient_type = "none"

    orient_value = safe_nonneg_number(item.get("orient_value"), 0.0)
    finality = normalize_text(item.get("finality")).lower()
    if finality not in {"final", "guide", "none"}:
        finality = "none"

    if orient_type == "none" or orient_value <= 0:
        orient_type = "none"
        orient_value = 0.0
        finality = "none"

    value_text = normalize_text(item.get("value_text"))[:240]
    evidence_text = normalize_text(item.get("evidence_text"))[:800]
    source_article_ref = normalize_text(item.get("source_article_ref"))[:120]
    source_url = normalize_text(item.get("source_url"))[:500]
    source_title = normalize_text(item.get("source_title"))[:500]
    confidence = safe_float_01(item.get("confidence"), 0.5)
    reason_short = truncate_words(normalize_text(item.get("reason_short")), 25)[:300]

    return {
        "excel_row": excel_row,
        "orient_type": orient_type,
        "orient_value": orient_value,
        "finality": finality,
        "value_text": value_text,
        "evidence_text": evidence_text,
        "source_article_ref": source_article_ref,
        "source_url": source_url,
        "source_title": source_title,
        "confidence": confidence,
        "reason_short": reason_short,
    }



def generate_structured(
    client: GeminiClient,
    model_name: str,
    prompt: str,
    max_retries: int,
    max_output_tokens: int,
    min_interval: float,
) -> CallResult:
    last_exc: Optional[Exception] = None
    for attempt in range(1, max_retries + 1):
        try:
            sleep_to_respect_min_interval(min_interval)
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=SYSTEM_PROMPT,
                    temperature=0,
                    max_output_tokens=max_output_tokens,
                    response_mime_type="application/json",
                    response_json_schema=RESULTS_JSON_SCHEMA,
                ),
            )
            mark_request_sent()
            raw_text = getattr(response, "text", "") or ""
            parsed = getattr(response, "parsed", None)
            usage = getattr(response, "usage_metadata", None)
            version = getattr(response, "model_version", None)
            return CallResult(
                parsed=parsed,
                raw_text=raw_text,
                response_model_version=version,
                usage_metadata=usage,
            )
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
            wait = extract_retry_delay_seconds(exc)
            if wait is None:
                wait = min(5 * (2 ** (attempt - 1)), 60)
            if attempt < max_retries:
                print(f"  retry {attempt}/{max_retries} after {wait:.1f}s: {exc}")
                time.sleep(wait)
    raise RuntimeError(str(last_exc) if last_exc else "Unknown API error")



def classify_batch(
    client: GeminiClient,
    model_name: str,
    batch_groups: List[BatchGroup],
    group_lookup: Dict[str, GroupInfo],
    max_retries: int,
    max_output_tokens: int,
    min_interval: float,
) -> Dict[str, Dict[str, Any]]:
    prompt = build_user_prompt(batch_groups)
    expected_issue_rows: Dict[int, str] = {}
    for batch_group in batch_groups:
        group = group_lookup[batch_group.group_uid]
        for issue in group.issues:
            expected_issue_rows[issue.excel_row] = batch_group.group_uid

    last_raw_excerpt: Optional[str] = None
    try:
        call_result = generate_structured(
            client=client,
            model_name=model_name,
            prompt=prompt,
            max_retries=max_retries,
            max_output_tokens=max_output_tokens,
            min_interval=min_interval,
        )
        prompt_tokens, output_tokens, total_tokens = usage_counts(call_result.usage_metadata)
        parsed_items = parse_model_items(call_result.parsed if call_result.parsed is not None else call_result.raw_text)
        last_raw_excerpt = (call_result.raw_text or "")[:2500] if call_result.raw_text else None

        per_group_results: Dict[str, Dict[int, Dict[str, Any]]] = {bg.group_uid: {} for bg in batch_groups}
        for raw_item in parsed_items:
            norm = normalize_model_item(raw_item)
            if norm is None:
                continue
            row_idx = norm["excel_row"]
            g_uid = expected_issue_rows.get(row_idx)
            if not g_uid:
                continue
            per_group_results[g_uid][row_idx] = norm

        out: Dict[str, Dict[str, Any]] = {}
        for batch_group in batch_groups:
            group = group_lookup[batch_group.group_uid]
            found = per_group_results.get(batch_group.group_uid, {})
            results: List[Dict[str, Any]] = []
            status = "classified"
            missing_rows: List[int] = []

            for issue in group.issues:
                item = found.get(issue.excel_row)
                if item is None:
                    missing_rows.append(issue.excel_row)
                    results.append(
                        {
                            "excel_row": issue.excel_row,
                            "paper": issue.paper,
                            "isin": issue.isin,
                            "orient_type": "none",
                            "orient_value": 0.0,
                            "finality": "none",
                            "value_text": "",
                            "evidence_text": "",
                            "source_article_ref": "",
                            "source_url": "",
                            "source_title": "",
                            "confidence": 0.0,
                            "reason_short": "missing_item_in_response",
                        }
                    )
                    status = "missing_item_in_response"
                else:
                    enriched = dict(item)
                    enriched["paper"] = issue.paper
                    enriched["isin"] = issue.isin
                    results.append(enriched)

            out[batch_group.group_uid] = {
                "status": status,
                "results": results,
                "model_name": model_name,
                "model_version": call_result.response_model_version,
                "prompt_tokens": prompt_tokens,
                "output_tokens": output_tokens,
                "total_tokens": total_tokens,
                "raw_response_excerpt": last_raw_excerpt,
                "error_message": None if not missing_rows else f"Missing rows: {missing_rows}",
            }
        return out

    except Exception as exc:  # noqa: BLE001
        err = f"{type(exc).__name__}: {exc}"
        return {
            batch_group.group_uid: {
                "status": "api_error",
                "results": [
                    {
                        "excel_row": issue.excel_row,
                        "paper": issue.paper,
                        "isin": issue.isin,
                        "orient_type": "none",
                        "orient_value": 0.0,
                        "finality": "none",
                        "value_text": "",
                        "evidence_text": "",
                        "source_article_ref": "",
                        "source_url": "",
                        "source_title": "",
                        "confidence": 0.0,
                        "reason_short": "api_error",
                    }
                    for issue in group_lookup[batch_group.group_uid].issues
                ],
                "model_name": model_name,
                "model_version": None,
                "prompt_tokens": None,
                "output_tokens": None,
                "total_tokens": None,
                "raw_response_excerpt": last_raw_excerpt,
                "error_message": err,
            }
            for batch_group in batch_groups
        }


# -----------------------------------------------------------------------------
# Reporting
# -----------------------------------------------------------------------------


def cache_summary(cache: Dict[str, Any]) -> Dict[str, int]:
    summary: Dict[str, int] = {}
    for payload in cache.values():
        status = "unknown"
        if isinstance(payload, dict):
            status = str(payload.get("status") or "unknown")
        summary[status] = summary.get(status, 0) + 1
    return summary



def print_summary(cache: Dict[str, Any]) -> None:
    summary = cache_summary(cache)
    print("\n" + "=" * 72)
    print("SUMMARY")
    print("=" * 72)
    for key in sorted(summary):
        print(f"{key}: {summary[key]}")

    issue_count = 0
    extracted_count = 0
    coupon_count = 0
    spread_count = 0
    final_count = 0
    for payload in cache.values():
        if not isinstance(payload, dict):
            continue
        for item in payload.get("results") or []:
            issue_count += 1
            orient_type = item.get("orient_type")
            orient_value = safe_nonneg_number(item.get("orient_value"), 0.0)
            if orient_type in {"coupon", "spread"} and orient_value > 0:
                extracted_count += 1
                if orient_type == "coupon":
                    coupon_count += 1
                if orient_type == "spread":
                    spread_count += 1
            if item.get("finality") == "final":
                final_count += 1

    print(f"issue_results: {issue_count}")
    print(f"nonzero_extracted: {extracted_count}")
    print(f"coupon: {coupon_count}")
    print(f"spread: {spread_count}")
    print(f"final: {final_count}")


# -----------------------------------------------------------------------------
# CLI / main
# -----------------------------------------------------------------------------


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Extract coupon/spread guides from Cbonds news with Gemini.")
    p.add_argument("--input", type=Path, default=Path("bonds_ext_enriched_rusbonds_orientir.xlsx"))
    p.add_argument("--output", type=Path, default=Path("bonds_ext_enriched_rusbonds_orientir_gemini.xlsx"))
    p.add_argument("--db", type=Path, default=Path("cbonds_news.sqlite3"))
    p.add_argument("--cache", type=Path, default=Path("gemini_orientir_cache.json"))
    p.add_argument("--model", type=str, default=DEFAULT_MODEL)
    p.add_argument("--max-groups-per-batch", type=int, default=DEFAULT_MAX_GROUPS_PER_BATCH)
    p.add_argument("--max-batch-chars", type=int, default=DEFAULT_MAX_BATCH_CHARS)
    p.add_argument("--max-requests", type=int, default=DEFAULT_MAX_REQUESTS)
    p.add_argument("--delay", type=float, default=DEFAULT_DELAY)
    p.add_argument("--max-retries", type=int, default=DEFAULT_MAX_RETRIES)
    p.add_argument("--max-output-tokens", type=int, default=DEFAULT_MAX_OUTPUT_TOKENS)
    p.add_argument("--save-every", type=int, default=DEFAULT_SAVE_EVERY)
    p.add_argument("--write-every", type=int, default=DEFAULT_WRITE_EVERY)
    p.add_argument("--api-key", type=str, default=os.getenv("GEMINI_API_KEY", ""))
    p.add_argument("--no-confirm", action="store_true")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--only-blank-pb-tizer", action="store_true", default=True)
    p.add_argument("--include-nonblank-pb-tizer", action="store_true")
    return p



def main() -> None:
    args = build_arg_parser().parse_args()
    if args.include_nonblank_pb_tizer:
        args.only_blank_pb_tizer = False

    wb, ws, excel_issues = load_excel_issues(args.input, only_blank_pb_tizer=args.only_blank_pb_tizer)
    groups = load_groups_from_db(args.db, excel_issues)
    group_lookup = {g.group_uid: g for g in groups}

    cache = load_json(args.cache)
    pending_groups = prepare_pending_groups(groups, cache)
    batches = make_batches(pending_groups, args.max_groups_per_batch, args.max_batch_chars)
    if len(batches) > args.max_requests:
        batches = batches[: args.max_requests]
        print(f"⚠ Limited to first {args.max_requests} requests due to max_requests.")

    total_issues = sum(len(g.issues) for g in groups)
    total_articles = sum(len(g.articles) for g in groups)

    print("=" * 72)
    print("Gemini extraction of coupon/spread guides from Cbonds news")
    print("=" * 72)
    print(f"Input workbook: {args.input}")
    print(f"Output workbook: {args.output}")
    print(f"DB: {args.db}")
    print(f"Cache: {args.cache}")
    print(f"Model: {args.model}")
    print(f"Workbook rows eligible: {len(excel_issues):,}")
    print(f"Groups with news in DB: {len(groups):,}")
    print(f"Issues with news in DB: {total_issues:,}")
    print(f"Articles linked in DB: {total_articles:,}")
    print(f"Cached groups already present: {len(cache):,}")
    print(f"Pending groups: {len(pending_groups):,}")
    print(f"Planned batches this run: {len(batches):,}")
    print(f"Delay between requests: {args.delay:.1f}s")

    if batches:
        batch_sizes = [len(b) for b in batches]
        batch_chars = [sum(x.estimated_chars for x in b) for b in batches]
        print(
            "Batch size stats (groups): "
            f"min={min(batch_sizes)}, median={int(median(batch_sizes))}, max={max(batch_sizes)}"
        )
        print(
            "Batch char stats:          "
            f"min={min(batch_chars)}, median={int(median(batch_chars))}, max={max(batch_chars)}"
        )

    if args.dry_run:
        save_partial_workbook(wb, ws, cache, args.output)
        print("\nDry run complete. Workbook written from existing cache only.")
        return

    if genai is None or types is None:
        raise ImportError("Не найден пакет google-genai. Установи его: pip install google-genai")
    if not args.api_key:
        raise ValueError("API key not found. Pass --api-key or set GEMINI_API_KEY.")

    if batches and not args.no_confirm:
        answer = input(f"Send {len(batches)} requests to Gemini? (y/n): ").strip().lower()
        if answer != "y":
            print("Cancelled by user.")
            return

    client = genai.Client(api_key=args.api_key)

    try:
        for idx, batch in enumerate(batches, start=1):
            chars_now = sum(x.estimated_chars for x in batch)
            print(f"\nBatch {idx}/{len(batches)} | groups={len(batch)} | est_chars={chars_now}")
            batch_results = classify_batch(
                client=client,
                model_name=args.model,
                batch_groups=batch,
                group_lookup=group_lookup,
                max_retries=args.max_retries,
                max_output_tokens=args.max_output_tokens,
                min_interval=args.delay,
            )
            cache.update(batch_results)

            ok_now = sum(1 for x in batch_results.values() if x.get("status") == "classified")
            print(f"  classified_groups: {ok_now} | non-classified: {len(batch_results) - ok_now}")

            if idx % args.save_every == 0:
                save_json(args.cache, cache)
                print(f"  cache saved: {args.cache}")
            if idx % args.write_every == 0:
                save_partial_workbook(wb, ws, cache, args.output)
                print(f"  workbook saved: {args.output}")

    except KeyboardInterrupt:
        print("\nInterrupted by user. Saving progress...")

    finally:
        save_json(args.cache, cache)
        save_partial_workbook(wb, ws, cache, args.output)
        print(f"\nProgress saved to cache: {args.cache}")
        print(f"Progress saved to workbook: {args.output}")

    print_summary(cache)


if __name__ == "__main__":
    main()
