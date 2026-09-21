#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Извлечение ориентира из набора новостей Cbonds вокруг начала букбилдинга.

Логика для каждого выпуска:
1. Берём datetime начала букбилдинга из файла bonds_offer_updated_book_period.xlsx,
   колонка `cbonds_book_start_datetime`.
2. Собираем все новости по эмитенту/выпуску из БД за месяц до начала
   букбилдинга, включая календарную дату букбилдинга, где has_coupon_spread_term = 1.
3. Вместе с каждой новостью передаём Gemini дату/время, заголовок, текст и
   спарсенную эмиссию из article_emissions, если она есть.
4. Gemini сам выбирает самый поздний ориентир до начала букбилдинга; если
   подходящего нет — самый ранний ориентир после начала букбилдинга.
5. Если Gemini сомневается, что ориентир относится к данной эмиссии, он должен
   вернуть needs_attention=true.

Результат пишется в копию Excel с новыми колонками и в SQLite-таблицу
issue_orientir_bookbuilding. Промежуточный кэш сохраняется в JSON, поэтому
скрипт можно безопасно перезапускать.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sqlite3
import time
from dataclasses import dataclass
from datetime import date, datetime, time as dt_time
from pathlib import Path
from typing import Any, Optional

import openpyxl

try:
    from google import genai
    from google.genai import types
except ImportError:  # pragma: no cover
    genai = None
    types = None


DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite-preview")
DEFAULT_DELAY = 5.0
DEFAULT_MAX_RETRIES = 5
DEFAULT_MAX_OUTPUT_TOKENS = 8192
DEFAULT_SAVE_EVERY = 10
DEFAULT_MAX_BATCH_CHARS = 25000
DEFAULT_MAX_ARTICLE_CHARS = 0
DEFAULT_MAX_ISSUE_NEWS = 0

INPUT_FILE = "bonds_offer_updated_book_period.xlsx"
OUTPUT_FILE = "bonds_offer_updated_cbonds_orientir_gemini.xlsx"
DB_FILE = "cbonds_news.sqlite3"
CACHE_FILE = "gemini_bookbuilding_coupon_spread_news_cache.json"
RESULTS_TABLE = "issue_orientir_bookbuilding"

REQUIRED_HEADERS = [
    "Бумага",
    "ISIN",
    "Эмитент",
    "cbonds_book_start_datetime",
    "pb_tizer",
]
OPTIONAL_FLOATING_HEADERS = ["Floating_rate", "floating_rate"]

OUTPUT_COLUMNS = [
    "bb_gemini_orient_type",
    "bb_gemini_orient_value",
    "bb_gemini_orient_finality",
    "bb_gemini_coupon_guide",
    "bb_gemini_spread_guide",
    "bb_gemini_source_timing",
    "bb_gemini_source_title",
    "bb_gemini_source_url",
    "bb_gemini_source_article_id",
    "bb_gemini_source_datetime",
    "bb_gemini_emission_cbonds_id",
    "bb_gemini_emission_name",
    "bb_gemini_needs_attention",
    "bb_gemini_attention_reason",
    "bb_gemini_value_text",
    "bb_gemini_evidence_text",
    "bb_gemini_reason_short",
    "bb_gemini_confidence",
    "bb_gemini_model_name",
    "bb_gemini_model_version",
    "bb_gemini_status",
]

SYSTEM_PROMPT = """Ты аналитик российского первичного долгового рынка.

Тебе передают ОДИН выпуск облигаций и СПИСОК новостей Cbonds за окно вокруг начала букбилдинга.

Задача:
1. Определи, какие новости относятся именно к указанному выпуску/эмиссии.
2. Среди релевантных новостей выбери самый поздний ориентир ДО начала букбилдинга.
3. Если до начала букбилдинга релевантного ориентира нет, выбери самый ранний ориентир ПОСЛЕ начала букбилдинга.
4. Верни ориентир купона или спреда/маржи именно для указанного выпуска.

ЖЁСТКОЕ ПРАВИЛО ВЫБОРА ВРЕМЕНИ:
- Сначала проверь все релевантные новости с published_at < bookbuilding_start.
- Если среди них есть ориентир для целевой эмиссии, выбери самый поздний такой ориентир.
- Если среди них НЕТ ориентира для целевой эмиссии, проверь новости с published_at >= bookbuilding_start.
- В этом случае выбери САМЫЙ РАННИЙ ориентир после bookbuilding_start.
- НЕ выбирай более поздний "финальный ориентир" после bookbuilding_start, если уже есть более ранний релевантный ориентир после bookbuilding_start.
- Слово "финальный" НЕ имеет приоритета над правилом времени.

КРИТИЧЕСКИ ВАЖНО:
- Нужен только ОРИЕНТИР / guidance / маркетируемый уровень, но слово "ориентир" может отсутствовать.
- Ищи смысловой ориентир, а не только буквальное слово "ориентир".
- Подходят предварительные маркетинговые значения для открытия/ведения книги заявок: "не выше", "не более", "диапазон", "составит", "предварительный уровень", "КС + спред не выше".
- Не возвращай фактическую ставку купона по итогам размещения/закрытия книги, если это именно результат размещения, а не маркетинговый уровень для книги.
- Если в новости обсуждаются несколько выпусков, выбери значение, которое относится к целевой эмиссии из ISSUE.paper.
- Поле parsed_emissions у новости — это эмиссия/эмиссии, спарсенные со страницы. Оно может отсутствовать или содержать только одну эмиссию, даже если в тексте обсуждаются две.
- Не полагайся только на parsed_emissions: обязательно читай заголовок и текст новости.
- Если ориентир найден, но связь с целевой эмиссией неоднозначна, верни значение и поставь needs_attention=true.
- Если нельзя понять, что ориентир относится к целевой эмиссии, верни orient_type="none".

Подходящие формулировки:
- "ориентир ставки 1 купона"
- "ориентир купона"
- "ориентир ставки купона"
- "ориентир спреда"
- "ориентир маржи"
- "тип купона ... фиксированный. ... не выше 18% годовых"
- "тип купона ... переменный ... КС Банка России + спред не выше 450 б.п."
- "ставка купона ... не выше 18% годовых" в новости об открытии/предварительном открытии книги заявок
- "спред ... не выше 450 б.п." в новости об открытии/предварительном открытии книги заявок
- "не выше 17.5% годовых"
- "диапазон ставки купона ..."
- "ориентир установлен ..."

Не подходит:
- "ставка 1 купона установлена на уровне 20%"
- "купон установлен на уровне 20%"
- "по итогам букбилдинга ставка составила 20%"
- "значение спреда установлено ..."
если это явно итог/результат размещения или закрытия книги, а не guidance.

Тип:
- coupon: значение в процентах.
- spread: значение в базисных пунктах.
- none: ничего не найдено.
- Если ориентир задан диапазоном, например "7.75–8%" или "275–300 б.п.",
  orient_value должен быть ВЕРХНЕЙ границей диапазона: 8 или 300.
- Если формулировка "не выше / не более", orient_value равен указанному потолку.

finality:
- final: финальный ориентир.
- guide: обычный/предварительный/новый ориентир.
- none: ничего не найдено.
- Это поле только диагностическое. Оно НЕ должно влиять на выбор новости.

source_timing:
- pre: выбранная новость раньше bookbuilding_start.
- post: выбранная новость в bookbuilding_start или позже.
- none: если ничего не найдено.

source_article_id:
- верни точный integer article_id выбранной новости из списка ARTICLES.
- если ориентир не найден, верни 0.

Верни только JSON-массив объектов без markdown. Один объект на каждый выпуск из ISSUES.
"""

RESULT_SCHEMA = {
    "type": "array",
    "items": {
        "type": "object",
        "properties": {
            "excel_row": {"type": "integer", "minimum": 2},
            "orient_type": {"type": "string", "enum": ["coupon", "spread", "none"]},
            "orient_value": {"type": "number", "minimum": 0},
            "finality": {"type": "string", "enum": ["final", "guide", "none"]},
            "value_text": {"type": "string"},
            "evidence_text": {"type": "string"},
            "needs_attention": {"type": "boolean"},
            "attention_reason": {"type": "string"},
            "source_article_id": {"type": "integer"},
            "source_timing": {"type": "string", "enum": ["pre", "post", "none"]},
            "source_emission_name": {"type": "string"},
            "confidence": {"type": "number", "minimum": 0, "maximum": 1},
            "reason_short": {"type": "string"},
        },
        "required": [
            "excel_row",
            "orient_type",
            "orient_value",
            "finality",
            "value_text",
            "evidence_text",
            "needs_attention",
            "attention_reason",
            "source_article_id",
            "source_timing",
            "source_emission_name",
            "confidence",
            "reason_short",
        ],
        "additionalProperties": False,
    },
}


LAST_REQUEST_TS = 0.0


@dataclass
class IssueInfo:
    excel_row: int
    paper: str
    isin: str
    emitent: str
    bookbuilding_start: datetime
    floating_rate: str
    pb_tizer: str
    check_status: str


@dataclass
class CandidateArticle:
    article_row_id: int
    article_id: int | None
    url: str
    title: str
    date_time: str
    published_at: datetime
    text: str
    emission_cbonds_id: int | None
    emission_name: str
    emission_url: str
    timing: str


@dataclass
class BatchItem:
    excel_row: int
    prompt_text: str
    estimated_chars: int


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def normalize_key(value: Any) -> str:
    text = normalize_text(value).casefold().replace("ё", "е")
    text = re.sub(r"\b(пао|ао|ооо|зао|оао)\b", " ", text)
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def safe_float_01(value: Any, default: float = 0.5) -> float:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return default
    if math.isnan(x):
        return default
    return max(0.0, min(1.0, x))


def safe_nonneg_number(value: Any, default: float = 0.0) -> float:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return default
    if math.isnan(x):
        return default
    return max(0.0, x)


def parse_date(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_text(value)
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def parse_time(value: Any) -> dt_time:
    if value is None or normalize_text(value) == "":
        return dt_time(0, 0, 0)
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, dt_time):
        return value.replace(microsecond=0)
    text = normalize_text(value)
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            pass
    return dt_time(0, 0, 0)


def parse_article_datetime(value: Any) -> Optional[datetime]:
    text = normalize_text(value)
    if not text:
        return None
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text)
    except Exception:
        return None


def parse_bookbuilding_start(value: Any) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    text = normalize_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text)
    except Exception:
        return None


def load_json(path: Path) -> dict[str, Any]:
    if path.exists():
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_json(path: Path, payload: dict[str, Any]) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


def build_header_index(ws) -> dict[str, int]:
    return {normalize_text(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}


def find_optional_column(headers: dict[str, int], variants: list[str]) -> int | None:
    normalized = {normalize_key(k): v for k, v in headers.items()}
    for variant in variants:
        idx = normalized.get(normalize_key(variant))
        if idx:
            return idx
    return None


def find_target_sheet(wb: openpyxl.Workbook, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Лист {sheet_name!r} не найден. Доступные листы: {', '.join(wb.sheetnames)}")
        return wb[sheet_name]
    required = {normalize_key(x) for x in REQUIRED_HEADERS}
    for ws in wb.worksheets:
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        if required.issubset({normalize_key(h) for h in headers if h is not None}):
            return ws
    raise RuntimeError(f"Не найден лист с колонками: {REQUIRED_HEADERS}")


def strip_workbook_images(wb: openpyxl.Workbook) -> None:
    for ws in wb.worksheets:
        if hasattr(ws, "_images"):
            ws._images = []


def load_excel_issues(input_path: Path, sheet_name: str | None, only_nonblank_pb_tizer: bool) -> tuple[openpyxl.Workbook, Any, dict[int, IssueInfo]]:
    wb = openpyxl.load_workbook(input_path)
    strip_workbook_images(wb)
    ws = find_target_sheet(wb, sheet_name)
    headers = build_header_index(ws)
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        raise RuntimeError(f"Не найдены колонки на листе {ws.title!r}: {missing}")

    floating_col = find_optional_column(headers, OPTIONAL_FLOATING_HEADERS)
    check_status_col = find_optional_column(headers, ["check_status", "статус проверки"])
    issues: dict[int, IssueInfo] = {}
    for row_idx in range(2, ws.max_row + 1):
        paper = normalize_text(ws.cell(row_idx, headers["Бумага"]).value)
        isin = normalize_text(ws.cell(row_idx, headers["ISIN"]).value)
        emitent = normalize_text(ws.cell(row_idx, headers["Эмитент"]).value)
        bookbuilding_start = parse_bookbuilding_start(ws.cell(row_idx, headers["cbonds_book_start_datetime"]).value)
        pb_tizer = normalize_text(ws.cell(row_idx, headers["pb_tizer"]).value)
        floating_rate = normalize_text(ws.cell(row_idx, floating_col).value) if floating_col else ""
        check_status = normalize_text(ws.cell(row_idx, check_status_col).value) if check_status_col else ""

        if only_nonblank_pb_tizer and not pb_tizer:
            continue
        if not paper or not emitent or bookbuilding_start is None:
            continue

        issues[row_idx] = IssueInfo(
            excel_row=row_idx,
            paper=paper,
            isin=isin,
            emitent=emitent,
            bookbuilding_start=bookbuilding_start,
            floating_rate=floating_rate,
            pb_tizer=pb_tizer,
            check_status=check_status,
        )

    return wb, ws, issues


def ensure_output_columns(ws) -> dict[str, int]:
    headers = build_header_index(ws)
    for col_name in OUTPUT_COLUMNS:
        if col_name not in headers:
            col_idx = ws.max_column + 1
            ws.cell(1, col_idx, col_name)
            headers[col_name] = col_idx
    return headers


def paper_matches_emission(paper: str, emission_name: str) -> bool:
    paper_key = normalize_key(paper)
    emission_key = normalize_key(emission_name)
    if not paper_key or not emission_key:
        return False
    return paper_key == emission_key or paper_key in emission_key or emission_key in paper_key


def month_before(value: datetime) -> datetime:
    year = value.year
    month = value.month - 1
    if month == 0:
        month = 12
        year -= 1
    day = min(value.day, [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month - 1])
    return value.replace(year=year, month=month, day=day)


def load_article_candidates(db_path: Path, issues: dict[int, IssueInfo]) -> dict[int, list[CandidateArticle]]:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        """
        SELECT
            ia.excel_row,
            a.id AS article_row_id,
            a.article_id,
            a.url,
            a.title,
            a.date_time,
            a.content_text,
            a.introtext,
            GROUP_CONCAT(DISTINCT ae.emission_cbonds_id) AS emission_cbonds_ids,
            GROUP_CONCAT(DISTINCT ae.emission_name) AS emission_names,
            GROUP_CONCAT(DISTINCT ae.emission_url) AS emission_urls
        FROM issue_articles ia
        JOIN articles a ON a.article_id = ia.article_id
        LEFT JOIN article_emissions ae ON ae.article_row_id = a.id
        WHERE a.has_coupon_spread_term = 1
        GROUP BY
            ia.excel_row,
            a.id,
            a.article_id,
            a.url,
            a.title,
            a.date_time,
            a.content_text,
            a.introtext
        ORDER BY ia.excel_row, a.date_time, a.id
        """
    ).fetchall()
    conn.close()

    by_issue: dict[int, dict[int, CandidateArticle]] = {}
    for row in rows:
        excel_row = int(row["excel_row"])
        issue = issues.get(excel_row)
        if issue is None:
            continue

        published_at = parse_article_datetime(row["date_time"])
        if published_at is None:
            continue

        window_start = month_before(issue.bookbuilding_start)
        window_end = issue.bookbuilding_start.replace(hour=23, minute=59, second=59)
        timing = "pre" if window_start <= published_at < issue.bookbuilding_start else "post" if issue.bookbuilding_start <= published_at <= window_end else "outside"
        if timing == "outside":
            continue

        emission_names = normalize_text(row["emission_names"])
        emission_urls = normalize_text(row["emission_urls"])
        emission_ids_text = normalize_text(row["emission_cbonds_ids"])
        emission_cbonds_id = None
        if emission_ids_text and "," not in emission_ids_text:
            try:
                emission_cbonds_id = int(emission_ids_text)
            except ValueError:
                emission_cbonds_id = None

        article_row_id = int(row["article_row_id"])
        text = normalize_text(" ".join([normalize_text(row["introtext"]), normalize_text(row["content_text"])]))
        candidate = CandidateArticle(
            article_row_id=article_row_id,
            article_id=None if row["article_id"] is None else int(row["article_id"]),
            url=normalize_text(row["url"]),
            title=normalize_text(row["title"]),
            date_time=normalize_text(row["date_time"]),
            published_at=published_at,
            text=text,
            emission_cbonds_id=emission_cbonds_id,
            emission_name=emission_names,
            emission_url=emission_urls,
            timing=timing,
        )

        by_issue.setdefault(excel_row, {})[article_row_id] = candidate

    return {row_id: sorted(items.values(), key=lambda x: x.published_at) for row_id, items in by_issue.items()}


def validate_sql_identifier(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value or ""):
        raise ValueError(f"Некорректное имя SQL-таблицы: {value!r}")
    return value


def ensure_results_table(conn: sqlite3.Connection, table_name: str = RESULTS_TABLE) -> None:
    table_name = validate_sql_identifier(table_name)
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {table_name} (
            excel_row INTEGER PRIMARY KEY,
            paper TEXT,
            isin TEXT,
            emitent TEXT,
            bookbuilding_start TEXT,
            orient_type TEXT,
            orient_value REAL,
            finality TEXT,
            coupon_guide REAL,
            spread_guide REAL,
            source_timing TEXT,
            source_title TEXT,
            source_url TEXT,
            source_article_id INTEGER,
            source_datetime TEXT,
            emission_cbonds_id INTEGER,
            emission_name TEXT,
            needs_attention INTEGER NOT NULL DEFAULT 0,
            attention_reason TEXT,
            value_text TEXT,
            evidence_text TEXT,
            reason_short TEXT,
            confidence REAL,
            model_name TEXT,
            model_version TEXT,
            status TEXT,
            updated_at TEXT NOT NULL
        )
        """
    )
    conn.commit()


def write_cache_to_db(
    db_path: Path,
    issues: dict[int, IssueInfo],
    cache: dict[str, Any],
    table_name: str = RESULTS_TABLE,
) -> None:
    conn = sqlite3.connect(db_path)
    table_name = validate_sql_identifier(table_name)
    ensure_results_table(conn, table_name)
    now = datetime.now().isoformat(timespec="seconds")

    rows = []
    for key, result in cache.items():
        try:
            row_idx = int(key)
        except ValueError:
            continue
        issue = issues.get(row_idx)
        if issue is None or not isinstance(result, dict):
            continue
        rows.append(
            (
                row_idx,
                issue.paper,
                issue.isin,
                issue.emitent,
                issue.bookbuilding_start.isoformat(sep=" "),
                result.get("orient_type"),
                result.get("orient_value"),
                result.get("finality"),
                result.get("coupon_guide"),
                result.get("spread_guide"),
                result.get("source_timing"),
                result.get("source_title"),
                result.get("source_url"),
                result.get("source_article_id"),
                result.get("source_datetime"),
                result.get("emission_cbonds_id"),
                result.get("emission_name"),
                int(result.get("needs_attention") or 0),
                result.get("attention_reason"),
                result.get("value_text"),
                result.get("evidence_text"),
                result.get("reason_short"),
                result.get("confidence"),
                result.get("model_name"),
                result.get("model_version"),
                result.get("status"),
                now,
            )
        )

    conn.executemany(
        f"""
        INSERT INTO {table_name} (
            excel_row, paper, isin, emitent, bookbuilding_start,
            orient_type, orient_value, finality, coupon_guide, spread_guide,
            source_timing, source_title, source_url, source_article_id, source_datetime,
            emission_cbonds_id, emission_name, needs_attention, attention_reason,
            value_text, evidence_text, reason_short, confidence,
            model_name, model_version, status, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(excel_row) DO UPDATE SET
            paper = excluded.paper,
            isin = excluded.isin,
            emitent = excluded.emitent,
            bookbuilding_start = excluded.bookbuilding_start,
            orient_type = excluded.orient_type,
            orient_value = excluded.orient_value,
            finality = excluded.finality,
            coupon_guide = excluded.coupon_guide,
            spread_guide = excluded.spread_guide,
            source_timing = excluded.source_timing,
            source_title = excluded.source_title,
            source_url = excluded.source_url,
            source_article_id = excluded.source_article_id,
            source_datetime = excluded.source_datetime,
            emission_cbonds_id = excluded.emission_cbonds_id,
            emission_name = excluded.emission_name,
            needs_attention = excluded.needs_attention,
            attention_reason = excluded.attention_reason,
            value_text = excluded.value_text,
            evidence_text = excluded.evidence_text,
            reason_short = excluded.reason_short,
            confidence = excluded.confidence,
            model_name = excluded.model_name,
            model_version = excluded.model_version,
            status = excluded.status,
            updated_at = excluded.updated_at
        """,
        rows,
    )
    conn.commit()
    conn.close()


def build_issue_prompt_block(
    issue: IssueInfo,
    candidates: list[CandidateArticle],
    max_article_chars: int = DEFAULT_MAX_ARTICLE_CHARS,
    max_issue_news: int = DEFAULT_MAX_ISSUE_NEWS,
    post_only_retry: bool = False,
) -> str:
    article_blocks: list[str] = []
    pre = [c for c in candidates if c.timing == "pre"]
    post = [c for c in candidates if c.timing == "post"]
    if max_issue_news and max_issue_news > 0:
        selected_candidates = sorted(pre, key=lambda x: x.published_at, reverse=True)[:max_issue_news]
        if len(selected_candidates) < max_issue_news:
            selected_candidates.extend(sorted(post, key=lambda x: x.published_at)[: max_issue_news - len(selected_candidates)])
        selected_candidates = sorted({c.article_row_id: c for c in selected_candidates}.values(), key=lambda x: x.published_at)
    else:
        selected_candidates = sorted(candidates, key=lambda x: x.published_at)

    for idx, candidate in enumerate(selected_candidates, start=1):
        text = candidate.text
        if max_article_chars and len(text) > max_article_chars:
            text = text[:max_article_chars]
        article_blocks.append(
            "\n".join(
                [
                    f"[A{idx}]",
                    f"article_id: {candidate.article_id or 'NA'}",
                    f"published_at: {candidate.date_time}",
                    f"timing_relative_to_book_start: {candidate.timing}",
                    f"url: {candidate.url}",
                    f"title: {candidate.title}",
                    f"parsed_emissions: {candidate.emission_name or 'NA'}",
                    f"parsed_emission_urls: {candidate.emission_url or 'NA'}",
                    "text:",
                    text or "NA",
                ]
            )
        )

    if post_only_retry:
        selection_rule = [
            "Это retry-прогон: в ARTICLES переданы только новости после bookbuilding_start.",
            "Выбери самый ранний релевантный ориентир купона или спреда после bookbuilding_start.",
            "Не выбирай более поздний новый/финальный/уточненный ориентир, если уже есть более ранний релевантный ориентир после bookbuilding_start.",
            "Слово 'ориентир' может отсутствовать: учитывай смысловые формулировки вроде 'не выше', 'не более', 'диапазон', 'КС + спред', 'ставка купона ... %'.",
            "Поле finality только диагностическое и не имеет приоритета над правилом времени.",
            "Для каждой новости учитывай parsed_emissions, но не полагайся только на него: в тексте может обсуждаться несколько выпусков.",
        ]
    else:
        selection_rule = [
            "Сначала выбери самый поздний релевантный ориентир до bookbuilding_start.",
            "Если до bookbuilding_start релевантного ориентира нет, выбери самый ранний релевантный ориентир после bookbuilding_start в списке.",
            "Не выбирай более поздний финальный ориентир после bookbuilding_start, если уже есть более ранний релевантный ориентир после bookbuilding_start.",
            "Слово 'ориентир' может отсутствовать: учитывай смысловые формулировки вроде 'не выше', 'не более', 'диапазон', 'КС + спред', 'ставка купона ... %'.",
            "Поле finality только диагностическое и не имеет приоритета над правилом времени.",
            "Для каждой новости учитывай parsed_emissions, но не полагайся только на него: в тексте может обсуждаться несколько выпусков.",
        ]

    return "\n".join(
        [
            "ISSUE:",
            f"excel_row: {issue.excel_row}",
            f"paper: {issue.paper}",
            f"isin: {issue.isin or 'NA'}",
            f"emitent: {issue.emitent}",
            f"bookbuilding_start: {issue.bookbuilding_start.isoformat(sep=' ')}",
            f"floating_rate: {issue.floating_rate or 'NA'}",
            "",
            "SELECTION_RULE:",
            *selection_rule,
            "",
            "ARTICLES:",
            "\n\n".join(article_blocks) if article_blocks else "NA",
        ]
    )


def build_batch_prompt(batch: list[BatchItem]) -> str:
    return "\n\n".join(
        [
            "Ниже несколько выпусков. Для каждого блока ISSUE верни ровно один объект в JSON-массиве.",
            "Ключ сопоставления результата: excel_row.",
            "Для каждого выпуска применяй правило выбора новости независимо от остальных выпусков.",
            *(item.prompt_text for item in batch),
        ]
    )


def make_batches(items: list[BatchItem], max_batch_chars: int) -> list[list[BatchItem]]:
    batches: list[list[BatchItem]] = []
    batch_chars: list[int] = []
    for item in sorted(items, key=lambda x: x.estimated_chars, reverse=True):
        placed = False
        for idx, current_chars in enumerate(batch_chars):
            if current_chars + item.estimated_chars <= max_batch_chars:
                batches[idx].append(item)
                batch_chars[idx] += item.estimated_chars
                placed = True
                break
        if not placed:
            batches.append([item])
            batch_chars.append(item.estimated_chars)
    return batches


def extract_json_object(raw_text: str) -> str:
    raw_text = (raw_text or "").strip()
    if raw_text.startswith("```"):
        raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text, flags=re.IGNORECASE)
        raw_text = re.sub(r"\s*```$", "", raw_text).strip()
    start = raw_text.find("{")
    end = raw_text.rfind("}")
    if start != -1 and end != -1 and end > start:
        return raw_text[start : end + 1]
    raise ValueError("JSON object not found in model output")


def extract_json_array(raw_text: str) -> str:
    raw_text = (raw_text or "").strip()
    if raw_text.startswith("```"):
        raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text, flags=re.IGNORECASE)
        raw_text = re.sub(r"\s*```$", "", raw_text).strip()
    start = raw_text.find("[")
    end = raw_text.rfind("]")
    if start != -1 and end != -1 and end > start:
        return raw_text[start : end + 1]
    raise ValueError("JSON array not found in model output")


def parse_model_items(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        parsed = payload
    elif isinstance(payload, str):
        parsed = json.loads(extract_json_array(payload))
    else:
        raise ValueError("Unsupported model payload type")
    if not isinstance(parsed, list):
        raise ValueError("Model output is not a list")
    return [item for item in parsed if isinstance(item, dict)]


def normalize_model_item(
    item: dict[str, Any],
    candidates: list[CandidateArticle],
    model_name: str,
    model_version: str | None,
) -> dict[str, Any]:
    orient_type = normalize_text(item.get("orient_type")).lower()
    if orient_type not in {"coupon", "spread", "none"}:
        orient_type = "none"
    orient_value = safe_nonneg_number(item.get("orient_value"), 0.0)
    finality = normalize_text(item.get("finality")).lower()
    if finality not in {"final", "guide", "none"}:
        finality = "none"
    if orient_type == "none" or orient_value <= 0:
        orient_type = "none"
        orient_value = 0.0
        finality = "none"

    source_article_id = None
    try:
        raw_source_article_id = int(item.get("source_article_id") or 0)
        source_article_id = raw_source_article_id or None
    except (TypeError, ValueError):
        source_article_id = None

    source_candidate = None
    if source_article_id is not None:
        for candidate in candidates:
            if candidate.article_id == source_article_id:
                source_candidate = candidate
                break

    source_timing = normalize_text(item.get("source_timing")).lower()
    if source_timing not in {"pre", "post", "none"}:
        source_timing = source_candidate.timing if source_candidate else "none"

    needs_attention = bool(item.get("needs_attention"))
    attention_reason = normalize_text(item.get("attention_reason")) if needs_attention else ""
    if orient_type != "none" and source_article_id is not None and source_candidate is None:
        needs_attention = True
        attention_reason = " ".join(
            x
            for x in [
                attention_reason,
                "Gemini вернул source_article_id, которого нет в переданном списке новостей.",
            ]
            if x
        ).strip()

    return {
        "orient_type": orient_type,
        "orient_value": orient_value,
        "finality": finality,
        "coupon_guide": orient_value if orient_type == "coupon" and orient_value > 0 else None,
        "spread_guide": orient_value if orient_type == "spread" and orient_value > 0 else None,
        "source_timing": source_timing,
        "source_title": source_candidate.title if source_candidate else "",
        "source_url": source_candidate.url if source_candidate else "",
        "source_article_id": source_article_id,
        "source_datetime": source_candidate.date_time if source_candidate else "",
        "emission_cbonds_id": source_candidate.emission_cbonds_id if source_candidate else None,
        "emission_name": normalize_text(item.get("source_emission_name")) or (source_candidate.emission_name if source_candidate else ""),
        "needs_attention": int(needs_attention),
        "attention_reason": attention_reason,
        "value_text": normalize_text(item.get("value_text"))[:240],
        "evidence_text": normalize_text(item.get("evidence_text"))[:900],
        "reason_short": normalize_text(item.get("reason_short"))[:300],
        "confidence": safe_float_01(item.get("confidence"), 0.5),
        "model_name": model_name,
        "model_version": model_version,
        "status": "classified",
    }


def has_extracted_orientir(result: dict[str, Any]) -> bool:
    return (
        result.get("orient_type") in {"coupon", "spread"}
        and safe_nonneg_number(result.get("orient_value"), 0.0) > 0
    )


def sleep_to_respect_min_interval(min_interval: float) -> None:
    global LAST_REQUEST_TS
    elapsed = time.time() - LAST_REQUEST_TS
    if LAST_REQUEST_TS and elapsed < min_interval:
        time.sleep(min_interval - elapsed)


def mark_request_sent() -> None:
    global LAST_REQUEST_TS
    LAST_REQUEST_TS = time.time()


def extract_retry_delay_seconds(exc: Exception) -> float | None:
    text = str(exc)
    for pattern in (r"retry in\s+([0-9]+(?:\.[0-9]+)?)s", r"Retry in\s+([0-9]+(?:\.[0-9]+)?)s"):
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return float(match.group(1))
    return None


def call_gemini(client: Any, model_name: str, prompt: str, max_retries: int, max_output_tokens: int, delay: float) -> tuple[Any, str | None]:
    last_exc: Exception | None = None
    for attempt in range(1, max_retries + 1):
        try:
            sleep_to_respect_min_interval(delay)
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=SYSTEM_PROMPT,
                    temperature=0,
                    max_output_tokens=max_output_tokens,
                    response_mime_type="application/json",
                    response_json_schema=RESULT_SCHEMA,
                ),
            )
            mark_request_sent()
            parsed = getattr(response, "parsed", None)
            raw_text = getattr(response, "text", "") or ""
            version = getattr(response, "model_version", None)
            return (parsed if parsed is not None else raw_text), version
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
            wait = extract_retry_delay_seconds(exc)
            if wait is None:
                wait = min(5 * (2 ** (attempt - 1)), 60)
            if attempt < max_retries:
                print(f"  retry {attempt}/{max_retries} after {wait:.1f}s: {exc}")
                time.sleep(wait)
    raise RuntimeError(str(last_exc) if last_exc else "Unknown API error")


def empty_result(status: str, reason: str) -> dict[str, Any]:
    return {
        "orient_type": "none",
        "orient_value": 0.0,
        "finality": "none",
        "coupon_guide": None,
        "spread_guide": None,
        "source_timing": "",
        "source_title": "",
        "source_url": "",
        "source_article_id": None,
        "source_datetime": "",
        "emission_cbonds_id": None,
        "emission_name": "",
        "needs_attention": 1 if status != "no_candidate" else 0,
        "attention_reason": reason,
        "value_text": "",
        "evidence_text": "",
        "reason_short": reason,
        "confidence": 0.0,
        "model_name": "",
        "model_version": "",
        "status": status,
    }


def apply_cache_to_workbook(ws, headers: dict[str, int], cache: dict[str, Any]) -> None:
    mapping = {
        "bb_gemini_orient_type": "orient_type",
        "bb_gemini_orient_value": "orient_value",
        "bb_gemini_orient_finality": "finality",
        "bb_gemini_coupon_guide": "coupon_guide",
        "bb_gemini_spread_guide": "spread_guide",
        "bb_gemini_source_timing": "source_timing",
        "bb_gemini_source_title": "source_title",
        "bb_gemini_source_url": "source_url",
        "bb_gemini_source_article_id": "source_article_id",
        "bb_gemini_source_datetime": "source_datetime",
        "bb_gemini_emission_cbonds_id": "emission_cbonds_id",
        "bb_gemini_emission_name": "emission_name",
        "bb_gemini_needs_attention": "needs_attention",
        "bb_gemini_attention_reason": "attention_reason",
        "bb_gemini_value_text": "value_text",
        "bb_gemini_evidence_text": "evidence_text",
        "bb_gemini_reason_short": "reason_short",
        "bb_gemini_confidence": "confidence",
        "bb_gemini_model_name": "model_name",
        "bb_gemini_model_version": "model_version",
        "bb_gemini_status": "status",
    }
    for key, result in cache.items():
        try:
            row_idx = int(key)
        except ValueError:
            continue
        if not isinstance(result, dict) or row_idx < 2 or row_idx > ws.max_row:
            continue
        for col_name, result_key in mapping.items():
            ws.cell(row_idx, headers[col_name], result.get(result_key))


def save_workbook(wb: openpyxl.Workbook, ws, cache: dict[str, Any], output_path: Path) -> None:
    headers = ensure_output_columns(ws)
    apply_cache_to_workbook(ws, headers, cache)
    try:
        wb.save(output_path)
    except ValueError as exc:
        if "closed file" not in str(exc).lower():
            raise
        strip_workbook_images(wb)
        wb.save(output_path)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Extract Cbonds orientir from all candidate bookbuilding news with Gemini.")
    parser.add_argument("--input", type=Path, default=Path(INPUT_FILE))
    parser.add_argument("--output", type=Path, default=Path(OUTPUT_FILE))
    parser.add_argument("--db", type=Path, default=Path(DB_FILE))
    parser.add_argument(
        "--results-table",
        type=str,
        default=RESULTS_TABLE,
        help=f"Таблица SQLite для результатов Gemini, по умолчанию {RESULTS_TABLE}.",
    )
    parser.add_argument("--cache", type=Path, default=Path(CACHE_FILE))
    parser.add_argument("--sheet", type=str, default=None)
    parser.add_argument("--model", type=str, default=DEFAULT_MODEL)
    parser.add_argument("--api-key", type=str, default=os.getenv("GEMINI_API_KEY", ""))
    parser.add_argument("--delay", type=float, default=DEFAULT_DELAY)
    parser.add_argument("--max-retries", type=int, default=DEFAULT_MAX_RETRIES)
    parser.add_argument("--max-output-tokens", type=int, default=DEFAULT_MAX_OUTPUT_TOKENS)
    parser.add_argument("--max-requests", type=int, default=None)
    parser.add_argument("--max-batch-chars", type=int, default=DEFAULT_MAX_BATCH_CHARS)
    parser.add_argument("--max-article-chars", type=int, default=DEFAULT_MAX_ARTICLE_CHARS)
    parser.add_argument("--max-issue-news", type=int, default=DEFAULT_MAX_ISSUE_NEWS)
    parser.add_argument("--issue-char-threshold", type=int, default=DEFAULT_MAX_BATCH_CHARS)
    parser.add_argument(
        "--issue-size-mode",
        choices=["all", "under", "over"],
        default="all",
        help="all: все выпуски; under: только выпуски <= threshold; over: только выпуски > threshold.",
    )
    parser.add_argument("--one-issue-per-batch", action="store_true", help="Не объединять выпуски в батчи.")
    parser.add_argument("--save-every", type=int, default=DEFAULT_SAVE_EVERY)
    parser.add_argument("--include-blank-pb-tizer", action="store_true")
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument(
        "--retry-check-status",
        type=str,
        default="",
        help="Обработать только строки с указанным check_status; для таких строк кеш игнорируется.",
    )
    parser.add_argument(
        "--retry-post-only",
        action="store_true",
        help="Для retry брать только новости после bookbuilding_start и выбирать самый ранний релевантный ориентир.",
    )
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-confirm", action="store_true")
    return parser


def main() -> None:
    args = build_arg_parser().parse_args()
    retry_check_status = normalize_text(args.retry_check_status).casefold()
    retry_post_only = bool(retry_check_status) or args.retry_post_only
    wb, ws, issues = load_excel_issues(
        args.input,
        sheet_name=args.sheet,
        only_nonblank_pb_tizer=not args.include_blank_pb_tizer,
    )
    loaded_issue_count = len(issues)
    if retry_check_status:
        issues = {
            row_idx: issue
            for row_idx, issue in issues.items()
            if normalize_text(issue.check_status).casefold() == retry_check_status
        }
    candidates_by_issue = load_article_candidates(args.db, issues)
    if retry_post_only:
        candidates_by_issue = {
            row_idx: [candidate for candidate in candidates if candidate.timing == "post"]
            for row_idx, candidates in candidates_by_issue.items()
        }
    cache = load_json(args.cache)

    no_candidate_count = 0
    for row_idx, issue in issues.items():
        if not candidates_by_issue.get(row_idx):
            no_candidate_count += 1
            reason = (
                "Не найдена новость с has_coupon_spread_term=1 после старта букбилдинга для retry."
                if retry_post_only
                else "Не найдена новость с has_coupon_spread_term=1 вокруг старта букбилдинга."
            )
            if retry_check_status:
                cache[str(row_idx)] = empty_result("no_candidate", reason)
            else:
                cache.setdefault(str(row_idx), empty_result("no_candidate", reason))

    rows_with_candidates = [row_idx for row_idx in sorted(issues) if candidates_by_issue.get(row_idx)]
    if retry_check_status:
        pending_rows = rows_with_candidates
    else:
        pending_rows = [row_idx for row_idx in rows_with_candidates if args.refresh or str(row_idx) not in cache]
    pending_items: list[BatchItem] = []
    for row_idx in pending_rows:
        prompt_text = build_issue_prompt_block(
            issues[row_idx],
            candidates_by_issue[row_idx],
            max_article_chars=args.max_article_chars,
            max_issue_news=args.max_issue_news,
            post_only_retry=retry_post_only,
        )
        item = BatchItem(
            excel_row=row_idx,
            prompt_text=prompt_text,
            estimated_chars=len(prompt_text) + 350,
        )
        if args.issue_size_mode == "under" and item.estimated_chars > args.issue_char_threshold:
            continue
        if args.issue_size_mode == "over" and item.estimated_chars <= args.issue_char_threshold:
            continue
        pending_items.append(item)

    batches = [[item] for item in pending_items] if args.one_issue_per_batch else make_batches(pending_items, args.max_batch_chars)
    if args.max_requests is not None:
        batches = batches[: args.max_requests]

    print("=" * 72)
    print("Bookbuilding orientir extraction from all candidate Cbonds news")
    print("=" * 72)
    print(f"Input: {args.input}")
    print(f"Output: {args.output}")
    print(f"DB: {args.db}")
    print(f"Results table: {args.results_table}")
    print(f"Sheet: {ws.title}")
    print(f"Issues loaded: {loaded_issue_count:,}")
    if retry_check_status:
        print(f"Retry check_status: {args.retry_check_status!r}")
    print(f"Issues eligible: {len(issues):,}")
    print(f"Issues with candidate news: {len(rows_with_candidates):,}")
    print(f"No candidate: {no_candidate_count:,}")
    print(f"Cache rows: {len(cache):,}")
    print(f"Pending issues: {len(pending_rows):,}")
    print(f"Pending issues after size filter: {len(pending_items):,}")
    print(f"Planned batches: {len(batches):,}")
    print(f"Max batch chars: {args.max_batch_chars:,}")
    print(f"Max article chars: {args.max_article_chars:,}")
    print(f"Max issue news: {args.max_issue_news:,}")
    print(f"Issue size mode: {args.issue_size_mode} | threshold: {args.issue_char_threshold:,}")
    print(f"One issue per batch: {args.one_issue_per_batch}")
    print(f"Retry post-only: {retry_post_only}")

    total_candidate_news = sum(len(items) for items in candidates_by_issue.values())
    pre_news_count = sum(1 for items in candidates_by_issue.values() for c in items if c.timing == "pre")
    post_news_count = sum(1 for items in candidates_by_issue.values() for c in items if c.timing == "post")
    print(f"Candidate news total: {total_candidate_news:,}")
    print(f"Candidate news pre-start: {pre_news_count:,}")
    print(f"Candidate news post-start: {post_news_count:,}")
    if batches:
        batch_chars = [len(build_batch_prompt(batch)) for batch in batches]
        print(f"Batch chars avg: {sum(batch_chars) / len(batch_chars):,.1f}")
        print(f"Batch chars max: {max(batch_chars):,}")

    if args.dry_run:
        save_workbook(wb, ws, cache, args.output)
        write_cache_to_db(args.db, issues, cache, table_name=args.results_table)
        print("\nDry run complete. Workbook written from existing cache only.")
        return

    if genai is None or types is None:
        raise ImportError("Не найден пакет google-genai. Установи его: pip install google-genai")
    if not args.api_key:
        raise ValueError("API key not found. Pass --api-key or set GEMINI_API_KEY.")

    if batches and not args.no_confirm:
        answer = input(f"Send {len(batches)} batched requests to Gemini? (y/n): ").strip().lower()
        if answer != "y":
            print("Cancelled by user.")
            save_workbook(wb, ws, cache, args.output)
            write_cache_to_db(args.db, issues, cache, table_name=args.results_table)
            return

    client = genai.Client(api_key=args.api_key)
    try:
        for idx, batch in enumerate(batches, start=1):
            batch_rows = [item.excel_row for item in batch]
            print(f"\nBatch {idx}/{len(batches)} | issues={len(batch)} | rows={batch_rows[:8]}{'...' if len(batch_rows) > 8 else ''}")
            try:
                payload, model_version = call_gemini(
                    client=client,
                    model_name=args.model,
                    prompt=build_batch_prompt(batch),
                    max_retries=args.max_retries,
                    max_output_tokens=args.max_output_tokens,
                    delay=args.delay,
                )
                parsed_items = parse_model_items(payload)
                results_by_row: dict[int, dict[str, Any]] = {}
                for item in parsed_items:
                    try:
                        item_row = int(item.get("excel_row"))
                    except (TypeError, ValueError):
                        continue
                    if item_row not in batch_rows:
                        continue
                    results_by_row[item_row] = normalize_model_item(
                        item,
                        candidates_by_issue[item_row],
                        args.model,
                        model_version,
                    )

                for row_idx in batch_rows:
                    result = results_by_row.get(row_idx)
                    if result is None:
                        result = empty_result("missing_item_in_response", "Gemini не вернул объект для этой строки в батче.")
                    cache[str(row_idx)] = result

                extracted_now = sum(1 for r in results_by_row.values() if has_extracted_orientir(r))
                print(f"  -> returned={len(results_by_row)}/{len(batch_rows)} extracted={extracted_now}")
            except Exception as exc:  # noqa: BLE001
                for row_idx in batch_rows:
                    cache[str(row_idx)] = empty_result("api_error", f"{type(exc).__name__}: {exc}")
                print(f"  [error] {exc}")

            if idx % args.save_every == 0:
                save_json(args.cache, cache)
                save_workbook(wb, ws, cache, args.output)
                write_cache_to_db(args.db, issues, cache, table_name=args.results_table)
                print(f"  saved: {args.cache}, {args.output}")

    except KeyboardInterrupt:
        print("\nInterrupted. Saving progress...")
    finally:
        save_json(args.cache, cache)
        save_workbook(wb, ws, cache, args.output)
        write_cache_to_db(args.db, issues, cache, table_name=args.results_table)
        print(f"\nSaved cache: {args.cache}")
        print(f"Saved workbook: {args.output}")
        print(f"Saved DB table: {args.results_table}")

    extracted = sum(
        1
        for item in cache.values()
        if isinstance(item, dict) and item.get("orient_type") in {"coupon", "spread"} and safe_nonneg_number(item.get("orient_value")) > 0
    )
    attention = sum(1 for item in cache.values() if isinstance(item, dict) and int(item.get("needs_attention") or 0) == 1)
    print(f"Extracted nonzero: {extracted:,}")
    print(f"Needs attention: {attention:,}")


if __name__ == "__main__":
    main()
