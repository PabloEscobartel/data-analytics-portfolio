#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Флаг строк, где pre evidence text по объему содержит признаки общего/совокупного объема.

Ищет в `volume_gemini_pre_evidence_text` подстроки:
- "совокуп"
- "общ"

Добавляет/обновляет колонки:
- volume_pre_aggregate_term_flag
- volume_pre_aggregate_term_match
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any

import openpyxl


DEFAULT_WORKBOOK = Path("bonds_offer_before2018_placement_volume_gemini.xlsx")
DEFAULT_SHEET = "place_before_2018"
TEXT_COLUMN = "volume_gemini_pre_evidence_text"
FLAG_COLUMN = "volume_pre_aggregate_term_flag"
MATCH_COLUMN = "volume_pre_aggregate_term_match"
PATTERN = re.compile(r"(?iu)(совокуп\w*|общ\w*)")


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def build_header_index(ws) -> dict[str, int]:
    return {normalize_text(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}


def ensure_column(ws, headers: dict[str, int], column: str) -> int:
    if column not in headers:
        col_idx = ws.max_column + 1
        ws.cell(1, col_idx, column)
        headers[column] = col_idx
    return headers[column]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Flag pre evidence texts with aggregate/common volume terms.")
    parser.add_argument("--input", type=Path, default=DEFAULT_WORKBOOK, help=f"Workbook to check, default: {DEFAULT_WORKBOOK}")
    parser.add_argument("--output", type=Path, default=None, help="Output workbook. If omitted, updates input file in place.")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help=f"Sheet name, default: {DEFAULT_SHEET!r}")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_path = args.output or args.input

    wb = openpyxl.load_workbook(args.input)
    if args.sheet not in wb.sheetnames:
        raise RuntimeError(f"Лист {args.sheet!r} не найден. Доступные листы: {', '.join(wb.sheetnames)}")
    ws = wb[args.sheet]

    headers = build_header_index(ws)
    if TEXT_COLUMN not in headers:
        raise RuntimeError(f"Не найдена колонка: {TEXT_COLUMN!r}")

    flag_col = ensure_column(ws, headers, FLAG_COLUMN)
    match_col = ensure_column(ws, headers, MATCH_COLUMN)
    text_col = headers[TEXT_COLUMN]

    checked = 0
    flagged = 0
    for row_idx in range(2, ws.max_row + 1):
        text = normalize_text(ws.cell(row_idx, text_col).value)
        matches = sorted({m.group(0).casefold() for m in PATTERN.finditer(text)})
        is_flagged = bool(matches)
        checked += 1
        flagged += int(is_flagged)
        ws.cell(row_idx, flag_col, int(is_flagged))
        ws.cell(row_idx, match_col, ", ".join(matches))

    wb.save(output_path)
    print(f"Workbook: {output_path}")
    print(f"Checked rows: {checked}")
    print(f"Flagged rows: {flagged}")


if __name__ == "__main__":
    main()
