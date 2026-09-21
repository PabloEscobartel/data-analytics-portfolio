#!/usr/bin/env python3
"""
Сбор новостей Cbonds вокруг периода букбилдинга/размещения для выпусков из Excel.

Что делает:
1. Читает Excel-файл bonds_offer_before2018_all_orientirs_gemini_cbonds.xlsx.
2. Берёт выпуски с листа place_before_2018, где эмитент, cbonds_book_start_datetime
   и Окончание размещения заполнены.
3. Объединяет выпуски в группы (Эмитент, окно дат), чтобы не дублировать запросы.
4. Для каждой группы ищет новости Cbonds по эмитенту за период:
   [cbonds_book_start_datetime минус 1 календарный месяц; Окончание размещения плюс 1 неделя].
5. Сохраняет статьи и связи выпуск-статья в существующую SQLite-базу cbonds_news.sqlite3.

Cookie берётся так же, как в cbonds_news_scrape.py:
- из переменной окружения CBONDS_COOKIE;
- или из файла cbonds_cookie.txt рядом со скриптом.

Запуск:
    python cbonds_news_offer_scrape.py

Полезные опции:
    python cbonds_news_offer_scrape.py --sheet place_before_2018
    python cbonds_news_offer_scrape.py --input bonds_offer_before2018_all_orientirs_gemini_cbonds.xlsx --book-date-column cbonds_book_start_datetime --placement-end-column "Окончание размещения"
    python cbonds_news_offer_scrape.py --limit-groups 10
    python cbonds_news_offer_scrape.py --dry-run
    python cbonds_news_offer_scrape.py --no-cookie
    python cbonds_news_offer_scrape.py --no-cache
"""

from __future__ import annotations

import argparse
import calendar
import json
import os
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote

import openpyxl

import cbonds_news_scrape as base


INPUT_FILE = "bonds_offer_before2018_all_orientirs_gemini_cbonds.xlsx"
DB_FILE = base.DB_FILE
BASE_CACHE_FILE = base.CACHE_FILE
OFFER_CACHE_FILE = "cbonds_news_offer_before2018_book_period_cache.json"

DEFAULT_SHEET = "place_before_2018"
DEFAULT_PAPER_COLUMN = "Бумага"
DEFAULT_ISIN_COLUMN = "ISIN"
DEFAULT_ISSUER_COLUMN = "Эмитент"
DEFAULT_BOOK_DATE_COLUMN = "cbonds_book_start_datetime"
DEFAULT_PLACEMENT_END_COLUMN = "Окончание размещения"
PB_TIZER_COLUMN = "pb_tizer"


@dataclass(frozen=True)
class OfferIssueRow:
    excel_row: int
    paper: str | None
    isin: str | None
    emitent: str
    bookbuilding_date: str  # YYYY-MM-DD
    placement_date_end: str  # YYYY-MM-DD
    search_date_min: str  # YYYY-MM-DD
    search_date_max: str  # YYYY-MM-DD


def is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def is_veb_rf(value: Any) -> bool:
    return base.normalize_text(value) in {"вэб.рф", "веб.рф"}


def subtract_calendar_month(value: date) -> date:
    year = value.year
    month = value.month - 1
    if month == 0:
        month = 12
        year -= 1

    day = min(value.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def date_range_for_bookbuilding(event_date: str) -> tuple[str, str]:
    end_date = datetime.strptime(event_date, "%Y-%m-%d").date()
    start_date = subtract_calendar_month(end_date)
    return start_date.isoformat(), end_date.isoformat()


def date_range_for_offer(book_date: str, placement_date_end: str) -> tuple[str, str]:
    book_day = datetime.strptime(book_date, "%Y-%m-%d").date()
    placement_end_day = datetime.strptime(placement_date_end, "%Y-%m-%d").date()
    start_date = subtract_calendar_month(book_day)
    end_date = placement_end_day + timedelta(days=7)
    return start_date.isoformat(), end_date.isoformat()


def range_group_key(emitent_id: int, date_min: str, date_max: str) -> str:
    return f"{emitent_id}||{date_min}||{date_max}"


def required_headers(
    paper_column: str,
    isin_column: str,
    issuer_column: str,
    book_date_column: str,
    placement_end_column: str,
    require_pb_tizer: bool,
) -> list[str]:
    headers = [paper_column, isin_column, issuer_column, book_date_column, placement_end_column]
    if require_pb_tizer:
        headers.append(PB_TIZER_COLUMN)
    return headers


def header_index(headers: list[Any], column_name: str) -> int:
    normalized_target = base.normalize_text(column_name)
    for idx, header in enumerate(headers):
        if base.normalize_text(header) == normalized_target:
            return idx
    raise RuntimeError(f"Не найдена колонка: {column_name!r}")


def find_target_sheet(
    wb: openpyxl.Workbook,
    sheet_name: str | None = None,
    paper_column: str = DEFAULT_PAPER_COLUMN,
    isin_column: str = DEFAULT_ISIN_COLUMN,
    issuer_column: str = DEFAULT_ISSUER_COLUMN,
    book_date_column: str = DEFAULT_BOOK_DATE_COLUMN,
    placement_end_column: str = DEFAULT_PLACEMENT_END_COLUMN,
    require_pb_tizer: bool = True,
):
    required = required_headers(
        paper_column,
        isin_column,
        issuer_column,
        book_date_column,
        placement_end_column,
        require_pb_tizer,
    )
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            available = ", ".join(wb.sheetnames)
            raise RuntimeError(f"Лист {sheet_name!r} не найден. Доступные листы: {available}")
        ws = wb[sheet_name]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        normalized_headers = {base.normalize_text(h) for h in headers if h is not None}
        missing = [h for h in required if base.normalize_text(h) not in normalized_headers]
        if missing:
            raise RuntimeError(f"На листе {sheet_name!r} не найдены колонки: {missing}")
        return ws

    normalized_required = {base.normalize_text(x) for x in required}
    for ws in wb.worksheets:
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        normalized_headers = {base.normalize_text(h) for h in headers if h is not None}
        if normalized_required.issubset(normalized_headers):
            return ws

    available = ", ".join(wb.sheetnames)
    raise RuntimeError(f"Не найден лист с колонками {required}. Доступные листы: {available}")


def read_target_rows(
    input_file: str,
    sheet_name: str | None = None,
    paper_column: str = DEFAULT_PAPER_COLUMN,
    isin_column: str = DEFAULT_ISIN_COLUMN,
    issuer_column: str = DEFAULT_ISSUER_COLUMN,
    book_date_column: str = DEFAULT_BOOK_DATE_COLUMN,
    placement_end_column: str = DEFAULT_PLACEMENT_END_COLUMN,
    require_pb_tizer: bool = False,
    filter_column: str | None = None,
    filter_value: str | None = None,
) -> tuple[str, list[OfferIssueRow]]:
    wb = openpyxl.load_workbook(input_file, data_only=True)
    ws = find_target_sheet(
        wb,
        sheet_name=sheet_name,
        paper_column=paper_column,
        isin_column=isin_column,
        issuer_column=issuer_column,
        book_date_column=book_date_column,
        placement_end_column=placement_end_column,
        require_pb_tizer=require_pb_tizer,
    )
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    paper_idx = header_index(headers, paper_column)
    isin_idx = header_index(headers, isin_column)
    emitent_idx = header_index(headers, issuer_column)
    book_date_idx = header_index(headers, book_date_column)
    placement_end_idx = header_index(headers, placement_end_column)
    pb_tizer_idx = None
    if any(base.normalize_text(h) == base.normalize_text(PB_TIZER_COLUMN) for h in headers):
        pb_tizer_idx = header_index(headers, PB_TIZER_COLUMN)
    filter_idx = None
    normalized_filter_value = None
    if filter_column:
        filter_idx = header_index(headers, filter_column)
        normalized_filter_value = base.normalize_text(filter_value).casefold()

    result: list[OfferIssueRow] = []

    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if filter_idx is not None:
            current_filter_value = base.normalize_text(row[filter_idx]).casefold()
            if current_filter_value != normalized_filter_value:
                continue

        pb_tizer = row[pb_tizer_idx] if pb_tizer_idx is not None else None
        emitent = row[emitent_idx]
        book_date_value = row[book_date_idx]
        placement_end_value = row[placement_end_idx]

        if require_pb_tizer and is_blank(pb_tizer):
            continue
        if is_blank(emitent) or is_blank(book_date_value) or is_blank(placement_end_value):
            continue

        book_date_str = base.ensure_date_string(book_date_value)
        placement_end_str = base.ensure_date_string(placement_end_value)
        if not book_date_str or not placement_end_str:
            continue
        date_min, date_max = date_range_for_offer(book_date_str, placement_end_str)

        result.append(
            OfferIssueRow(
                excel_row=excel_row,
                paper=None if is_blank(row[paper_idx]) else str(row[paper_idx]).strip(),
                isin=None if is_blank(row[isin_idx]) else str(row[isin_idx]).strip(),
                emitent=str(emitent).strip(),
                bookbuilding_date=book_date_str,
                placement_date_end=placement_end_str,
                search_date_min=date_min,
                search_date_max=date_max,
            )
        )

    return ws.title, result


def load_offer_cache(cache_file: str) -> dict:
    cache: dict = {"emitents": {}, "search_results": {}}

    if Path(BASE_CACHE_FILE).exists():
        with open(BASE_CACHE_FILE, "r", encoding="utf-8") as f:
            base_cache = json.load(f)
        cache["emitents"].update(base_cache.get("emitents", {}))

    if Path(cache_file).exists():
        with open(cache_file, "r", encoding="utf-8") as f:
            offer_cache = json.load(f)
        cache["emitents"].update(offer_cache.get("emitents", {}))
        cache["search_results"].update(offer_cache.get("search_results", {}))

    return cache


def save_offer_cache(cache: dict, cache_file: str) -> None:
    with open(cache_file, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


def existing_columns(conn: Any, table_name: str) -> set[str]:
    rows = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    return {str(row[1]) for row in rows}


def ensure_offer_schema(conn: Any) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS offer_searches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            emitent_name TEXT NOT NULL,
            emitent_id INTEGER,
            bookbuilding_date TEXT NOT NULL,
            placement_date_end TEXT NOT NULL,
            date_min TEXT NOT NULL,
            date_max TEXT NOT NULL,
            search_url TEXT,
            total_items INTEGER,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(emitent_name, bookbuilding_date, placement_date_end, date_min, date_max)
        );
        """
    )

    issue_article_columns = existing_columns(conn, "issue_articles")
    additions = [
        ("placement_date_end", "TEXT"),
        ("search_date_min", "TEXT"),
        ("search_date_max", "TEXT"),
    ]
    for column_name, column_sql in additions:
        if column_name not in issue_article_columns:
            conn.execute(f"ALTER TABLE issue_articles ADD COLUMN {column_name} {column_sql}")

    conn.commit()


def upsert_offer_search(
    conn: Any,
    emitent_name: str,
    emitent_id: int | None,
    bookbuilding_date: str,
    placement_date_end: str,
    date_min: str,
    date_max: str,
    search_url: str,
    total_items: int | None,
) -> None:
    conn.execute(
        """
        INSERT INTO offer_searches (
            emitent_name, emitent_id, bookbuilding_date, placement_date_end,
            date_min, date_max, search_url, total_items
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(emitent_name, bookbuilding_date, placement_date_end, date_min, date_max)
        DO UPDATE SET
            emitent_id = excluded.emitent_id,
            search_url = excluded.search_url,
            total_items = excluded.total_items
        """,
        (
            emitent_name,
            emitent_id,
            bookbuilding_date,
            placement_date_end,
            date_min,
            date_max,
            search_url,
            total_items,
        ),
    )
    conn.commit()


def link_offer_issue_to_article(
    conn: Any,
    issue: OfferIssueRow,
    emitent_id: int | None,
    article_id: int,
    article_url: str,
) -> None:
    conn.execute(
        """
        INSERT INTO issue_articles (
            excel_row, paper, isin, emitent_name, emitent_id, bookbuilding_date,
            placement_date_end, search_date_min, search_date_max, article_id, article_url
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(excel_row, article_id) DO UPDATE SET
            paper = excluded.paper,
            isin = excluded.isin,
            emitent_name = excluded.emitent_name,
            emitent_id = excluded.emitent_id,
            bookbuilding_date = excluded.bookbuilding_date,
            placement_date_end = excluded.placement_date_end,
            search_date_min = excluded.search_date_min,
            search_date_max = excluded.search_date_max,
            article_url = excluded.article_url
        """,
        (
            issue.excel_row,
            issue.paper,
            issue.isin,
            issue.emitent,
            emitent_id,
            issue.bookbuilding_date,
            issue.placement_date_end,
            issue.search_date_min,
            issue.search_date_max,
            article_id,
            article_url,
        ),
    )
    conn.commit()


def get_emitent_id(emitent_name: str, cache: dict, cache_file: str) -> base.EmitentMatch | None:
    cached = cache.setdefault("emitents", {}).get(emitent_name)
    if cached:
        return base.EmitentMatch(
            emitent_id=int(cached["emitent_id"]),
            ttl=str(cached.get("ttl") or emitent_name),
            status_id=None if cached.get("status_id") is None else str(cached.get("status_id")),
        )

    payload = {
        "term": quote(emitent_name, safe=""),
        "emitent_stop_statuses_ids": [],
    }

    resp = base.request_with_retry(
        "POST",
        base.SUGGEST_COMPANIES_URL,
        json=payload,
        headers={
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": base.SITE_BASE_URL,
            "Referer": f"{base.SITE_BASE_URL}/",
        },
    )
    data = base.parse_json_response(resp)
    items = data.get("items", [])
    match = base.choose_best_emitent_match(emitent_name, items)

    if match:
        cache.setdefault("emitents", {})[emitent_name] = {
            "emitent_id": match.emitent_id,
            "ttl": match.ttl,
            "status_id": match.status_id,
        }
        save_offer_cache(cache, cache_file)

    return match


def build_search_url(emitent_id: int, emitent_name: str, date_min: str, date_max: str, page: int = 1) -> str:
    query = "&".join(
        [
            f"emitent_id[]={quote(str(emitent_id), safe='')}",
            f"emitent_id_name[]={quote(emitent_name, safe='')}",
            f"date_min={quote(date_min, safe='')}",
            f"date_max={quote(date_max, safe='')}",
            f"page={page}",
        ]
    )
    return f"{base.NEWS_SEARCH_BASE_URL}?{query}"


def build_news_api_payload(emitent_id: int, date_min: str, date_max: str, page: int) -> dict:
    return {
        "filters": [
            {"field": "show_only_loans", "operator": "eq", "value": 0},
            {"field": "emitent_id", "operator": "in", "value": [str(emitent_id)]},
            {"field": "date", "operator": "ge", "value": date_min},
            {"field": "date", "operator": "le", "value": date_max},
        ],
        "lang": "rus",
        "quantity": {
            "offset": (page - 1) * base.NEWS_PAGE_LIMIT,
            "limit": base.NEWS_PAGE_LIMIT,
            "page": page,
        },
        "sorting": [],
    }


def fetch_news_items_for_range(
    emitent_name: str,
    emitent_id: int,
    date_min: str,
    date_max: str,
    cache: dict,
    cache_file: str,
    use_cache: bool,
) -> list[dict]:
    key = range_group_key(emitent_id, date_min, date_max)
    cached_items = cache.setdefault("search_results", {}).get(key)
    if use_cache and cached_items is not None:
        return cached_items

    all_items: list[dict] = []
    total_expected: int | None = None

    for page in range(1, base.MAX_API_PAGES + 1):
        search_url = build_search_url(emitent_id, emitent_name, date_min, date_max, page=page)
        payload = build_news_api_payload(emitent_id, date_min, date_max, page)

        resp = base.request_with_retry(
            "POST",
            base.NEWS_API_URL,
            json=payload,
            headers={
                "Accept": "application/json, text/plain, */*",
                "Content-Type": "application/json",
                "Origin": base.SITE_BASE_URL,
                "Referer": search_url,
            },
        )

        try:
            data = base.parse_json_response(resp)
        except Exception as exc:  # noqa: BLE001
            debug_path = base.dump_debug_json(
                f"failed_parse_{emitent_id}_{date_min}_{date_max}_page{page}",
                {
                    "error": str(exc),
                    "status_code": resp.status_code,
                    "headers": dict(resp.headers),
                    "text": resp.text,
                    "request_payload": payload,
                    "search_url": search_url,
                },
            )
            raise RuntimeError(f"Не удалось распарсить JSON API news. Debug: {debug_path}") from exc

        error_block = data.get("error") or {}
        err_no = error_block.get("err_no", 0)
        err_str = str(error_block.get("err_str") or "").strip()
        if err_no or err_str:
            debug_path = base.dump_debug_json(f"api_error_{emitent_id}_{date_min}_{date_max}_page{page}", data)
            raise RuntimeError(f"API /api/news/ вернул ошибку err_no={err_no} err_str={err_str!r}. Debug: {debug_path}")

        response_block = data.get("response") or {}
        items = response_block.get("items") or []

        if total_expected is None:
            total_expected = response_block.get("total")
            if total_expected is not None:
                total_expected = int(total_expected)

        if not items:
            break

        all_items.extend(items)
        base.maybe_sleep()

        if total_expected is not None and len(all_items) >= total_expected:
            break
        if len(items) < base.NEWS_PAGE_LIMIT:
            break

    cache.setdefault("search_results", {})[key] = all_items
    save_offer_cache(cache, cache_file)
    return all_items


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Собрать новости Cbonds по эмитенту в индивидуальном окне дат для каждого выпуска."
    )
    parser.add_argument("--input", default=INPUT_FILE, help=f"Входной Excel-файл, по умолчанию {INPUT_FILE}")
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=f"Имя листа. Если не указано, по умолчанию {DEFAULT_SHEET!r}.",
    )
    parser.add_argument(
        "--paper-column",
        default=DEFAULT_PAPER_COLUMN,
        help=f"Колонка названия выпуска, по умолчанию {DEFAULT_PAPER_COLUMN!r}.",
    )
    parser.add_argument(
        "--isin-column",
        default=DEFAULT_ISIN_COLUMN,
        help=f"Колонка ISIN, по умолчанию {DEFAULT_ISIN_COLUMN!r}.",
    )
    parser.add_argument(
        "--issuer-column",
        default=DEFAULT_ISSUER_COLUMN,
        help=f"Колонка эмитента, по умолчанию {DEFAULT_ISSUER_COLUMN!r}.",
    )
    parser.add_argument(
        "--book-date-column",
        default=DEFAULT_BOOK_DATE_COLUMN,
        help=f"Колонка даты начала букбилдинга, по умолчанию {DEFAULT_BOOK_DATE_COLUMN!r}.",
    )
    parser.add_argument(
        "--date-column",
        dest="legacy_date_column",
        default=None,
        help="Устаревший алиас для --book-date-column.",
    )
    parser.add_argument(
        "--placement-end-column",
        default=DEFAULT_PLACEMENT_END_COLUMN,
        help=f"Колонка даты окончания размещения, по умолчанию {DEFAULT_PLACEMENT_END_COLUMN!r}.",
    )
    parser.add_argument(
        "--include-blank-pb-tizer",
        action="store_true",
        help="Не отбрасывать строки с пустым pb_tizer, если включён --require-pb-tizer.",
    )
    parser.add_argument(
        "--require-pb-tizer",
        action="store_true",
        help="Обрабатывать только строки с непустым pb_tizer, если такая колонка нужна для другого файла.",
    )
    parser.add_argument("--filter-column", default=None, help="Колонка для дополнительного фильтра строк.")
    parser.add_argument("--filter-value", default=None, help="Значение дополнительного фильтра строк.")
    parser.add_argument("--db", default=DB_FILE, help=f"SQLite-база, по умолчанию {DB_FILE}")
    parser.add_argument("--cache", default=OFFER_CACHE_FILE, help=f"Кэш диапазонных запросов, по умолчанию {OFFER_CACHE_FILE}")
    parser.add_argument("--limit-groups", type=int, default=None, help="Ограничить число групп для тестового запуска.")
    parser.add_argument("--dry-run", action="store_true", help="Только показать рассчитанные группы без запросов к Cbonds.")
    parser.add_argument("--no-cookie", action="store_true", help="Не подключать Cookie даже из env/файла.")
    parser.add_argument("--skip-veb-rf", action="store_true", help="Пропустить выпуски эмитента ВЭБ.РФ / ВЭБ.РФ.")
    parser.add_argument("--no-cache", action="store_true", help="Не использовать кэш результатов поиска новостей.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    base_dir = Path(__file__).resolve().parent
    os.chdir(base_dir)
    if args.legacy_date_column:
        args.book_date_column = args.legacy_date_column

    if not Path(args.input).exists():
        raise FileNotFoundError(f"Не найден входной файл: {base_dir / args.input}")

    print(f"Excel: {args.input}")
    print(f"SQLite: {args.db}")

    require_pb_tizer = args.require_pb_tizer and not args.include_blank_pb_tizer
    if bool(args.filter_column) != bool(args.filter_value):
        raise ValueError("Параметры --filter-column и --filter-value нужно указывать вместе.")
    sheet_title, target_rows = read_target_rows(
        args.input,
        sheet_name=args.sheet,
        paper_column=args.paper_column,
        isin_column=args.isin_column,
        issuer_column=args.issuer_column,
        book_date_column=args.book_date_column,
        placement_end_column=args.placement_end_column,
        require_pb_tizer=require_pb_tizer,
        filter_column=args.filter_column,
        filter_value=args.filter_value,
    )
    skipped_veb = 0
    if args.skip_veb_rf:
        before_skip = len(target_rows)
        target_rows = [row for row in target_rows if not is_veb_rf(row.emitent)]
        skipped_veb = before_skip - len(target_rows)
    print(f"Лист: {sheet_title}")
    print(f"Колонка выпуска: {args.paper_column}")
    print(f"Колонка эмитента: {args.issuer_column}")
    print(f"Колонка начала букбилдинга: {args.book_date_column}")
    print(f"Колонка окончания размещения: {args.placement_end_column}")
    print(f"Фильтр pb_tizer непустой: {'да' if require_pb_tizer else 'нет'}")
    if args.filter_column:
        print(f"Дополнительный фильтр: {args.filter_column} = {args.filter_value}")
    if args.skip_veb_rf:
        print(f"ВЭБ.РФ пропущено строк: {skipped_veb}")
    print(f"Строк к обработке: {len(target_rows)}")

    grouped: dict[tuple[str, str, str, str, str], list[OfferIssueRow]] = defaultdict(list)
    for issue in target_rows:
        grouped[
            (
                issue.emitent,
                issue.bookbuilding_date,
                issue.placement_date_end,
                issue.search_date_min,
                issue.search_date_max,
            )
        ].append(issue)

    grouped_items = sorted(grouped.items())
    if args.limit_groups is not None:
        grouped_items = grouped_items[: args.limit_groups]

    total_groups = len(grouped_items)
    print(f"Уникальных групп (Эмитент + окно дат): {len(grouped)}")
    if args.limit_groups is not None:
        print(f"Тестовый лимит групп: {total_groups}")
    print()

    if args.dry_run:
        preview_limit = total_groups if args.limit_groups is not None else min(total_groups, 20)
        for idx, ((emitent_name, bookbuilding_date, placement_date_end, date_min, date_max), issues) in enumerate(
            grouped_items[:preview_limit],
            start=1,
        ):
            print(
                f"[{idx}/{total_groups}] {emitent_name} | book={bookbuilding_date} "
                f"| placement_end={placement_date_end} | news={date_min}..{date_max} "
                f"| выпусков: {len(issues)}"
            )
        if preview_limit < total_groups:
            print(f"... показаны первые {preview_limit} групп из {total_groups}.")
        print("\nDry-run: запросы к Cbonds и запись в SQLite не выполнялись.")
        return

    if args.no_cookie:
        base.session.headers.pop("Cookie", None)
    elif "Cookie" not in base.session.headers:
        cookie_header = base.load_cookie_header()
        if cookie_header:
            base.session.headers["Cookie"] = cookie_header

    print(f"Cookie: {'загружены' if 'Cookie' in base.session.headers else 'не заданы'}")

    cache = load_offer_cache(args.cache)
    conn = base.init_db(args.db)
    ensure_offer_schema(conn)

    for idx, ((emitent_name, bookbuilding_date, placement_date_end, date_min, date_max), issues) in enumerate(
        grouped_items,
        start=1,
    ):
        print(
            f"[{idx}/{total_groups}] {emitent_name} | book={bookbuilding_date} "
            f"| placement_end={placement_date_end} | news={date_min}..{date_max} "
            f"| выпусков: {len(issues)}"
        )

        try:
            match = get_emitent_id(emitent_name, cache, args.cache)
            if not match:
                print("    [skip] emitent_id не найден")
                continue

            print(f"    emitent_id={match.emitent_id} | ttl={match.ttl}")
            base.maybe_sleep()

            search_url = build_search_url(match.emitent_id, match.ttl, date_min, date_max, page=1)
            news_items = fetch_news_items_for_range(
                match.ttl,
                match.emitent_id,
                date_min,
                date_max,
                cache,
                args.cache,
                use_cache=not args.no_cache,
            )
            upsert_offer_search(
                conn,
                emitent_name,
                match.emitent_id,
                bookbuilding_date,
                placement_date_end,
                date_min,
                date_max,
                search_url,
                len(news_items),
            )

            print(f"    найдено новостей: {len(news_items)}")
            if not news_items:
                continue

            for item in news_items:
                article_id = base.save_article(conn, item)
                article_url = str(item.get("cb_link") or f"{base.SITE_BASE_URL}/news/{article_id}/")
                for issue in issues:
                    link_offer_issue_to_article(conn, issue, match.emitent_id, article_id, article_url)

        except Exception as exc:  # noqa: BLE001
            print(f"    [error] {exc}")

    conn.close()
    save_offer_cache(cache, args.cache)
    print(f"\nГотово. Результаты сохранены в {base_dir / args.db}")


if __name__ == "__main__":
    main()
