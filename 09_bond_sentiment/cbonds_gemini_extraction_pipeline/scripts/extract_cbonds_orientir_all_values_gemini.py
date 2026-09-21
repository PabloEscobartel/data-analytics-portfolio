#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Извлечение всех возможных ориентиров купона/спреда из новостей Cbonds.

Сценарий: для выпусков, у которых на Cbonds не найден период книги заявок
(`cbonds_book_status = not_found`), новости собраны за месяц до начала
размещения. Даты букбилдинга нет, поэтому Gemini не выбирает один "правильный"
ориентир по правилу времени, а возвращает все возможные найденные ориентиры
с ссылками на новости и кратким описанием контекста.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sqlite3
import time
from dataclasses import dataclass
from datetime import date, datetime, time as dt_time
from pathlib import Path
from typing import Any, Optional

import openpyxl

try:
    from google import genai
    from google.genai import types
except ImportError:  # pragma: no cover
    genai = None
    types = None


DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite-preview")
DEFAULT_DELAY = 5.0
DEFAULT_MAX_RETRIES = 5
DEFAULT_MAX_OUTPUT_TOKENS = 8192
DEFAULT_SAVE_EVERY = 10
DEFAULT_MAX_BATCH_CHARS = 25000
DEFAULT_MAX_ARTICLE_CHARS = 0
DEFAULT_MAX_ISSUE_NEWS = 0

INPUT_FILE = "bonds_offer_before2018_cbonds_orientir_gemini.xlsx"
OUTPUT_FILE = "bonds_offer_before2018_all_orientirs_gemini.xlsx"
DB_FILE = "cbonds_news.sqlite3"
CACHE_FILE = "gemini_all_orientirs_before2018_cache.json"
RESULTS_TABLE = "issue_orientir_all_values_before2018"

REQUIRED_HEADERS = [
    "Бумага",
    "ISIN",
    "Эмитент",
    "Начало размещения",
    "cbonds_book_status",
]
OPTIONAL_FLOATING_HEADERS = ["Floating_rate", "floating_rate"]

OUTPUT_COLUMNS = [
    "all_orientirs_count",
    "all_orientirs_coupon_values",
    "all_orientirs_spread_values",
    "all_orientirs_sources",
    "all_orientirs_json",
    "all_orientirs_needs_attention",
    "all_orientirs_status",
    "all_orientirs_model_name",
    "all_orientirs_model_version",
]

SYSTEM_PROMPT = """Ты аналитик российского первичного долгового рынка.

Тебе передают ОДИН выпуск облигаций и СПИСОК новостей Cbonds за месяц до начала размещения.
Для этих выпусков дата букбилдинга неизвестна, поэтому НЕ надо выбирать один ориентир по правилу времени.

Задача:
1. Найди ВСЕ возможные ориентиры купона, доходности, спреда или маржи, которые относятся к указанному выпуску/эмиссии.
2. Для каждого найденного ориентира верни тип, числовое значение, источник-новость и краткое описание контекста.
3. Если в новости обсуждаются несколько выпусков, выбери только значения, которые относятся к целевой эмиссии из ISSUE.paper.
4. Если связь значения с целевой эмиссией вероятна, но неоднозначна, сохрани значение и поставь needs_attention=true.

КРИТИЧЕСКИ ВАЖНО:
- Слово "ориентир" может отсутствовать. Ищи смысловой ориентир, а не только буквальное слово.
- Подходят формулировки: "не выше", "не более", "диапазон", "предварительный уровень", "КС + спред", "ставка купона ... %", если это маркетируемый уровень/ориентир для первичного размещения.
- Не возвращай фактическую ставку купона по итогам размещения, если это явно результат закрытия книги/размещения, а не guidance.
- Если ориентир задан диапазоном, например "7.75-8%" или "275-300 б.п.", orient_value должен быть ВЕРХНЕЙ границей диапазона: 8 или 300.
- Если формулировка "не выше / не более", orient_value равен указанному потолку.
- parsed_emissions у новости может отсутствовать или содержать только одну эмиссию, даже если в тексте обсуждаются две. Обязательно читай заголовок и текст.

Тип:
- coupon: значение в процентах годовых.
- spread: значение в базисных пунктах.

source_article_id:
- верни точный integer article_id новости из списка ARTICLES.

Верни только JSON-массив объектов без markdown. Один объект на каждый выпуск из ISSUES.
"""

RESULT_SCHEMA = {
    "type": "array",
    "items": {
        "type": "object",
        "properties": {
            "excel_row": {"type": "integer", "minimum": 2},
            "status": {"type": "string", "enum": ["found", "none"]},
            "orientirs": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "orient_type": {"type": "string", "enum": ["coupon", "spread"]},
                        "orient_value": {"type": "number", "minimum": 0},
                        "value_text": {"type": "string"},
                        "evidence_text": {"type": "string"},
                        "source_article_id": {"type": "integer"},
                        "source_emission_name": {"type": "string"},
                        "short_description": {"type": "string"},
                        "needs_attention": {"type": "boolean"},
                        "attention_reason": {"type": "string"},
                        "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                    },
                    "required": [
                        "orient_type",
                        "orient_value",
                        "value_text",
                        "evidence_text",
                        "source_article_id",
                        "source_emission_name",
                        "short_description",
                        "needs_attention",
                        "attention_reason",
                        "confidence",
                    ],
                    "additionalProperties": False,
                },
            },
            "reason_short": {"type": "string"},
        },
        "required": ["excel_row", "status", "orientirs", "reason_short"],
        "additionalProperties": False,
    },
}

LAST_REQUEST_TS = 0.0


@dataclass
class IssueInfo:
    excel_row: int
    paper: str
    isin: str
    emitent: str
    event_date: date
    floating_rate: str
    cbonds_book_status: str


@dataclass
class CandidateArticle:
    article_row_id: int
    article_id: int | None
    url: str
    title: str
    date_time: str
    published_at: datetime
    text: str
    emission_cbonds_id: int | None
    emission_name: str
    emission_url: str


@dataclass
class BatchItem:
    excel_row: int
    prompt_text: str
    estimated_chars: int


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def normalize_key(value: Any) -> str:
    text = normalize_text(value).casefold().replace("ё", "е")
    text = re.sub(r"\b(пао|ао|ооо|зао|оао)\b", " ", text)
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def safe_float_01(value: Any, default: float = 0.5) -> float:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return default
    if math.isnan(x):
        return default
    return max(0.0, min(1.0, x))


def safe_nonneg_number(value: Any, default: float = 0.0) -> float:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return default
    if math.isnan(x):
        return default
    return max(0.0, x)


def parse_date(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_text(value)
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def parse_article_datetime(value: Any) -> Optional[datetime]:
    text = normalize_text(value)
    if not text:
        return None
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text)
    except Exception:
        return None


def month_before(value: date) -> date:
    year = value.year
    month = value.month - 1
    if month == 0:
        month = 12
        year -= 1
    days = [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    return date(year, month, min(value.day, days[month - 1]))


def load_json(path: Path) -> dict[str, Any]:
    if path.exists():
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_json(path: Path, payload: dict[str, Any]) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


def build_header_index(ws) -> dict[str, int]:
    return {normalize_text(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}


def find_optional_column(headers: dict[str, int], variants: list[str]) -> int | None:
    normalized = {normalize_key(k): v for k, v in headers.items()}
    for variant in variants:
        idx = normalized.get(normalize_key(variant))
        if idx:
            return idx
    return None


def find_target_sheet(wb: openpyxl.Workbook, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Лист {sheet_name!r} не найден. Доступные листы: {', '.join(wb.sheetnames)}")
        return wb[sheet_name]
    required = {normalize_key(x) for x in REQUIRED_HEADERS}
    for ws in wb.worksheets:
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        if required.issubset({normalize_key(h) for h in headers if h is not None}):
            return ws
    raise RuntimeError(f"Не найден лист с колонками: {REQUIRED_HEADERS}")


def strip_workbook_images(wb: openpyxl.Workbook) -> None:
    for ws in wb.worksheets:
        if hasattr(ws, "_images"):
            ws._images = []


def load_excel_issues(
    input_path: Path,
    sheet_name: str | None,
    status_column: str,
    status_value: str,
    date_column: str,
) -> tuple[openpyxl.Workbook, Any, dict[int, IssueInfo]]:
    wb = openpyxl.load_workbook(input_path)
    strip_workbook_images(wb)
    ws = find_target_sheet(wb, sheet_name)
    headers = build_header_index(ws)
    required = ["Бумага", "ISIN", "Эмитент", date_column, status_column]
    missing = [name for name in required if name not in headers]
    if missing:
        raise RuntimeError(f"Не найдены колонки на листе {ws.title!r}: {missing}")

    floating_col = find_optional_column(headers, OPTIONAL_FLOATING_HEADERS)
    target_status = normalize_text(status_value).casefold()
    issues: dict[int, IssueInfo] = {}
    for row_idx in range(2, ws.max_row + 1):
        status = normalize_text(ws.cell(row_idx, headers[status_column]).value)
        if status.casefold() != target_status:
            continue

        paper = normalize_text(ws.cell(row_idx, headers["Бумага"]).value)
        isin = normalize_text(ws.cell(row_idx, headers["ISIN"]).value)
        emitent = normalize_text(ws.cell(row_idx, headers["Эмитент"]).value)
        event_date = parse_date(ws.cell(row_idx, headers[date_column]).value)
        floating_rate = normalize_text(ws.cell(row_idx, floating_col).value) if floating_col else ""
        if not paper or not emitent or event_date is None:
            continue

        issues[row_idx] = IssueInfo(
            excel_row=row_idx,
            paper=paper,
            isin=isin,
            emitent=emitent,
            event_date=event_date,
            floating_rate=floating_rate,
            cbonds_book_status=status,
        )

    return wb, ws, issues


def load_article_candidates(
    db_path: Path,
    issues: dict[int, IssueInfo],
    match_link_date: bool = True,
) -> dict[int, list[CandidateArticle]]:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        """
        SELECT
            ia.excel_row,
            ia.bookbuilding_date AS link_date,
            a.id AS article_row_id,
            a.article_id,
            a.url,
            a.title,
            a.date_time,
            a.content_text,
            a.introtext,
            GROUP_CONCAT(DISTINCT ae.emission_cbonds_id) AS emission_cbonds_ids,
            GROUP_CONCAT(DISTINCT ae.emission_name) AS emission_names,
            GROUP_CONCAT(DISTINCT ae.emission_url) AS emission_urls
        FROM issue_articles ia
        JOIN articles a ON a.article_id = ia.article_id
        LEFT JOIN article_emissions ae ON ae.article_row_id = a.id
        WHERE a.has_coupon_spread_term = 1
        GROUP BY
            ia.excel_row,
            ia.bookbuilding_date,
            a.id,
            a.article_id,
            a.url,
            a.title,
            a.date_time,
            a.content_text,
            a.introtext
        ORDER BY ia.excel_row, a.date_time, a.id
        """
    ).fetchall()
    conn.close()

    by_issue: dict[int, dict[int, CandidateArticle]] = {}
    for row in rows:
        excel_row = int(row["excel_row"])
        issue = issues.get(excel_row)
        if issue is None:
            continue

        if match_link_date and normalize_text(row["link_date"]) != issue.event_date.isoformat():
            continue

        published_at = parse_article_datetime(row["date_time"])
        if published_at is None:
            continue

        window_start = datetime.combine(month_before(issue.event_date), dt_time(0, 0, 0))
        window_end = datetime.combine(issue.event_date, dt_time(23, 59, 59))
        if not (window_start <= published_at <= window_end):
            continue

        emission_names = normalize_text(row["emission_names"])
        emission_urls = normalize_text(row["emission_urls"])
        emission_ids_text = normalize_text(row["emission_cbonds_ids"])
        emission_cbonds_id = None
        if emission_ids_text and "," not in emission_ids_text:
            try:
                emission_cbonds_id = int(emission_ids_text)
            except ValueError:
                emission_cbonds_id = None

        article_row_id = int(row["article_row_id"])
        text = normalize_text(" ".join([normalize_text(row["introtext"]), normalize_text(row["content_text"])]))
        candidate = CandidateArticle(
            article_row_id=article_row_id,
            article_id=None if row["article_id"] is None else int(row["article_id"]),
            url=normalize_text(row["url"]),
            title=normalize_text(row["title"]),
            date_time=normalize_text(row["date_time"]),
            published_at=published_at,
            text=text,
            emission_cbonds_id=emission_cbonds_id,
            emission_name=emission_names,
            emission_url=emission_urls,
        )
        by_issue.setdefault(excel_row, {})[article_row_id] = candidate

    return {row_id: sorted(items.values(), key=lambda x: x.published_at) for row_id, items in by_issue.items()}


def build_issue_prompt_block(
    issue: IssueInfo,
    candidates: list[CandidateArticle],
    max_article_chars: int = DEFAULT_MAX_ARTICLE_CHARS,
    max_issue_news: int = DEFAULT_MAX_ISSUE_NEWS,
) -> str:
    selected_candidates = sorted(candidates, key=lambda x: x.published_at)
    if max_issue_news and max_issue_news > 0:
        selected_candidates = selected_candidates[:max_issue_news]

    article_blocks: list[str] = []
    for idx, candidate in enumerate(selected_candidates, start=1):
        text = candidate.text
        if max_article_chars and len(text) > max_article_chars:
            text = text[:max_article_chars]
        article_blocks.append(
            "\n".join(
                [
                    f"[A{idx}]",
                    f"article_id: {candidate.article_id or 'NA'}",
                    f"published_at: {candidate.date_time}",
                    f"url: {candidate.url}",
                    f"title: {candidate.title}",
                    f"parsed_emissions: {candidate.emission_name or 'NA'}",
                    f"parsed_emission_urls: {candidate.emission_url or 'NA'}",
                    "text:",
                    text or "NA",
                ]
            )
        )

    return "\n".join(
        [
            "ISSUE:",
            f"excel_row: {issue.excel_row}",
            f"paper: {issue.paper}",
            f"isin: {issue.isin or 'NA'}",
            f"emitent: {issue.emitent}",
            f"placement_start_date: {issue.event_date.isoformat()}",
            f"floating_rate: {issue.floating_rate or 'NA'}",
            "",
            "TASK:",
            "Найди все возможные ориентиры купона/доходности/спреда/маржи для этой эмиссии во всех ARTICLES.",
            "Не выбирай один лучший ориентир: верни список всех разных значений и источников, если они относятся к целевой эмиссии.",
            "Если одинаковое значение повторяется в нескольких новостях, можно вернуть несколько источников только если контекст существенно отличается; иначе оставь самый информативный источник.",
            "Если ориентир относится к другому выпуску в той же новости, не включай его.",
            "",
            "ARTICLES:",
            "\n\n".join(article_blocks) if article_blocks else "NA",
        ]
    )


def build_batch_prompt(batch: list[BatchItem]) -> str:
    return "\n\n".join(
        [
            "Ниже несколько выпусков. Для каждого блока ISSUE верни ровно один объект в JSON-массиве.",
            "Ключ сопоставления результата: excel_row.",
            "Для каждого выпуска анализируй только его ARTICLES.",
            *(item.prompt_text for item in batch),
        ]
    )


def make_batches(items: list[BatchItem], max_batch_chars: int) -> list[list[BatchItem]]:
    batches: list[list[BatchItem]] = []
    batch_chars: list[int] = []
    for item in sorted(items, key=lambda x: x.estimated_chars, reverse=True):
        placed = False
        for idx, current_chars in enumerate(batch_chars):
            if current_chars + item.estimated_chars <= max_batch_chars:
                batches[idx].append(item)
                batch_chars[idx] += item.estimated_chars
                placed = True
                break
        if not placed:
            batches.append([item])
            batch_chars.append(item.estimated_chars)
    return batches


def extract_json_array(raw_text: str) -> str:
    raw_text = (raw_text or "").strip()
    if raw_text.startswith("```"):
        raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text, flags=re.IGNORECASE)
        raw_text = re.sub(r"\s*```$", "", raw_text).strip()
    start = raw_text.find("[")
    end = raw_text.rfind("]")
    if start != -1 and end != -1 and end > start:
        return raw_text[start : end + 1]
    raise ValueError("JSON array not found in model output")


def parse_model_items(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        parsed = payload
    elif isinstance(payload, str):
        parsed = json.loads(extract_json_array(payload))
    else:
        raise ValueError("Unsupported model payload type")
    if not isinstance(parsed, list):
        raise ValueError("Model output is not a list")
    return [item for item in parsed if isinstance(item, dict)]


def normalize_model_item(
    item: dict[str, Any],
    candidates: list[CandidateArticle],
    model_name: str,
    model_version: str | None,
) -> dict[str, Any]:
    candidates_by_article_id = {c.article_id: c for c in candidates if c.article_id is not None}
    normalized_orientirs: list[dict[str, Any]] = []
    seen: set[tuple[str, float, int]] = set()

    for raw in item.get("orientirs") or []:
        if not isinstance(raw, dict):
            continue
        orient_type = normalize_text(raw.get("orient_type")).lower()
        if orient_type not in {"coupon", "spread"}:
            continue
        orient_value = safe_nonneg_number(raw.get("orient_value"), 0.0)
        if orient_value <= 0:
            continue
        try:
            source_article_id = int(raw.get("source_article_id") or 0)
        except (TypeError, ValueError):
            source_article_id = 0
        source_candidate = candidates_by_article_id.get(source_article_id)

        key = (orient_type, round(orient_value, 6), source_article_id)
        if key in seen:
            continue
        seen.add(key)

        needs_attention = bool(raw.get("needs_attention"))
        attention_reason = normalize_text(raw.get("attention_reason")) if needs_attention else ""
        if source_article_id and source_candidate is None:
            needs_attention = True
            attention_reason = " ".join(
                x
                for x in [
                    attention_reason,
                    "Gemini вернул source_article_id, которого нет в переданном списке новостей.",
                ]
                if x
            ).strip()

        normalized_orientirs.append(
            {
                "orient_type": orient_type,
                "orient_value": orient_value,
                "value_text": normalize_text(raw.get("value_text"))[:240],
                "evidence_text": normalize_text(raw.get("evidence_text"))[:900],
                "source_article_id": source_article_id or None,
                "source_title": source_candidate.title if source_candidate else "",
                "source_url": source_candidate.url if source_candidate else "",
                "source_datetime": source_candidate.date_time if source_candidate else "",
                "emission_cbonds_id": source_candidate.emission_cbonds_id if source_candidate else None,
                "source_emission_name": normalize_text(raw.get("source_emission_name")) or (source_candidate.emission_name if source_candidate else ""),
                "short_description": normalize_text(raw.get("short_description"))[:500],
                "needs_attention": int(needs_attention),
                "attention_reason": attention_reason,
                "confidence": safe_float_01(raw.get("confidence"), 0.5),
            }
        )

    status = "found" if normalized_orientirs else "none"
    return {
        "status": status,
        "orientirs": normalized_orientirs,
        "reason_short": normalize_text(item.get("reason_short"))[:300],
        "model_name": model_name,
        "model_version": model_version,
    }


def empty_result(status: str, reason: str) -> dict[str, Any]:
    return {
        "status": status,
        "orientirs": [],
        "reason_short": reason,
        "model_name": "",
        "model_version": "",
    }


def sleep_to_respect_min_interval(min_interval: float) -> None:
    global LAST_REQUEST_TS
    elapsed = time.time() - LAST_REQUEST_TS
    if LAST_REQUEST_TS and elapsed < min_interval:
        time.sleep(min_interval - elapsed)


def mark_request_sent() -> None:
    global LAST_REQUEST_TS
    LAST_REQUEST_TS = time.time()


def extract_retry_delay_seconds(exc: Exception) -> float | None:
    text = str(exc)
    for pattern in (r"retry in\s+([0-9]+(?:\.[0-9]+)?)s", r"Retry in\s+([0-9]+(?:\.[0-9]+)?)s"):
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return float(match.group(1))
    return None


def call_gemini(client: Any, model_name: str, prompt: str, max_retries: int, max_output_tokens: int, delay: float) -> tuple[Any, str | None]:
    last_exc: Exception | None = None
    for attempt in range(1, max_retries + 1):
        try:
            sleep_to_respect_min_interval(delay)
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=SYSTEM_PROMPT,
                    temperature=0,
                    max_output_tokens=max_output_tokens,
                    response_mime_type="application/json",
                    response_json_schema=RESULT_SCHEMA,
                ),
            )
            mark_request_sent()
            parsed = getattr(response, "parsed", None)
            raw_text = getattr(response, "text", "") or ""
            version = getattr(response, "model_version", None)
            return (parsed if parsed is not None else raw_text), version
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
            wait = extract_retry_delay_seconds(exc)
            if wait is None:
                wait = min(5 * (2 ** (attempt - 1)), 60)
            if attempt < max_retries:
                print(f"  retry {attempt}/{max_retries} after {wait:.1f}s: {exc}")
                time.sleep(wait)
    raise RuntimeError(str(last_exc) if last_exc else "Unknown API error")


def validate_sql_identifier(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value or ""):
        raise ValueError(f"Некорректное имя SQL-таблицы: {value!r}")
    return value


def ensure_results_table(conn: sqlite3.Connection, table_name: str) -> None:
    table_name = validate_sql_identifier(table_name)
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {table_name} (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            excel_row INTEGER NOT NULL,
            paper TEXT,
            isin TEXT,
            emitent TEXT,
            event_date TEXT,
            orient_index INTEGER NOT NULL,
            orient_type TEXT,
            orient_value REAL,
            value_text TEXT,
            evidence_text TEXT,
            source_article_id INTEGER,
            source_title TEXT,
            source_url TEXT,
            source_datetime TEXT,
            emission_cbonds_id INTEGER,
            source_emission_name TEXT,
            short_description TEXT,
            needs_attention INTEGER NOT NULL DEFAULT 0,
            attention_reason TEXT,
            confidence REAL,
            model_name TEXT,
            model_version TEXT,
            status TEXT,
            updated_at TEXT NOT NULL,
            UNIQUE(excel_row, orient_index)
        )
        """
    )
    conn.commit()


def write_cache_to_db(db_path: Path, issues: dict[int, IssueInfo], cache: dict[str, Any], table_name: str) -> None:
    conn = sqlite3.connect(db_path)
    table_name = validate_sql_identifier(table_name)
    ensure_results_table(conn, table_name)
    now = datetime.now().isoformat(timespec="seconds")

    row_ids = []
    rows = []
    for key, result in cache.items():
        try:
            row_idx = int(key)
        except ValueError:
            continue
        issue = issues.get(row_idx)
        if issue is None or not isinstance(result, dict):
            continue
        row_ids.append((row_idx,))
        orientirs = result.get("orientirs") or []
        if not orientirs:
            rows.append(
                (
                    row_idx,
                    issue.paper,
                    issue.isin,
                    issue.emitent,
                    issue.event_date.isoformat(),
                    0,
                    None,
                    None,
                    "",
                    "",
                    None,
                    "",
                    "",
                    "",
                    None,
                    "",
                    "",
                    0,
                    result.get("reason_short") or "",
                    0.0,
                    result.get("model_name") or "",
                    result.get("model_version") or "",
                    result.get("status") or "none",
                    now,
                )
            )
            continue
        for idx, orientir in enumerate(orientirs, start=1):
            rows.append(
                (
                    row_idx,
                    issue.paper,
                    issue.isin,
                    issue.emitent,
                    issue.event_date.isoformat(),
                    idx,
                    orientir.get("orient_type"),
                    orientir.get("orient_value"),
                    orientir.get("value_text"),
                    orientir.get("evidence_text"),
                    orientir.get("source_article_id"),
                    orientir.get("source_title"),
                    orientir.get("source_url"),
                    orientir.get("source_datetime"),
                    orientir.get("emission_cbonds_id"),
                    orientir.get("source_emission_name"),
                    orientir.get("short_description"),
                    int(orientir.get("needs_attention") or 0),
                    orientir.get("attention_reason"),
                    orientir.get("confidence"),
                    result.get("model_name"),
                    result.get("model_version"),
                    result.get("status"),
                    now,
                )
            )

    if row_ids:
        conn.executemany(f"DELETE FROM {table_name} WHERE excel_row = ?", row_ids)
    conn.executemany(
        f"""
        INSERT INTO {table_name} (
            excel_row, paper, isin, emitent, event_date, orient_index,
            orient_type, orient_value, value_text, evidence_text,
            source_article_id, source_title, source_url, source_datetime,
            emission_cbonds_id, source_emission_name, short_description,
            needs_attention, attention_reason, confidence,
            model_name, model_version, status, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )
    conn.commit()
    conn.close()


def ensure_output_columns(ws) -> dict[str, int]:
    headers = build_header_index(ws)
    for col_name in OUTPUT_COLUMNS:
        if col_name not in headers:
            col_idx = ws.max_column + 1
            ws.cell(1, col_idx, col_name)
            headers[col_name] = col_idx
    return headers


def summarize_values(orientirs: list[dict[str, Any]], orient_type: str) -> str:
    values = []
    for item in orientirs:
        if item.get("orient_type") == orient_type:
            value = item.get("orient_value")
            if value is not None:
                values.append(str(value))
    return "; ".join(dict.fromkeys(values))


def summarize_sources(orientirs: list[dict[str, Any]]) -> str:
    parts = []
    for item in orientirs:
        source = " | ".join(
            x
            for x in [
                normalize_text(item.get("source_datetime")),
                normalize_text(item.get("source_title")),
                normalize_text(item.get("source_url")),
                normalize_text(item.get("short_description")),
            ]
            if x
        )
        if source:
            parts.append(source)
    return "\n".join(parts)[:32000]


def apply_cache_to_workbook(ws, headers: dict[str, int], cache: dict[str, Any]) -> None:
    for key, result in cache.items():
        try:
            row_idx = int(key)
        except ValueError:
            continue
        if not isinstance(result, dict) or row_idx < 2 or row_idx > ws.max_row:
            continue
        orientirs = result.get("orientirs") or []
        ws.cell(row_idx, headers["all_orientirs_count"], len(orientirs))
        ws.cell(row_idx, headers["all_orientirs_coupon_values"], summarize_values(orientirs, "coupon"))
        ws.cell(row_idx, headers["all_orientirs_spread_values"], summarize_values(orientirs, "spread"))
        ws.cell(row_idx, headers["all_orientirs_sources"], summarize_sources(orientirs))
        ws.cell(row_idx, headers["all_orientirs_json"], json.dumps(orientirs, ensure_ascii=False))
        ws.cell(row_idx, headers["all_orientirs_needs_attention"], int(any(int(x.get("needs_attention") or 0) for x in orientirs)))
        ws.cell(row_idx, headers["all_orientirs_status"], result.get("status"))
        ws.cell(row_idx, headers["all_orientirs_model_name"], result.get("model_name"))
        ws.cell(row_idx, headers["all_orientirs_model_version"], result.get("model_version"))


def save_workbook(wb: openpyxl.Workbook, ws, cache: dict[str, Any], output_path: Path) -> None:
    headers = ensure_output_columns(ws)
    apply_cache_to_workbook(ws, headers, cache)
    try:
        wb.save(output_path)
    except ValueError as exc:
        if "closed file" not in str(exc).lower():
            raise
        strip_workbook_images(wb)
        wb.save(output_path)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Extract all possible Cbonds orientirs from candidate news with Gemini.")
    parser.add_argument("--input", type=Path, default=Path(INPUT_FILE))
    parser.add_argument("--output", type=Path, default=Path(OUTPUT_FILE))
    parser.add_argument("--sheet", type=str, default="place_before_2018")
    parser.add_argument("--db", type=Path, default=Path(DB_FILE))
    parser.add_argument("--results-table", type=str, default=RESULTS_TABLE)
    parser.add_argument("--cache", type=Path, default=Path(CACHE_FILE))
    parser.add_argument("--date-column", type=str, default="Начало размещения")
    parser.add_argument("--status-column", type=str, default="cbonds_book_status")
    parser.add_argument("--status-value", type=str, default="not_found")
    parser.add_argument("--no-match-link-date", action="store_true", help="Не требовать совпадения issue_articles.bookbuilding_date с date-column.")
    parser.add_argument("--model", type=str, default=DEFAULT_MODEL)
    parser.add_argument("--api-key", type=str, default=os.getenv("GEMINI_API_KEY", ""))
    parser.add_argument("--delay", type=float, default=DEFAULT_DELAY)
    parser.add_argument("--max-retries", type=int, default=DEFAULT_MAX_RETRIES)
    parser.add_argument("--max-output-tokens", type=int, default=DEFAULT_MAX_OUTPUT_TOKENS)
    parser.add_argument("--max-requests", type=int, default=None)
    parser.add_argument("--max-batch-chars", type=int, default=DEFAULT_MAX_BATCH_CHARS)
    parser.add_argument("--max-article-chars", type=int, default=DEFAULT_MAX_ARTICLE_CHARS)
    parser.add_argument("--max-issue-news", type=int, default=DEFAULT_MAX_ISSUE_NEWS)
    parser.add_argument("--one-issue-per-batch", action="store_true", help="Не объединять выпуски в батчи.")
    parser.add_argument("--save-every", type=int, default=DEFAULT_SAVE_EVERY)
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-confirm", action="store_true")
    return parser


def main() -> None:
    args = build_arg_parser().parse_args()
    wb, ws, issues = load_excel_issues(
        args.input,
        sheet_name=args.sheet,
        status_column=args.status_column,
        status_value=args.status_value,
        date_column=args.date_column,
    )
    candidates_by_issue = load_article_candidates(
        args.db,
        issues,
        match_link_date=not args.no_match_link_date,
    )
    cache = load_json(args.cache)

    no_candidate_count = 0
    for row_idx in issues:
        if not candidates_by_issue.get(row_idx):
            no_candidate_count += 1
            cache.setdefault(str(row_idx), empty_result("no_candidate", "Не найдена новость с has_coupon_spread_term=1 за окно до начала размещения."))

    rows_with_candidates = [row_idx for row_idx in sorted(issues) if candidates_by_issue.get(row_idx)]
    pending_rows = [row_idx for row_idx in rows_with_candidates if args.refresh or str(row_idx) not in cache]

    pending_items: list[BatchItem] = []
    for row_idx in pending_rows:
        prompt_text = build_issue_prompt_block(
            issues[row_idx],
            candidates_by_issue[row_idx],
            max_article_chars=args.max_article_chars,
            max_issue_news=args.max_issue_news,
        )
        pending_items.append(
            BatchItem(
                excel_row=row_idx,
                prompt_text=prompt_text,
                estimated_chars=len(prompt_text) + 500,
            )
        )

    batches = [[item] for item in pending_items] if args.one_issue_per_batch else make_batches(pending_items, args.max_batch_chars)
    if args.max_requests is not None:
        batches = batches[: args.max_requests]

    print("=" * 72)
    print("All possible orientirs extraction from Cbonds news")
    print("=" * 72)
    print(f"Input: {args.input}")
    print(f"Output: {args.output}")
    print(f"DB: {args.db}")
    print(f"Results table: {args.results_table}")
    print(f"Sheet: {ws.title}")
    print(f"Date column: {args.date_column}")
    print(f"Filter: {args.status_column} = {args.status_value}")
    print(f"Match link date: {not args.no_match_link_date}")
    print(f"Issues eligible: {len(issues):,}")
    print(f"Issues with candidate news: {len(rows_with_candidates):,}")
    print(f"No candidate: {no_candidate_count:,}")
    print(f"Cache rows: {len(cache):,}")
    print(f"Pending issues: {len(pending_rows):,}")
    print(f"Planned batches: {len(batches):,}")
    print(f"Max batch chars: {args.max_batch_chars:,}")
    print(f"Max article chars: {args.max_article_chars:,}")
    print(f"Max issue news: {args.max_issue_news:,}")
    total_candidate_news = sum(len(items) for items in candidates_by_issue.values())
    print(f"Candidate news total: {total_candidate_news:,}")
    if batches:
        batch_chars = [len(build_batch_prompt(batch)) for batch in batches]
        print(f"Batch chars avg: {sum(batch_chars) / len(batch_chars):,.1f}")
        print(f"Batch chars max: {max(batch_chars):,}")

    if args.dry_run:
        save_workbook(wb, ws, cache, args.output)
        write_cache_to_db(args.db, issues, cache, args.results_table)
        print("\nDry run complete. Workbook written from existing cache only.")
        return

    if genai is None or types is None:
        raise ImportError("Не найден пакет google-genai. Установи его: pip install google-genai")
    if not args.api_key:
        raise ValueError("API key not found. Pass --api-key or set GEMINI_API_KEY.")

    if batches and not args.no_confirm:
        answer = input(f"Send {len(batches)} batched requests to Gemini? (y/n): ").strip().lower()
        if answer != "y":
            print("Cancelled by user.")
            save_workbook(wb, ws, cache, args.output)
            write_cache_to_db(args.db, issues, cache, args.results_table)
            return

    client = genai.Client(api_key=args.api_key)
    try:
        for idx, batch in enumerate(batches, start=1):
            batch_rows = [item.excel_row for item in batch]
            print(f"\nBatch {idx}/{len(batches)} | issues={len(batch)} | rows={batch_rows[:8]}{'...' if len(batch_rows) > 8 else ''}")
            try:
                payload, model_version = call_gemini(
                    client=client,
                    model_name=args.model,
                    prompt=build_batch_prompt(batch),
                    max_retries=args.max_retries,
                    max_output_tokens=args.max_output_tokens,
                    delay=args.delay,
                )
                parsed_items = parse_model_items(payload)
                results_by_row: dict[int, dict[str, Any]] = {}
                for item in parsed_items:
                    try:
                        item_row = int(item.get("excel_row"))
                    except (TypeError, ValueError):
                        continue
                    if item_row not in batch_rows:
                        continue
                    results_by_row[item_row] = normalize_model_item(
                        item,
                        candidates_by_issue[item_row],
                        args.model,
                        model_version,
                    )

                for row_idx in batch_rows:
                    result = results_by_row.get(row_idx)
                    if result is None:
                        result = empty_result("missing_item_in_response", "Gemini не вернул объект для этой строки в батче.")
                    cache[str(row_idx)] = result

                extracted_now = sum(len(r.get("orientirs") or []) for r in results_by_row.values())
                print(f"  -> returned={len(results_by_row)}/{len(batch_rows)} orientirs={extracted_now}")
            except Exception as exc:  # noqa: BLE001
                for row_idx in batch_rows:
                    cache[str(row_idx)] = empty_result("api_error", f"{type(exc).__name__}: {exc}")
                print(f"  [error] {exc}")

            if idx % args.save_every == 0:
                save_json(args.cache, cache)
                save_workbook(wb, ws, cache, args.output)
                write_cache_to_db(args.db, issues, cache, args.results_table)
                print(f"  saved: {args.cache}, {args.output}")

    except KeyboardInterrupt:
        print("\nInterrupted. Saving progress...")
    finally:
        save_json(args.cache, cache)
        save_workbook(wb, ws, cache, args.output)
        write_cache_to_db(args.db, issues, cache, args.results_table)
        print(f"\nSaved cache: {args.cache}")
        print(f"Saved workbook: {args.output}")
        print(f"Saved DB table: {args.results_table}")

    extracted = sum(len(item.get("orientirs") or []) for item in cache.values() if isinstance(item, dict))
    attention = sum(
        1
        for item in cache.values()
        if isinstance(item, dict)
        for orientir in item.get("orientirs") or []
        if int(orientir.get("needs_attention") or 0) == 1
    )
    print(f"Extracted orientirs: {extracted:,}")
    print(f"Needs attention: {attention:,}")


if __name__ == "__main__":
    main()
