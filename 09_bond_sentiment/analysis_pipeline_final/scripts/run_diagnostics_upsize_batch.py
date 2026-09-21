#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Run upsize diagnostics for several sentiment-window/min-message datasets
and collect the outputs into one workbook.

Default datasets:
  regression_dataset_book_open_3/regression_ready.csv
  regression_dataset_book_open_5/regression_ready.csv
  regression_dataset_book_open_10/regression_ready.csv
  regression_dataset_book_close_3/regression_ready.csv
  regression_dataset_book_close_5/regression_ready.csv
  regression_dataset_book_close_10/regression_ready.csv

Usage:
  python run_diagnostics_upsize_batch.py
  python run_diagnostics_upsize_batch.py diagnostics_results_upsize_all_windows.xlsx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import run_diagnostics_upsize as diag


DEFAULT_OUTPUT_FILE = "diagnostics_results_upsize_all_windows.xlsx"
DEFAULT_DATASET_TEMPLATE = "regression_dataset_{window}_{min_messages}/regression_ready.csv"
WINDOWS = ["book_open", "book_close"]
MIN_MESSAGES = [3, 5, 10]
SAMPLE_MODES = ["all", "fixed", "floater"]
SAMPLE_LABELS = {
    "all": "All bonds",
    "fixed": "Fixed coupon bonds",
    "floater": "Floating-rate bonds",
}
ORIGINAL_BASE_VARS = list(diag.BASE_VARS)
ORIGINAL_MAIN_INTERACTION_VARS = list(diag.MAIN_INTERACTION_VARS)
ORIGINAL_NONLINEAR_VARS = list(diag.NONLINEAR_VARS)
ORIGINAL_INTERACTION_VARS = list(diag.INTERACTION_VARS)
ORIGINAL_EXTENSION_TERMS = list(diag.EXTENSION_TERMS)


def find_dataset(window: str, min_messages: int, dataset_template: str = DEFAULT_DATASET_TEMPLATE) -> Optional[Path]:
    candidates = [Path(dataset_template.format(window=window, min_messages=min_messages))]
    if dataset_template == DEFAULT_DATASET_TEMPLATE and min_messages == 3:
        candidates.append(Path(f"regression_dataset_{window}/regression_ready.csv"))
    for path in candidates:
        if path.exists():
            return path
    return None


def output_for_sample(output_path: Path, sample_mode: str) -> Path:
    if sample_mode == "all":
        return output_path
    return output_path.with_name(f"{output_path.stem}_{sample_mode}{output_path.suffix}")


def apply_sample_mode(full_sub: pd.DataFrame, baseline_sub: pd.DataFrame, sample_mode: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    if sample_mode == "all":
        return full_sub, baseline_sub
    if sample_mode not in {"fixed", "floater"}:
        raise ValueError(f"Unknown sample mode: {sample_mode}")

    target = 0 if sample_mode == "fixed" else 1

    def filt(df: pd.DataFrame) -> pd.DataFrame:
        out = df.loc[pd.to_numeric(df["is_floater"], errors="coerce").eq(target)].copy()
        if sample_mode == "floater":
            out["time_fe"] = np.where(
                pd.to_numeric(out["year"], errors="coerce").le(2021),
                "2018-2021",
                "2022-2025",
            )
        return out

    return filt(full_sub), filt(baseline_sub)


def configure_diagnostic_spec(sample_mode: str, history_spec: str = "baseline") -> None:
    base_vars = list(ORIGINAL_BASE_VARS)
    main_interaction_vars = list(ORIGINAL_MAIN_INTERACTION_VARS)
    extension_terms = list(ORIGINAL_EXTENSION_TERMS)
    nonlinear_vars = list(ORIGINAL_NONLINEAR_VARS)
    interaction_vars = list(ORIGINAL_INTERACTION_VARS)

    if history_spec in {"share", "baseline"}:
        base_vars = [
            "hist_reduction_share" if v == "history_repeat_prior_tightening" else v
            for v in base_vars
        ]
    elif history_spec == "dummy":
        base_vars = [
            "history_repeat_prior_tightening" if v == "hist_reduction_share" else v
            for v in base_vars
        ]
    else:
        raise ValueError(f"Unknown history specification: {history_spec}")

    if sample_mode in {"fixed", "floater"}:
        base_vars = [v for v in base_vars if v != "is_floater"]
        main_interaction_vars = []
        interaction_vars = [v for v in interaction_vars if v not in {"si_x_is_floater", "log_buzz_x_is_floater"}]
        extension_terms = [t for t in extension_terms if t[0] not in {"si_x_is_floater", "log_buzz_x_is_floater"}]
    else:
        interaction_vars = [v for v in interaction_vars if v not in main_interaction_vars]

    diag.BASE_VARS = base_vars
    diag.MAIN_INTERACTION_VARS = main_interaction_vars
    diag.EXPLANATORY_VARS = base_vars + ["si", "log_buzz"] + main_interaction_vars
    diag.INTERACTION_VARS = interaction_vars
    diag.EXTENSION_TERMS = extension_terms
    diag.RESET_SPECS = [
        ("Baseline", []),
        ("Baseline + nonlinear", nonlinear_vars),
        ("Baseline + interactions", interaction_vars),
        ("Baseline + nonlinear + interactions", nonlinear_vars + interaction_vars),
    ]


def parse_pipeline_summary(path: Path) -> Dict[str, Any]:
    summary_path = path.parent / "pipeline_summary.txt"
    out: Dict[str, Any] = {}
    if not summary_path.exists():
        return out
    for line in summary_path.read_text(encoding="utf-8").splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = key.strip().lower().replace(" ", "_")
        nums = re.findall(r"[\d,]+", value)
        out[key] = int(nums[0].replace(",", "")) if nums else value.strip()
    return out


def add_meta(df: pd.DataFrame, window: str, min_messages: int, dataset: Path, sample_mode: str) -> pd.DataFrame:
    out = df.copy()
    out.insert(0, "dataset", str(dataset))
    out.insert(0, "sample_mode", sample_mode)
    out.insert(0, "min_messages", min_messages)
    out.insert(0, "window", window)
    return out


def collect_sample_summary(
    window: str,
    min_messages: int,
    dataset: Path,
    sample_mode: str,
    full_sub: pd.DataFrame,
    baseline_sub: pd.DataFrame,
    vif_df: pd.DataFrame,
    reset_df: pd.DataFrame,
    infl_df: pd.DataFrame,
) -> Dict[str, Any]:
    pipe = parse_pipeline_summary(dataset)
    return {
        "window": window,
        "min_messages": min_messages,
        "sample_mode": sample_mode,
        "sample_label": SAMPLE_LABELS[sample_mode],
        "dataset": str(dataset),
        "pipeline_sentiment_rows": pipe.get("with_sentiment", np.nan),
        "aggregated_offerings": pipe.get("aggregated_offerings", np.nan),
        "baseline_N": len(baseline_sub),
        "baseline_coupon_N": int(baseline_sub["coupon_reduction_bp"].notna().sum()),
        "baseline_volume_N": int(baseline_sub["placement_vol_book"].notna().sum()),
        "full_N": len(full_sub),
        "issuers": int(baseline_sub["issuer"].nunique()),
        "max_VIF": float(vif_df["VIF"].max()) if not vif_df.empty else np.nan,
        "min_RESET_p": float(pd.to_numeric(reset_df["p_value"], errors="coerce").min()) if not reset_df.empty else np.nan,
        "max_Cook_D": float(infl_df["max_cooks_d"].max()) if not infl_df.empty else np.nan,
    }


def autosize(ws) -> None:
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 44)


def write_df_section(ws, title: str, df: pd.DataFrame, start_row: int) -> int:
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=13, name="Arial")
    row = start_row + 2
    if df.empty:
        ws.cell(row=row, column=1, value="No rows").font = Font(name="Arial", size=10)
        return row + 2

    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=j, value=col)
        cell.font = Font(bold=True, name="Arial", size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    row += 1

    for _, rec in df.iterrows():
        for j, col in enumerate(df.columns, start=1):
            val = rec[col]
            if isinstance(val, np.integer):
                val = int(val)
            elif isinstance(val, np.floating):
                val = float(val)
            ws.cell(row=row, column=j, value=val).font = Font(name="Arial", size=10)
        row += 1
    return row + 2


def write_scenario_sheet(ws, item: Dict[str, Any]) -> None:
    window = item["window"]
    min_messages = item["min_messages"]
    sample_mode = item["sample_mode"]
    dataset = item["dataset"]
    full_sub = item["full_sub"]
    baseline_sub = item["baseline_sub"]

    ws["A1"] = f"{window}, min_messages >= {min_messages}, sample: {SAMPLE_LABELS.get(sample_mode, sample_mode)}"
    ws["A1"].font = Font(bold=True, size=15, name="Arial")
    ws["A2"] = f"Dataset: {dataset}"
    ws["A3"] = (
        f"Baseline N={len(baseline_sub)}, coupon N={baseline_sub['coupon_reduction_bp'].notna().sum()}, "
        f"volume N={baseline_sub['placement_vol_book'].notna().sum()}"
    )

    row = 5
    row = write_df_section(ws, "Sample Filters", item["sample_df"], row)
    row = write_df_section(ws, "Data Quality", item["quality_df"], row)
    row = write_df_section(ws, "VIF", item["vif_df"], row)
    row = write_df_section(ws, "VIF with Fixed Effects", item["vif_fe_df"], row)
    row = write_df_section(ws, "High Correlations |r| > 0.3", item["high_corr_df"], row)
    row = write_df_section(ws, "Heteroskedasticity", item["het_df"], row)
    row = write_df_section(ws, "Normality", item["norm_df"], row)
    row = write_df_section(ws, "RESET", item["reset_df"], row)
    row = write_df_section(ws, "Influence / Cook's D", item["infl_df"], row)
    row = write_df_section(ws, "Robustness Snapshot", item["robustness_df"], row)

    sig_ext = item["extensions_df"][pd.to_numeric(item["extensions_df"]["p_value"], errors="coerce") < 0.10].copy()
    if not sig_ext.empty:
        sig_ext = sig_ext[["dependent", "term_label", "type", "N", "coef", "std_error", "p_value", "evidence"]]
    write_df_section(ws, "Nonlinear and Interaction Terms with p < 0.10", sig_ext, row)
    autosize(ws)


def format_machine_sheets(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for ws in wb.worksheets:
        if ws.title == "Summary" or ws.title.startswith("open_") or ws.title.startswith("close_"):
            continue
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        autosize(ws)
    wb.save(path)


def run_one(window: str, min_messages: int, dataset: Path, sample_mode: str) -> Dict[str, Any]:
    print(f"\n=== {window}, min_messages={min_messages}, sample={sample_mode} ===")
    print(f"Input: {dataset}")
    full_sub, baseline_sub, quality_df, sample_df, anomalies_df = diag.load_data(str(dataset))
    full_sub, baseline_sub = apply_sample_mode(full_sub, baseline_sub, sample_mode)
    vif_df = diag.test_vif(baseline_sub)
    vif_fe_df = diag.test_vif_with_fe(baseline_sub)
    corr_matrix, high_corr_df = diag.test_correlations(baseline_sub)
    het_df = diag.test_heteroskedasticity(baseline_sub)
    norm_df = diag.test_normality(baseline_sub)
    reset_df = diag.test_reset(baseline_sub)
    infl_df = diag.test_influence(baseline_sub)
    robustness_df = diag.build_robustness_report(full_sub, baseline_sub)
    extensions_df = diag.test_extension_terms(baseline_sub)

    return {
        "window": window,
        "min_messages": min_messages,
        "sample_mode": sample_mode,
        "dataset": dataset,
        "full_sub": full_sub,
        "baseline_sub": baseline_sub,
        "quality_df": quality_df,
        "sample_df": sample_df,
        "anomalies_df": anomalies_df,
        "vif_df": vif_df,
        "vif_fe_df": vif_fe_df,
        "corr_matrix": corr_matrix,
        "high_corr_df": high_corr_df,
        "het_df": het_df,
        "norm_df": norm_df,
        "reset_df": reset_df,
        "infl_df": infl_df,
        "robustness_df": robustness_df,
        "extensions_df": extensions_df,
    }


def build_workbook(runs: List[Dict[str, Any]], sample_summary: pd.DataFrame, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_df_section(ws, "Sample and Diagnostic Summary", sample_summary, 1)
    autosize(ws)

    for item in runs:
        short_window = "open" if item["window"] == "book_open" else "close"
        ws = wb.create_sheet(f"{short_window}_{item['min_messages']}")
        write_scenario_sheet(ws, item)

    wb.save(output_path)


def build_output(
    output_path: Path,
    sample_mode: str,
    history_spec: str = "baseline",
    dataset_template: str = DEFAULT_DATASET_TEMPLATE,
) -> None:
    configure_diagnostic_spec(sample_mode, history_spec=history_spec)
    final_output_path = output_for_sample(output_path, sample_mode)

    runs = []
    sample_rows = []
    missing_rows = []

    for window in WINDOWS:
        for min_messages in MIN_MESSAGES:
            dataset = find_dataset(window, min_messages, dataset_template=dataset_template)
            if dataset is None:
                missing_rows.append({
                    "window": window,
                    "min_messages": min_messages,
                    "sample_mode": sample_mode,
                    "status": "missing regression_ready.csv",
                })
                continue
            item = run_one(window, min_messages, dataset, sample_mode)
            runs.append(item)
            sample_rows.append(collect_sample_summary(
                window,
                min_messages,
                dataset,
                sample_mode,
                item["full_sub"],
                item["baseline_sub"],
                item["vif_df"],
                item["reset_df"],
                item["infl_df"],
            ))

    if not runs:
        raise FileNotFoundError("No regression_ready.csv files found for the default book_open/book_close min-message specs.")

    sample_summary = pd.DataFrame(sample_rows)
    build_workbook(runs, sample_summary, final_output_path)

    vif_all = pd.concat([add_meta(item["vif_df"], item["window"], item["min_messages"], item["dataset"], sample_mode) for item in runs], ignore_index=True)
    vif_fe_all = pd.concat([add_meta(item["vif_fe_df"], item["window"], item["min_messages"], item["dataset"], sample_mode) for item in runs], ignore_index=True)
    reset_all = pd.concat([add_meta(item["reset_df"], item["window"], item["min_messages"], item["dataset"], sample_mode) for item in runs], ignore_index=True)
    robust_all = pd.concat([add_meta(item["robustness_df"], item["window"], item["min_messages"], item["dataset"], sample_mode) for item in runs], ignore_index=True)
    high_corr_all = pd.concat([add_meta(item["high_corr_df"], item["window"], item["min_messages"], item["dataset"], sample_mode) for item in runs], ignore_index=True)
    extensions_all = pd.concat([add_meta(item["extensions_df"], item["window"], item["min_messages"], item["dataset"], sample_mode) for item in runs], ignore_index=True)
    extensions_p10 = extensions_all[pd.to_numeric(extensions_all["p_value"], errors="coerce") < 0.10].copy()
    missing_df = pd.DataFrame(missing_rows)

    with pd.ExcelWriter(final_output_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        vif_all.to_excel(writer, index=False, sheet_name="VIF All")
        vif_fe_all.to_excel(writer, index=False, sheet_name="VIF FE All")
        reset_all.to_excel(writer, index=False, sheet_name="RESET All")
        robust_all.to_excel(writer, index=False, sheet_name="Robustness All")
        high_corr_all.to_excel(writer, index=False, sheet_name="High Corr All")
        extensions_p10.to_excel(writer, index=False, sheet_name="Ext p10")
        extensions_all.to_excel(writer, index=False, sheet_name="Extensions All")
        if not missing_df.empty:
            missing_df.to_excel(writer, index=False, sheet_name="Missing Inputs")

    format_machine_sheets(final_output_path)
    print(f"\nSaved: {final_output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Batch upsize diagnostics across sentiment windows and min-message thresholds.")
    parser.add_argument("output", nargs="?", default=DEFAULT_OUTPUT_FILE)
    parser.add_argument("--sample", choices=SAMPLE_MODES, default="all", help="Sample restriction for all scenarios.")
    parser.add_argument("--all-samples", action="store_true", help="Create separate all/fixed/floater workbooks.")
    parser.add_argument(
        "--history-spec",
        choices=["share", "dummy", "baseline"],
        default="share",
        help="Issuer-history controls: default share uses hist_reduction_share; dummy restores repeat-prior-tightening dummy.",
    )
    parser.add_argument(
        "--dataset-template",
        default=DEFAULT_DATASET_TEMPLATE,
        help=(
            "Template for regression_ready.csv paths. Available fields: {window}, {min_messages}. "
            "Default: regression_dataset_{window}_{min_messages}/regression_ready.csv"
        ),
    )
    args = parser.parse_args()
    output_path = Path(args.output)

    sample_modes = SAMPLE_MODES if args.all_samples else [args.sample]
    for sample_mode in sample_modes:
        build_output(
            output_path,
            sample_mode,
            history_spec=args.history_spec,
            dataset_template=args.dataset_template,
        )


if __name__ == "__main__":
    main()
