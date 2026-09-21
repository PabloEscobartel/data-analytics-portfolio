#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Regression analysis for primary bond placements.

Main outcomes:
1) coupon_reduction_bp              — coupon tightening in bp
2) ihs_coupon_reduction             — inverse hyperbolic sine of coupon tightening
3) coupon_reduced_dummy             — 1 if coupon was tightened
4) ihs_coupon_positive              — IHS coupon tightening conditional on tightening > 0
5) log_placement_vol_book           — log(final placed volume / pre-book teaser volume)
6) upsize_dummy                     — 1 if final placed volume exceeds teaser volume by >1%
7) joint_success_upsize             — strict success: coupon tightened and volume upsized
8) success_type                     — multinomial outcome: neither / price_only / volume_only / dual_success

Usage:
    python run_regressions_upsize.py
    python run_regressions_upsize.py regression_ready.csv
    python run_regressions_upsize.py regression_ready.csv results.xlsx

Requirements:
    pip install pandas numpy statsmodels scipy openpyxl
"""

from __future__ import annotations

import os
import sys
import warnings
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import statsmodels.api as sm
from statsmodels.tools.sm_exceptions import ConvergenceWarning
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

warnings.filterwarnings("default")

# ============================================================
# SETTINGS
# ============================================================

DEFAULT_INPUT_CANDIDATES = [
    # "regression_dataset_book_close/regression_ready.csv",
    "regression_dataset_book_open/regression_ready.csv",
]
DEFAULT_OUTPUT_FILE = "regression_results_upsize.xlsx"
YEAR_FROM = 2018

BASELINE_EXCLUDE_ISSUERS = {"ВЭБ.РФ"}
BASELINE_DROP_ZERO_ORGANIZERS = True
BASELINE_DROP_BOOK_AFTER_PLACEMENT = True

USE_YEAR_FE = True
USE_QUARTER_FE = False
USE_RATING_BUCKET_FE = False
USE_INDUSTRY_FE = False
ESTIMATE_MNLOGIT = True
INDUSTRY_FE_MIN_N = 20

WINSOR_LOWER = 0.01
WINSOR_UPPER = 0.99
COOK_DROP_Q = 0.99

# For upsize: avoid treating tiny rounding deviations as true increases.
UPSIZE_THRESHOLD = 1.01
FULL_VOLUME_THRESHOLD = 0.99

RATING_MAP = {
    "AAA": 15, "AA+": 14, "AA": 13, "AA-": 12,
    "A+": 11, "A": 10, "A-": 9,
    "BBB+": 8, "BBB": 7, "BBB-": 6,
    "BB+": 5, "BB": 4, "BB-": 3,
    "B+": 2, "B": 1, "B-": 0,
    "C": -1,
}

VAR_LABELS = {
    "const": "Constant",
    "coupon_reduced_dummy": "Coupon reduction dummy",
    "coupon_success": "Coupon success",
    "volume_success": "Volume success",
    "success_score": "Success score: coupon + volume",
    "any_success": "Any success: coupon or volume",
    "dual_success": "Dual success: coupon and volume",
    "joint_success_upsize": "Strict dual success",
    "joint_success_full": "Coupon tightening and full volume",
    "success_type_code": "Success type",
    "success_type": "Success type",
    "ihs_coupon_reduction": "IHS coupon reduction",
    "ihs_coupon_positive": "IHS coupon reduction | coupon > 0",
    "num_organizers": "Number of organizers",
    "log_num_organizers": "ln(1 + organizers)",
    "issuer_history_group": "Issuer history group",
    "history_debut": "Debut issuer",
    "history_repeat_prior_tightening": "Repeat issuer, prior tightening",
    "hist_reduction_share": "Hist. share of prior tightenings",
    "log_full_issue_number": "ln(Full issue number)",
    "hist_ever_reduced": "Prior tightening (dummy)",
    "hist_avg_volume_ratio": "Hist. avg volume ratio",
    "rating_num": "Credit rating (ordinal)",
    "rating_bucket": "Rating bucket",
    "industry_fe": "Industry FE",
    "log_dur": "ln(Term to exit, days)",
    "has_put": "Has put option (dummy)",
    "is_floater": "Floating rate (dummy)",
    "si": "Sentiment Index (SI)",
    "si_positive": "Positive SI",
    "si_negative": "Negative SI intensity",
    "si_pos": "Positive SI",
    "si_neg": "Negative SI intensity",
    "si_sq": "SI squared",
    "log_buzz": "ln(Buzz)",
    "log_buzz_sq": "ln(Buzz) squared",
    "share_pos": "Share positive",
    "share_neg": "Share negative",
    "ofz_yield": "OFZ matched yield, %",
    "ofz_yield_within_year": "OFZ yield, within-year deviation",
    "rvi": "RVI (volatility index)",
    "rating_num_sq": "Credit rating squared",
    "si_x_log_buzz": "SI x ln(Buzz)",
    "si_x_is_floater": "SI x Floater",
    "log_buzz_x_is_floater": "ln(Buzz) x Floater",
    "si_x_rating_num": "SI x Rating",
    "log_buzz_x_rating_num": "ln(Buzz) x Rating",
    "num_organizers_x_si": "Organizers x SI",
}

SUCCESS_TYPE_CODES = {
    "neither": 0,
    "price_only": 1,
    "volume_only": 2,
    "dual_success": 3,
}
SUCCESS_TYPE_LABELS = {
    0: "Neither",
    1: "Price only",
    2: "Volume only",
    3: "Dual success",
}

BASE_VARS = [
    "log_num_organizers",
    "history_debut",
    "hist_reduction_share",
    "log_full_issue_number",
    "hist_avg_volume_ratio",
    "rating_num",
    "ofz_yield_within_year",
    "rvi",
    "log_dur",
    "has_put",
    "is_floater",
]
MAIN_INTERACTION_VARS: List[str] = []
SENT_VARS_SI = BASE_VARS + ["si", "log_buzz"]
SENT_VARS_SHARES = BASE_VARS + ["share_pos", "share_neg", "log_buzz"]
SENT_VARS_SPLIT = BASE_VARS + ["si_positive", "si_negative", "log_buzz"]
NONLINEAR_VARS = ["si_sq", "log_buzz_sq", "rating_num_sq"]
INTERACTION_VARS = [
    "si_x_log_buzz",
    "si_x_is_floater",
    "log_buzz_x_is_floater",
    "si_x_rating_num",
    "log_buzz_x_rating_num",
    "num_organizers_x_si",
]
EXTENSION_TERMS = [
    ("si_sq", "nonlinear", "negative", "Strong sentiment has diminishing/asymmetric effect"),
    ("log_buzz_sq", "nonlinear", "negative", "Media coverage has diminishing returns"),
    ("rating_num_sq", "nonlinear", "any", "Rating effect is nonlinear"),
    ("si_x_log_buzz", "interaction", "positive", "Sentiment matters more when discussion is broader"),
    ("si_x_is_floater", "interaction", "any", "Sentiment differs for floaters"),
    ("log_buzz_x_is_floater", "interaction", "any", "Media coverage effect differs for floaters"),
    ("si_x_rating_num", "interaction", "negative", "Sentiment matters more for lower-rated issuers"),
    ("log_buzz_x_rating_num", "interaction", "negative", "Coverage matters more for lower-rated issuers"),
    ("num_organizers_x_si", "interaction", "any", "Organizer syndicate changes the sentiment effect"),
]


def make_sentiment_var_sets(base_vars: List[str]) -> Tuple[List[str], List[str], List[str]]:
    return (
        base_vars + ["si", "log_buzz"],
        base_vars + ["share_pos", "share_neg", "log_buzz"],
        base_vars + ["si_positive", "si_negative", "log_buzz"],
    )


# ============================================================
# HELPERS
# ============================================================

@dataclass
class SampleBundle:
    full: pd.DataFrame
    baseline: pd.DataFrame


def pick_input_file() -> str:
    if len(sys.argv) >= 2:
        return sys.argv[1]
    for p in DEFAULT_INPUT_CANDIDATES:
        if os.path.exists(p):
            return p
    raise FileNotFoundError(
        "No input file found. Pass path explicitly, e.g. python run_regressions_upsize.py regression_ready.csv"
    )


def pick_output_file() -> str:
    if len(sys.argv) >= 3:
        return sys.argv[2]
    return DEFAULT_OUTPUT_FILE


def first_existing_col(df: pd.DataFrame, candidates: List[str]) -> str:
    for col in candidates:
        if col in df.columns:
            return col
    raise KeyError(f"Не найдена ни одна из колонок: {', '.join(candidates)}")


def optional_existing_col(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    for col in candidates:
        if col in df.columns:
            return col
    return None


def normalize_rating(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace("–", "-", regex=False)
        .str.replace("—", "-", regex=False)
        .str.strip()
    )


def rating_bucket_from_num(value: float) -> str:
    if pd.isna(value):
        return "Missing rating"
    if value >= 12:
        return "High: AAA to AA-"
    if value >= 6:
        return "Middle: A+ to BBB-"
    return "HY/ВДО: BB+ and below"


def add_industry_fe(full_sub: pd.DataFrame, baseline_sub: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame]:
    if "industry" not in full_sub.columns:
        full_sub = full_sub.copy()
        baseline_sub = baseline_sub.copy()
        full_sub["industry_fe"] = "Missing industry"
        baseline_sub["industry_fe"] = "Missing industry"
        return full_sub, baseline_sub

    counts = baseline_sub["industry"].fillna("Missing industry").astype(str).value_counts()
    keep = set(counts[counts >= INDUSTRY_FE_MIN_N].index)

    def collapse(series: pd.Series) -> pd.Series:
        labels = series.fillna("Missing industry").astype(str)
        return labels.where(labels.isin(keep), "Other / rare industries")

    full_sub = full_sub.copy()
    baseline_sub = baseline_sub.copy()
    full_sub["industry_fe"] = collapse(full_sub["industry"])
    baseline_sub["industry_fe"] = collapse(baseline_sub["industry"])
    return full_sub, baseline_sub


def safe_numeric(df: pd.DataFrame, col: Optional[str], default: float = np.nan) -> pd.Series:
    if col is None:
        return pd.Series(default, index=df.index, dtype="float64")
    return pd.to_numeric(df[col], errors="coerce")


def winsorize_series(s: pd.Series, lower: float = 0.01, upper: float = 0.99) -> pd.Series:
    s2 = s.dropna()
    if s2.empty:
        return s.copy()
    lo, hi = s2.quantile(lower), s2.quantile(upper)
    return s.clip(lower=lo, upper=hi)


def add_extension_terms(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    sub["si_sq"] = sub["si"] ** 2
    sub["log_buzz_sq"] = sub["log_buzz"] ** 2
    sub["rating_num_sq"] = sub["rating_num"] ** 2
    sub["si_x_log_buzz"] = sub["si"] * sub["log_buzz"]
    sub["si_x_is_floater"] = sub["si"] * sub["is_floater"]
    sub["log_buzz_x_is_floater"] = sub["log_buzz"] * sub["is_floater"]
    sub["si_x_rating_num"] = sub["si"] * sub["rating_num"]
    sub["log_buzz_x_rating_num"] = sub["log_buzz"] * sub["rating_num"]
    sub["num_organizers_x_si"] = sub["num_organizers"] * sub["si"]
    return sub


def add_issuer_history_controls(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    valid = sub["full_issue_number"].notna() & sub["hist_ever_reduced"].notna()
    debut = valid & sub["full_issue_number"].eq(1)
    repeat = valid & sub["full_issue_number"].gt(1)
    prior = sub["hist_ever_reduced"].eq(1)
    no_prior = sub["hist_ever_reduced"].eq(0)

    sub["issuer_history_group"] = np.select(
        [debut, repeat & prior, repeat & no_prior],
        ["debut", "repeat_prior_tightening", "repeat_no_prior_tightening"],
        default=None,
    )
    sub["history_debut"] = pd.Series(np.where(valid, debut.astype(int), np.nan), index=sub.index).astype("Int64")
    sub["history_repeat_prior_tightening"] = pd.Series(
        np.where(valid, (repeat & prior).astype(int), np.nan),
        index=sub.index,
    ).astype("Int64")
    return sub


def evidence_label(coef: float, p_value: float, expected_sign: str) -> str:
    if pd.isna(coef) or pd.isna(p_value):
        return "not estimated"
    if p_value >= 0.10:
        return "no evidence"
    if expected_sign == "negative":
        return "supports" if coef < 0 else "opposite sign"
    if expected_sign == "positive":
        return "supports" if coef > 0 else "opposite sign"
    return "significant"


def sig_stars(p: float) -> str:
    if p is None or pd.isna(p):
        return ""
    if p < 0.01:
        return "***"
    if p < 0.05:
        return "**"
    if p < 0.10:
        return "*"
    return ""


def add_quality_flags(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    sub["placement_dt"] = pd.to_datetime(sub["placement_date"], errors="coerce")
    sub["book_dt"] = pd.to_datetime(sub["book_date"], errors="coerce")
    sub["book_after_placement"] = (
        sub["book_dt"].notna()
        & sub["placement_dt"].notna()
        & (sub["book_dt"].dt.date > sub["placement_dt"].dt.date)
    )
    sub["is_special_issuer"] = sub["issuer"].astype(str).isin(BASELINE_EXCLUDE_ISSUERS)
    sub["zero_organizers"] = sub["num_organizers"].eq(0)
    if USE_QUARTER_FE:
        sub["time_fe"] = sub["placement_dt"].dt.to_period("Q").astype(str)
    else:
        sub["time_fe"] = sub["year"].astype("Int64").astype(str)
    return sub


def add_ofz_within_year(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    sub["ofz_yield_within_year"] = (
        sub["ofz_yield"]
        - sub.groupby("year")["ofz_yield"].transform("mean")
    )
    return sub


def apply_baseline_filters(sub: pd.DataFrame) -> pd.DataFrame:
    mask = pd.Series(True, index=sub.index)
    if BASELINE_DROP_BOOK_AFTER_PLACEMENT:
        mask &= ~sub["book_after_placement"]
    if BASELINE_EXCLUDE_ISSUERS:
        mask &= ~sub["is_special_issuer"]
    if BASELINE_DROP_ZERO_ORGANIZERS:
        mask &= ~sub["zero_organizers"]
    return sub.loc[mask].copy()


def build_X(data: pd.DataFrame, variables: List[str], add_fe: bool = True) -> pd.DataFrame:
    X = data[variables].copy()
    if add_fe:
        fe_parts = []
        if USE_QUARTER_FE:
            fe_parts.append(pd.get_dummies(data["time_fe"], prefix="q", drop_first=True, dtype=float))
        elif USE_YEAR_FE:
            fe_parts.append(pd.get_dummies(data["time_fe"], prefix="year", drop_first=True, dtype=float))
        if USE_RATING_BUCKET_FE and "rating_bucket" in data.columns:
            fe_parts.append(pd.get_dummies(data["rating_bucket"], prefix="rating_bucket", drop_first=True, dtype=float))
        if USE_INDUSTRY_FE and "industry_fe" in data.columns:
            fe_parts.append(pd.get_dummies(data["industry_fe"], prefix="industry", drop_first=True, dtype=float))
        fe_parts = [fe for fe in fe_parts if not fe.empty]
        if fe_parts:
            X = pd.concat([X] + fe_parts, axis=1)
    X = X.apply(pd.to_numeric, errors="coerce")
    constant_cols = [c for c in X.columns if X[c].nunique(dropna=True) <= 1]
    if constant_cols:
        X = X.drop(columns=constant_cols)
    X = sm.add_constant(X, has_constant="add")
    return X.astype(float)


def subset_for_model(data: pd.DataFrame, y_col: str, variables: List[str]) -> pd.DataFrame:
    needed = [y_col] + variables + ["issuer_id", "time_fe"]
    needed = [c for c in needed if c in data.columns]
    return data.dropna(subset=needed).copy()


def fit_ols(y: pd.Series, X: pd.DataFrame, cov_type: str = "HC3", groups: Optional[pd.Series] = None):
    model = sm.OLS(y, X, missing="drop")
    if cov_type == "cluster":
        if groups is None or pd.Series(groups).nunique() < 2:
            return model.fit(cov_type="HC3")
        return model.fit(cov_type="cluster", cov_kwds={"groups": groups})
    return model.fit(cov_type=cov_type)


def fit_glm_binomial(y: pd.Series, X: pd.DataFrame, groups: Optional[pd.Series] = None):
    model = sm.GLM(y, X, family=sm.families.Binomial(), missing="drop")
    if groups is not None and pd.Series(groups).nunique() >= 2:
        try:
            return model.fit(cov_type="cluster", cov_kwds={"groups": groups}, maxiter=200, disp=0)
        except (np.linalg.LinAlgError, ValueError, sm.tools.sm_exceptions.PerfectSeparationError) as e:
            warnings.warn(f"GLM clustered SE failed ({e}); falling back to HC1.", RuntimeWarning)
    try:
        return model.fit(cov_type="HC1", maxiter=200, disp=0)
    except Exception as e:
        warnings.warn(f"GLM HC1 failed ({e}); returning conventional covariance.", RuntimeWarning)
        return model.fit(maxiter=200, disp=0)


def fit_mnlogit(y: pd.Series, X: pd.DataFrame, groups: Optional[pd.Series] = None):
    y = pd.to_numeric(y, errors="coerce").astype("Int64")
    valid = y.notna()
    y = y.loc[valid].astype(int)
    X = X.loc[valid]
    groups = pd.Series(groups, index=X.index).loc[valid] if groups is not None else None
    if y.nunique() < 2 or 0 not in set(y.unique()):
        warnings.warn("MNLogit skipped: need at least two categories and 'neither' as the baseline.", RuntimeWarning)
        return None
    model = sm.MNLogit(y, X, missing="drop")

    def valid_result(res: Any) -> bool:
        if res is None or not hasattr(res, "params"):
            return False
        vals = np.asarray(res.params, dtype=float)
        return np.isfinite(vals).all()

    cov_specs = []
    if groups is not None and groups.nunique() >= 2:
        cov_specs.append(("cluster", {"groups": groups}))
    cov_specs.append(("HC1", {}))

    for cov_type, cov_kwds in cov_specs:
        for method in ["newton", "bfgs", "powell"]:
            try:
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore", ConvergenceWarning)
                    warnings.simplefilter("ignore", RuntimeWarning)
                    res = model.fit(
                        method=method,
                        maxiter=500,
                        disp=False,
                        cov_type=cov_type,
                        cov_kwds=cov_kwds,
                    )
                if valid_result(res):
                    return res
            except (np.linalg.LinAlgError, ValueError):
                continue
    warnings.warn("MNLogit failed for all optimizers/covariance choices.", RuntimeWarning)
    return None


def get_csp(res: Any, var: str, colnames: List[str]) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    if hasattr(res, "params") and var in res.params.index:
        return float(res.params[var]), float(res.bse[var]), float(res.pvalues[var])
    return None, None, None


def mnlogit_col_for_category(res: Any, category_code: int) -> Optional[int]:
    ynames = getattr(getattr(res, "model", None), "_ynames_map", None)
    if not ynames:
        return category_code - 1 if category_code > 0 else None
    for internal_code, label in ynames.items():
        if str(label) == str(category_code):
            if internal_code == 0:
                return None
            return internal_code - 1
    return None


def get_mnlogit_csp(res: Any, var: str, category_code: int) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    if res is None or not hasattr(res, "params") or category_code == 0:
        return None, None, None
    col = mnlogit_col_for_category(res, category_code)
    if col is None:
        return None, None, None
    if var not in res.params.index or col not in res.params.columns:
        return None, None, None
    return float(res.params.loc[var, col]), float(res.bse.loc[var, col]), float(res.pvalues.loc[var, col])


def collect_mnlogit_rows(res: Any, variables: List[str], nobs: int) -> pd.DataFrame:
    rows = []
    for code in [1, 2, 3]:
        for var in variables:
            coef, se, p_value = get_mnlogit_csp(res, var, code)
            rows.append({
                "outcome": SUCCESS_TYPE_LABELS[code],
                "baseline": SUCCESS_TYPE_LABELS[0],
                "variable": var,
                "label": VAR_LABELS.get(var, var),
                "N": nobs,
                "coef": round(coef, 6) if coef is not None and pd.notna(coef) else np.nan,
                "std_error": round(se, 6) if se is not None and pd.notna(se) else np.nan,
                "p_value": round(p_value, 6) if p_value is not None and pd.notna(p_value) else np.nan,
                "stars": sig_stars(p_value) if p_value is not None else "",
            })
    return pd.DataFrame(rows)


def model_stat_value(res: Any, stat_name: str) -> str:
    if res is None:
        return "—"
    if stat_name == "N":
        return str(int(res.nobs)) if hasattr(res, "nobs") else "—"
    if stat_name == "R²":
        return f"{res.rsquared:.4f}" if hasattr(res, "rsquared") else "—"
    if stat_name == "Adj. R²":
        return f"{res.rsquared_adj:.4f}" if hasattr(res, "rsquared_adj") else "—"
    if stat_name == "Pseudo R²":
        try:
            return f"{res.pseudo_rsquared(kind='cs'):.4f}"
        except Exception:
            return "—"
    if stat_name == "Log-lik":
        return f"{res.llf:.2f}" if hasattr(res, "llf") else "—"
    if stat_name == "AIC":
        return f"{res.aic:.2f}" if hasattr(res, "aic") else "—"
    return "—"


def average_marginal_effect(res: Any, X: pd.DataFrame, var: str) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    if var not in X.columns or not hasattr(res, "params"):
        return None, None, None
    params = res.params.reindex(X.columns).fillna(0)
    beta = float(params[var])
    p = pd.Series(res.predict(X), index=X.index).clip(1e-9, 1 - 1e-9)
    w = p * (1 - p)
    ame = beta * float(w.mean())

    cov = res.cov_params()
    if not isinstance(cov, pd.DataFrame):
        cov = pd.DataFrame(cov, index=X.columns, columns=X.columns)
    cov = cov.reindex(index=X.columns, columns=X.columns).fillna(0)

    grad = pd.Series(0.0, index=X.columns)
    grad[var] = float(w.mean())
    x_weighted = X.mul((1 - 2 * p) * w, axis=0).mean(axis=0)
    grad = grad.add(beta * x_weighted, fill_value=0)
    var_ame = float(grad.to_numpy() @ cov.to_numpy() @ grad.to_numpy())
    se = float(np.sqrt(max(var_ame, 0)))
    z = ame / se if se > 0 else np.nan
    p_value = float(2 * (1 - stats.norm.cdf(abs(z)))) if pd.notna(z) else np.nan
    return ame, se, p_value


# ============================================================
# DATA PREP
# ============================================================

def load_and_prepare(path: str, year_from: int) -> SampleBundle:
    df = pd.read_csv(path, low_memory=False)
    df["placement_dt"] = pd.to_datetime(df["placement_date"], errors="coerce")
    df["year"] = df["placement_dt"].dt.year
    sub = df[df["year"] >= year_from].copy()

    # Sentiment.
    si_col = first_existing_col(sub, ["si_relevant", "si"])
    buzz_col = first_existing_col(sub, ["log_buzz_relevant", "log_buzz_all", "log_buzz"])
    share_pos_col = optional_existing_col(sub, ["share_positive_relevant", "share_positive"])
    share_neg_col = optional_existing_col(sub, ["share_negative_relevant", "share_negative"])

    si_raw = safe_numeric(sub, si_col)
    log_buzz_raw = safe_numeric(sub, buzz_col)
    share_pos_raw = safe_numeric(sub, share_pos_col)
    share_neg_raw = safe_numeric(sub, share_neg_col)

    sub["si"] = si_raw.fillna(0)
    sub["si_positive"] = sub["si"].clip(lower=0)
    sub["si_negative"] = (-sub["si"]).clip(lower=0)
    sub["si_pos"] = sub["si_positive"]
    sub["si_neg"] = sub["si_negative"]
    sub["log_buzz"] = log_buzz_raw.fillna(0)
    sub["share_pos"] = share_pos_raw.fillna(0)
    sub["share_neg"] = share_neg_raw.fillna(0)

    # Outcomes.
    sub["coupon_reduction_bp"] = pd.to_numeric(sub["coupon_reduction_bp"], errors="coerce")
    sub["coupon_reduced_dummy"] = np.nan
    m_coupon = sub["coupon_reduction_bp"].notna()
    sub.loc[m_coupon, "coupon_reduced_dummy"] = (sub.loc[m_coupon, "coupon_reduction_bp"] > 0).astype(int)
    sub["coupon_success"] = sub["coupon_reduced_dummy"]
    sub["ihs_coupon_reduction"] = np.arcsinh(sub["coupon_reduction_bp"])
    sub["ihs_coupon_positive"] = np.where(
        sub["coupon_reduction_bp"] > 0,
        np.arcsinh(sub["coupon_reduction_bp"]),
        np.nan,
    )
    sub["placement_vol_book"] = pd.to_numeric(sub["placement_vol_book"], errors="coerce")

    # Guard against impossible non-positive ratios.
    sub.loc[sub["placement_vol_book"] <= 0, "placement_vol_book"] = np.nan
    sub["log_placement_vol_book"] = np.log(sub["placement_vol_book"])

    sub["upsize_dummy"] = np.nan
    m_vol = sub["placement_vol_book"].notna()
    sub.loc[m_vol, "upsize_dummy"] = (sub.loc[m_vol, "placement_vol_book"] > UPSIZE_THRESHOLD).astype(int)
    sub["volume_success"] = sub["upsize_dummy"]

    sub["full_or_more_dummy"] = np.nan
    sub.loc[m_vol, "full_or_more_dummy"] = (sub.loc[m_vol, "placement_vol_book"] >= FULL_VOLUME_THRESHOLD).astype(int)

    m_joint = sub["coupon_reduction_bp"].notna() & sub["placement_vol_book"].notna()
    sub["joint_success_upsize"] = np.nan
    sub.loc[m_joint, "joint_success_upsize"] = (
        (sub.loc[m_joint, "coupon_reduction_bp"] > 0)
        & (sub.loc[m_joint, "placement_vol_book"] > UPSIZE_THRESHOLD)
    ).astype(int)
    sub["dual_success"] = sub["joint_success_upsize"]

    sub["joint_success_full"] = np.nan
    sub.loc[m_joint, "joint_success_full"] = (
        (sub.loc[m_joint, "coupon_reduction_bp"] > 0)
        & (sub.loc[m_joint, "placement_vol_book"] >= FULL_VOLUME_THRESHOLD)
    ).astype(int)
    sub["any_success"] = np.nan
    sub.loc[m_joint, "any_success"] = (
        sub.loc[m_joint, "coupon_success"].eq(1)
        | sub.loc[m_joint, "volume_success"].eq(1)
    ).astype(int)
    sub["success_score"] = np.nan
    sub.loc[m_joint, "success_score"] = (
        sub.loc[m_joint, "coupon_success"].astype(float)
        + sub.loc[m_joint, "volume_success"].astype(float)
    )
    sub["success_type"] = pd.Series(np.nan, index=sub.index, dtype="object")
    sub.loc[m_joint, "success_type"] = "neither"
    sub.loc[m_joint & sub["coupon_success"].eq(1) & sub["volume_success"].eq(0), "success_type"] = "price_only"
    sub.loc[m_joint & sub["coupon_success"].eq(0) & sub["volume_success"].eq(1), "success_type"] = "volume_only"
    sub.loc[m_joint & sub["coupon_success"].eq(1) & sub["volume_success"].eq(1), "success_type"] = "dual_success"
    sub["success_type_code"] = sub["success_type"].map(SUCCESS_TYPE_CODES)

    # Controls.
    sub["num_organizers"] = pd.to_numeric(sub["num_organizers"], errors="coerce")
    sub["log_num_organizers"] = np.log1p(sub["num_organizers"])
    if "full_issue_number" not in sub.columns:
        raise KeyError("Не найдена обязательная колонка: full_issue_number")
    sub["full_issue_number"] = pd.to_numeric(sub["full_issue_number"], errors="coerce")
    sub.loc[sub["full_issue_number"] <= 0, "full_issue_number"] = np.nan
    sub["log_full_issue_number"] = np.log(sub["full_issue_number"])
    for col in ["hist_ever_reduced", "has_put", "hist_avg_volume_ratio"]:
        if col not in sub.columns:
            raise KeyError(f"Не найдена обязательная колонка: {col}")
        sub[col] = pd.to_numeric(sub[col], errors="coerce")
    if "hist_reduction_share" in sub.columns:
        sub["hist_reduction_share"] = pd.to_numeric(sub["hist_reduction_share"], errors="coerce")
    sub = add_issuer_history_controls(sub)
    if "log_term" in sub.columns:
        sub["log_dur"] = pd.to_numeric(sub["log_term"], errors="coerce")
    elif "term" in sub.columns:
        term = pd.to_numeric(sub["term"], errors="coerce")
        sub["log_dur"] = np.log(term.clip(lower=1))
    else:
        raise KeyError("Нужна колонка log_term или term")

    sub["final_rating"] = normalize_rating(sub["final_rating"])
    sub["rating_num"] = sub["final_rating"].map(RATING_MAP)
    sub["rating_bucket"] = sub["rating_num"].map(rating_bucket_from_num)
    sub["ofz_yield"] = pd.to_numeric(sub["ofz_matched_yield"], errors="coerce")
    sub = add_ofz_within_year(sub)
    sub["rvi"] = pd.to_numeric(sub["rvi_close"], errors="coerce")
    sub["is_floater"] = pd.to_numeric(sub.get("is_floater", 0), errors="coerce")
    sub = add_extension_terms(sub)

    sub = add_quality_flags(sub)
    sub["issuer_id"] = sub["issuer"].astype(str).astype("category").cat.codes

    common_drop_vars = list(dict.fromkeys(["issuer_id"] + SENT_VARS_SI + ["share_pos", "share_neg"]))
    full_sub = sub.dropna(subset=common_drop_vars).copy()
    baseline_sub = apply_baseline_filters(full_sub)
    full_sub = add_ofz_within_year(full_sub)
    baseline_sub = add_ofz_within_year(baseline_sub)
    full_sub, baseline_sub = add_industry_fe(full_sub, baseline_sub)
    return SampleBundle(full=full_sub, baseline=baseline_sub)


# ============================================================
# ESTIMATION
# ============================================================

def estimate_models(sub: pd.DataFrame, sample_name: str = "Baseline", base_vars: Optional[List[str]] = None) -> Dict[str, Any]:
    R: Dict[str, Any] = {"_sample_name": sample_name}
    ame_rows = []
    base_vars = BASE_VARS if base_vars is None else base_vars
    sent_vars_si, sent_vars_shares, sent_vars_split = make_sentiment_var_sets(base_vars)

    def add_ame(model_name: str, res: Any, X: pd.DataFrame, variables: List[str]) -> None:
        for var in variables:
            ame, se, p_value = average_marginal_effect(res, X, var)
            if ame is None:
                continue
            ame_rows.append({
                "sample": sample_name,
                "model": model_name,
                "variable": var,
                "label": VAR_LABELS.get(var, var),
                "AME": round(ame, 6),
                "std_error": round(se, 6) if se is not None and pd.notna(se) else np.nan,
                "p_value": round(p_value, 6) if p_value is not None and pd.notna(p_value) else np.nan,
            })

    # Coupon: raw and winsorized.
    d_coupon = subset_for_model(sub, "coupon_reduction_bp", sent_vars_si)
    X_coupon = build_X(d_coupon, sent_vars_si, add_fe=True)
    y_coupon = d_coupon["coupon_reduction_bp"].astype(float)
    R["coupon_ols_hc3"] = fit_ols(y_coupon, X_coupon, cov_type="HC3")
    R["coupon_ols_cluster"] = fit_ols(y_coupon, X_coupon, cov_type="cluster", groups=d_coupon["issuer_id"])
    R["coupon_ihs_cluster"] = fit_ols(
        d_coupon["ihs_coupon_reduction"].astype(float), X_coupon, cov_type="cluster", groups=d_coupon["issuer_id"]
    )

    d_coupon_bin = subset_for_model(sub, "coupon_reduced_dummy", sent_vars_si)
    X_coupon_bin = build_X(d_coupon_bin, sent_vars_si, add_fe=True)
    y_coupon_bin = d_coupon_bin["coupon_reduced_dummy"].astype(float)
    R["coupon_reduced_logit_cluster"] = fit_glm_binomial(y_coupon_bin, X_coupon_bin, groups=d_coupon_bin["issuer_id"])
    R["coupon_reduced_lpm_cluster"] = fit_ols(
        y_coupon_bin, X_coupon_bin, cov_type="cluster", groups=d_coupon_bin["issuer_id"]
    )
    add_ame("coupon_reduced_logit_cluster", R["coupon_reduced_logit_cluster"], X_coupon_bin, ["si", "log_buzz"])

    d_coupon_pos = subset_for_model(sub.loc[sub["coupon_reduction_bp"] > 0], "ihs_coupon_positive", sent_vars_si)
    X_coupon_pos = build_X(d_coupon_pos, sent_vars_si, add_fe=True)
    y_coupon_pos = d_coupon_pos["ihs_coupon_positive"].astype(float)
    R["coupon_positive_ihs_cluster"] = fit_ols(
        y_coupon_pos, X_coupon_pos, cov_type="cluster", groups=d_coupon_pos["issuer_id"]
    )

    d_coupon_bin_split = subset_for_model(sub, "coupon_reduced_dummy", sent_vars_split)
    X_coupon_bin_split = build_X(d_coupon_bin_split, sent_vars_split, add_fe=True)
    R["coupon_reduced_split_logit"] = fit_glm_binomial(
        d_coupon_bin_split["coupon_reduced_dummy"].astype(float), X_coupon_bin_split, groups=d_coupon_bin_split["issuer_id"]
    )
    add_ame(
        "coupon_reduced_split_logit",
        R["coupon_reduced_split_logit"],
        X_coupon_bin_split,
        ["si_positive", "si_negative", "log_buzz"],
    )

    d_coupon_pos_split = subset_for_model(sub.loc[sub["coupon_reduction_bp"] > 0], "ihs_coupon_positive", sent_vars_split)
    X_coupon_pos_split = build_X(d_coupon_pos_split, sent_vars_split, add_fe=True)
    R["coupon_positive_split_ihs"] = fit_ols(
        d_coupon_pos_split["ihs_coupon_positive"].astype(float),
        X_coupon_pos_split,
        cov_type="cluster",
        groups=d_coupon_pos_split["issuer_id"],
    )

    d_coupon["coupon_reduction_bp_w"] = winsorize_series(d_coupon["coupon_reduction_bp"], WINSOR_LOWER, WINSOR_UPPER)
    R["coupon_winsor_cluster"] = fit_ols(
        d_coupon["coupon_reduction_bp_w"].astype(float), X_coupon, cov_type="cluster", groups=d_coupon["issuer_id"]
    )

    X_coupon_sh = build_X(d_coupon, sent_vars_shares, add_fe=True)
    R["coupon_shares_cluster"] = fit_ols(y_coupon, X_coupon_sh, cov_type="cluster", groups=d_coupon["issuer_id"])

    # Volume: log ratio and raw ratio robustness.
    d_vol = subset_for_model(sub, "log_placement_vol_book", sent_vars_si)
    X_vol = build_X(d_vol, sent_vars_si, add_fe=True)
    y_logvol = d_vol["log_placement_vol_book"].astype(float)
    R["logvol_ols_hc3"] = fit_ols(y_logvol, X_vol, cov_type="HC3")
    R["logvol_ols_cluster"] = fit_ols(y_logvol, X_vol, cov_type="cluster", groups=d_vol["issuer_id"])

    d_vol["placement_vol_book_w"] = winsorize_series(d_vol["placement_vol_book"], WINSOR_LOWER, WINSOR_UPPER)
    R["volratio_winsor_cluster"] = fit_ols(
        d_vol["placement_vol_book_w"].astype(float), X_vol, cov_type="cluster", groups=d_vol["issuer_id"]
    )

    X_vol_sh = build_X(d_vol, sent_vars_shares, add_fe=True)
    R["logvol_shares_cluster"] = fit_ols(y_logvol, X_vol_sh, cov_type="cluster", groups=d_vol["issuer_id"])

    d_vol_split = subset_for_model(sub, "log_placement_vol_book", sent_vars_split)
    X_vol_split = build_X(d_vol_split, sent_vars_split, add_fe=True)
    R["logvol_split_cluster"] = fit_ols(
        d_vol_split["log_placement_vol_book"].astype(float),
        X_vol_split,
        cov_type="cluster",
        groups=d_vol_split["issuer_id"],
    )

    # Binary outcomes.
    d_up = subset_for_model(sub, "upsize_dummy", sent_vars_si)
    X_up = build_X(d_up, sent_vars_si, add_fe=True)
    y_up = d_up["upsize_dummy"].astype(float)
    R["upsize_logit_cluster"] = fit_glm_binomial(y_up, X_up, groups=d_up["issuer_id"])
    R["upsize_lpm_cluster"] = fit_ols(y_up, X_up, cov_type="cluster", groups=d_up["issuer_id"])
    add_ame("upsize_logit_cluster", R["upsize_logit_cluster"], X_up, ["si", "log_buzz"])

    d_up_split = subset_for_model(sub, "upsize_dummy", sent_vars_split)
    X_up_split = build_X(d_up_split, sent_vars_split, add_fe=True)
    R["upsize_split_logit"] = fit_glm_binomial(
        d_up_split["upsize_dummy"].astype(float), X_up_split, groups=d_up_split["issuer_id"]
    )
    add_ame("upsize_split_logit", R["upsize_split_logit"], X_up_split, ["si_positive", "si_negative", "log_buzz"])

    d_joint = subset_for_model(sub, "joint_success_upsize", sent_vars_si)
    X_joint = build_X(d_joint, sent_vars_si, add_fe=True)
    y_joint = d_joint["joint_success_upsize"].astype(float)
    R["joint_logit_cluster"] = fit_glm_binomial(y_joint, X_joint, groups=d_joint["issuer_id"])
    R["joint_lpm_cluster"] = fit_ols(y_joint, X_joint, cov_type="cluster", groups=d_joint["issuer_id"])
    add_ame("joint_logit_cluster", R["joint_logit_cluster"], X_joint, ["si", "log_buzz"])

    d_joint_split = subset_for_model(sub, "joint_success_upsize", sent_vars_split)
    X_joint_split = build_X(d_joint_split, sent_vars_split, add_fe=True)
    R["joint_split_logit"] = fit_glm_binomial(
        d_joint_split["joint_success_upsize"].astype(float), X_joint_split, groups=d_joint_split["issuer_id"]
    )
    add_ame("joint_split_logit", R["joint_split_logit"], X_joint_split, ["si_positive", "si_negative", "log_buzz"])

    # Multinomial placement outcome: neither is the baseline category.
    # Keep it on the main baseline sample; smaller robustness cuts often create sparse category/FE cells.
    d_success_type = subset_for_model(sub, "success_type_code", sent_vars_si)
    X_success_type = build_X(d_success_type, sent_vars_si, add_fe=True)
    if sample_name == "Baseline" and ESTIMATE_MNLOGIT:
        y_success_type = d_success_type["success_type_code"].astype(int)
        R["success_type_mnlogit"] = fit_mnlogit(y_success_type, X_success_type, groups=d_success_type["issuer_id"])
        R["_mnlogit_rows"] = collect_mnlogit_rows(
            R["success_type_mnlogit"],
            ["si", "log_buzz"],
            int(len(d_success_type)),
        )
    else:
        R["success_type_mnlogit"] = None
        R["_mnlogit_rows"] = pd.DataFrame()

    # Metadata.
    R["_n_base"] = int(len(sub))
    R["_n_coupon"] = int(len(d_coupon))
    R["_n_coupon_binary"] = int(len(d_coupon_bin))
    R["_n_coupon_positive"] = int(len(d_coupon_pos))
    R["_n_volume"] = int(len(d_vol))
    R["_n_upsize"] = int(len(d_up))
    R["_n_joint"] = int(len(d_joint))
    R["_n_success_type"] = int(len(d_success_type))
    R["_success_type_counts"] = d_success_type["success_type"].value_counts().to_dict()
    R["_n_issuers"] = int(sub["issuer"].nunique())
    R["_cn_si"] = list(X_coupon.columns)
    R["_cn_si_coupon_binary"] = list(X_coupon_bin.columns)
    R["_cn_si_coupon_positive"] = list(X_coupon_pos.columns)
    R["_cn_si_volume"] = list(X_vol.columns)
    R["_cn_si_upsize"] = list(X_up.columns)
    R["_cn_si_joint"] = list(X_joint.columns)
    R["_cn_success_type"] = list(X_success_type.columns)
    R["_cn_sh_coupon"] = list(X_coupon_sh.columns)
    R["_cn_sh_volume"] = list(X_vol_sh.columns)
    R["_cn_split_coupon_binary"] = list(X_coupon_bin_split.columns)
    R["_cn_split_coupon_positive"] = list(X_coupon_pos_split.columns)
    R["_cn_split_volume"] = list(X_vol_split.columns)
    R["_cn_split_upsize"] = list(X_up_split.columns)
    R["_cn_split_joint"] = list(X_joint_split.columns)
    R["_ame_rows"] = ame_rows

    return R


def drop_top_cooks_d(sub: pd.DataFrame, y_col: str, q: float = 0.99) -> pd.DataFrame:
    d = subset_for_model(sub, y_col, SENT_VARS_SI)
    if len(d) < 30:
        return sub.copy()
    X = build_X(d, SENT_VARS_SI, add_fe=True)
    y = d[y_col].astype(float)
    cooks = sm.OLS(y, X).fit().get_influence().cooks_distance[0]
    cutoff = np.nanquantile(cooks, q)
    keep_index = d.index[cooks <= cutoff]
    return sub.loc[sub.index.isin(keep_index)].copy()


def make_robustness_samples(bundle: SampleBundle) -> Dict[str, pd.DataFrame]:
    full = bundle.full
    base = bundle.baseline
    samples = {
        "Baseline": base,
        "Full cleaned sample": full,
        "No ВЭБ.РФ only": full.loc[~full["is_special_issuer"]].copy(),
        "No zero organizers only": full.loc[~full["zero_organizers"]].copy(),
        "Rating A+ to BB only": base.loc[base["rating_num"].between(4, 11)].copy(),
        f"Baseline without top {int((1 - COOK_DROP_Q) * 100)}% Cook D volume": drop_top_cooks_d(
            base, "log_placement_vol_book", COOK_DROP_Q
        ),
    }
    return {k: v for k, v in samples.items() if len(v) >= 30}


def make_control_robustness_specs() -> Dict[str, List[str]]:
    return {
        "Baseline without hist avg volume ratio": [v for v in BASE_VARS if v != "hist_avg_volume_ratio"],
        "Baseline without ln issue number": [v for v in BASE_VARS if v != "log_full_issue_number"],
        "Baseline without both issue-count controls": [
            v for v in BASE_VARS if v not in {"hist_avg_volume_ratio", "log_full_issue_number"}
        ],
    }


def estimate_all_samples(bundle: SampleBundle) -> Dict[str, Dict[str, Any]]:
    out: Dict[str, Dict[str, Any]] = {}
    for name, df_s in make_robustness_samples(bundle).items():
        out[name] = estimate_models(df_s, sample_name=name)
    for name, base_vars in make_control_robustness_specs().items():
        out[name] = estimate_models(bundle.baseline, sample_name=name, base_vars=base_vars)
    return out


def collect_ame_rows(all_results: Dict[str, Dict[str, Any]]) -> pd.DataFrame:
    rows = []
    for result in all_results.values():
        rows.extend(result.get("_ame_rows", []))
    return pd.DataFrame(rows)


def estimate_extension_terms(sub: pd.DataFrame) -> pd.DataFrame:
    rows = []
    outcome_specs = [
        ("coupon_reduction_bp", "Coupon reduction, bp", "OLS cluster", "coupon"),
        ("ihs_coupon_reduction", "IHS coupon reduction", "OLS cluster", "coupon"),
        ("coupon_reduced_dummy", "Coupon reduction dummy", "Logit cluster", "coupon_binary"),
        ("ihs_coupon_positive", "IHS coupon reduction | coupon > 0", "OLS cluster", "coupon_positive"),
        ("log_placement_vol_book", "ln(volume / teaser)", "OLS cluster", "volume"),
        ("upsize_dummy", "Upsize dummy", "Logit cluster", "upsize"),
        ("joint_success_upsize", "Joint success", "Logit cluster", "joint"),
    ]
    for y_col, y_label, model_type, outcome_group in outcome_specs:
        for term, term_type, expected_sign, hypothesis in EXTENSION_TERMS:
            variables = SENT_VARS_SI + [term]
            d = subset_for_model(sub, y_col, variables)
            x_vars = [v for v in variables if d[v].nunique(dropna=False) > 1]
            if len(d) < len(x_vars) + 5 or term not in x_vars:
                rows.append({
                    "dependent": y_col,
                    "label": y_label,
                    "model": model_type,
                    "term": term,
                    "term_label": VAR_LABELS.get(term, term),
                    "type": term_type,
                    "expected_sign": expected_sign,
                    "hypothesis": hypothesis,
                    "N": len(d),
                    "coef": np.nan,
                    "std_error": np.nan,
                    "p_value": np.nan,
                    "evidence": "too few observations or constant term",
                })
                continue
            X = build_X(d, x_vars, add_fe=True)
            y = d[y_col].astype(float)
            if outcome_group in {"coupon", "coupon_positive", "volume"}:
                res = fit_ols(y, X, cov_type="cluster", groups=d["issuer_id"])
            else:
                res = fit_glm_binomial(y, X, groups=d["issuer_id"])
            coef, se, p_value = get_csp(res, term, list(X.columns))
            rows.append({
                "dependent": y_col,
                "label": y_label,
                "model": model_type,
                "term": term,
                "term_label": VAR_LABELS.get(term, term),
                "type": term_type,
                "expected_sign": expected_sign,
                "hypothesis": hypothesis,
                "N": int(res.nobs) if hasattr(res, "nobs") else len(d),
                "coef": round(coef, 6) if coef is not None and pd.notna(coef) else np.nan,
                "std_error": round(se, 6) if se is not None and pd.notna(se) else np.nan,
                "p_value": round(p_value, 6) if p_value is not None and pd.notna(p_value) else np.nan,
                "evidence": evidence_label(coef, p_value, expected_sign) if coef is not None and p_value is not None else "not estimated",
            })
    return pd.DataFrame(rows)


# ============================================================
# EXCEL OUTPUT
# ============================================================

def write_model_table(
    ws,
    title: str,
    model_specs: List[Tuple[str, str, List[str]]],
    results: Dict[str, Any],
    start_row: int = 1,
) -> int:
    hf = Font(bold=True, size=10, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    sfill = PatternFill("solid", fgColor="E2EFDA")
    n_models = len(model_specs)

    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=13, name="Arial")
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=1 + n_models)
    row = start_row + 2
    ws.row_dimensions[row].height = 92
    ws.cell(row=row, column=1, value="Variable").font = hf
    ws.cell(row=row, column=1).fill = hfill
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="center")

    for j, (_, ttl, _) in enumerate(model_specs, start=2):
        c = ws.cell(row=row, column=j, value=ttl)
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    row += 1

    all_vars: List[str] = []
    for _, _, cn in model_specs:
        for v in cn:
            if v.startswith("year_") or v.startswith("q_") or v.startswith("rating_bucket_") or v.startswith("industry_"):
                continue
            if v not in all_vars:
                all_vars.append(v)

    time_fe_vars: List[str] = []
    for _, _, cn in model_specs:
        for v in cn:
            if not (v.startswith("year_") or v.startswith("q_")):
                continue
            if v not in time_fe_vars:
                time_fe_vars.append(v)

    def variable_label(var: str) -> str:
        if var.startswith("year_"):
            value = var.replace("year_", "")
            return f"Regime {value}" if "-" in value else f"Year {value}"
        if var.startswith("q_"):
            return f"Quarter {var.replace('q_', '')}"
        return VAR_LABELS.get(var, var)

    def write_coef_row(var: str) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=variable_label(var)).font = Font(name="Arial", size=10)
        for j, (key, _, cn) in enumerate(model_specs, start=2):
            res = results[key]
            coef, se, p = get_csp(res, var, cn)
            if coef is None:
                ws.cell(row=row, column=j, value="—").font = Font(name="Arial", size=10, color="999999")
                continue
            coef_text = f"{coef:.4f}{sig_stars(p)}"
            coef_cell = ws.cell(row=row, column=j, value=coef_text)
            coef_cell.alignment = Alignment(horizontal="right")
            coef_cell.font = Font(name="Arial", size=10, bold=(p is not None and p < 0.10))
            if p is not None and p < 0.10:
                coef_cell.fill = sfill
            se_cell = ws.cell(row=row + 1, column=j, value=f"({se:.4f})" if se is not None and pd.notna(se) else "(—)")
            se_cell.font = Font(name="Arial", size=9, color="666666")
            se_cell.alignment = Alignment(horizontal="right")
        row += 2

    for var in all_vars:
        write_coef_row(var)

    if time_fe_vars:
        row += 1
        ws.cell(row=row, column=1, value="Time fixed effects").font = Font(bold=True, name="Arial", size=10)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1 + n_models)
        row += 1
        for var in time_fe_vars:
            write_coef_row(var)

    row += 1
    ws.cell(row=row, column=1, value="Model statistics").font = Font(bold=True, name="Arial")
    row += 1
    for stat_name in ["N", "R²", "Adj. R²", "Pseudo R²", "Log-lik", "AIC"]:
        ws.cell(row=row, column=1, value=stat_name).font = Font(name="Arial", size=10, italic=True)
        for j, (key, _, _) in enumerate(model_specs, start=2):
            ws.cell(row=row, column=j, value=model_stat_value(results[key], stat_name)).font = Font(name="Arial", size=10)
        row += 1

    row += 1
    notes = [
        "*** p<0.01, ** p<0.05, * p<0.1. Standard errors in parentheses.",
        "OLS HC3 = heteroskedasticity-robust SE; cluster = SE clustered by issuer.",
        f"Upsize dummy = 1 if placement_vol_book > {UPSIZE_THRESHOLD:.2f}.",
        "log volume outcome = ln(placement_volume / teaser_volume).",
        "strict/dual success = coupon_reduction_bp > 0 and placement_vol_book > upsize threshold.",
        "multinomial success type uses neither as the baseline category.",
        "Missing SI/buzz/shares values are set to 0.",
        "Baseline excludes ВЭБ.РФ, num_organizers=0, and book_date > placement_date.",
        "Year fixed effects included." if USE_YEAR_FE and not USE_QUARTER_FE else "Quarter fixed effects included." if USE_QUARTER_FE else "No time fixed effects included.",
        "Rating-bucket fixed effects included." if USE_RATING_BUCKET_FE else "Rating-bucket fixed effects not included.",
        "When rating_num and rating-bucket FE are both included, rating_num is interpreted within rating buckets." if USE_RATING_BUCKET_FE and "rating_num" in BASE_VARS else "",
        f"Industry fixed effects included; industries with baseline N < {INDUSTRY_FE_MIN_N} are grouped." if USE_INDUSTRY_FE else "Industry fixed effects not included.",
    ]
    for note in notes:
        if not note:
            continue
        ws.cell(row=row, column=1, value=note).font = Font(name="Arial", size=9, italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1 + n_models)
        row += 1

    ws.freeze_panes = ws["B4"]
    ws.column_dimensions["A"].width = 36
    for i in range(2, 2 + n_models):
        ws.column_dimensions[get_column_letter(i)].width = 24
    return row


def write_robustness_snapshot(ws, all_results: Dict[str, Dict[str, Any]]) -> None:
    hf = Font(bold=True, size=10, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    ws["A1"] = "Robustness Snapshot: SI and Buzz Coefficients"
    ws["A1"].font = Font(bold=True, size=13, name="Arial")
    headers = [
        "Sample", "N coupon", "N coupon > 0", "N volume", "N upsize", "N joint",
        "SI coupon", "p", "SI IHS coupon", "p", "SI coupon logit", "p", "SI IHS coupon > 0", "p",
        "SI log volume", "p", "SI upsize logit", "p", "SI joint logit", "p",
        "Buzz coupon", "p", "Buzz IHS coupon", "p", "Buzz coupon logit", "p", "Buzz IHS coupon > 0", "p",
        "Buzz log volume", "p", "Buzz upsize logit", "p",
    ]
    row = 3
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=j, value=h)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    row += 1

    def pair(res, var, cn):
        coef, _, p = get_csp(res, var, cn)
        if coef is None:
            return "—", "—"
        return round(coef, 4), round(p, 4) if p is not None and pd.notna(p) else "—"

    for sample_name, R in all_results.items():
        vals = [sample_name, R["_n_coupon"], R["_n_coupon_positive"], R["_n_volume"], R["_n_upsize"], R["_n_joint"]]
        specs = [
            ("coupon_ols_cluster", "si", R["_cn_si"]),
            ("coupon_ihs_cluster", "si", R["_cn_si"]),
            ("coupon_reduced_logit_cluster", "si", R["_cn_si_coupon_binary"]),
            ("coupon_positive_ihs_cluster", "si", R["_cn_si_coupon_positive"]),
            ("logvol_ols_cluster", "si", R["_cn_si_volume"]),
            ("upsize_logit_cluster", "si", R["_cn_si_upsize"]),
            ("joint_logit_cluster", "si", R["_cn_si_joint"]),
            ("coupon_ols_cluster", "log_buzz", R["_cn_si"]),
            ("coupon_ihs_cluster", "log_buzz", R["_cn_si"]),
            ("coupon_reduced_logit_cluster", "log_buzz", R["_cn_si_coupon_binary"]),
            ("coupon_positive_ihs_cluster", "log_buzz", R["_cn_si_coupon_positive"]),
            ("logvol_ols_cluster", "log_buzz", R["_cn_si_volume"]),
            ("upsize_logit_cluster", "log_buzz", R["_cn_si_upsize"]),
        ]
        for key, var, cn in specs:
            vals.extend(pair(R[key], var, cn))
        for j, v in enumerate(vals, start=1):
            ws.cell(row=row, column=j, value=v).font = Font(name="Arial", size=10)
        row += 1

    ws.column_dimensions["A"].width = 34
    for i in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15


def write_descriptive_stats(ws, sub: pd.DataFrame, title: str) -> None:
    hf = Font(bold=True, size=10, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, name="Arial")
    headers = ["Variable", "N", "Mean", "Std", "Min", "p25", "Median", "p75", "Max"]
    row = 3
    for j, h in enumerate(headers, start=1):
        ws.cell(row=row, column=j, value=h).font = hf
        ws.cell(row=row, column=j).fill = hfill
    row += 1
    vars_to_show = [
        "coupon_reduction_bp", "coupon_reduced_dummy", "ihs_coupon_reduction", "ihs_coupon_positive",
        "placement_vol_book", "log_placement_vol_book", "upsize_dummy",
        "full_or_more_dummy", "joint_success_upsize", "joint_success_full", "any_success",
        "success_score", "success_type_code", "si", "log_buzz",
        "share_pos", "share_neg", "si_positive", "si_negative",
        "num_organizers", "log_num_organizers", "history_debut", "history_repeat_prior_tightening",
        "hist_reduction_share",
        "log_full_issue_number", "hist_ever_reduced", "hist_avg_volume_ratio", "rating_num",
        "ofz_yield", "ofz_yield_within_year", "rvi", "log_dur", "has_put",
        "si_x_is_floater", "log_buzz_x_is_floater",
    ]
    for var in vars_to_show:
        if var not in sub.columns:
            continue
        s = pd.to_numeric(sub[var], errors="coerce")
        if s.notna().sum() == 0:
            continue
        vals = [int(s.notna().sum()), s.mean(), s.std(), s.min(), s.quantile(.25), s.median(), s.quantile(.75), s.max()]
        ws.cell(row=row, column=1, value=VAR_LABELS.get(var, var)).font = Font(name="Arial", size=10)
        for j, val in enumerate(vals, start=2):
            ws.cell(row=row, column=j, value=round(float(val), 4) if pd.notna(val) else "—").number_format = "0.0000"
        row += 1
    ws.column_dimensions["A"].width = 34
    for i in range(2, 10):
        ws.column_dimensions[get_column_letter(i)].width = 12


def write_multinomial_table(ws, title: str, R: Dict[str, Any]) -> None:
    mn_rows = R.get("_mnlogit_rows", pd.DataFrame()).copy()
    counts = R.get("_success_type_counts", {})
    count_rows = pd.DataFrame(
        [{"success_type": label, "N": int(counts.get(name, 0))} for name, label in [
            ("neither", "Neither"),
            ("price_only", "Price only"),
            ("volume_only", "Volume only"),
            ("dual_success", "Dual success"),
        ]]
    )
    if not mn_rows.empty:
        mn_rows = mn_rows[["outcome", "baseline", "label", "N", "coef", "std_error", "p_value", "stars"]]
    write_dataframe_sheet(ws, title, mn_rows)
    start = ws.max_row + 3
    ws.cell(row=start, column=1, value="Outcome Counts").font = Font(bold=True, size=13, name="Arial")
    header_row = start + 2
    hfill = PatternFill("solid", fgColor="D9E1F2")
    for j, col in enumerate(count_rows.columns, start=1):
        cell = ws.cell(row=header_row, column=j, value=col)
        cell.font = Font(bold=True, name="Arial", size=10)
        cell.fill = hfill
    for i, (_, rec) in enumerate(count_rows.iterrows(), start=header_row + 1):
        ws.cell(row=i, column=1, value=rec["success_type"]).font = Font(name="Arial", size=10)
        ws.cell(row=i, column=2, value=int(rec["N"])).font = Font(name="Arial", size=10)


def regressor_inventory(columns: List[str]) -> pd.DataFrame:
    rows = []
    for col in columns:
        if col == "const":
            group = "constant"
        elif col.startswith("year_") or col.startswith("q_"):
            group = "time fixed effect"
        elif col.startswith("rating_bucket_"):
            group = "rating-bucket fixed effect"
        elif col.startswith("industry_"):
            group = "industry fixed effect"
        elif col in {"si", "log_buzz", "si_positive", "si_negative", "share_pos", "share_neg", *MAIN_INTERACTION_VARS}:
            group = "sentiment / attention"
        elif col in BASE_VARS:
            group = "control"
        else:
            group = "other"
        rows.append({
            "regressor": col,
            "label": VAR_LABELS.get(col, col),
            "group": group,
        })
    return pd.DataFrame(rows)


def write_dataframe_sheet(ws, title: str, df_data: pd.DataFrame) -> None:
    hf = Font(bold=True, size=10, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, name="Arial")
    row = 3
    if df_data.empty:
        ws.cell(row=row, column=1, value="No rows")
        return
    for j, h in enumerate(df_data.columns, start=1):
        cell = ws.cell(row=row, column=j, value=h)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(wrap_text=True)
    row += 1
    for _, r in df_data.iterrows():
        for j, col in enumerate(df_data.columns, start=1):
            val = r[col]
            ws.cell(row=row, column=j, value=float(val) if isinstance(val, np.floating) else int(val) if isinstance(val, np.integer) else val)
        row += 1
    for i in range(1, len(df_data.columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    ws.column_dimensions["H"].width = 44
    ws.freeze_panes = "A4"


def build_excel(all_results: Dict[str, Dict[str, Any]], bundle: SampleBundle, extension_df: pd.DataFrame, ame_df: pd.DataFrame, path: str) -> None:
    wb = Workbook()
    Rb = all_results["Baseline"]
    fe_label = "quarter" if USE_QUARTER_FE else "year" if USE_YEAR_FE else "none"

    ws1 = wb.active
    ws1.title = "Main Results"
    main_models = [
        ("coupon_ols_cluster", f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={Rb['_n_coupon']}", Rb["_cn_si"]),
        ("coupon_ihs_cluster", f"Y: IHS coupon reduction\nDV: asinh(coupon bp)\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={Rb['_n_coupon']}", Rb["_cn_si"]),
        ("coupon_reduced_logit_cluster", f"Y: coupon reduction dummy\nDV: 1 if coupon>0\nModel: logit/GLM\nSE: clustered by issuer\nFE: {fe_label}\nN={Rb['_n_coupon_binary']}", Rb["_cn_si_coupon_binary"]),
        ("coupon_positive_ihs_cluster", f"Y: IHS coupon reduction\nDV: asinh(coupon bp), coupon>0 only\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={Rb['_n_coupon_positive']}", Rb["_cn_si_coupon_positive"]),
        ("logvol_ols_cluster", f"Y: ln(volume/teaser)\nDV: log placement_vol_book\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={Rb['_n_volume']}", Rb["_cn_si_volume"]),
        ("upsize_logit_cluster", f"Y: upsize dummy\nDV: 1 if ratio>{UPSIZE_THRESHOLD:.2f}\nModel: logit/GLM\nSE: clustered by issuer\nFE: {fe_label}\nN={Rb['_n_upsize']}", Rb["_cn_si_upsize"]),
        ("joint_logit_cluster", f"Y: strict dual success\nDV: coupon>0 and upsize\nModel: logit/GLM\nSE: clustered by issuer\nFE: {fe_label}\nN={Rb['_n_joint']}", Rb["_cn_si_joint"]),
    ]
    write_model_table(ws1, f"Table 1. Main Results — Baseline Sample ({YEAR_FROM}+)", main_models, Rb)

    ws2 = wb.create_sheet("Robustness Snapshot")
    write_robustness_snapshot(ws2, all_results)

    ws3 = wb.create_sheet("Coupon Robustness")
    coupon_models = [
        ("coupon_ols_hc3", f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: HC3\nFE: {fe_label}\nN={Rb['_n_coupon']}", Rb["_cn_si"]),
        ("coupon_ols_cluster", f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon']}", Rb["_cn_si"]),
        ("coupon_ihs_cluster", f"Y: IHS coupon reduction\nDV: asinh(coupon bp)\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon']}", Rb["_cn_si"]),
        ("coupon_reduced_logit_cluster", f"Two-part step 1\nY: coupon reduction dummy\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon_binary']}", Rb["_cn_si_coupon_binary"]),
        ("coupon_positive_ihs_cluster", f"Two-part step 2\nY: asinh(coupon bp)\nSample: coupon>0\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon_positive']}", Rb["_cn_si_coupon_positive"]),
        ("coupon_winsor_cluster", f"Y: coupon reduction, bp\nDV: winsor 1/99%\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon']}", Rb["_cn_si"]),
        ("coupon_shares_cluster", f"Y: coupon reduction, bp\nX: pos/neg shares\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon']}", Rb["_cn_sh_coupon"]),
    ]
    write_model_table(ws3, "Table 2. Coupon Robustness — Baseline Sample", coupon_models, Rb)

    ws4 = wb.create_sheet("Volume Robustness")
    volume_models = [
        ("logvol_ols_hc3", f"Y: ln(volume/teaser)\nModel: OLS\nSE: HC3\nFE: {fe_label}\nN={Rb['_n_volume']}", Rb["_cn_si_volume"]),
        ("logvol_ols_cluster", f"Y: ln(volume/teaser)\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_volume']}", Rb["_cn_si_volume"]),
        ("volratio_winsor_cluster", f"Y: placement_vol_book\nDV: winsor 1/99%\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_volume']}", Rb["_cn_si_volume"]),
        ("upsize_lpm_cluster", f"Y: upsize dummy\nModel: LPM\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_upsize']}", Rb["_cn_si_upsize"]),
        ("logvol_shares_cluster", f"Y: ln(volume/teaser)\nX: pos/neg shares\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_volume']}", Rb["_cn_sh_volume"]),
    ]
    write_model_table(ws4, "Table 3. Volume Robustness — Baseline Sample", volume_models, Rb)

    ws5 = wb.create_sheet("Nonlinear Interactions")
    write_dataframe_sheet(ws5, "Nonlinear and Interaction Terms — Baseline Sample", extension_df)

    ws6 = wb.create_sheet("Split SI")
    split_models = [
        ("coupon_reduced_split_logit", f"Y: coupon reduction dummy\nX: positive/negative SI\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon_binary']}", Rb["_cn_split_coupon_binary"]),
        ("coupon_positive_split_ihs", f"Y: asinh(coupon bp)\nSample: coupon>0\nX: positive/negative SI\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_coupon_positive']}", Rb["_cn_split_coupon_positive"]),
        ("logvol_split_cluster", f"Y: ln(volume/teaser)\nX: positive/negative SI\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_volume']}", Rb["_cn_split_volume"]),
        ("upsize_split_logit", f"Y: upsize dummy\nX: positive/negative SI\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_upsize']}", Rb["_cn_split_upsize"]),
        ("joint_split_logit", f"Y: strict dual success\nX: positive/negative SI\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={Rb['_n_joint']}", Rb["_cn_split_joint"]),
    ]
    write_model_table(ws6, "Table 4. Positive vs Negative Sentiment — Baseline Sample", split_models, Rb)

    ws7 = wb.create_sheet("Marginal Effects")
    write_dataframe_sheet(ws7, "Average Marginal Effects for Logit Models", ame_df)

    ws8 = wb.create_sheet("Multinomial Success")
    write_multinomial_table(ws8, "Table 5b. Multinomial Success Type — Baseline Sample", Rb)

    ws9 = wb.create_sheet("Regressors")
    write_dataframe_sheet(ws9, "Actual Regressors in Main SI Specification", regressor_inventory(Rb["_cn_si"]))

    ws10 = wb.create_sheet("Descriptive Baseline")
    write_descriptive_stats(ws10, bundle.baseline, f"Descriptive Statistics — Baseline Sample (N={len(bundle.baseline)})")
    ws11 = wb.create_sheet("Descriptive Full")
    write_descriptive_stats(ws11, bundle.full, f"Descriptive Statistics — Full Cleaned Sample (N={len(bundle.full)})")

    wb.save(path)


# ============================================================
# CONSOLE OUTPUT
# ============================================================

def print_key_result(res: Any, var: str, cn: List[str]) -> str:
    coef, se, p = get_csp(res, var, cn)
    if coef is None:
        return "—"
    return f"{coef:.4f} ({se:.4f}), p={p:.4f} {sig_stars(p)}"


def print_sample_summary(bundle: SampleBundle) -> None:
    full = bundle.full
    base = bundle.baseline
    print("\n" + "=" * 70)
    print("SAMPLE / DATA QUALITY")
    print("=" * 70)
    print(f"Full cleaned sample       N={len(full)}, issuers={full['issuer'].nunique()}")
    print(f"Baseline sample           N={len(base)}, issuers={base['issuer'].nunique()}")
    print(f"Baseline coupon N         {base['coupon_reduction_bp'].notna().sum()}")
    print(f"Baseline coupon > 0 N     {(base['coupon_reduction_bp'] > 0).sum()}")
    print(f"Baseline volume N         {base['placement_vol_book'].notna().sum()}")
    print(f"Baseline joint N          {(base['coupon_reduction_bp'].notna() & base['placement_vol_book'].notna()).sum()}")
    print(f"Upsize cases              {(base['placement_vol_book'] > UPSIZE_THRESHOLD).sum()}")
    print(f"Full/more volume cases    {(base['placement_vol_book'] >= FULL_VOLUME_THRESHOLD).sum()}")
    if "success_type" in base.columns:
        print("Success type:")
        for name, n in base["success_type"].value_counts().items():
            print(f"  {name:28s} {int(n)}")
    print(f"book_date > placement     {int(full['book_after_placement'].sum())}")
    print(f"Special issuer ВЭБ.РФ     {int(full['is_special_issuer'].sum())}")
    print(f"num_organizers = 0        {int(full['zero_organizers'].sum())}")
    if "rating_bucket" in base.columns:
        print("Rating buckets:")
        for name, n in base["rating_bucket"].value_counts().items():
            print(f"  {name:28s} {int(n)}")
    if "issuer_history_group" in base.columns:
        print("Issuer history groups:")
        for name, n in base["issuer_history_group"].value_counts().items():
            print(f"  {name:28s} {int(n)}")
    if "industry_fe" in base.columns:
        print(f"Industry FE categories    {base['industry_fe'].nunique()} (min N={INDUSTRY_FE_MIN_N}; rare grouped)")
    if "volume_increased_dummy" in full.columns:
        old = pd.to_numeric(full["volume_increased_dummy"], errors="coerce")
        print(f"Old volume_increased_dummy sum {int(old.fillna(0).sum())} — do not use for new upsize logic")


def print_results(all_results: Dict[str, Dict[str, Any]]) -> None:
    print("\n" + "=" * 70)
    print("REGRESSION SNAPSHOT")
    print("=" * 70)
    for sample_name, R in all_results.items():
        print(f"\n--- {sample_name} ---")
        print(f"N coupon={R['_n_coupon']}, N coupon>0={R['_n_coupon_positive']}, N volume={R['_n_volume']}, N upsize={R['_n_upsize']}, N joint={R['_n_joint']}")
        print(f"SI coupon cluster:     {print_key_result(R['coupon_ols_cluster'], 'si', R['_cn_si'])}")
        print(f"SI IHS coupon:         {print_key_result(R['coupon_ihs_cluster'], 'si', R['_cn_si'])}")
        print(f"SI coupon logit:       {print_key_result(R['coupon_reduced_logit_cluster'], 'si', R['_cn_si_coupon_binary'])}")
        print(f"SI IHS coupon > 0:     {print_key_result(R['coupon_positive_ihs_cluster'], 'si', R['_cn_si_coupon_positive'])}")
        print(f"SI log volume cluster: {print_key_result(R['logvol_ols_cluster'], 'si', R['_cn_si_volume'])}")
        print(f"SI upsize logit:       {print_key_result(R['upsize_logit_cluster'], 'si', R['_cn_si_upsize'])}")
        print(f"SI joint logit:        {print_key_result(R['joint_logit_cluster'], 'si', R['_cn_si_joint'])}")
        if R.get("success_type_mnlogit") is not None:
            for code in [1, 2, 3]:
                coef, se, p = get_mnlogit_csp(R["success_type_mnlogit"], "si", code)
                if coef is not None:
                    print(f"SI MNLogit {SUCCESS_TYPE_LABELS[code]:12s} vs Neither: {coef:.4f} ({se:.4f}), p={p:.4f} {sig_stars(p)}")
        print(f"Buzz coupon cluster:   {print_key_result(R['coupon_ols_cluster'], 'log_buzz', R['_cn_si'])}")
        print(f"Buzz IHS coupon:       {print_key_result(R['coupon_ihs_cluster'], 'log_buzz', R['_cn_si'])}")
        print(f"Buzz coupon logit:     {print_key_result(R['coupon_reduced_logit_cluster'], 'log_buzz', R['_cn_si_coupon_binary'])}")
        print(f"Buzz IHS coupon > 0:   {print_key_result(R['coupon_positive_ihs_cluster'], 'log_buzz', R['_cn_si_coupon_positive'])}")
        print(f"Buzz log volume:       {print_key_result(R['logvol_ols_cluster'], 'log_buzz', R['_cn_si_volume'])}")


def print_extension_results(extension_df: pd.DataFrame) -> None:
    print("\n" + "=" * 70)
    print("NONLINEAR / INTERACTION TERMS (Baseline, p < 0.10)")
    print("=" * 70)
    sig = extension_df[pd.to_numeric(extension_df["p_value"], errors="coerce") < 0.10]
    if sig.empty:
        print("none")
        return
    for _, r in sig.iterrows():
        print(f"{r['dependent']:25s} {r['term_label']:28s} coef={r['coef']}, p={r['p_value']}, {r['evidence']}")


def print_ame_results(ame_df: pd.DataFrame) -> None:
    print("\n" + "=" * 70)
    print("AVERAGE MARGINAL EFFECTS (Baseline)")
    print("=" * 70)
    if ame_df.empty:
        print("none")
        return
    base = ame_df[ame_df["sample"] == "Baseline"]
    for _, r in base.iterrows():
        print(f"{r['model']:32s} {r['label']:18s} AME={r['AME']}, p={r['p_value']}")


def main() -> None:
    input_file = pick_input_file()
    output_file = pick_output_file()
    print(f"Input:  {input_file}")
    print(f"Output: {output_file}")
    print("Загрузка и подготовка данных...")
    bundle = load_and_prepare(input_file, YEAR_FROM)
    print_sample_summary(bundle)
    print("\nОценка моделей...")
    all_results = estimate_all_samples(bundle)
    ame_df = collect_ame_rows(all_results)
    extension_df = estimate_extension_terms(bundle.baseline)
    print_results(all_results)
    print_extension_results(extension_df)
    print_ame_results(ame_df)
    print("\nСохранение Excel...")
    build_excel(all_results, bundle, extension_df, ame_df, output_file)
    print(f"\n✅ Таблицы сохранены: {output_file}")


if __name__ == "__main__":
    main()
