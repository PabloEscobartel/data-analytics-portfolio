#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Diagnostics for primary bond placement regressions with upsize outcome.

Checks:
- sample construction and data quality
- VIF and correlations
- heteroskedasticity, residual normality, RESET, Cook's D for OLS outcomes
- quick robustness snapshot for coupon and volume outcomes

Usage:
    python run_diagnostics_upsize.py
    python run_diagnostics_upsize.py regression_ready.csv
    python run_diagnostics_upsize.py regression_ready.csv diagnostics_results_upsize.xlsx

Requirements:
    pip install pandas numpy statsmodels scipy openpyxl
"""

from __future__ import annotations

import os
import sys
import warnings
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats
from statsmodels.stats.diagnostic import het_breuschpagan
from statsmodels.stats.outliers_influence import variance_inflation_factor
from statsmodels.stats.stattools import jarque_bera

try:
    os.environ.setdefault("MPLCONFIGDIR", "/private/tmp/matplotlib")
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    HAS_MATPLOTLIB = True
except Exception:
    HAS_MATPLOTLIB = False

warnings.filterwarnings("default")

# ============================================================
# SETTINGS
# ============================================================

DEFAULT_INPUT_CANDIDATES = [
    # "regression_dataset_book_close/regression_ready.csv",
    "regression_dataset_book_open/regression_ready.csv",
]
DEFAULT_OUTPUT_FILE = "diagnostics_results_upsize.xlsx"
YEAR_FROM = 2018

BASELINE_EXCLUDE_ISSUERS = {"ВЭБ.РФ"}
BASELINE_DROP_ZERO_ORGANIZERS = True
BASELINE_DROP_BOOK_AFTER_PLACEMENT = True
USE_RATING_BUCKET_FE = False
USE_INDUSTRY_FE = False
INDUSTRY_FE_MIN_N = 20

UPSIZE_THRESHOLD = 1.01
FULL_VOLUME_THRESHOLD = 0.99
WINSOR_LOWER = 0.01
WINSOR_UPPER = 0.99

RATING_MAP = {
    "AAA": 15, "AA+": 14, "AA": 13, "AA-": 12,
    "A+": 11, "A": 10, "A-": 9,
    "BBB+": 8, "BBB": 7, "BBB-": 6,
    "BB+": 5, "BB": 4, "BB-": 3,
    "B+": 2, "B": 1, "B-": 0,
    "C": -1,
}

VAR_LABELS = {
    "const": "Constant",
    "coupon_reduced_dummy": "Coupon reduction dummy",
    "ihs_coupon_reduction": "IHS coupon reduction",
    "ihs_coupon_positive": "IHS coupon reduction | coupon > 0",
    "num_organizers": "Number of organizers",
    "log_num_organizers": "ln(1 + organizers)",
    "issuer_history_group": "Issuer history group",
    "history_debut": "Debut issuer",
    "history_repeat_prior_tightening": "Repeat issuer, prior tightening",
    "hist_reduction_share": "Hist. share of prior tightenings",
    "log_full_issue_number": "ln(Full issue number)",
    "hist_ever_reduced": "Prior tightening",
    "hist_avg_volume_ratio": "Hist. avg volume ratio",
    "rating_num": "Credit rating (ordinal)",
    "rating_bucket": "Rating bucket",
    "industry_fe": "Industry FE",
    "log_dur": "ln(Term to exit, days)",
    "has_put": "Has put option",
    "is_floater": "Floating rate (dummy)",
    "si": "Sentiment Index (SI)",
    "si_positive": "Positive SI",
    "si_negative": "Negative SI intensity",
    "si_pos": "Positive SI",
    "si_neg": "Negative SI intensity",
    "si_sq": "SI squared",
    "log_buzz": "ln(Buzz)",
    "log_buzz_sq": "ln(Buzz) squared",
    "ofz_yield": "OFZ matched yield, %",
    "ofz_yield_within_year": "OFZ yield, within-year deviation",
    "rvi": "RVI (volatility index)",
    "rating_num_sq": "Credit rating squared",
    "si_x_log_buzz": "SI x ln(Buzz)",
    "si_x_is_floater": "SI x Floater",
    "log_buzz_x_is_floater": "ln(Buzz) x Floater",
    "si_x_rating_num": "SI x Rating",
    "log_buzz_x_rating_num": "ln(Buzz) x Rating",
    "num_organizers_x_si": "Organizers x SI",
}

BASE_VARS = [
    "log_num_organizers", "history_debut", "hist_reduction_share", "log_full_issue_number", "hist_avg_volume_ratio",
    "rating_num", "ofz_yield_within_year", "rvi", "log_dur", "has_put", "is_floater",
]
MAIN_INTERACTION_VARS: List[str] = []
EXPLANATORY_VARS = BASE_VARS + ["si", "log_buzz"]
OLS_OUTCOMES = [
    ("coupon_reduction_bp", "Coupon reduction, bp"),
    ("ihs_coupon_reduction", "IHS coupon reduction"),
    ("ihs_coupon_positive", "IHS coupon reduction | coupon > 0"),
    ("log_placement_vol_book", "ln(placement volume / teaser volume)"),
]
NONLINEAR_VARS = ["si_sq", "log_buzz_sq", "rating_num_sq"]
INTERACTION_VARS = [
    "si_x_log_buzz",
    "si_x_is_floater",
    "log_buzz_x_is_floater",
    "si_x_rating_num",
    "log_buzz_x_rating_num",
    "num_organizers_x_si",
]
EXTENSION_TERMS = [
    ("si_sq", "nonlinear", "negative", "Strong sentiment has diminishing/asymmetric effect"),
    ("log_buzz_sq", "nonlinear", "negative", "Media coverage has diminishing returns"),
    ("rating_num_sq", "nonlinear", "any", "Rating effect is nonlinear"),
    ("si_x_log_buzz", "interaction", "positive", "Sentiment matters more when discussion is broader"),
    ("si_x_is_floater", "interaction", "any", "Sentiment differs for floaters"),
    ("log_buzz_x_is_floater", "interaction", "any", "Media coverage effect differs for floaters"),
    ("si_x_rating_num", "interaction", "negative", "Sentiment matters more for lower-rated issuers"),
    ("log_buzz_x_rating_num", "interaction", "negative", "Coverage matters more for lower-rated issuers"),
    ("num_organizers_x_si", "interaction", "any", "Organizer syndicate changes the sentiment effect"),
]
RESET_SPECS = [
    ("Baseline", []),
    ("Baseline + nonlinear", NONLINEAR_VARS),
    ("Baseline + interactions", INTERACTION_VARS),
    ("Baseline + nonlinear + interactions", NONLINEAR_VARS + INTERACTION_VARS),
]


def make_explanatory_vars(base_vars: List[str]) -> List[str]:
    return base_vars + ["si", "log_buzz"]


def make_control_robustness_specs() -> Dict[str, List[str]]:
    return {
        "Baseline without hist avg volume ratio": [v for v in BASE_VARS if v != "hist_avg_volume_ratio"],
        "Baseline without ln issue number": [v for v in BASE_VARS if v != "log_full_issue_number"],
        "Baseline without both issue-count controls": [
            v for v in BASE_VARS if v not in {"hist_avg_volume_ratio", "log_full_issue_number"}
        ],
    }
DISTRIBUTION_NUMERIC_VARS = [
    "coupon_reduction_bp",
    "coupon_reduced_dummy",
    "ihs_coupon_reduction",
    "ihs_coupon_positive",
    "placement_vol_book",
    "log_placement_vol_book",
    "upsize_dummy",
    "joint_success_upsize",
    "si",
    "si_positive",
    "si_negative",
    "si_pos",
    "si_neg",
    "log_buzz",
    "num_organizers",
    "log_num_organizers",
    "history_debut",
    "history_repeat_prior_tightening",
    "hist_reduction_share",
    "log_full_issue_number",
    "rating_num",
    "ofz_yield",
    "ofz_yield_within_year",
    "rvi",
    "log_dur",
    "hist_avg_volume_ratio",
]
DISTRIBUTION_CATEGORICAL_VARS = [
    "final_rating",
    "rating_bucket",
    "issuer_history_group",
    "industry_fe",
    "industry",
    "year",
]


# ============================================================
# HELPERS
# ============================================================

def pick_input_file() -> str:
    if len(sys.argv) >= 2:
        return sys.argv[1]
    for p in DEFAULT_INPUT_CANDIDATES:
        if os.path.exists(p):
            return p
    raise FileNotFoundError("No input file found. Pass path explicitly.")


def pick_output_file() -> str:
    if len(sys.argv) >= 3:
        return sys.argv[2]
    return DEFAULT_OUTPUT_FILE


def first_existing_col(df: pd.DataFrame, candidates: List[str]) -> str:
    for col in candidates:
        if col in df.columns:
            return col
    raise KeyError(f"Не найдена ни одна из колонок: {', '.join(candidates)}")


def optional_existing_col(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    for col in candidates:
        if col in df.columns:
            return col
    return None


def normalize_rating(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.replace("–", "-", regex=False)
        .str.replace("—", "-", regex=False)
        .str.strip()
    )


def rating_bucket_from_num(value: float) -> str:
    if pd.isna(value):
        return "Missing rating"
    if value >= 12:
        return "High: AAA to AA-"
    if value >= 6:
        return "Middle: A+ to BBB-"
    return "HY/ВДО: BB+ and below"


def add_industry_fe(full_sub: pd.DataFrame, baseline_sub: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame]:
    if "industry" not in full_sub.columns:
        full_sub = full_sub.copy()
        baseline_sub = baseline_sub.copy()
        full_sub["industry_fe"] = "Missing industry"
        baseline_sub["industry_fe"] = "Missing industry"
        return full_sub, baseline_sub

    counts = baseline_sub["industry"].fillna("Missing industry").astype(str).value_counts()
    keep = set(counts[counts >= INDUSTRY_FE_MIN_N].index)

    def collapse(series: pd.Series) -> pd.Series:
        labels = series.fillna("Missing industry").astype(str)
        return labels.where(labels.isin(keep), "Other / rare industries")

    full_sub = full_sub.copy()
    baseline_sub = baseline_sub.copy()
    full_sub["industry_fe"] = collapse(full_sub["industry"])
    baseline_sub["industry_fe"] = collapse(baseline_sub["industry"])
    return full_sub, baseline_sub


def safe_numeric(df: pd.DataFrame, col: Optional[str], default: float = np.nan) -> pd.Series:
    if col is None:
        return pd.Series(default, index=df.index, dtype="float64")
    return pd.to_numeric(df[col], errors="coerce")


def winsorize_series(s: pd.Series, lower: float = 0.01, upper: float = 0.99) -> pd.Series:
    s2 = s.dropna()
    if s2.empty:
        return s.copy()
    return s.clip(lower=s2.quantile(lower), upper=s2.quantile(upper))


def add_extension_terms(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    sub["si_sq"] = sub["si"] ** 2
    sub["log_buzz_sq"] = sub["log_buzz"] ** 2
    sub["rating_num_sq"] = sub["rating_num"] ** 2
    sub["si_x_log_buzz"] = sub["si"] * sub["log_buzz"]
    sub["si_x_is_floater"] = sub["si"] * sub["is_floater"]
    sub["log_buzz_x_is_floater"] = sub["log_buzz"] * sub["is_floater"]
    sub["si_x_rating_num"] = sub["si"] * sub["rating_num"]
    sub["log_buzz_x_rating_num"] = sub["log_buzz"] * sub["rating_num"]
    sub["num_organizers_x_si"] = sub["num_organizers"] * sub["si"]
    return sub


def add_issuer_history_controls(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    valid = sub["full_issue_number"].notna() & sub["hist_ever_reduced"].notna()
    debut = valid & sub["full_issue_number"].eq(1)
    repeat = valid & sub["full_issue_number"].gt(1)
    prior = sub["hist_ever_reduced"].eq(1)
    no_prior = sub["hist_ever_reduced"].eq(0)

    sub["issuer_history_group"] = np.select(
        [debut, repeat & prior, repeat & no_prior],
        ["debut", "repeat_prior_tightening", "repeat_no_prior_tightening"],
        default=None,
    )
    sub["history_debut"] = pd.Series(np.where(valid, debut.astype(int), np.nan), index=sub.index).astype("Int64")
    sub["history_repeat_prior_tightening"] = pd.Series(
        np.where(valid, (repeat & prior).astype(int), np.nan),
        index=sub.index,
    ).astype("Int64")
    return sub


def evidence_label(coef: float, p_value: float, expected_sign: str) -> str:
    if pd.isna(coef) or pd.isna(p_value):
        return "not estimated"
    if p_value >= 0.10:
        return "no evidence"
    if expected_sign == "negative":
        return "supports" if coef < 0 else "opposite sign"
    if expected_sign == "positive":
        return "supports" if coef > 0 else "opposite sign"
    return "significant"


def add_quality_flags(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    sub["placement_dt"] = pd.to_datetime(sub["placement_date"], errors="coerce")
    sub["book_dt"] = pd.to_datetime(sub["book_date"], errors="coerce")
    sub["book_after_placement"] = (
        sub["book_dt"].notna()
        & sub["placement_dt"].notna()
        & (sub["book_dt"].dt.date > sub["placement_dt"].dt.date)
    )
    sub["is_special_issuer"] = sub["issuer"].astype(str).isin(BASELINE_EXCLUDE_ISSUERS)
    sub["zero_organizers"] = sub["num_organizers"].eq(0)
    sub["time_fe"] = sub["year"].astype("Int64").astype(str)
    return sub


def add_ofz_within_year(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    sub["ofz_yield_within_year"] = (
        sub["ofz_yield"]
        - sub.groupby("year")["ofz_yield"].transform("mean")
    )
    return sub


def apply_baseline_filters(sub: pd.DataFrame) -> pd.DataFrame:
    mask = pd.Series(True, index=sub.index)
    if BASELINE_DROP_BOOK_AFTER_PLACEMENT:
        mask &= ~sub["book_after_placement"]
    if BASELINE_EXCLUDE_ISSUERS:
        mask &= ~sub["is_special_issuer"]
    if BASELINE_DROP_ZERO_ORGANIZERS:
        mask &= ~sub["zero_organizers"]
    return sub.loc[mask].copy()


def build_X(data: pd.DataFrame, variables: List[str], add_fe: bool = True) -> pd.DataFrame:
    X = data[variables].copy()
    if add_fe and "time_fe" in data.columns:
        fe_parts = [pd.get_dummies(data["time_fe"], prefix="year", drop_first=True, dtype=float)]
        if USE_RATING_BUCKET_FE and "rating_bucket" in data.columns:
            fe_parts.append(pd.get_dummies(data["rating_bucket"], prefix="rating_bucket", drop_first=True, dtype=float))
        if USE_INDUSTRY_FE and "industry_fe" in data.columns:
            fe_parts.append(pd.get_dummies(data["industry_fe"], prefix="industry", drop_first=True, dtype=float))
        fe_parts = [fe for fe in fe_parts if not fe.empty]
        if fe_parts:
            X = pd.concat([X] + fe_parts, axis=1)
    X = sm.add_constant(X, has_constant="add")
    return X.apply(pd.to_numeric, errors="coerce").astype(float)


def fit_ols_summary(sample_name: str, dep_label: str, data: pd.DataFrame, y_col: str, variables: Optional[List[str]] = None) -> dict:
    variables = EXPLANATORY_VARS if variables is None else variables
    d = data.dropna(subset=[y_col] + variables).copy()
    x_vars = [v for v in variables if d[v].nunique(dropna=False) > 1]
    if len(d) < len(x_vars) + 5:
        return {
            "sample": sample_name, "dependent": dep_label, "N": len(d), "R2": np.nan,
            "si_coef": np.nan, "si_p": np.nan,
            "log_buzz_coef": np.nan, "log_buzz_p": np.nan,
            "note": "Too few observations",
        }
    X = build_X(d, x_vars, add_fe=True)
    res = sm.OLS(d[y_col], X).fit(cov_type="HC3")
    return {
        "sample": sample_name,
        "dependent": dep_label,
        "N": int(res.nobs),
        "R2": round(res.rsquared, 4),
        "si_coef": round(res.params.get("si", np.nan), 4),
        "si_p": round(res.pvalues.get("si", np.nan), 4),
        "log_buzz_coef": round(res.params.get("log_buzz", np.nan), 4),
        "log_buzz_p": round(res.pvalues.get("log_buzz", np.nan), 4),
        "note": "OLS HC3 with year FE; constant regressors omitted",
    }


def remove_top_cooks(data: pd.DataFrame, y_col: str) -> pd.DataFrame:
    d = data.dropna(subset=[y_col] + EXPLANATORY_VARS).copy()
    if len(d) < len(EXPLANATORY_VARS) + 5:
        return d
    X = build_X(d, EXPLANATORY_VARS, add_fe=True)
    cooks = sm.OLS(d[y_col], X).fit().get_influence().cooks_distance[0]
    cutoff = np.nanquantile(cooks, 0.99)
    return d.loc[cooks <= cutoff].copy()


# ============================================================
# LOAD DATA
# ============================================================

def load_data(path: str) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    df = pd.read_csv(path, low_memory=False)
    df["placement_dt"] = pd.to_datetime(df["placement_date"], errors="coerce")
    df["year"] = df["placement_dt"].dt.year
    sub = df[df["year"] >= YEAR_FROM].copy()

    # Sentiment.
    si_col = first_existing_col(sub, ["si_relevant", "si"])
    buzz_col = first_existing_col(sub, ["log_buzz_relevant", "log_buzz_all", "log_buzz"])
    si_raw = safe_numeric(sub, si_col)
    log_buzz_raw = safe_numeric(sub, buzz_col)

    sub["si"] = si_raw.fillna(0)
    sub["si_positive"] = sub["si"].clip(lower=0)
    sub["si_negative"] = (-sub["si"]).clip(lower=0)
    sub["si_pos"] = sub["si_positive"]
    sub["si_neg"] = sub["si_negative"]
    sub["log_buzz"] = log_buzz_raw.fillna(0)

    # Outcomes.
    sub["coupon_reduction_bp"] = pd.to_numeric(sub["coupon_reduction_bp"], errors="coerce")
    sub["coupon_reduced_dummy"] = np.nan
    m_coupon = sub["coupon_reduction_bp"].notna()
    sub.loc[m_coupon, "coupon_reduced_dummy"] = (sub.loc[m_coupon, "coupon_reduction_bp"] > 0).astype(int)
    sub["ihs_coupon_reduction"] = np.arcsinh(sub["coupon_reduction_bp"])
    sub["ihs_coupon_positive"] = np.where(
        sub["coupon_reduction_bp"] > 0,
        np.arcsinh(sub["coupon_reduction_bp"]),
        np.nan,
    )
    sub["placement_vol_book"] = pd.to_numeric(sub["placement_vol_book"], errors="coerce")
    sub.loc[sub["placement_vol_book"] <= 0, "placement_vol_book"] = np.nan
    sub["log_placement_vol_book"] = np.log(sub["placement_vol_book"])
    sub["upsize_dummy"] = np.nan
    m_vol = sub["placement_vol_book"].notna()
    sub.loc[m_vol, "upsize_dummy"] = (sub.loc[m_vol, "placement_vol_book"] > UPSIZE_THRESHOLD).astype(int)
    sub["full_or_more_dummy"] = np.nan
    sub.loc[m_vol, "full_or_more_dummy"] = (sub.loc[m_vol, "placement_vol_book"] >= FULL_VOLUME_THRESHOLD).astype(int)
    m_joint = sub["coupon_reduction_bp"].notna() & sub["placement_vol_book"].notna()
    sub["joint_success_upsize"] = np.nan
    sub.loc[m_joint, "joint_success_upsize"] = (
        (sub.loc[m_joint, "coupon_reduction_bp"] > 0)
        & (sub.loc[m_joint, "placement_vol_book"] > UPSIZE_THRESHOLD)
    ).astype(int)

    # Controls.
    sub["num_organizers"] = pd.to_numeric(sub["num_organizers"], errors="coerce")
    sub["log_num_organizers"] = np.log1p(sub["num_organizers"])
    if "full_issue_number" not in sub.columns:
        raise KeyError("Не найдена обязательная колонка: full_issue_number")
    sub["full_issue_number"] = pd.to_numeric(sub["full_issue_number"], errors="coerce")
    sub.loc[sub["full_issue_number"] <= 0, "full_issue_number"] = np.nan
    sub["log_full_issue_number"] = np.log(sub["full_issue_number"])
    for col in ["hist_ever_reduced", "has_put", "hist_avg_volume_ratio"]:
        if col not in sub.columns:
            raise KeyError(f"Не найдена обязательная колонка: {col}")
        sub[col] = pd.to_numeric(sub[col], errors="coerce")
    if "hist_reduction_share" in sub.columns:
        sub["hist_reduction_share"] = pd.to_numeric(sub["hist_reduction_share"], errors="coerce")
    sub = add_issuer_history_controls(sub)
    sub["rating_had_long_dash"] = sub["final_rating"].astype(str).str.contains("–|—", regex=True, na=False)
    sub["final_rating"] = normalize_rating(sub["final_rating"])
    sub["rating_num"] = sub["final_rating"].map(RATING_MAP)
    sub["rating_bucket"] = sub["rating_num"].map(rating_bucket_from_num)
    if "log_term" in sub.columns:
        sub["log_dur"] = pd.to_numeric(sub["log_term"], errors="coerce")
    elif "term" in sub.columns:
        sub["log_dur"] = np.log(pd.to_numeric(sub["term"], errors="coerce").clip(lower=1))
    else:
        raise KeyError("Нужна колонка log_term или term")
    sub["term"] = pd.to_numeric(sub["term"], errors="coerce") if "term" in sub.columns else np.nan
    sub["ofz_yield"] = pd.to_numeric(sub["ofz_matched_yield"], errors="coerce")
    sub = add_ofz_within_year(sub)
    sub["rvi"] = pd.to_numeric(sub["rvi_close"], errors="coerce")
    sub["is_floater"] = pd.to_numeric(sub.get("is_floater", 0), errors="coerce")
    sub = add_extension_terms(sub)

    sub = add_quality_flags(sub)
    common_drop = EXPLANATORY_VARS.copy()
    full_sub = sub.dropna(subset=common_drop).copy()
    baseline_sub = apply_baseline_filters(full_sub)
    full_sub = add_ofz_within_year(full_sub)
    baseline_sub = add_ofz_within_year(baseline_sub)
    full_sub, baseline_sub = add_industry_fe(full_sub, baseline_sub)

    quality_df, sample_df, anomalies_df = build_quality_reports(sub, full_sub, baseline_sub)
    return full_sub, baseline_sub, quality_df, sample_df, anomalies_df


def build_quality_reports(raw_sub: pd.DataFrame, full_sub: pd.DataFrame, baseline_sub: pd.DataFrame):
    rows = []

    def add(issue, n, action):
        rows.append({"issue": issue, "N": int(n), "action": action})

    add("Ratings with en/em dash before normalization", raw_sub["rating_had_long_dash"].sum(), "Normalized before rating mapping")
    add("Missing/unmapped rating after normalization", raw_sub["rating_num"].isna().sum(), "Dropped by regression dropna if still missing")
    add("book_date later than placement_date", raw_sub["book_after_placement"].sum(), "Excluded from baseline; check manually")
    add("Special/development institution issuer: ВЭБ.РФ", raw_sub["is_special_issuer"].sum(), "Excluded from baseline; kept for full robustness")
    add("num_organizers = 0", raw_sub["zero_organizers"].sum(), "Excluded from baseline")
    add("Debut issue with prior tightening flag", ((raw_sub["full_issue_number"] == 1) & (raw_sub["hist_ever_reduced"] == 1)).sum(), "Should normally be zero; check history coding if positive")
    add("placement_vol_book missing", raw_sub["placement_vol_book"].isna().sum(), "Excluded only from volume/upsize models")
    add("placement_vol_book > threshold", (raw_sub["placement_vol_book"] > UPSIZE_THRESHOLD).sum(), "Upsize cases")
    add("placement_vol_book around 1", raw_sub["placement_vol_book"].between(0.99, 1.01, inclusive="both").sum(), "No material volume change")
    add("placement_vol_book below 0.99", (raw_sub["placement_vol_book"] < 0.99).sum(), "Placed below teaser volume")
    if "volume_increased_dummy" in raw_sub.columns:
        add("Old volume_increased_dummy = 1", pd.to_numeric(raw_sub["volume_increased_dummy"], errors="coerce").fillna(0).sum(), "Do not use old dummy for new upsize logic")

    for col, label in [("coupon_reduction_bp", "Coupon reduction p01/p99"), ("placement_vol_book", "placement_vol_book p01/p99"), ("rvi", "RVI p99/max")]:
        if col in raw_sub.columns:
            s = pd.to_numeric(raw_sub[col], errors="coerce").dropna()
            if len(s):
                if col == "rvi":
                    action = f"p99={s.quantile(.99):.2f}, max={s.max():.2f}; do not auto-drop crisis dates"
                else:
                    action = f"p01={s.quantile(.01):.2f}, p99={s.quantile(.99):.2f}; winsorize only as robustness"
                add(label, len(s), action)

    if "term" in raw_sub.columns:
        add("term < 30 days", (raw_sub["term"] < 30).sum(), "Review as possible technical placements")
        add("term > 3650 days", (raw_sub["term"] > 3650).sum(), "Kept; log_term dampens tail")

    sample_df = pd.DataFrame([
        {"sample": "Full cleaned sample", "N": len(full_sub), "coupon_N": full_sub["coupon_reduction_bp"].notna().sum(), "volume_N": full_sub["placement_vol_book"].notna().sum(), "joint_N": (full_sub["coupon_reduction_bp"].notna() & full_sub["placement_vol_book"].notna()).sum(), "issuers": full_sub["issuer"].nunique()},
        {"sample": "Baseline sample", "N": len(baseline_sub), "coupon_N": baseline_sub["coupon_reduction_bp"].notna().sum(), "volume_N": baseline_sub["placement_vol_book"].notna().sum(), "joint_N": (baseline_sub["coupon_reduction_bp"].notna() & baseline_sub["placement_vol_book"].notna()).sum(), "issuers": baseline_sub["issuer"].nunique()},
    ])

    anomaly_cols = [c for c in ["issuer", "bond_name", "ISIN", "placement_date", "book_date", "final_rating", "num_organizers", "placement_vol_book", "coupon_reduction_bp"] if c in raw_sub.columns]
    anomalies_df = raw_sub.loc[raw_sub["book_after_placement"], anomaly_cols].copy()
    return pd.DataFrame(rows), sample_df, anomalies_df


# ============================================================
# TESTS
# ============================================================

def test_vif(sub: pd.DataFrame) -> pd.DataFrame:
    X = sub[EXPLANATORY_VARS].apply(pd.to_numeric, errors="coerce").astype(float)
    X = sm.add_constant(X, has_constant="add")
    rows = []
    for i, var in enumerate(X.columns):
        if var == "const":
            continue
        vif = variance_inflation_factor(X.values, i)
        rows.append({"variable": var, "label": VAR_LABELS.get(var, var), "VIF": round(vif, 2), "status": "ВЫСОКИЙ" if vif > 10 else "умеренный" if vif > 5 else "ОК"})
    return pd.DataFrame(rows)


def regressor_group(var: str) -> str:
    if var == "const":
        return "constant"
    if var.startswith("year_") or var.startswith("q_"):
        return "time fixed effect"
    if var.startswith("rating_bucket_"):
        return "rating-bucket fixed effect"
    if var.startswith("industry_"):
        return "industry fixed effect"
    if var in {"si", "log_buzz", *MAIN_INTERACTION_VARS}:
        return "sentiment / attention"
    if var in BASE_VARS:
        return "control"
    return "other"


def test_vif_with_fe(sub: pd.DataFrame) -> pd.DataFrame:
    d = sub.dropna(subset=EXPLANATORY_VARS).copy()
    X = build_X(d, EXPLANATORY_VARS, add_fe=True)
    rows = []
    for i, var in enumerate(X.columns):
        if var == "const":
            continue
        try:
            vif = variance_inflation_factor(X.values, i)
        except Exception:
            vif = np.nan
        rows.append({
            "variable": var,
            "label": VAR_LABELS.get(var, var),
            "group": regressor_group(var),
            "VIF": round(float(vif), 2) if pd.notna(vif) else np.nan,
            "status": "ВЫСОКИЙ" if pd.notna(vif) and vif > 10 else "умеренный" if pd.notna(vif) and vif > 5 else "ОК",
        })
    return pd.DataFrame(rows)


def test_correlations(sub: pd.DataFrame):
    vars_corr = ["coupon_reduction_bp", "coupon_reduced_dummy", "ihs_coupon_reduction", "ihs_coupon_positive", "placement_vol_book", "log_placement_vol_book", "upsize_dummy", "joint_success_upsize"] + EXPLANATORY_VARS
    vars_corr = [v for v in vars_corr if v in sub.columns]
    corr = sub[vars_corr].corr()
    rows = []
    for i in range(len(corr.columns)):
        for j in range(i + 1, len(corr.columns)):
            r = corr.iloc[i, j]
            if pd.notna(r) and abs(r) > 0.3:
                rows.append({"var1": corr.columns[i], "var2": corr.columns[j], "correlation": round(r, 3), "status": "ВЫСОКАЯ" if abs(r) > 0.7 else "умеренная"})
    return corr, pd.DataFrame(rows)


def test_heteroskedasticity(sub: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for y_col, label in OLS_OUTCOMES:
        d = sub.dropna(subset=[y_col] + EXPLANATORY_VARS).copy()
        X = build_X(d, EXPLANATORY_VARS, add_fe=True)
        ols = sm.OLS(d[y_col], X).fit()
        bp_stat, bp_p, _, _ = het_breuschpagan(ols.resid, X)
        rows.append({"dependent": y_col, "label": label, "N": len(d), "BP_statistic": round(bp_stat, 2), "BP_p_value": round(bp_p, 4), "heteroskedasticity": "Обнаружена" if bp_p < 0.05 else "Не обнаружена", "solution": "HC3 / clustered SE" if bp_p < 0.05 else "—"})
    return pd.DataFrame(rows)


def test_normality(sub: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for y_col, label in OLS_OUTCOMES:
        d = sub.dropna(subset=[y_col] + EXPLANATORY_VARS).copy()
        X = build_X(d, EXPLANATORY_VARS, add_fe=True)
        ols = sm.OLS(d[y_col], X).fit()
        jb_stat, jb_p, skew, kurt = jarque_bera(ols.resid)
        sw_stat, sw_p = stats.shapiro(ols.resid[:5000]) if len(ols.resid) > 5000 else stats.shapiro(ols.resid)
        rows.append({"dependent": y_col, "label": label, "N": len(d), "JB_statistic": round(jb_stat, 1), "JB_p_value": round(jb_p, 6), "skewness": round(skew, 2), "kurtosis": round(kurt, 2), "SW_statistic": round(sw_stat, 4), "SW_p_value": round(sw_p, 6), "normality": "Отвергается" if jb_p < 0.05 else "Не отвергается", "note": "Not critical for OLS with robust/clustered SE; inspect outliers/robustness"})
    return pd.DataFrame(rows)


def test_reset(sub: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for y_col, label in OLS_OUTCOMES:
        for spec_name, extra_vars in RESET_SPECS:
            variables = list(dict.fromkeys(EXPLANATORY_VARS + extra_vars))
            d = sub.dropna(subset=[y_col] + variables).copy()
            x_vars = [v for v in variables if d[v].nunique(dropna=False) > 1]
            if len(d) < len(x_vars) + 5:
                rows.append({
                    "dependent": y_col,
                    "label": label,
                    "reset_spec": spec_name,
                    "N": len(d),
                    "k_variables": len(x_vars),
                    "F_statistic": np.nan,
                    "p_value": np.nan,
                    "specification": "Too few observations",
                    "note": "",
                })
                continue
            X = build_X(d, x_vars, add_fe=True)
            ols = sm.OLS(d[y_col], X).fit()
            X_reset = X.copy()
            X_reset["y_hat2"] = ols.fittedvalues ** 2
            X_reset["y_hat3"] = ols.fittedvalues ** 3
            res2 = sm.OLS(d[y_col], X_reset).fit()
            r_matrix = np.zeros((2, len(res2.params)))
            r_matrix[0, -2] = 1
            r_matrix[1, -1] = 1
            f_test = res2.f_test(r_matrix)
            p_value = float(f_test.pvalue)
            rows.append({
                "dependent": y_col,
                "label": label,
                "reset_spec": spec_name,
                "N": len(d),
                "k_variables": len(x_vars),
                "F_statistic": round(float(f_test.fvalue), 2),
                "p_value": round(p_value, 4),
                "specification": "Проверить спецификацию" if p_value < 0.05 else "Адекватна",
                "note": "RESET still rejects" if p_value < 0.05 else "RESET no longer rejects",
            })
    return pd.DataFrame(rows)


def test_influence(sub: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for y_col, label in OLS_OUTCOMES:
        d = sub.dropna(subset=[y_col] + EXPLANATORY_VARS).copy()
        X = build_X(d, EXPLANATORY_VARS, add_fe=True)
        ols = sm.OLS(d[y_col], X).fit()
        cooks = ols.get_influence().cooks_distance[0]
        threshold = 4 / len(d)
        n_inf = int((cooks > threshold).sum())
        rows.append({"dependent": y_col, "label": label, "N": len(d), "threshold_4_N": round(threshold, 4), "n_influential": n_inf, "pct_influential": round(n_inf / len(d) * 100, 1), "max_cooks_d": round(float(np.nanmax(cooks)), 4), "status": "Много" if n_inf > len(d) * 0.05 else "В норме"})
    return pd.DataFrame(rows)


def test_extension_terms(sub: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for y_col, label in OLS_OUTCOMES:
        for term, term_type, expected_sign, hypothesis in EXTENSION_TERMS:
            variables = EXPLANATORY_VARS + [term]
            d = sub.dropna(subset=[y_col] + variables).copy()
            x_vars = [v for v in variables if d[v].nunique(dropna=False) > 1]
            if len(d) < len(x_vars) + 5 or term not in x_vars:
                rows.append({
                    "dependent": y_col,
                    "label": label,
                    "term": term,
                    "term_label": VAR_LABELS.get(term, term),
                    "type": term_type,
                    "expected_sign": expected_sign,
                    "hypothesis": hypothesis,
                    "N": len(d),
                    "coef": np.nan,
                    "std_error": np.nan,
                    "p_value": np.nan,
                    "evidence": "too few observations or constant term",
                })
                continue
            X = build_X(d, x_vars, add_fe=True)
            res = sm.OLS(d[y_col], X).fit(cov_type="HC3")
            coef = float(res.params.get(term, np.nan))
            se = float(res.bse.get(term, np.nan))
            p_value = float(res.pvalues.get(term, np.nan))
            rows.append({
                "dependent": y_col,
                "label": label,
                "term": term,
                "term_label": VAR_LABELS.get(term, term),
                "type": term_type,
                "expected_sign": expected_sign,
                "hypothesis": hypothesis,
                "N": int(res.nobs),
                "coef": round(coef, 6),
                "std_error": round(se, 6),
                "p_value": round(p_value, 6),
                "evidence": evidence_label(coef, p_value, expected_sign),
            })
    return pd.DataFrame(rows)


def build_robustness_report(full_sub: pd.DataFrame, baseline_sub: pd.DataFrame) -> pd.DataFrame:
    rows = []
    samples = [
        ("Baseline", baseline_sub),
        ("Full cleaned sample", full_sub),
        ("No ВЭБ.РФ only", full_sub[~full_sub["is_special_issuer"]]),
        ("No zero organizers only", full_sub[~full_sub["zero_organizers"]]),
        ("Rating A+ to BB only", baseline_sub[baseline_sub["rating_num"].between(4, 11)]),
        ("Baseline without top 1% Cook D volume", remove_top_cooks(baseline_sub, "log_placement_vol_book")),
    ]
    for name, data in samples:
        rows.append(fit_ols_summary(name, "coupon_reduction_bp", data, "coupon_reduction_bp"))
        rows.append(fit_ols_summary(name, "ihs_coupon_reduction", data, "ihs_coupon_reduction"))
        rows.append(fit_ols_summary(name, "ihs_coupon_positive", data, "ihs_coupon_positive"))
        rows.append(fit_ols_summary(name, "log_placement_vol_book", data, "log_placement_vol_book"))
    for name, base_vars in make_control_robustness_specs().items():
        variables = make_explanatory_vars(base_vars)
        rows.append(fit_ols_summary(name, "coupon_reduction_bp", baseline_sub, "coupon_reduction_bp", variables))
        rows.append(fit_ols_summary(name, "ihs_coupon_reduction", baseline_sub, "ihs_coupon_reduction", variables))
        rows.append(fit_ols_summary(name, "ihs_coupon_positive", baseline_sub, "ihs_coupon_positive", variables))
        rows.append(fit_ols_summary(name, "log_placement_vol_book", baseline_sub, "log_placement_vol_book", variables))
    d = baseline_sub.dropna(subset=["coupon_reduction_bp"] + EXPLANATORY_VARS).copy()
    if len(d):
        d["coupon_reduction_bp_w"] = winsorize_series(d["coupon_reduction_bp"], WINSOR_LOWER, WINSOR_UPPER)
        rows.append(fit_ols_summary("Baseline winsorized 1/99%", "coupon_reduction_bp_w", d, "coupon_reduction_bp_w"))
    d = baseline_sub.dropna(subset=["placement_vol_book"] + EXPLANATORY_VARS).copy()
    if len(d):
        d["placement_vol_book_w"] = winsorize_series(d["placement_vol_book"], WINSOR_LOWER, WINSOR_UPPER)
        rows.append(fit_ols_summary("Baseline winsorized 1/99%", "placement_vol_book_w", d, "placement_vol_book_w"))
    return pd.DataFrame(rows)


# ============================================================
# DISTRIBUTION PLOTS
# ============================================================

def safe_plot_filename(name: str) -> str:
    keep = []
    for ch in name:
        keep.append(ch if ch.isalnum() or ch in {"_", "-"} else "_")
    return "".join(keep).strip("_") or "plot"


def distribution_plot_dir(output_file: str) -> Path:
    out = Path(output_file)
    parent = out.parent if str(out.parent) else Path(".")
    return parent / f"{out.stem}_plots"


def plot_numeric_distribution(data: pd.DataFrame, var: str, path: Path) -> dict:
    s = pd.to_numeric(data[var], errors="coerce")
    clean = s.dropna()
    stats_row = {
        "sample": "Baseline",
        "variable": var,
        "label": VAR_LABELS.get(var, var),
        "type": "numeric",
        "N": int(clean.size),
        "missing": int(s.isna().sum()),
        "mean": np.nan,
        "std": np.nan,
        "min": np.nan,
        "p01": np.nan,
        "p25": np.nan,
        "median": np.nan,
        "p75": np.nan,
        "p99": np.nan,
        "max": np.nan,
        "skew": np.nan,
        "kurtosis": np.nan,
        "top_value": "",
        "top_count": np.nan,
        "plot_file": str(path),
    }
    if clean.empty:
        return stats_row

    stats_row.update({
        "mean": round(float(clean.mean()), 6),
        "std": round(float(clean.std()), 6),
        "min": round(float(clean.min()), 6),
        "p01": round(float(clean.quantile(0.01)), 6),
        "p25": round(float(clean.quantile(0.25)), 6),
        "median": round(float(clean.median()), 6),
        "p75": round(float(clean.quantile(0.75)), 6),
        "p99": round(float(clean.quantile(0.99)), 6),
        "max": round(float(clean.max()), 6),
        "skew": round(float(clean.skew()), 6),
        "kurtosis": round(float(clean.kurtosis()), 6),
    })

    fig, axes = plt.subplots(1, 2, figsize=(11, 4), gridspec_kw={"width_ratios": [3, 1]})
    unique_count = clean.nunique(dropna=True)
    if unique_count <= 10:
        counts = clean.value_counts().sort_index()
        axes[0].bar([str(x) for x in counts.index], counts.values, color="#4C78A8")
        axes[0].set_xlabel(VAR_LABELS.get(var, var))
        axes[0].set_ylabel("Count")
        axes[0].tick_params(axis="x", rotation=45)
    else:
        axes[0].hist(clean, bins=40, color="#4C78A8", edgecolor="white", alpha=0.9)
        axes[0].axvline(clean.median(), color="#F58518", linewidth=2, label="median")
        axes[0].axvline(clean.mean(), color="#54A24B", linewidth=2, label="mean")
        axes[0].legend()
        axes[0].set_xlabel(VAR_LABELS.get(var, var))
        axes[0].set_ylabel("Count")

    try:
        axes[1].boxplot(clean, orientation="vertical", patch_artist=True, boxprops={"facecolor": "#A0CBE8"})
    except TypeError:
        axes[1].boxplot(clean, vert=True, patch_artist=True, boxprops={"facecolor": "#A0CBE8"})
    axes[1].set_xticks([])
    axes[1].set_ylabel(VAR_LABELS.get(var, var))
    fig.suptitle(f"{VAR_LABELS.get(var, var)} | N={len(clean)}, missing={stats_row['missing']}")
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return stats_row


def plot_categorical_distribution(data: pd.DataFrame, var: str, path: Path) -> dict:
    s = data[var].fillna("Missing").astype(str)
    counts = s.value_counts()
    top = counts.head(25).sort_values()
    stats_row = {
        "sample": "Baseline",
        "variable": var,
        "label": VAR_LABELS.get(var, var),
        "type": "categorical",
        "N": int(s.size),
        "missing": int(data[var].isna().sum()),
        "mean": np.nan,
        "std": np.nan,
        "min": np.nan,
        "p01": np.nan,
        "p25": np.nan,
        "median": np.nan,
        "p75": np.nan,
        "p99": np.nan,
        "max": np.nan,
        "skew": np.nan,
        "kurtosis": np.nan,
        "top_value": counts.index[0] if len(counts) else "",
        "top_count": int(counts.iloc[0]) if len(counts) else np.nan,
        "plot_file": str(path),
    }

    fig, ax = plt.subplots(figsize=(10, max(4, 0.28 * len(top))))
    ax.barh(top.index, top.values, color="#4C78A8")
    ax.set_xlabel("Count")
    ax.set_title(f"{VAR_LABELS.get(var, var)} | categories={counts.size}, missing={stats_row['missing']}")
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return stats_row


def build_distribution_plots(baseline_sub: pd.DataFrame, output_file: str) -> pd.DataFrame:
    rows = []
    plot_dir = distribution_plot_dir(output_file)

    if not HAS_MATPLOTLIB:
        return pd.DataFrame([{
            "sample": "Baseline",
            "variable": "",
            "label": "matplotlib unavailable",
            "type": "error",
            "N": 0,
            "missing": 0,
            "plot_file": "",
        }])

    plot_dir.mkdir(parents=True, exist_ok=True)

    for var in DISTRIBUTION_NUMERIC_VARS:
        if var not in baseline_sub.columns:
            continue
        path = plot_dir / f"{safe_plot_filename(var)}.png"
        rows.append(plot_numeric_distribution(baseline_sub, var, path))

    for var in DISTRIBUTION_CATEGORICAL_VARS:
        if var not in baseline_sub.columns:
            continue
        path = plot_dir / f"{safe_plot_filename(var)}.png"
        rows.append(plot_categorical_distribution(baseline_sub, var, path))

    return pd.DataFrame(rows)


# ============================================================
# EXCEL
# ============================================================

def write_dataframe_sheet(wb: Workbook, sheet_name: str, df_data: pd.DataFrame) -> None:
    hf = Font(bold=True, size=10, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    ws = wb.create_sheet(sheet_name)
    ws["A1"] = sheet_name
    ws["A1"].font = Font(bold=True, size=13, name="Arial")
    row = 3
    if df_data.empty:
        ws.cell(row=row, column=1, value="No rows")
        return
    for j, h in enumerate(df_data.columns, start=1):
        ws.cell(row=row, column=j, value=h).font = hf
        ws.cell(row=row, column=j).fill = hfill
    row += 1
    for _, r in df_data.iterrows():
        for j, col in enumerate(df_data.columns, start=1):
            val = r[col]
            ws.cell(row=row, column=j, value=float(val) if isinstance(val, np.floating) else int(val) if isinstance(val, np.integer) else val)
        row += 1
    for i in range(1, len(df_data.columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24
    ws.freeze_panes = "A4"


def build_excel(path: str, vif_df, vif_fe_df, corr_matrix, high_corr_df, het_df, norm_df, reset_df, infl_df, quality_df, sample_df, anomalies_df, robustness_df, extensions_df, distributions_df, n_base, n_full):
    wb = Workbook()
    hf = Font(bold=True, size=10, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    ok_fill = PatternFill("solid", fgColor="E2EFDA")
    warn_fill = PatternFill("solid", fgColor="FCE4EC")

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Diagnostic Tests Summary (Baseline N={n_base}, Full cleaned N={n_full}, Period: {YEAR_FROM}+)"
    ws["A1"].font = Font(bold=True, size=14, name="Arial")
    ws.merge_cells("A1:E1")
    summary_rows = [
        ("1. Multicollinearity (VIF)", f"Max VIF = {vif_df['VIF'].max():.2f}", "ОК" if vif_df["VIF"].max() < 5 else "Review"),
        ("2. Heteroskedasticity (BP)", f"min p = {het_df['BP_p_value'].min():.4f}", "Use HC3 / clustered SE"),
        ("3. Normality (JB)", f"max p = {norm_df['JB_p_value'].max():.6f}", "Not critical for OLS; use robustness"),
        ("4. Specification (RESET)", f"min p = {reset_df['p_value'].min():.4f}", "Check nonlinearities if significant"),
        ("5. Influential obs (Cook)", f"max Cook D = {infl_df['max_cooks_d'].max():.4f}", "Report robustness"),
    ]
    row = 3
    for test, stat, verdict in summary_rows:
        ws.cell(row=row, column=1, value=test).font = Font(name="Arial", size=11, bold=True)
        ws.cell(row=row, column=2, value=stat).font = Font(name="Arial", size=11)
        c = ws.cell(row=row, column=3, value=verdict)
        c.font = Font(name="Arial", size=11)
        c.fill = ok_fill if "ОК" in verdict or "Use" in verdict else warn_fill
        row += 1
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 44

    write_dataframe_sheet(wb, "Sample Filters", sample_df)
    write_dataframe_sheet(wb, "Data Quality", quality_df)
    write_dataframe_sheet(wb, "Date Anomalies", anomalies_df)
    write_dataframe_sheet(wb, "Distributions", distributions_df)
    write_dataframe_sheet(wb, "Robustness", robustness_df)
    write_dataframe_sheet(wb, "Nonlinear Interactions", extensions_df)
    write_dataframe_sheet(wb, "VIF with FE", vif_fe_df)

    ws_vif = wb.create_sheet("VIF")
    ws_vif["A1"] = "Variance Inflation Factors"
    ws_vif["A1"].font = Font(bold=True, size=13, name="Arial")
    row = 3
    for j, h in enumerate(["Variable", "VIF", "Status"], start=1):
        ws_vif.cell(row=row, column=j, value=h).font = hf
        ws_vif.cell(row=row, column=j).fill = hfill
    row += 1
    for _, r in vif_df.iterrows():
        ws_vif.cell(row=row, column=1, value=r["label"])
        ws_vif.cell(row=row, column=2, value=r["VIF"]).number_format = "0.00"
        ws_vif.cell(row=row, column=3, value=r["status"]).fill = ok_fill if r["status"] == "ОК" else warn_fill
        row += 1
    ws_vif.column_dimensions["A"].width = 34

    ws_corr = wb.create_sheet("Correlations")
    ws_corr["A1"] = "Correlation Matrix"
    ws_corr["A1"].font = Font(bold=True, size=13, name="Arial")
    cols = list(corr_matrix.columns)
    row = 3
    for j, col in enumerate(cols, start=2):
        c = ws_corr.cell(row=row, column=j, value=VAR_LABELS.get(col, col))
        c.font = Font(bold=True, size=8, name="Arial")
        c.fill = hfill
        c.alignment = Alignment(text_rotation=60, wrap_text=True)
    row += 1
    for var in cols:
        ws_corr.cell(row=row, column=1, value=VAR_LABELS.get(var, var)).fill = hfill
        for j, var2 in enumerate(cols, start=2):
            val = corr_matrix.loc[var, var2]
            c = ws_corr.cell(row=row, column=j, value=round(float(val), 3) if pd.notna(val) else "—")
            if var != var2 and pd.notna(val) and abs(val) > 0.5:
                c.fill = warn_fill
                c.font = Font(bold=True, size=9)
        row += 1
    ws_corr.column_dimensions["A"].width = 30

    write_dataframe_sheet(wb, "High Correlations", high_corr_df)
    write_dataframe_sheet(wb, "Heteroskedasticity", het_df)
    write_dataframe_sheet(wb, "Normality", norm_df)
    write_dataframe_sheet(wb, "RESET", reset_df)
    write_dataframe_sheet(wb, "Influence", infl_df)

    wb.save(path)


# ============================================================
# CONSOLE
# ============================================================

def print_results(vif_df, high_corr_df, het_df, norm_df, reset_df, infl_df, quality_df, sample_df, robustness_df, extensions_df, distributions_df, n_base, n_full):
    print("\n" + "=" * 70)
    print(f"ДИАГНОСТИЧЕСКИЕ ТЕСТЫ (baseline N={n_base}, full cleaned N={n_full}, period {YEAR_FROM}+)")
    print("=" * 70)
    print("\n0. SAMPLE / DATA QUALITY")
    for _, r in sample_df.iterrows():
        print(f"  {r['sample']:25s} N={int(r['N'])}, coupon_N={int(r['coupon_N'])}, volume_N={int(r['volume_N'])}, joint_N={int(r['joint_N'])}, issuers={int(r['issuers'])}")
    for _, r in quality_df.iterrows():
        if r["N"] > 0:
            print(f"  {r['issue']:45s} N={int(r['N'])} → {r['action']}")
    if not distributions_df.empty and "plot_file" in distributions_df.columns:
        plot_dir = Path(str(distributions_df["plot_file"].dropna().iloc[0])).parent if distributions_df["plot_file"].dropna().size else None
        if plot_dir:
            print(f"  Distribution plots: {plot_dir}")

    print("\n1. VIF")
    for _, r in vif_df.iterrows():
        print(f"  {r['label']:35s} {r['VIF']:6.2f} {r['status']}")
    print(f"  Max VIF: {vif_df['VIF'].max():.2f}")

    print("\n2. HIGH CORRELATIONS |r| > 0.3")
    if high_corr_df.empty:
        print("  none")
    else:
        for _, r in high_corr_df.iterrows():
            print(f"  {VAR_LABELS.get(r['var1'], r['var1']):30s} × {VAR_LABELS.get(r['var2'], r['var2']):30s} = {r['correlation']:.3f}")

    print("\n3. HETEROSKEDASTICITY")
    for _, r in het_df.iterrows():
        print(f"  {r['dependent']:25s} N={int(r['N'])}, BP={r['BP_statistic']:.2f}, p={r['BP_p_value']:.4f} → {r['solution']}")

    print("\n4. NORMALITY")
    for _, r in norm_df.iterrows():
        print(f"  {r['dependent']:25s} N={int(r['N'])}, JB={r['JB_statistic']:.1f}, skew={r['skewness']:.2f}, kurt={r['kurtosis']:.2f}")

    print("\n5. RESET")
    for _, r in reset_df.iterrows():
        spec = r["reset_spec"] if "reset_spec" in r.index else "Baseline"
        print(f"  {r['dependent']:25s} {spec:38s} N={int(r['N'])}, F={r['F_statistic']:.2f}, p={r['p_value']:.4f}, {r['specification']}")

    print("\n6. COOK'S D")
    for _, r in infl_df.iterrows():
        print(f"  {r['dependent']:25s} N={int(r['N'])}, influential={int(r['n_influential'])} ({r['pct_influential']}%), max={r['max_cooks_d']:.4f}")

    print("\n7. ROBUSTNESS SNAPSHOT")
    for _, r in robustness_df.iterrows():
        print(f"  {r['sample']:35s} {r['dependent']:28s} N={int(r['N']):4d}, SI={r['si_coef']}, p={r['si_p']}, buzz={r['log_buzz_coef']}, p={r['log_buzz_p']}")

    print("\n8. NONLINEAR / INTERACTION TERMS (p < 0.10)")
    sig = extensions_df[pd.to_numeric(extensions_df["p_value"], errors="coerce") < 0.10]
    if sig.empty:
        print("  none")
    else:
        for _, r in sig.iterrows():
            print(f"  {r['dependent']:25s} {r['term_label']:28s} coef={r['coef']}, p={r['p_value']}, {r['evidence']}")


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    input_file = pick_input_file()
    output_file = pick_output_file()
    print(f"Input:  {input_file}")
    print(f"Output: {output_file}")
    print("Загрузка данных...")
    full_sub, baseline_sub, quality_df, sample_df, anomalies_df = load_data(input_file)
    n_base = len(baseline_sub)
    n_full = len(full_sub)

    print("Запуск тестов...")
    distributions_df = build_distribution_plots(baseline_sub, output_file)
    vif_df = test_vif(baseline_sub)
    vif_fe_df = test_vif_with_fe(baseline_sub)
    corr_matrix, high_corr_df = test_correlations(baseline_sub)
    het_df = test_heteroskedasticity(baseline_sub)
    norm_df = test_normality(baseline_sub)
    reset_df = test_reset(baseline_sub)
    infl_df = test_influence(baseline_sub)
    robustness_df = build_robustness_report(full_sub, baseline_sub)
    extensions_df = test_extension_terms(baseline_sub)

    print_results(vif_df, high_corr_df, het_df, norm_df, reset_df, infl_df, quality_df, sample_df, robustness_df, extensions_df, distributions_df, n_base, n_full)
    build_excel(output_file, vif_df, vif_fe_df, corr_matrix, high_corr_df, het_df, norm_df, reset_df, infl_df, quality_df, sample_df, anomalies_df, robustness_df, extensions_df, distributions_df, n_base, n_full)
    print(f"\n✅ Результаты сохранены: {output_file}")


if __name__ == "__main__":
    main()
