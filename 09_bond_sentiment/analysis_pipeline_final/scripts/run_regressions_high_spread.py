#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
High-spread / high-yield proxy heterogeneity checks.

This script does not rebuild the core pipeline. It reads an existing
regression_ready.csv, adds high-spread proxy indicators from teaser terms, and
estimates:
  1) interaction models on the baseline sample;
  2) baseline sentiment models on high-spread-only subsamples.

Definitions:
  - fixed coupon bonds: teaser coupon - matched OFZ yield >= threshold pp
  - floaters: teaser spread >= threshold bp
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import numpy as np
import pandas as pd
from statsmodels.stats.outliers_influence import variance_inflation_factor
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import run_regressions_upsize as reg


DEFAULT_INPUT = "regression_dataset_book_open_5/regression_ready.csv"
DEFAULT_OUTPUT = "regression_results_high_spread.xlsx"

THRESHOLDS = [
    {"suffix": "4", "label": ">=4 pp / >=400 bp", "fixed_pp": 4.0, "floater_bp": 400.0},
    {"suffix": "5", "label": ">=5 pp / >=500 bp", "fixed_pp": 5.0, "floater_bp": 500.0},
]

OUTCOME_SPECS = [
    ("coupon_ols", "coupon_reduction_bp", "Coupon reduction, bp", "OLS cluster", "continuous"),
    ("coupon_ihs", "ihs_coupon_reduction", "IHS coupon reduction", "OLS cluster", "continuous"),
    ("coupon_logit", "coupon_reduced_dummy", "Coupon reduction dummy", "Logit/GLM cluster", "binary"),
    ("log_volume_ols", "log_placement_vol_book", "ln(volume / teaser)", "OLS cluster", "continuous"),
    ("upsize_logit", "upsize_dummy", "Upsize dummy", "Logit/GLM cluster", "binary"),
    ("joint_logit", "joint_success_upsize", "Strict dual success", "Logit/GLM cluster", "binary"),
]

BASE_INTEREST_VARS = ["si", "log_buzz"]


def sig_stars(p_value: Optional[float]) -> str:
    return reg.sig_stars(p_value)


def first_existing_col(df: pd.DataFrame, candidates: Iterable[str]) -> str:
    for col in candidates:
        if col in df.columns:
            return col
    raise KeyError(f"None of these columns exists: {', '.join(candidates)}")


def add_high_spread_terms(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    teaser_col = first_existing_col(out, ["teaser_value", "tizer_value"])
    ofz_col = first_existing_col(out, ["ofz_matched_yield", "ofz_yield"])

    out["teaser_value_for_hs"] = pd.to_numeric(out[teaser_col], errors="coerce")
    out["ofz_yield_for_hs"] = pd.to_numeric(out[ofz_col], errors="coerce")
    out["is_floater"] = pd.to_numeric(out["is_floater"], errors="coerce")
    out["fixed_teaser_spread_to_ofz_pp"] = np.where(
        out["is_floater"].eq(0),
        out["teaser_value_for_hs"] - out["ofz_yield_for_hs"],
        np.nan,
    )
    out["floater_teaser_spread_bp"] = np.where(
        out["is_floater"].eq(1),
        out["teaser_value_for_hs"],
        np.nan,
    )

    for item in THRESHOLDS:
        suffix = item["suffix"]
        hs = f"high_spread_{suffix}"
        si_x = f"si_x_high_spread_{suffix}"
        buzz_x = f"log_buzz_x_high_spread_{suffix}"

        fixed = out["is_floater"].eq(0) & out["fixed_teaser_spread_to_ofz_pp"].ge(item["fixed_pp"])
        floater = out["is_floater"].eq(1) & out["floater_teaser_spread_bp"].ge(item["floater_bp"])
        out[hs] = (fixed | floater).astype(int)
        out[si_x] = out["si"] * out[hs]
        out[buzz_x] = out["log_buzz"] * out[hs]

        reg.VAR_LABELS[hs] = f"High-spread proxy ({item['label']})"
        reg.VAR_LABELS[si_x] = f"SI x High-spread ({item['label']})"
        reg.VAR_LABELS[buzz_x] = f"ln(Buzz) x High-spread ({item['label']})"

    reg.VAR_LABELS["teaser_value_for_hs"] = "Teaser coupon/rate value"
    reg.VAR_LABELS["fixed_teaser_spread_to_ofz_pp"] = "Fixed teaser spread to OFZ, pp"
    reg.VAR_LABELS["floater_teaser_spread_bp"] = "Floater teaser spread, bp"
    return out


def model_variables_for_interaction(suffix: str) -> List[str]:
    return list(dict.fromkeys(
        reg.BASE_VARS
        + ["si", "log_buzz", f"high_spread_{suffix}", f"si_x_high_spread_{suffix}", f"log_buzz_x_high_spread_{suffix}"]
    ))


def model_variables_for_subsample() -> List[str]:
    return list(dict.fromkeys(reg.BASE_VARS + ["si", "log_buzz"]))


def fit_one_model(df: pd.DataFrame, y_col: str, variables: List[str], outcome_type: str) -> Tuple[Any, List[str], int]:
    d = reg.subset_for_model(df, y_col, variables)
    x_vars = [v for v in variables if v in d.columns and d[v].nunique(dropna=False) > 1]
    X = reg.build_X(d, x_vars, add_fe=True)
    y = d[y_col].astype(float)
    if outcome_type == "binary":
        res = reg.fit_glm_binomial(y, X, groups=d["issuer_id"])
    else:
        res = reg.fit_ols(y, X, cov_type="cluster", groups=d["issuer_id"])
    return res, list(X.columns), len(d)


def collect_model_rows(
    res: Any,
    colnames: List[str],
    variables: List[str],
    meta: Dict[str, Any],
) -> List[Dict[str, Any]]:
    rows = []
    for var in variables:
        coef, se, p_value = reg.get_csp(res, var, colnames)
        rows.append({
            **meta,
            "variable": var,
            "label": reg.VAR_LABELS.get(var, var),
            "coef": round(coef, 6) if coef is not None and pd.notna(coef) else np.nan,
            "std_error": round(se, 6) if se is not None and pd.notna(se) else np.nan,
            "p_value": round(p_value, 6) if p_value is not None and pd.notna(p_value) else np.nan,
            "stars": sig_stars(p_value) if p_value is not None else "",
        })
    return rows


def linear_combination(res: Any, terms: List[Tuple[str, float]]) -> Tuple[float, float, float]:
    params = res.params
    cov = res.cov_params()
    if not isinstance(cov, pd.DataFrame):
        cov = pd.DataFrame(cov, index=params.index, columns=params.index)

    coef = 0.0
    weights = pd.Series(0.0, index=params.index)
    for var, weight in terms:
        if var in params.index:
            coef += float(weight * params[var])
            weights.loc[var] = weight
    var = float(weights.to_numpy() @ cov.reindex(index=params.index, columns=params.index).fillna(0).to_numpy() @ weights.to_numpy())
    se = float(np.sqrt(max(var, 0.0)))
    z = coef / se if se > 0 else np.nan
    p_value = float(2 * (1 - reg.stats.norm.cdf(abs(z)))) if pd.notna(z) else np.nan
    return coef, se, p_value


def group_ame_for_logit(res: Any, X: pd.DataFrame, base_var: str, interaction_var: str, high_spread_col: str, group_value: int) -> float:
    if base_var not in X.columns or high_spread_col not in X.columns:
        return np.nan
    params = res.params.reindex(X.columns).fillna(0)
    beta = float(params[base_var])
    if interaction_var in X.columns:
        beta += group_value * float(params[interaction_var])
    mask = X[high_spread_col].round().astype(int).eq(group_value)
    if not mask.any():
        return np.nan
    p = pd.Series(res.predict(X), index=X.index).clip(1e-9, 1 - 1e-9)
    return float((p.loc[mask] * (1 - p.loc[mask]) * beta).mean())


def estimate_interactions(df: pd.DataFrame, threshold: Dict[str, Any]) -> Tuple[pd.DataFrame, pd.DataFrame]:
    suffix = threshold["suffix"]
    hs = f"high_spread_{suffix}"
    variables = model_variables_for_interaction(suffix)
    coef_rows: List[Dict[str, Any]] = []
    effect_rows: List[Dict[str, Any]] = []

    for model_key, y_col, y_label, model_type, outcome_type in OUTCOME_SPECS:
        res, colnames, n_model = fit_one_model(df, y_col, variables, outcome_type)
        meta = {
            "threshold": threshold["label"],
            "model_key": model_key,
            "dependent": y_col,
            "dependent_label": y_label,
            "model_type": model_type,
            "N": int(res.nobs) if hasattr(res, "nobs") else n_model,
        }
        interest = ["si", "log_buzz", hs, f"si_x_high_spread_{suffix}", f"log_buzz_x_high_spread_{suffix}"]
        coef_rows.extend(collect_model_rows(res, colnames, interest, meta))

        X = res.model.exog
        X = pd.DataFrame(X, columns=colnames)
        for base_var, int_var in [
            ("si", f"si_x_high_spread_{suffix}"),
            ("log_buzz", f"log_buzz_x_high_spread_{suffix}"),
        ]:
            for group_value in [0, 1]:
                terms = [(base_var, 1.0), (int_var, float(group_value))]
                coef, se, p_value = linear_combination(res, terms)
                row = {
                    **meta,
                    "effect_variable": base_var,
                    "high_spread": group_value,
                    "high_spread_label": "High-spread" if group_value else "Non-high-spread",
                    "linear_effect": round(coef, 6),
                    "linear_effect_se": round(se, 6) if pd.notna(se) else np.nan,
                    "linear_effect_p_value": round(p_value, 6) if pd.notna(p_value) else np.nan,
                    "linear_effect_stars": sig_stars(p_value),
                    "logit_AME": np.nan,
                    "note": "OLS coefficient for continuous models; logit index coefficient for binary models.",
                }
                if outcome_type == "binary":
                    ame = group_ame_for_logit(res, X, base_var, int_var, hs, group_value)
                    row["logit_AME"] = round(ame, 6) if pd.notna(ame) else np.nan
                    row["note"] = "For binary models, logit_AME is the average derivative within the group."
                effect_rows.append(row)

    return pd.DataFrame(coef_rows), pd.DataFrame(effect_rows)


def estimate_high_spread_subsample(df: pd.DataFrame, threshold: Dict[str, Any]) -> pd.DataFrame:
    suffix = threshold["suffix"]
    hs = f"high_spread_{suffix}"
    sub = df.loc[df[hs].eq(1)].copy()
    variables = model_variables_for_subsample()
    rows: List[Dict[str, Any]] = []

    for model_key, y_col, y_label, model_type, outcome_type in OUTCOME_SPECS:
        res, colnames, n_model = fit_one_model(sub, y_col, variables, outcome_type)
        meta = {
            "threshold": threshold["label"],
            "sample": "high_spread_only",
            "model_key": model_key,
            "dependent": y_col,
            "dependent_label": y_label,
            "model_type": model_type,
            "N": int(res.nobs) if hasattr(res, "nobs") else n_model,
        }
        rows.extend(collect_model_rows(res, colnames, ["si", "log_buzz"], meta))

    return pd.DataFrame(rows)


def sample_summary(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for threshold in THRESHOLDS:
        suffix = threshold["suffix"]
        hs = f"high_spread_{suffix}"
        sub = df.loc[df[hs].eq(1)]
        rows.append({
            "threshold": threshold["label"],
            "fixed_rule": f"teaser_value - OFZ >= {threshold['fixed_pp']:g} pp",
            "floater_rule": f"teaser_value >= {threshold['floater_bp']:g} bp",
            "baseline_N": len(df),
            "high_spread_N": len(sub),
            "non_high_spread_N": int(df[hs].eq(0).sum()),
            "fixed_high_spread_N": int(sub["is_floater"].eq(0).sum()),
            "floater_high_spread_N": int(sub["is_floater"].eq(1).sum()),
            "issuers_high_spread_N": int(sub["issuer"].nunique()) if "issuer" in sub.columns else np.nan,
            "coupon_N": int(sub["coupon_reduction_bp"].notna().sum()),
            "volume_N": int(sub["placement_vol_book"].notna().sum()),
            "joint_N": int((sub["coupon_reduction_bp"].notna() & sub["placement_vol_book"].notna()).sum()),
        })
    return pd.DataFrame(rows)


def compute_vif(X: pd.DataFrame) -> pd.DataFrame:
    X = X.replace([np.inf, -np.inf], np.nan).dropna(axis=0)
    constant_cols = [c for c in X.columns if X[c].nunique(dropna=True) <= 1]
    X = X.drop(columns=constant_cols, errors="ignore")
    rows = []
    if X.empty or len(X) <= X.shape[1]:
        return pd.DataFrame(columns=["variable", "VIF", "status"])

    for i, col in enumerate(X.columns):
        try:
            vif = float(variance_inflation_factor(X.to_numpy(dtype=float), i))
        except Exception:
            vif = np.nan
        if pd.isna(vif):
            status = "not estimated"
        elif vif >= 10:
            status = "high"
        elif vif >= 5:
            status = "moderate"
        else:
            status = "OK"
        rows.append({"variable": col, "VIF": round(vif, 4) if pd.notna(vif) else np.nan, "status": status})
    return pd.DataFrame(rows)


def multicollinearity_diagnostics(df: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    vif_rows = []
    corr_rows = []
    condition_rows = []

    for threshold in THRESHOLDS:
        suffix = threshold["suffix"]
        hs = f"high_spread_{suffix}"
        si_x = f"si_x_high_spread_{suffix}"
        buzz_x = f"log_buzz_x_high_spread_{suffix}"
        variables = model_variables_for_interaction(suffix)
        needed = variables + ["issuer_id", "time_fe"]
        d = df.dropna(subset=[c for c in needed if c in df.columns]).copy()

        for include_fe, label in [(False, "without FE"), (True, "with year FE")]:
            X = reg.build_X(d, variables, add_fe=include_fe)
            vif_df = compute_vif(X.drop(columns=["const"], errors="ignore"))
            if not vif_df.empty:
                vif_df.insert(0, "threshold", threshold["label"])
                vif_df.insert(1, "design", label)
                vif_df["label"] = vif_df["variable"].map(lambda v: reg.VAR_LABELS.get(v, v))
                vif_rows.append(vif_df)

            Xn = X.drop(columns=["const"], errors="ignore").replace([np.inf, -np.inf], np.nan).dropna()
            if not Xn.empty and len(Xn) > Xn.shape[1]:
                try:
                    condition_number = float(np.linalg.cond(Xn.to_numpy(dtype=float)))
                except Exception:
                    condition_number = np.nan
            else:
                condition_number = np.nan
            condition_rows.append({
                "threshold": threshold["label"],
                "design": label,
                "N": len(Xn),
                "k_regressors_ex_const": Xn.shape[1] if not Xn.empty else 0,
                "condition_number": round(condition_number, 4) if pd.notna(condition_number) else np.nan,
            })

        corr_vars = ["rating_num", hs, si_x, buzz_x, "si", "log_buzz"]
        corr_vars = [v for v in corr_vars if v in d.columns]
        corr = d[corr_vars].apply(pd.to_numeric, errors="coerce").corr()
        for left in corr_vars:
            for right in corr_vars:
                if left >= right:
                    continue
                corr_rows.append({
                    "threshold": threshold["label"],
                    "left": left,
                    "left_label": reg.VAR_LABELS.get(left, left),
                    "right": right,
                    "right_label": reg.VAR_LABELS.get(right, right),
                    "correlation": round(float(corr.loc[left, right]), 4) if pd.notna(corr.loc[left, right]) else np.nan,
                    "focus_pair": int("rating_num" in {left, right}),
                })

    vif_all = pd.concat(vif_rows, ignore_index=True) if vif_rows else pd.DataFrame()
    corr_all = pd.DataFrame(corr_rows)
    condition_all = pd.DataFrame(condition_rows)

    if not vif_all.empty:
        focus_prefixes = ("rating_num", "high_spread_", "si_x_high_spread_", "log_buzz_x_high_spread_", "si", "log_buzz")
        vif_all["focus_variable"] = vif_all["variable"].map(lambda v: int(v.startswith(focus_prefixes)))
        vif_all = vif_all.sort_values(["threshold", "design", "focus_variable", "VIF"], ascending=[True, True, False, False])
    if not corr_all.empty:
        corr_all = corr_all.sort_values(["threshold", "focus_pair", "correlation"], ascending=[True, False, False])
    return vif_all, corr_all, condition_all


def write_workbook(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, name="Arial")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 42)
    wb.save(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Estimate high-spread sentiment heterogeneity checks.")
    parser.add_argument("input", nargs="?", default=DEFAULT_INPUT, help="Path to regression_ready.csv.")
    parser.add_argument("output", nargs="?", default=DEFAULT_OUTPUT, help="Output xlsx path.")
    parser.add_argument("--year-from", type=int, default=reg.YEAR_FROM)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    bundle = reg.load_and_prepare(args.input, args.year_from)
    baseline = add_high_spread_terms(bundle.baseline)

    interaction_frames = []
    group_effect_frames = []
    subsample_frames = []
    for threshold in THRESHOLDS:
        coefs, effects = estimate_interactions(baseline, threshold)
        interaction_frames.append(coefs)
        group_effect_frames.append(effects)
        subsample_frames.append(estimate_high_spread_subsample(baseline, threshold))
    vif_df, corr_df, condition_df = multicollinearity_diagnostics(baseline)

    sheets = {
        "Summary": sample_summary(baseline),
        "Interaction Coefs": pd.concat(interaction_frames, ignore_index=True),
        "Group Effects": pd.concat(group_effect_frames, ignore_index=True),
        "HighSpread Subsample": pd.concat(subsample_frames, ignore_index=True),
        "VIF Diagnostics": vif_df,
        "Correlation Diagnostics": corr_df,
        "Condition Numbers": condition_df,
        "Variable Notes": pd.DataFrame([
            {
                "item": "Preferred interpretation",
                "note": "Use Interaction Coefs for significance of heterogeneity; use Group Effects for SI/log_buzz effects by high-spread group.",
            },
            {
                "item": "High-spread proxy",
                "note": "This is a spread-based high-yield proxy, not a pure rating-based VDO classification.",
            },
            {
                "item": "Rating control",
                "note": "rating_num remains in the control set, so interactions are conditional on formal credit rating.",
            },
            {
                "item": "Binary models",
                "note": "linear_effect is a logit-index coefficient; logit_AME is the within-group average derivative.",
            },
            {
                "item": "Multicollinearity diagnostics",
                "note": "VIF Diagnostics reports the full interaction design with and without year FE. Correlation Diagnostics highlights pairwise correlations involving rating_num.",
            },
        ]),
    }
    write_workbook(Path(args.output), sheets)
    print(f"Saved: {args.output}")
    print(sheets["Summary"].to_string(index=False))


if __name__ == "__main__":
    main()
