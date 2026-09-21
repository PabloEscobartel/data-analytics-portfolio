#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Run upsize regressions for several sentiment-window/min-message datasets
and collect the key outputs into one workbook.

Default datasets:
  regression_dataset_book_open_3/regression_ready.csv
  regression_dataset_book_open_5/regression_ready.csv
  regression_dataset_book_open_10/regression_ready.csv
  regression_dataset_book_close_3/regression_ready.csv
  regression_dataset_book_close_5/regression_ready.csv
  regression_dataset_book_close_10/regression_ready.csv

Usage:
  python run_regressions_upsize_batch.py
  python run_regressions_upsize_batch.py regression_results_upsize_all_windows.xlsx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import run_regressions_upsize as reg


DEFAULT_OUTPUT_FILE = "regression_results_upsize_all_windows.xlsx"
DEFAULT_DATASET_TEMPLATE = "regression_dataset_{window}_{min_messages}/regression_ready.csv"
WINDOWS = ["book_open", "book_close"]
MIN_MESSAGES = [3, 5, 10]
SAMPLE_MODES = ["all", "fixed", "floater"]
SAMPLE_LABELS = {
    "all": "All bonds",
    "fixed": "Fixed coupon bonds",
    "floater": "Floating-rate bonds",
}
ORIGINAL_BASE_VARS = list(reg.BASE_VARS)
ORIGINAL_MAIN_INTERACTION_VARS = list(reg.MAIN_INTERACTION_VARS)
ORIGINAL_EXTENSION_TERMS = list(reg.EXTENSION_TERMS)


MODEL_SPECS = [
    ("coupon_ols_cluster", "Coupon reduction, bp", "_cn_si"),
    ("coupon_ihs_cluster", "IHS coupon reduction", "_cn_si"),
    ("coupon_reduced_logit_cluster", "Coupon reduction dummy", "_cn_si_coupon_binary"),
    ("coupon_positive_ihs_cluster", "IHS coupon reduction | coupon > 0", "_cn_si_coupon_positive"),
    ("logvol_ols_cluster", "ln(volume/teaser)", "_cn_si_volume"),
    ("upsize_logit_cluster", "Upsize dummy", "_cn_si_upsize"),
    ("joint_logit_cluster", "Strict dual success", "_cn_si_joint"),
]


SPLIT_MODEL_SPECS = [
    ("coupon_reduced_split_logit", "Coupon reduction dummy", "_cn_split_coupon_binary"),
    ("coupon_positive_split_ihs", "IHS coupon reduction | coupon > 0", "_cn_split_coupon_positive"),
    ("logvol_split_cluster", "ln(volume/teaser)", "_cn_split_volume"),
    ("upsize_split_logit", "Upsize dummy", "_cn_split_upsize"),
    ("joint_split_logit", "Strict dual success", "_cn_split_joint"),
]


def find_dataset(window: str, min_messages: int, dataset_template: str = DEFAULT_DATASET_TEMPLATE) -> Optional[Path]:
    candidates = [
        Path(dataset_template.format(window=window, min_messages=min_messages)),
    ]
    if dataset_template == DEFAULT_DATASET_TEMPLATE and min_messages == 3:
        candidates.append(Path(f"regression_dataset_{window}/regression_ready.csv"))
    for path in candidates:
        if path.exists():
            return path
    return None


def output_for_sample(output_path: Path, sample_mode: str) -> Path:
    if sample_mode == "all":
        return output_path
    return output_path.with_name(f"{output_path.stem}_{sample_mode}{output_path.suffix}")


def apply_sample_mode(bundle: reg.SampleBundle, sample_mode: str) -> reg.SampleBundle:
    if sample_mode == "all":
        return bundle
    if sample_mode not in {"fixed", "floater"}:
        raise ValueError(f"Unknown sample mode: {sample_mode}")

    target = 0 if sample_mode == "fixed" else 1

    def filt(df: pd.DataFrame) -> pd.DataFrame:
        out = df.loc[pd.to_numeric(df["is_floater"], errors="coerce").eq(target)].copy()
        if sample_mode == "floater":
            out["time_fe"] = np.where(
                pd.to_numeric(out["year"], errors="coerce").le(2021),
                "2018-2021",
                "2022-2025",
            )
        return out

    return reg.SampleBundle(full=filt(bundle.full), baseline=filt(bundle.baseline))


def configure_regression_spec(sample_mode: str, history_spec: str = "baseline") -> None:
    base_vars = list(ORIGINAL_BASE_VARS)
    main_interaction_vars = list(ORIGINAL_MAIN_INTERACTION_VARS)
    extension_terms = list(ORIGINAL_EXTENSION_TERMS)

    if history_spec in {"share", "baseline"}:
        base_vars = [
            "hist_reduction_share" if v == "history_repeat_prior_tightening" else v
            for v in base_vars
        ]
    elif history_spec == "dummy":
        base_vars = [
            "history_repeat_prior_tightening" if v == "hist_reduction_share" else v
            for v in base_vars
        ]
    else:
        raise ValueError(f"Unknown history specification: {history_spec}")

    if sample_mode in {"fixed", "floater"}:
        base_vars = [v for v in base_vars if v != "is_floater"]
        main_interaction_vars = []
        extension_terms = [t for t in extension_terms if t[0] not in {"si_x_is_floater", "log_buzz_x_is_floater"}]

    reg.BASE_VARS = base_vars
    reg.MAIN_INTERACTION_VARS = main_interaction_vars
    reg.SENT_VARS_SI, reg.SENT_VARS_SHARES, reg.SENT_VARS_SPLIT = reg.make_sentiment_var_sets(base_vars)
    reg.EXTENSION_TERMS = extension_terms
    reg.ESTIMATE_MNLOGIT = sample_mode != "floater"


def parse_pipeline_summary(path: Path) -> Dict[str, Any]:
    summary_path = path.parent / "pipeline_summary.txt"
    out: Dict[str, Any] = {}
    if not summary_path.exists():
        return out
    for line in summary_path.read_text(encoding="utf-8").splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = key.strip().lower().replace(" ", "_")
        nums = re.findall(r"[\d,]+", value)
        out[key] = int(nums[0].replace(",", "")) if nums else value.strip()
    return out


def get_coef_row(
    res: Any,
    colnames: List[str],
    var: str,
    meta: Dict[str, Any],
    model_key: str,
    model_label: str,
) -> Dict[str, Any]:
    coef, se, p_value = reg.get_csp(res, var, colnames)
    return {
        **meta,
        "model": model_key,
        "dependent": model_label,
        "variable": var,
        "label": reg.VAR_LABELS.get(var, var),
        "coef": coef,
        "std_error": se,
        "p_value": p_value,
        "stars": reg.sig_stars(p_value) if p_value is not None else "",
    }


def model_n(result: Dict[str, Any], model_key: str) -> int:
    if model_key in {"coupon_ols_cluster", "coupon_ihs_cluster"}:
        return result["_n_coupon"]
    if model_key == "coupon_reduced_logit_cluster":
        return result["_n_coupon_binary"]
    if model_key == "coupon_positive_ihs_cluster":
        return result["_n_coupon_positive"]
    if model_key == "logvol_ols_cluster":
        return result["_n_volume"]
    if model_key == "upsize_logit_cluster":
        return result["_n_upsize"]
    if model_key == "joint_logit_cluster":
        return result["_n_joint"]
    return result.get("_n_base", np.nan)


def collect_key_results(
    window: str,
    min_messages: int,
    dataset: Path,
    sample_mode: str,
    all_results: Dict[str, Dict[str, Any]],
) -> pd.DataFrame:
    rows = []
    key_vars = ["si", "log_buzz"]
    for sample_name, result in all_results.items():
        for model_key, model_label, cn_key in MODEL_SPECS:
            meta = {
                "window": window,
                "min_messages": min_messages,
                "sample_mode": sample_mode,
                "dataset": str(dataset),
                "sample": sample_name,
                "N": model_n(result, model_key),
            }
            for var in key_vars:
                rows.append(get_coef_row(result[model_key], result[cn_key], var, meta, model_key, model_label))
    return pd.DataFrame(rows)


def collect_split_results(window: str, min_messages: int, dataset: Path, sample_mode: str, result: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    for model_key, model_label, cn_key in SPLIT_MODEL_SPECS:
        meta = {
            "window": window,
            "min_messages": min_messages,
            "sample_mode": sample_mode,
            "dataset": str(dataset),
            "sample": "Baseline",
            "N": result.get("_n_base", np.nan),
        }
        for var in ["si_positive", "si_negative", "log_buzz"]:
            rows.append(get_coef_row(result[model_key], result[cn_key], var, meta, model_key, model_label))
    return pd.DataFrame(rows)


def collect_mnlogit_results(
    window: str,
    min_messages: int,
    dataset: Path,
    sample_mode: str,
    all_results: Dict[str, Dict[str, Any]],
) -> pd.DataFrame:
    frames = []
    for sample_name, result in all_results.items():
        df = result.get("_mnlogit_rows", pd.DataFrame()).copy()
        if df.empty:
            continue
        df.insert(0, "sample", sample_name)
        df.insert(0, "dataset", str(dataset))
        df.insert(0, "sample_mode", sample_mode)
        df.insert(0, "min_messages", min_messages)
        df.insert(0, "window", window)
        frames.append(df)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def collect_sample_summary(window: str, min_messages: int, dataset: Path, sample_mode: str, bundle: reg.SampleBundle) -> Dict[str, Any]:
    pipe = parse_pipeline_summary(dataset)
    base = bundle.baseline
    full = bundle.full
    return {
        "window": window,
        "min_messages": min_messages,
        "sample_mode": sample_mode,
        "sample_label": SAMPLE_LABELS[sample_mode],
        "dataset": str(dataset),
        "pipeline_sentiment_rows": pipe.get("with_sentiment", np.nan),
        "aggregated_offerings": pipe.get("aggregated_offerings", np.nan),
        "baseline_N": len(base),
        "baseline_coupon_N": int(base["coupon_reduction_bp"].notna().sum()),
        "baseline_volume_N": int(base["placement_vol_book"].notna().sum()),
        "baseline_success_neither": int(base["success_type"].eq("neither").sum()) if "success_type" in base else np.nan,
        "baseline_success_price_only": int(base["success_type"].eq("price_only").sum()) if "success_type" in base else np.nan,
        "baseline_success_volume_only": int(base["success_type"].eq("volume_only").sum()) if "success_type" in base else np.nan,
        "baseline_success_dual": int(base["success_type"].eq("dual_success").sum()) if "success_type" in base else np.nan,
        "full_N": len(full),
        "issuers": int(base["issuer"].nunique()),
    }


def autosize(ws) -> None:
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 42)


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    header_font = Font(bold=True, name="Arial")
    for ws in wb.worksheets:
        if ws.title == "Summary" or ws.title.startswith("open_") or ws.title.startswith("close_"):
            continue
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        autosize(ws)
    wb.save(path)


def write_small_dataframe(ws, title: str, df: pd.DataFrame, start_row: int) -> int:
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=13, name="Arial")
    row = start_row + 2
    if df.empty:
        ws.cell(row=row, column=1, value="No rows")
        return row + 2

    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=j, value=col)
        cell.font = Font(bold=True, name="Arial", size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    row += 1

    for _, rec in df.iterrows():
        for j, col in enumerate(df.columns, start=1):
            val = rec[col]
            if isinstance(val, np.integer):
                val = int(val)
            elif isinstance(val, np.floating):
                val = float(val)
            ws.cell(row=row, column=j, value=val).font = Font(name="Arial", size=10)
        row += 1

    for j in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(ws.column_dimensions[get_column_letter(j)].width or 12, 16), 42)
    return row + 2


def write_scenario_tables(
    ws,
    window: str,
    min_messages: int,
    dataset: Path,
    bundle: reg.SampleBundle,
    all_results: Dict[str, Dict[str, Any]],
    ame_df: pd.DataFrame,
    extension_df: pd.DataFrame,
) -> None:
    rb = all_results["Baseline"]
    sample_mode = rb.get("_sample_mode", "all")
    if sample_mode == "floater" and reg.USE_YEAR_FE and not reg.USE_QUARTER_FE:
        fe_label = "regime: 2018-2021 / 2022-2025"
    else:
        fe_label = "quarter" if reg.USE_QUARTER_FE else "year" if reg.USE_YEAR_FE else "none"
    ws["A1"] = f"{window}, min_messages >= {min_messages}, sample: {SAMPLE_LABELS.get(sample_mode, sample_mode)}"
    ws["A1"].font = Font(bold=True, size=15, name="Arial")
    ws["A2"] = f"Dataset: {dataset}"
    ws["A3"] = (
        f"Baseline N={len(bundle.baseline)}, coupon N={rb['_n_coupon']}, "
        f"volume N={rb['_n_volume']}"
    )
    row = 5

    row = write_small_dataframe(
        ws,
        "Table 0. Actual Regressors in Main SI Specification",
        reg.regressor_inventory(rb["_cn_si"]),
        row,
    )

    main_models = [
        ("coupon_ols_cluster", f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={rb['_n_coupon']}", rb["_cn_si"]),
        ("coupon_ihs_cluster", f"Y: IHS coupon reduction\nDV: asinh(coupon bp)\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={rb['_n_coupon']}", rb["_cn_si"]),
        ("coupon_reduced_logit_cluster", f"Y: coupon reduction dummy\nDV: 1 if coupon>0\nModel: logit/GLM\nSE: clustered by issuer\nFE: {fe_label}\nN={rb['_n_coupon_binary']}", rb["_cn_si_coupon_binary"]),
        ("coupon_positive_ihs_cluster", f"Y: IHS coupon reduction\nDV: asinh(coupon bp), coupon>0 only\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={rb['_n_coupon_positive']}", rb["_cn_si_coupon_positive"]),
        ("logvol_ols_cluster", f"Y: ln(volume/teaser)\nDV: log placement_vol_book\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={rb['_n_volume']}", rb["_cn_si_volume"]),
        ("upsize_logit_cluster", f"Y: upsize dummy\nDV: 1 if ratio>{reg.UPSIZE_THRESHOLD:.2f}\nModel: logit/GLM\nSE: clustered by issuer\nFE: {fe_label}\nN={rb['_n_upsize']}", rb["_cn_si_upsize"]),
        ("joint_logit_cluster", f"Y: strict dual success\nDV: coupon>0 and upsize\nModel: logit/GLM\nSE: clustered by issuer\nFE: {fe_label}\nN={rb['_n_joint']}", rb["_cn_si_joint"]),
    ]
    row = reg.write_model_table(ws, "Table 1. Main Results — Baseline Sample", main_models, rb, start_row=row) + 2

    coupon_models = [
        ("coupon_ols_hc3", f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: HC3\nFE: {fe_label}\nN={rb['_n_coupon']}", rb["_cn_si"]),
        ("coupon_ols_cluster", f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon']}", rb["_cn_si"]),
        ("coupon_ihs_cluster", f"Y: IHS coupon reduction\nDV: asinh(coupon bp)\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon']}", rb["_cn_si"]),
        ("coupon_reduced_logit_cluster", f"Two-part step 1\nY: coupon reduction dummy\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon_binary']}", rb["_cn_si_coupon_binary"]),
        ("coupon_positive_ihs_cluster", f"Two-part step 2\nY: asinh(coupon bp)\nSample: coupon>0\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon_positive']}", rb["_cn_si_coupon_positive"]),
        ("coupon_winsor_cluster", f"Y: coupon reduction, bp\nDV: winsor 1/99%\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon']}", rb["_cn_si"]),
        ("coupon_shares_cluster", f"Y: coupon reduction, bp\nX: pos/neg shares\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon']}", rb["_cn_sh_coupon"]),
    ]
    row = reg.write_model_table(ws, "Table 2. Coupon Robustness — Baseline Sample", coupon_models, rb, start_row=row) + 2

    volume_models = [
        ("logvol_ols_hc3", f"Y: ln(volume/teaser)\nModel: OLS\nSE: HC3\nFE: {fe_label}\nN={rb['_n_volume']}", rb["_cn_si_volume"]),
        ("logvol_ols_cluster", f"Y: ln(volume/teaser)\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_volume']}", rb["_cn_si_volume"]),
        ("volratio_winsor_cluster", f"Y: placement_vol_book\nDV: winsor 1/99%\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_volume']}", rb["_cn_si_volume"]),
        ("upsize_lpm_cluster", f"Y: upsize dummy\nModel: LPM\nSE: clustered\nFE: {fe_label}\nN={rb['_n_upsize']}", rb["_cn_si_upsize"]),
        ("logvol_shares_cluster", f"Y: ln(volume/teaser)\nX: pos/neg shares\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_volume']}", rb["_cn_sh_volume"]),
    ]
    row = reg.write_model_table(ws, "Table 3. Volume Robustness — Baseline Sample", volume_models, rb, start_row=row) + 2

    split_models = [
        ("coupon_reduced_split_logit", f"Y: coupon reduction dummy\nX: positive/negative SI\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon_binary']}", rb["_cn_split_coupon_binary"]),
        ("coupon_positive_split_ihs", f"Y: asinh(coupon bp)\nSample: coupon>0\nX: positive/negative SI\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_coupon_positive']}", rb["_cn_split_coupon_positive"]),
        ("logvol_split_cluster", f"Y: ln(volume/teaser)\nX: positive/negative SI\nModel: OLS\nSE: clustered\nFE: {fe_label}\nN={rb['_n_volume']}", rb["_cn_split_volume"]),
        ("upsize_split_logit", f"Y: upsize dummy\nX: positive/negative SI\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={rb['_n_upsize']}", rb["_cn_split_upsize"]),
        ("joint_split_logit", f"Y: strict dual success\nX: positive/negative SI\nModel: logit/GLM\nSE: clustered\nFE: {fe_label}\nN={rb['_n_joint']}", rb["_cn_split_joint"]),
    ]
    row = reg.write_model_table(ws, "Table 4. Positive vs Negative Sentiment — Baseline Sample", split_models, rb, start_row=row) + 2

    baseline_ame = ame_df[ame_df["sample"].eq("Baseline")].copy() if not ame_df.empty else pd.DataFrame()
    if not baseline_ame.empty:
        baseline_ame = baseline_ame[["model", "label", "AME", "std_error", "p_value"]]
    row = write_small_dataframe(ws, "Table 5. Average Marginal Effects — Baseline Logit Models", baseline_ame, row)

    sig_ext = extension_df[pd.to_numeric(extension_df.get("p_value"), errors="coerce") < 0.10].copy()
    if not sig_ext.empty:
        sig_ext = sig_ext[["dependent", "term_label", "type", "N", "coef", "std_error", "p_value", "evidence"]]
    row = write_small_dataframe(ws, "Table 6. Nonlinear and Interaction Terms with p < 0.10", sig_ext, row)

    mn_df = rb.get("_mnlogit_rows", pd.DataFrame()).copy()
    if not mn_df.empty:
        mn_df = mn_df[["outcome", "baseline", "label", "N", "coef", "std_error", "p_value", "stars"]]
    write_small_dataframe(ws, "Table 7. Multinomial Success Type — Baseline Category: Neither", mn_df, row)


def build_table_workbook(table_runs: List[Dict[str, Any]], sample_df: pd.DataFrame, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_small_dataframe(ws, "Sample Summary", sample_df, 1)
    ws.freeze_panes = "A4"

    for item in table_runs:
        short_window = "open" if item["window"] == "book_open" else "close"
        ws = wb.create_sheet(f"{short_window}_{item['min_messages']}")
        write_scenario_tables(
            ws,
            item["window"],
            item["min_messages"],
            item["dataset"],
            item["bundle"],
            item["all_results"],
            item["ame_df"],
            item["extension_df"],
        )
    wb.save(output_path)


def run_one(window: str, min_messages: int, dataset: Path):
    print(f"\n=== {window}, min_messages={min_messages} ===")
    print(f"Input: {dataset}")
    bundle = reg.load_and_prepare(str(dataset), reg.YEAR_FROM)
    all_results = reg.estimate_all_samples(bundle)
    ame_df = reg.collect_ame_rows(all_results)
    extension_df = reg.estimate_extension_terms(bundle.baseline)
    return bundle, all_results, ame_df, extension_df


def build_output(
    output_path: Path,
    sample_mode: str,
    history_spec: str = "baseline",
    dataset_template: str = DEFAULT_DATASET_TEMPLATE,
) -> None:
    configure_regression_spec(sample_mode, history_spec=history_spec)
    final_output_path = output_for_sample(output_path, sample_mode)

    key_frames = []
    split_frames = []
    mnlogit_frames = []
    ame_frames = []
    extension_frames = []
    sample_rows = []
    missing_rows = []
    table_runs = []

    for window in WINDOWS:
        for min_messages in MIN_MESSAGES:
            dataset = find_dataset(window, min_messages, dataset_template=dataset_template)
            if dataset is None:
                missing_rows.append({
                    "window": window,
                    "min_messages": min_messages,
                    "sample_mode": sample_mode,
                    "status": "missing regression_ready.csv",
                })
                continue
            print(f"\n=== {window}, min_messages={min_messages}, sample={sample_mode}, history={history_spec} ===")
            print(f"Input: {dataset}")
            bundle = reg.load_and_prepare(str(dataset), reg.YEAR_FROM)
            bundle = apply_sample_mode(bundle, sample_mode)
            all_results = reg.estimate_all_samples(bundle)
            ame_df = reg.collect_ame_rows(all_results)
            extension_df = reg.estimate_extension_terms(bundle.baseline)
            for result in all_results.values():
                result["_sample_mode"] = sample_mode
            baseline_result = all_results["Baseline"]
            table_runs.append({
                "window": window,
                "min_messages": min_messages,
                "sample_mode": sample_mode,
                "dataset": dataset,
                "bundle": bundle,
                "all_results": all_results,
                "ame_df": ame_df,
                "extension_df": extension_df,
            })

            key_frames.append(collect_key_results(window, min_messages, dataset, sample_mode, all_results))
            split_frames.append(collect_split_results(window, min_messages, dataset, sample_mode, baseline_result))
            mnlogit_df = collect_mnlogit_results(window, min_messages, dataset, sample_mode, all_results)
            if not mnlogit_df.empty:
                mnlogit_frames.append(mnlogit_df)
            sample_rows.append(collect_sample_summary(window, min_messages, dataset, sample_mode, bundle))

            if not ame_df.empty:
                ame_df = ame_df.copy()
                ame_df.insert(0, "min_messages", min_messages)
                ame_df.insert(0, "window", window)
                ame_df.insert(2, "sample_mode", sample_mode)
                ame_df.insert(2, "dataset", str(dataset))
                ame_frames.append(ame_df)
            if not extension_df.empty:
                extension_df = extension_df.copy()
                extension_df.insert(0, "min_messages", min_messages)
                extension_df.insert(0, "window", window)
                extension_df.insert(2, "sample_mode", sample_mode)
                extension_df.insert(2, "dataset", str(dataset))
                extension_frames.append(extension_df)

    if not key_frames:
        raise FileNotFoundError("No regression_ready.csv files found for the default book_open/book_close min-message specs.")

    key_df = pd.concat(key_frames, ignore_index=True)
    split_df = pd.concat(split_frames, ignore_index=True) if split_frames else pd.DataFrame()
    mnlogit_all = pd.concat(mnlogit_frames, ignore_index=True) if mnlogit_frames else pd.DataFrame()
    ame_all = pd.concat(ame_frames, ignore_index=True) if ame_frames else pd.DataFrame()
    ext_all = pd.concat(extension_frames, ignore_index=True) if extension_frames else pd.DataFrame()
    sample_df = pd.DataFrame(sample_rows)
    missing_df = pd.DataFrame(missing_rows)
    regressor_df = pd.concat([
        reg.regressor_inventory(item["all_results"]["Baseline"]["_cn_si"]).assign(
            window=item["window"],
            min_messages=item["min_messages"],
            sample_mode=sample_mode,
            dataset=str(item["dataset"]),
        )
        for item in table_runs
    ], ignore_index=True)

    baseline_key = key_df[key_df["sample"].eq("Baseline")].copy()
    baseline_ame = ame_all[ame_all["sample"].eq("Baseline")].copy() if not ame_all.empty else pd.DataFrame()
    baseline_mnlogit = mnlogit_all[mnlogit_all["sample"].eq("Baseline")].copy() if not mnlogit_all.empty else pd.DataFrame()
    sig_extensions = ext_all[pd.to_numeric(ext_all.get("p_value"), errors="coerce") < 0.10].copy() if not ext_all.empty else pd.DataFrame()

    build_table_workbook(table_runs, sample_df, final_output_path)

    # Keep machine-readable extracts at the end of the same workbook for quick filtering.
    with pd.ExcelWriter(final_output_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        baseline_key.to_excel(writer, index=False, sheet_name="Baseline Key")
        key_df.to_excel(writer, index=False, sheet_name="All Key")
        split_df.to_excel(writer, index=False, sheet_name="Split Key")
        regressor_df.to_excel(writer, index=False, sheet_name="Regressors")
        baseline_mnlogit.to_excel(writer, index=False, sheet_name="MNLogit Key")
        if not mnlogit_all.empty:
            mnlogit_all.to_excel(writer, index=False, sheet_name="MNLogit All")
        baseline_ame.to_excel(writer, index=False, sheet_name="AME Key")
        sig_extensions.to_excel(writer, index=False, sheet_name="Ext p10")
        if not missing_df.empty:
            missing_df.to_excel(writer, index=False, sheet_name="Missing Inputs")

    format_workbook(final_output_path)
    print(f"\nSaved: {final_output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Batch upsize regressions across sentiment windows and min-message thresholds.")
    parser.add_argument("output", nargs="?", default=DEFAULT_OUTPUT_FILE)
    parser.add_argument("--sample", choices=SAMPLE_MODES, default="all", help="Sample restriction for all scenarios.")
    parser.add_argument("--all-samples", action="store_true", help="Create separate all/fixed/floater workbooks.")
    parser.add_argument(
        "--history-spec",
        choices=["share", "dummy", "baseline"],
        default="share",
        help="Issuer-history controls: default share uses hist_reduction_share; dummy restores repeat-prior-tightening dummy.",
    )
    parser.add_argument(
        "--dataset-template",
        default=DEFAULT_DATASET_TEMPLATE,
        help=(
            "Template for regression_ready.csv paths. Available fields: {window}, {min_messages}. "
            "Default: regression_dataset_{window}_{min_messages}/regression_ready.csv"
        ),
    )
    args = parser.parse_args()
    output_path = Path(args.output)

    sample_modes = SAMPLE_MODES if args.all_samples else [args.sample]
    for sample_mode in sample_modes:
        build_output(
            output_path,
            sample_mode,
            history_spec=args.history_spec,
            dataset_template=args.dataset_template,
        )


if __name__ == "__main__":
    main()
