#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Build descriptive statistics and distribution plots for regression datasets.

Outputs:
  descriptive_statistics_all_windows.xlsx
  distribution_plots_all_windows.pdf

The workbook covers all available book_open/book_close windows, min-message
thresholds 3/5/10, and sample modes all/fixed/floater. The plot PDF focuses on
the main open_5 scenario plus threshold comparisons for sentiment variables.
"""

from __future__ import annotations

import argparse
import os
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import numpy as np
import pandas as pd

os.environ.setdefault("MPLCONFIGDIR", "/private/tmp/matplotlib")
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import run_diagnostics_upsize as diag


WINDOWS = ["book_open", "book_close"]
MIN_MESSAGES = [3, 5, 10]
DEFAULT_DATASET_TEMPLATE = "regression_dataset_{window}_{min_messages}/regression_ready.csv"
SAMPLE_MODES = ["all", "fixed", "floater"]
SAMPLE_LABELS = {
    "all": "All bonds",
    "fixed": "Fixed coupon bonds",
    "floater": "Floating-rate bonds",
}

EXTRA_LABELS = {
    "coupon_reduction_bp_w": "Coupon reduction, winsor 1/99%",
    "placement_vol_book_w": "Placement volume / teaser, winsor 1/99%",
    "log1p_n_messages_all": "ln(1 + all messages)",
    "log1p_n_messages_relevant": "ln(1 + relevant messages)",
    "log1p_views_sum": "ln(1 + total views)",
    "log1p_engagement_sum": "ln(1 + total engagement)",
    "log_placement_volume": "ln(1 + final placement volume)",
    "log_teaser_volume": "ln(1 + teaser volume)",
    "abs_si": "|Sentiment Index|",
    "coupon_positive_bp": "Coupon reduction, bp | coupon > 0",
}

NUMERIC_VARS = [
    # Outcomes.
    "coupon_reduction_bp",
    "coupon_reduced_dummy",
    "ihs_coupon_reduction",
    "ihs_coupon_positive",
    "coupon_reduction_bp_w",
    "coupon_positive_bp",
    "placement_vol_book",
    "placement_vol_book_w",
    "log_placement_vol_book",
    "upsize_dummy",
    "joint_success_upsize",
    # Final/teaser values used in prediction.
    "coupon_rate",
    "teaser_value",
    "placement_volume",
    "teaser_volume",
    "log_placement_volume",
    "log_teaser_volume",
    # Sentiment and attention.
    "si",
    "abs_si",
    "si_positive",
    "si_negative",
    "si_sq",
    "log_buzz",
    "log_buzz_sq",
    "n_messages_all",
    "n_messages_relevant",
    "log1p_n_messages_all",
    "log1p_n_messages_relevant",
    "n_sources",
    "share_positive",
    "share_negative",
    "share_relevant",
    "views_sum",
    "views_mean",
    "log1p_views_sum",
    "engagement_sum",
    "log1p_engagement_sum",
    # Controls.
    "num_organizers",
    "log_num_organizers",
    "history_debut",
    "hist_reduction_share",
    "history_repeat_prior_tightening",
    "full_issue_number",
    "log_full_issue_number",
    "hist_avg_volume_ratio",
    "rating_num",
    "rating_num_sq",
    "ofz_yield",
    "ofz_yield_within_year",
    "rvi",
    "term",
    "log_dur",
    "has_put",
    "is_floater",
    # Interactions used in nonlinear/extension specifications.
    "si_x_log_buzz",
    "si_x_rating_num",
    "log_buzz_x_rating_num",
    "num_organizers_x_si",
]

CATEGORICAL_VARS = [
    "year",
    "final_rating",
    "rating_bucket",
    "issuer_history_group",
    "industry_fe",
    "industry",
    "Currency",
    "status",
]

PLOT_NUMERIC_VARS = [
    "coupon_reduction_bp",
    "coupon_reduction_bp_w",
    "coupon_positive_bp",
    "ihs_coupon_reduction",
    "ihs_coupon_positive",
    "placement_vol_book",
    "placement_vol_book_w",
    "log_placement_vol_book",
    "coupon_rate",
    "teaser_value",
    "placement_volume",
    "log_placement_volume",
    "teaser_volume",
    "log_teaser_volume",
    "si",
    "abs_si",
    "si_positive",
    "si_negative",
    "si_sq",
    "log_buzz",
    "log_buzz_sq",
    "n_messages_all",
    "n_messages_relevant",
    "log1p_n_messages_all",
    "log1p_n_messages_relevant",
    "log1p_views_sum",
    "log1p_engagement_sum",
    "rating_num",
    "rating_num_sq",
    "ofz_yield_within_year",
    "rvi",
    "log_dur",
    "log_num_organizers",
    "hist_reduction_share",
    "log_full_issue_number",
    "hist_avg_volume_ratio",
    "si_x_log_buzz",
    "si_x_rating_num",
    "log_buzz_x_rating_num",
    "num_organizers_x_si",
]

PLOT_CATEGORICAL_VARS = [
    "year",
    "rating_bucket",
    "issuer_history_group",
    "industry_fe",
    "final_rating",
]


def var_label(var: str) -> str:
    return EXTRA_LABELS.get(var, diag.VAR_LABELS.get(var, var))


def find_dataset(window: str, min_messages: int, dataset_template: str = DEFAULT_DATASET_TEMPLATE) -> Optional[Path]:
    candidates = [Path(dataset_template.format(window=window, min_messages=min_messages))]
    if dataset_template == DEFAULT_DATASET_TEMPLATE and min_messages == 3:
        candidates.append(Path(f"regression_dataset_{window}/regression_ready.csv"))
    for path in candidates:
        if path.exists():
            return path
    return None


def scenario_label(window: str, min_messages: int) -> str:
    prefix = "open" if window == "book_open" else "close"
    return f"{prefix}_{min_messages}"


def apply_sample_mode(df: pd.DataFrame, sample_mode: str) -> pd.DataFrame:
    if sample_mode == "all":
        return df.copy()
    target = 0 if sample_mode == "fixed" else 1
    return df.loc[pd.to_numeric(df["is_floater"], errors="coerce").eq(target)].copy()


def add_prediction_terms(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "teaser_value" not in out.columns and "tizer_value" in out.columns:
        out["teaser_value"] = out["tizer_value"]
    if "teaser_volume" not in out.columns and "tizer_volume" in out.columns:
        out["teaser_volume"] = out["tizer_volume"]
    for col in [
        "coupon_reduction_bp",
        "placement_vol_book",
        "coupon_rate",
        "teaser_value",
        "placement_volume",
        "teaser_volume",
        "n_messages_all",
        "n_messages_relevant",
        "n_sources",
        "share_positive",
        "share_negative",
        "share_relevant",
        "views_sum",
        "views_mean",
        "engagement_sum",
    ]:
        if col in out.columns:
            out[col] = pd.to_numeric(out[col], errors="coerce")

    if "coupon_reduction_bp" in out.columns:
        out["coupon_reduction_bp_w"] = diag.winsorize_series(out["coupon_reduction_bp"], 0.01, 0.99)
        out["coupon_positive_bp"] = np.where(out["coupon_reduction_bp"] > 0, out["coupon_reduction_bp"], np.nan)
    else:
        out["coupon_reduction_bp_w"] = np.nan
        out["coupon_positive_bp"] = np.nan

    if "placement_vol_book" in out.columns:
        out["placement_vol_book_w"] = diag.winsorize_series(out["placement_vol_book"], 0.01, 0.99)
    else:
        out["placement_vol_book_w"] = np.nan

    if "placement_volume" in out.columns:
        vol = pd.to_numeric(out["placement_volume"], errors="coerce")
        out["log_placement_volume"] = np.where(vol > 0, np.log1p(vol), np.nan)
    else:
        out["log_placement_volume"] = np.nan

    if "teaser_volume" in out.columns:
        teaser = pd.to_numeric(out["teaser_volume"], errors="coerce")
        out["log_teaser_volume"] = np.where(teaser > 0, np.log1p(teaser), np.nan)
    else:
        out["log_teaser_volume"] = np.nan

    for source, target in [
        ("n_messages_all", "log1p_n_messages_all"),
        ("n_messages_relevant", "log1p_n_messages_relevant"),
        ("views_sum", "log1p_views_sum"),
        ("engagement_sum", "log1p_engagement_sum"),
    ]:
        if source in out.columns:
            values = pd.to_numeric(out[source], errors="coerce")
            out[target] = np.where(values >= 0, np.log1p(values), np.nan)
        else:
            out[target] = np.nan

    if "si" in out.columns:
        out["abs_si"] = pd.to_numeric(out["si"], errors="coerce").abs()
    else:
        out["abs_si"] = np.nan

    return out


def load_scenario(window: str, min_messages: int, dataset_template: str = DEFAULT_DATASET_TEMPLATE) -> Optional[Dict[str, Any]]:
    dataset = find_dataset(window, min_messages, dataset_template=dataset_template)
    if dataset is None:
        return None
    full_sub, baseline_sub, quality_df, sample_df, anomalies_df = diag.load_data(str(dataset))
    full_sub = add_prediction_terms(full_sub)
    baseline_sub = add_prediction_terms(baseline_sub)
    return {
        "window": window,
        "min_messages": min_messages,
        "scenario": scenario_label(window, min_messages),
        "dataset": dataset,
        "full": full_sub,
        "baseline": baseline_sub,
        "quality": quality_df,
        "sample": sample_df,
        "anomalies": anomalies_df,
    }


def numeric_stats(df: pd.DataFrame, variables: Iterable[str]) -> pd.DataFrame:
    rows = []
    for var in variables:
        if var not in df.columns:
            continue
        s = pd.to_numeric(df[var], errors="coerce")
        clean = s.dropna()
        row = {
            "variable": var,
            "label": var_label(var),
            "N": int(clean.size),
            "missing": int(s.isna().sum()),
            "missing_share": float(s.isna().mean()) if len(s) else np.nan,
            "mean": np.nan,
            "std": np.nan,
            "min": np.nan,
            "p01": np.nan,
            "p05": np.nan,
            "p25": np.nan,
            "median": np.nan,
            "p75": np.nan,
            "p95": np.nan,
            "p99": np.nan,
            "max": np.nan,
            "skew": np.nan,
            "kurtosis": np.nan,
        }
        if not clean.empty:
            row.update({
                "mean": float(clean.mean()),
                "std": float(clean.std()),
                "min": float(clean.min()),
                "p01": float(clean.quantile(0.01)),
                "p05": float(clean.quantile(0.05)),
                "p25": float(clean.quantile(0.25)),
                "median": float(clean.median()),
                "p75": float(clean.quantile(0.75)),
                "p95": float(clean.quantile(0.95)),
                "p99": float(clean.quantile(0.99)),
                "max": float(clean.max()),
                "skew": float(clean.skew()) if clean.size > 2 else np.nan,
                "kurtosis": float(clean.kurtosis()) if clean.size > 3 else np.nan,
            })
        rows.append(row)
    return pd.DataFrame(rows)


def categorical_counts(df: pd.DataFrame, variables: Iterable[str], top_n: int = 25) -> pd.DataFrame:
    rows = []
    for var in variables:
        if var not in df.columns:
            continue
        s = df[var]
        counts = s.fillna("Missing").astype(str).value_counts(dropna=False).head(top_n)
        total = len(s)
        for rank, (value, count) in enumerate(counts.items(), start=1):
            rows.append({
                "variable": var,
                "label": var_label(var),
                "rank": rank,
                "value": value,
                "count": int(count),
                "share": float(count / total) if total else np.nan,
                "total_N": total,
                "missing": int(s.isna().sum()),
            })
    return pd.DataFrame(rows)


def add_meta(df: pd.DataFrame, scenario: Dict[str, Any], sample_mode: str) -> pd.DataFrame:
    out = df.copy()
    out.insert(0, "dataset", str(scenario["dataset"]))
    out.insert(0, "sample_label", SAMPLE_LABELS[sample_mode])
    out.insert(0, "sample_mode", sample_mode)
    out.insert(0, "scenario", scenario["scenario"])
    out.insert(0, "min_messages", scenario["min_messages"])
    out.insert(0, "window", scenario["window"])
    return out


def count_sentiment_coverage(df: pd.DataFrame) -> int:
    """Rows with an attached sentiment aggregate, proxied by positive buzz."""
    if "log_buzz" not in df.columns:
        return 0
    return int(pd.to_numeric(df["log_buzz"], errors="coerce").gt(0).sum())


def count_si_nonzero(df: pd.DataFrame) -> int:
    if "si" not in df.columns:
        return 0
    si = pd.to_numeric(df["si"], errors="coerce")
    return int(si.fillna(0).ne(0).sum())


def outcome_sample_rows(scenario: Dict[str, Any]) -> List[Dict[str, Any]]:
    base = scenario["baseline"]
    outcome_specs = [
        ("Baseline", "baseline", pd.Series(True, index=base.index)),
        ("Coupon reduction", "coupon_reduction_bp_notna", base["coupon_reduction_bp"].notna()),
        ("Volume ratio", "placement_vol_book_notna", base["placement_vol_book"].notna()),
        (
            "Joint success",
            "coupon_and_volume_notna",
            base["coupon_reduction_bp"].notna() & base["placement_vol_book"].notna(),
        ),
    ]

    rows = []
    for label, filter_name, mask in outcome_specs:
        sub = base.loc[mask].copy()
        fixed = apply_sample_mode(sub, "fixed")
        floater = apply_sample_mode(sub, "floater")
        rows.append({
            "window": scenario["window"],
            "min_messages": scenario["min_messages"],
            "scenario": scenario["scenario"],
            "dataset": str(scenario["dataset"]),
            "row": label,
            "filter": filter_name,
            "all_N": len(sub),
            "fixed_N": len(fixed),
            "floater_N": len(floater),
            "issuers_N": int(sub["issuer"].nunique()) if "issuer" in sub.columns else np.nan,
            "fixed_issuers_N": int(fixed["issuer"].nunique()) if "issuer" in fixed.columns else np.nan,
            "floater_issuers_N": int(floater["issuer"].nunique()) if "issuer" in floater.columns else np.nan,
            "sentiment_coverage_N": count_sentiment_coverage(sub),
            "si_nonzero_N": count_si_nonzero(sub),
        })
    return rows


def build_outputs(
    scenarios: List[Dict[str, Any]],
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    summary_rows = []
    outcome_size_rows = []
    numeric_frames = []
    categorical_frames = []

    for scenario in scenarios:
        outcome_size_rows.extend(outcome_sample_rows(scenario))
        for sample_mode in SAMPLE_MODES:
            baseline = apply_sample_mode(scenario["baseline"], sample_mode)
            full = apply_sample_mode(scenario["full"], sample_mode)
            summary_rows.append({
                "window": scenario["window"],
                "min_messages": scenario["min_messages"],
                "scenario": scenario["scenario"],
                "sample_mode": sample_mode,
                "sample_label": SAMPLE_LABELS[sample_mode],
                "dataset": str(scenario["dataset"]),
                "baseline_N": len(baseline),
                "full_N": len(full),
                "issuers": int(baseline["issuer"].nunique()) if "issuer" in baseline.columns else np.nan,
                "coupon_N": int(baseline["coupon_reduction_bp"].notna().sum()) if "coupon_reduction_bp" in baseline else np.nan,
                "volume_N": int(baseline["placement_vol_book"].notna().sum()) if "placement_vol_book" in baseline else np.nan,
                "sentiment_coverage_N": count_sentiment_coverage(baseline),
                "si_nonzero_N": count_si_nonzero(baseline),
                "fixed_N": int(pd.to_numeric(baseline.get("is_floater", np.nan), errors="coerce").eq(0).sum()),
                "floater_N": int(pd.to_numeric(baseline.get("is_floater", np.nan), errors="coerce").eq(1).sum()),
            })

            numeric_frames.append(add_meta(numeric_stats(baseline, NUMERIC_VARS), scenario, sample_mode))
            categorical_frames.append(add_meta(categorical_counts(baseline, CATEGORICAL_VARS), scenario, sample_mode))

    numeric_all = pd.concat(numeric_frames, ignore_index=True) if numeric_frames else pd.DataFrame()
    categorical_all = pd.concat(categorical_frames, ignore_index=True) if categorical_frames else pd.DataFrame()
    summary = pd.DataFrame(summary_rows)
    outcome_sizes = pd.DataFrame(outcome_size_rows)

    main = numeric_all[
        (numeric_all["scenario"] == "open_5")
        & (numeric_all["sample_mode"] == "all")
    ].copy()
    return summary, outcome_sizes, main, numeric_all, categorical_all


def variable_dictionary() -> pd.DataFrame:
    rows = []
    for var in sorted(set(NUMERIC_VARS + CATEGORICAL_VARS)):
        rows.append({
            "variable": var,
            "label": var_label(var),
            "type": "numeric" if var in NUMERIC_VARS else "categorical",
            "included_in_plots": var in set(PLOT_NUMERIC_VARS + PLOT_CATEGORICAL_VARS),
        })
    return pd.DataFrame(rows)


def autosize(ws) -> None:
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 42)


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        autosize(ws)
    wb.save(path)


def write_workbook(
    path: Path,
    summary: pd.DataFrame,
    outcome_sizes: pd.DataFrame,
    main: pd.DataFrame,
    numeric_all: pd.DataFrame,
    categorical_all: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, index=False, sheet_name="Sample Summary")
        outcome_sizes.to_excel(writer, index=False, sheet_name="Outcome Sample Sizes")
        main.to_excel(writer, index=False, sheet_name="Main open_5 all")
        numeric_all.to_excel(writer, index=False, sheet_name="Numeric Stats")
        categorical_all.to_excel(writer, index=False, sheet_name="Categorical Counts")
        variable_dictionary().to_excel(writer, index=False, sheet_name="Variable Dictionary")
    format_workbook(path)


def percentile_xlim(series_list: List[pd.Series]) -> Optional[Tuple[float, float]]:
    clean = pd.concat([s.dropna() for s in series_list], ignore_index=True)
    if clean.empty:
        return None
    lo = float(clean.quantile(0.01))
    hi = float(clean.quantile(0.99))
    if not np.isfinite(lo) or not np.isfinite(hi) or lo == hi:
        return None
    pad = (hi - lo) * 0.05
    return lo - pad, hi + pad


def plot_numeric_by_sample(pdf: PdfPages, scenario: Dict[str, Any], var: str) -> None:
    frames = {mode: apply_sample_mode(scenario["baseline"], mode) for mode in SAMPLE_MODES}
    series = {
        mode: pd.to_numeric(df[var], errors="coerce").dropna()
        for mode, df in frames.items()
        if var in df.columns and pd.to_numeric(df[var], errors="coerce").notna().any()
    }
    if not series:
        return

    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.8), gridspec_kw={"width_ratios": [3, 1]})
    xlim = percentile_xlim(list(series.values()))
    for mode, s in series.items():
        data = s
        if xlim is not None:
            data = s[(s >= xlim[0]) & (s <= xlim[1])]
        axes[0].hist(data, bins=35, alpha=0.42, label=f"{SAMPLE_LABELS[mode]} (N={len(s)})")
    axes[0].set_title("Histogram, clipped to p1-p99 for readability")
    axes[0].set_xlabel(var_label(var))
    axes[0].set_ylabel("Count")
    axes[0].legend(fontsize=8)

    labels = []
    values = []
    for mode, s in series.items():
        labels.append(SAMPLE_LABELS[mode])
        values.append(s)
    axes[1].boxplot(values, tick_labels=labels, vert=True, patch_artist=True)
    axes[1].tick_params(axis="x", rotation=45)
    axes[1].set_title("Boxplot")

    fig.suptitle(f"{scenario['scenario']}: {var_label(var)}")
    fig.tight_layout()
    pdf.savefig(fig)
    plt.close(fig)


def plot_categorical(pdf: PdfPages, scenario: Dict[str, Any], var: str) -> None:
    df = scenario["baseline"]
    if var not in df.columns:
        return
    counts = df[var].fillna("Missing").astype(str).value_counts().head(20).sort_values()
    if counts.empty:
        return
    fig, ax = plt.subplots(figsize=(10.5, max(4.2, 0.32 * len(counts))))
    ax.barh(counts.index, counts.values, color="#4C78A8")
    ax.set_xlabel("Count")
    ax.set_title(f"{scenario['scenario']}: {var_label(var)}")
    fig.tight_layout()
    pdf.savefig(fig)
    plt.close(fig)


def plot_threshold_comparison(pdf: PdfPages, scenarios: List[Dict[str, Any]], var: str, window: str) -> None:
    selected = [s for s in scenarios if s["window"] == window]
    if not selected:
        return
    series = {}
    for s in selected:
        df = s["baseline"]
        if var in df.columns:
            clean = pd.to_numeric(df[var], errors="coerce").dropna()
            if not clean.empty:
                series[s["scenario"]] = clean
    if not series:
        return

    fig, ax = plt.subplots(figsize=(10.5, 5))
    xlim = percentile_xlim(list(series.values()))
    for label, s in sorted(series.items()):
        data = s
        if xlim is not None:
            data = s[(s >= xlim[0]) & (s <= xlim[1])]
        ax.hist(data, bins=35, alpha=0.38, label=f"{label} (N={len(s)})")
    ax.set_title(f"{window}: threshold comparison for {var_label(var)}")
    ax.set_xlabel(var_label(var))
    ax.set_ylabel("Count")
    ax.legend()
    fig.tight_layout()
    pdf.savefig(fig)
    plt.close(fig)


def write_plot_pdf(path: Path, scenarios: List[Dict[str, Any]]) -> None:
    main = next((s for s in scenarios if s["scenario"] == "open_5"), None)
    with PdfPages(path) as pdf:
        if main is not None:
            fig, ax = plt.subplots(figsize=(11.5, 4))
            ax.axis("off")
            ax.text(
                0.02,
                0.8,
                "Distribution plots\n\n"
                "Main section: open_5 baseline sample, compared across all/fixed/floater.\n"
                "Final section: sentiment/attention distributions by min-message threshold.",
                fontsize=14,
                va="top",
            )
            pdf.savefig(fig)
            plt.close(fig)

            for var in PLOT_NUMERIC_VARS:
                plot_numeric_by_sample(pdf, main, var)
            for var in PLOT_CATEGORICAL_VARS:
                plot_categorical(pdf, main, var)

        for window in WINDOWS:
            for var in ["si", "si_positive", "si_negative", "log_buzz", "n_messages_all", "n_messages_relevant"]:
                plot_threshold_comparison(pdf, scenarios, var, window)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build descriptive statistics and distribution plots.")
    parser.add_argument("--output-xlsx", default="descriptive_statistics_all_windows.xlsx")
    parser.add_argument("--output-plots", default="distribution_plots_all_windows.pdf")
    parser.add_argument(
        "--dataset-template",
        default=DEFAULT_DATASET_TEMPLATE,
        help=(
            "Template for regression_ready.csv paths. Available fields: {window}, {min_messages}. "
            "Default: regression_dataset_{window}_{min_messages}/regression_ready.csv"
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    scenarios = []
    for window in WINDOWS:
        for min_messages in MIN_MESSAGES:
            item = load_scenario(window, min_messages, dataset_template=args.dataset_template)
            if item is not None:
                scenarios.append(item)

    if not scenarios:
        raise FileNotFoundError("No regression datasets found.")

    summary, outcome_sizes, main, numeric_all, categorical_all = build_outputs(scenarios)
    write_workbook(Path(args.output_xlsx), summary, outcome_sizes, main, numeric_all, categorical_all)
    write_plot_pdf(Path(args.output_plots), scenarios)

    print(f"Saved descriptive statistics: {args.output_xlsx}")
    print(f"Saved distribution plots: {args.output_plots}")


if __name__ == "__main__":
    main()
