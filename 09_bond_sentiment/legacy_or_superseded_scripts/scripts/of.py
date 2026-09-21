from datetime import datetime, date
import re
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

INPUT_FILE = "la_finale/bonds_offer.xlsx"
OUTPUT_FILE = "bonds_offer_updated.xlsx"

MIN_VALID_YEAR = 2000

def normalize_text(x):
    if x is None:
        return ""
    return str(x).strip().lower().replace("ё", "е")

def find_col_by_header(ws, header_name, header_row=1):
    target = normalize_text(header_name)
    for cell in ws[header_row]:
        if normalize_text(cell.value) == target:
            return cell.column
    raise ValueError(f"На листе '{ws.title}' не найден столбец '{header_name}'")

def parse_date_from_cell(cell, epoch):
    v = cell.value

    if v is None:
        return None

    if isinstance(v, datetime):
        d = v.date()
        return d if d.year >= MIN_VALID_YEAR else None

    if isinstance(v, date):
        return v if v.year >= MIN_VALID_YEAR else None

    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None

        patterns = [
            ("%d.%m.%Y", s),
            ("%d.%m.%y", s),
            ("%Y-%m-%d", s),
            ("%Y.%m.%d", s),
            ("%d/%m/%Y", s),
            ("%d/%m/%y", s),
        ]

        for fmt, val in patterns:
            try:
                d = datetime.strptime(val, fmt).date()
                return d if d.year >= MIN_VALID_YEAR else None
            except ValueError:
                pass

        m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", s)
        if m:
            dd, mm, yyyy = m.groups()
            try:
                d = date(int(yyyy), int(mm), int(dd))
                return d if d.year >= MIN_VALID_YEAR else None
            except ValueError:
                return None

        return None

    if isinstance(v, (int, float)):
        if getattr(cell, "is_date", False):
            try:
                d = from_excel(v, epoch=epoch)
                if isinstance(d, datetime):
                    d = d.date()
                return d if d.year >= MIN_VALID_YEAR else None
            except Exception:
                return None
        return None

    return None

def build_earliest_map(ws, isin_col, date_col, epoch):
    result = {}

    for row in ws.iter_rows(min_row=2):
        isin_cell = row[isin_col - 1]
        date_cell = row[date_col - 1]

        isin = isin_cell.value
        if isin is None:
            continue

        isin = str(isin).strip()
        if not isin:
            continue

        d = parse_date_from_cell(date_cell, epoch)
        if d is None:
            continue

        prev = result.get(isin)
        if prev is None or d < prev:
            result[isin] = d

    return result

def update_main_column(ws_main, isin_col, target_col, source_map, epoch):
    filled_empty = 0
    replaced_with_earlier = 0
    unchanged = 0
    no_match = 0

    for row in ws_main.iter_rows(min_row=2):
        isin_cell = row[isin_col - 1]
        target_cell = row[target_col - 1]

        isin = isin_cell.value
        if isin is None:
            continue

        isin = str(isin).strip()
        if not isin:
            continue

        source_date = source_map.get(isin)
        if source_date is None:
            no_match += 1
            continue

        existing_date = parse_date_from_cell(target_cell, epoch)

        if existing_date is None:
            target_cell.value = source_date
            target_cell.number_format = "DD.MM.YYYY"
            filled_empty += 1
        else:
            if source_date < existing_date:
                target_cell.value = source_date
                target_cell.number_format = "DD.MM.YYYY"
                replaced_with_earlier += 1
            else:
                unchanged += 1

    return {
        "filled_empty": filled_empty,
        "replaced_with_earlier": replaced_with_earlier,
        "unchanged": unchanged,
        "no_match": no_match,
    }

def main():
    wb = load_workbook(INPUT_FILE)
    epoch = wb.epoch

    ws_main = wb["main"]
    ws_call = wb["call"]
    ws_put = wb["put"]

    main_isin_col = find_col_by_header(ws_main, "ISIN")
    main_call_col = find_col_by_header(ws_main, "Оферта (call)")
    main_put_col = find_col_by_header(ws_main, "Оферта (put)")

    call_isin_col = find_col_by_header(ws_call, "ISIN")
    put_isin_col = find_col_by_header(ws_put, "ISIN")

    call_date_col = find_col_by_header(ws_call, "Дата")
    put_date_col = find_col_by_header(ws_put, "Дата")

    call_map = build_earliest_map(ws_call, call_isin_col, call_date_col, epoch)
    put_map = build_earliest_map(ws_put, put_isin_col, put_date_col, epoch)

    call_stats = update_main_column(ws_main, main_isin_col, main_call_col, call_map, epoch)
    put_stats = update_main_column(ws_main, main_isin_col, main_put_col, put_map, epoch)

    wb.save(OUTPUT_FILE)

    print("Готово")
    print("Файл:", OUTPUT_FILE)
    print("CALL:", call_stats)
    print("PUT:", put_stats)

if __name__ == "__main__":
    main()
