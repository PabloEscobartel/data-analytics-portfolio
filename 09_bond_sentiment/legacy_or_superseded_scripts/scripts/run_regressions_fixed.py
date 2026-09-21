#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Регрессионный анализ успешности первичных размещений облигаций.

Что исправлено относительно предыдущей версии:
1. Скрипт возвращает и оценивает не только baseline, но и full/robustness samples.
2. Baseline исключает ВЭБ.РФ, num_organizers = 0 и book_date > placement_date.
3. SI/buzz/shares пропуски заполняются нулём.
5. Fractional logit не превращает нули в 0.001; нули остаются нулями.
6. Добавлены year fixed effects.
7. Добавлены HC3 и clustered SE by issuer.
8. Добавлен winsorized coupon reduction как robustness.
9. Tobit оставлен как appendix/robustness, а не как главный результат.
10. Warnings не подавляются полностью.

Требования:
    pip install pandas numpy statsmodels scipy openpyxl

Запуск:
    python run_regressions.py
"""

from __future__ import annotations

import warnings
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple, Any

import numpy as np
import pandas as pd
import statsmodels.api as sm
from scipy import optimize, stats
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("default")

# ============================================================
# SETTINGS
# ============================================================

INPUT_FILE = "regression_dataset/regression_ready.csv"
OUTPUT_FILE = "regression_results.xlsx"
YEAR_FROM = 2018

# Baseline cleaning rules
BASELINE_EXCLUDE_ISSUERS = {"ВЭБ.РФ"}
BASELINE_DROP_ZERO_ORGANIZERS = True
BASELINE_DROP_BOOK_AFTER_PLACEMENT = True

# Fixed effects
USE_YEAR_FE = True
USE_QUARTER_FE = False  # If True, year FE are ignored and quarter FE are used instead.

# Coupon robustness
WINSOR_LOWER = 0.01
WINSOR_UPPER = 0.99

# Robustness samples
COOK_DROP_Q = 0.99

RATING_MAP = {
    "AAA": 15, "AA+": 14, "AA": 13, "AA-": 12,
    "A+": 11, "A": 10, "A-": 9,
    "BBB+": 8, "BBB": 7, "BBB-": 6,
    "BB+": 5, "BB": 4, "BB-": 3,
    "B+": 2, "B": 1, "B-": 0,
    "C": -1,
}

VAR_LABELS = {
    "const": "Constant",
    "num_organizers": "Number of organizers",
    "is_debut": "Debut (dummy)",
    "hist_ever_reduced": "Prior tightening (dummy)",
    "hist_avg_volume_ratio": "Hist. avg volume ratio",
    "rating_num": "Credit rating (ordinal)",
    "log_dur": "ln(Term to exit, days)",
    "has_put": "Has put option (dummy)",
    "is_floater": "Floating rate (dummy)",
    "si": "Sentiment Index (SI)",
    "log_buzz": "ln(Buzz)",
    "share_pos": "Share positive",
    "share_neg": "Share negative",
    "ofz_yield": "OFZ matched yield, %",
    "rvi": "RVI (volatility index)",
}

BASE_VARS = [
    "num_organizers",
    "is_debut",
    "hist_ever_reduced",
    "hist_avg_volume_ratio",
    "rating_num",
    "ofz_yield",
    "rvi",
    "log_dur",
    "has_put",
    "is_floater",
]

SENT_VARS_SI = BASE_VARS + ["si", "log_buzz"]
SENT_VARS_SHARES = BASE_VARS + ["share_pos", "share_neg", "log_buzz"]


# ============================================================
# HELPERS
# ============================================================

@dataclass
class SampleBundle:
    full: pd.DataFrame
    baseline: pd.DataFrame


def first_existing_col(df: pd.DataFrame, candidates: List[str]) -> str:
    for col in candidates:
        if col in df.columns:
            return col
    raise KeyError(f"Не найдена ни одна из колонок: {', '.join(candidates)}")


def optional_existing_col(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    for col in candidates:
        if col in df.columns:
            return col
    return None


def normalize_rating(series: pd.Series) -> pd.Series:
    """Normalize en/em dashes in ratings such as BB– / B–."""
    return (
        series.astype(str)
        .str.replace("–", "-", regex=False)
        .str.replace("—", "-", regex=False)
        .str.strip()
    )


def safe_numeric(df: pd.DataFrame, col: Optional[str], default: float = np.nan) -> pd.Series:
    if col is None:
        return pd.Series(default, index=df.index, dtype="float64")
    return pd.to_numeric(df[col], errors="coerce")


def winsorize_series(s: pd.Series, lower: float = 0.01, upper: float = 0.99) -> pd.Series:
    s_nonmiss = s.dropna()
    if s_nonmiss.empty:
        return s.copy()
    lo, hi = s_nonmiss.quantile(lower), s_nonmiss.quantile(upper)
    return s.clip(lower=lo, upper=hi)


def sig_stars(p: float) -> str:
    if p is None or pd.isna(p):
        return ""
    if p < 0.01:
        return "***"
    if p < 0.05:
        return "**"
    if p < 0.10:
        return "*"
    return ""


def add_quality_flags(sub: pd.DataFrame) -> pd.DataFrame:
    sub = sub.copy()
    sub["placement_dt"] = pd.to_datetime(sub["placement_date"], errors="coerce")
    sub["book_dt"] = pd.to_datetime(sub["book_date"], errors="coerce")

    sub["book_after_placement"] = (
        sub["book_dt"].notna()
        & sub["placement_dt"].notna()
        & (sub["book_dt"].dt.date > sub["placement_dt"].dt.date)
    )

    sub["is_special_issuer"] = sub["issuer"].astype(str).isin(BASELINE_EXCLUDE_ISSUERS)
    sub["zero_organizers"] = sub["num_organizers"].eq(0)

    if USE_QUARTER_FE:
        sub["time_fe"] = sub["placement_dt"].dt.to_period("Q").astype(str)
    else:
        sub["time_fe"] = sub["year"].astype("Int64").astype(str)

    return sub


def apply_baseline_filters(sub: pd.DataFrame) -> pd.DataFrame:
    mask = pd.Series(True, index=sub.index)

    if BASELINE_DROP_BOOK_AFTER_PLACEMENT:
        mask &= ~sub["book_after_placement"]

    if BASELINE_EXCLUDE_ISSUERS:
        mask &= ~sub["is_special_issuer"]

    if BASELINE_DROP_ZERO_ORGANIZERS:
        mask &= ~sub["zero_organizers"]

    return sub.loc[mask].copy()


def build_X(data: pd.DataFrame, variables: List[str], add_fe: bool = True) -> pd.DataFrame:
    """Build numeric design matrix with constant and optional time fixed effects."""
    X = data[variables].copy()

    if add_fe:
        if USE_QUARTER_FE:
            fe = pd.get_dummies(data["time_fe"], prefix="q", drop_first=True, dtype=float)
        elif USE_YEAR_FE:
            fe = pd.get_dummies(data["time_fe"], prefix="year", drop_first=True, dtype=float)
        else:
            fe = pd.DataFrame(index=data.index)

        if not fe.empty:
            X = pd.concat([X, fe], axis=1)

    X = sm.add_constant(X, has_constant="add")
    X = X.apply(pd.to_numeric, errors="coerce").astype(float)
    return X


def get_csp(res: Any, var: str, colnames: List[str]) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    """Return coefficient, standard error, p-value for statsmodels or custom Tobit result."""
    if isinstance(res, dict):
        if var not in colnames:
            return None, None, None
        i = colnames.index(var)
        return res["coef"][i], res["se"][i], res["p"][i]

    if hasattr(res, "params") and var in res.params.index:
        return float(res.params[var]), float(res.bse[var]), float(res.pvalues[var])

    return None, None, None


def fit_ols(y: pd.Series, X: pd.DataFrame, cov_type: str = "HC3", groups: Optional[pd.Series] = None):
    model = sm.OLS(y, X, missing="drop")

    if cov_type == "cluster":
        if groups is None:
            raise ValueError("groups must be provided for clustered SE")
        if pd.Series(groups).nunique() < 2:
            # Fallback to HC3 if clustering is impossible.
            return model.fit(cov_type="HC3")
        return model.fit(cov_type="cluster", cov_kwds={"groups": groups})

    return model.fit(cov_type=cov_type)


def fit_frac_logit(y: pd.Series, X: pd.DataFrame, groups: Optional[pd.Series] = None):
    """
    Fractional response model via GLM Binomial quasi-likelihood.
    Keeps true zeros as zeros. Only clips accidental out-of-bounds values to [0, 1].

    Important:
    statsmodels can fail for GLM clustered covariance with "Singular matrix",
    especially when the dependent variable has many zeros and the model includes
    fixed effects. In that case we keep the fractional logit coefficients but
    fall back to HC1 robust covariance. If HC1 also fails, we return the
    conventional GLM covariance and print a warning.
    """
    y_frac = y.clip(lower=0, upper=1)
    model = sm.GLM(y_frac, X, family=sm.families.Binomial(), missing="drop")

    # 1) Preferred: clustered SE by issuer.
    if groups is not None and pd.Series(groups).nunique() >= 2:
        try:
            return model.fit(
                cov_type="cluster",
                cov_kwds={"groups": groups},
                maxiter=200,
                disp=0,
            )
        except (np.linalg.LinAlgError, ValueError) as e:
            warnings.warn(
                f"Fractional logit clustered SE failed ({e}). "
                "Falling back to HC1 robust SE for this model.",
                RuntimeWarning,
            )

    # 2) Fallback: heteroskedasticity-robust SE.
    try:
        return model.fit(cov_type="HC1", maxiter=200, disp=0)
    except (np.linalg.LinAlgError, ValueError) as e:
        warnings.warn(
            f"Fractional logit HC1 SE failed ({e}). "
            "Returning conventional GLM covariance for this model.",
            RuntimeWarning,
        )

    # 3) Last resort: conventional GLM covariance.
    return model.fit(maxiter=200, disp=0)


# ============================================================
# TOBIT MLE — APPENDIX / ROBUSTNESS ONLY
# ============================================================

def tobit_fit(y: np.ndarray, X: np.ndarray, lower: float = 0.0) -> Dict[str, Any]:
    """
    Simple left-censored Tobit MLE.
    Kept as robustness only. Does not implement clustered SE.
    """
    y = np.asarray(y, dtype=float)
    X = np.asarray(X, dtype=float)

    def neg_ll(params: np.ndarray) -> float:
        beta = params[:-1]
        log_sigma = params[-1]
        sigma = np.exp(log_sigma)
        xb = X @ beta

        uncens = y > lower
        cens = ~uncens

        ll = 0.0
        if uncens.any():
            r = (y[uncens] - xb[uncens]) / sigma
            ll += np.sum(-0.5 * np.log(2 * np.pi) - np.log(sigma) - 0.5 * r**2)

        if cens.any():
            z = (lower - xb[cens]) / sigma
            ll += np.sum(np.log(stats.norm.cdf(z) + 1e-12))

        return -ll

    ols = sm.OLS(y, X).fit()
    resid_std = np.std(ols.resid)
    if not np.isfinite(resid_std) or resid_std <= 0:
        resid_std = 1.0

    p0 = np.append(np.asarray(ols.params), np.log(resid_std))

    res = optimize.minimize(
        neg_ll,
        p0,
        method="BFGS",
        options={"maxiter": 20000, "gtol": 1e-6},
    )

    # If BFGS fails, try Nelder-Mead as fallback.
    if not res.success:
        res_nm = optimize.minimize(
            neg_ll,
            res.x,
            method="Nelder-Mead",
            options={"maxiter": 80000, "xatol": 1e-8, "fatol": 1e-8},
        )
        if res_nm.fun < res.fun:
            res = res_nm

    beta = res.x[:-1]
    sigma = float(np.exp(res.x[-1]))

    # Try covariance from inverse Hessian when available.
    se = np.full(len(res.x), np.nan)
    if hasattr(res, "hess_inv"):
        try:
            hess_inv = np.asarray(res.hess_inv)
            if hess_inv.shape == (len(res.x), len(res.x)):
                se = np.sqrt(np.abs(np.diag(hess_inv)))
        except Exception:
            pass

    # Fallback numerical Hessian.
    if np.isnan(se).all():
        eps = 1e-5
        n_p = len(res.x)
        H = np.zeros((n_p, n_p))
        for i in range(n_p):
            def fi(p, _i=i):
                return optimize.approx_fprime(p, neg_ll, eps)[_i]
            H[i] = optimize.approx_fprime(res.x, fi, eps)
        try:
            cov = np.linalg.inv(H)
            se = np.sqrt(np.abs(np.diag(cov)))
        except np.linalg.LinAlgError:
            pass

    se_beta = se[:-1]
    z_vals = np.where(se_beta > 0, beta / se_beta, np.nan)
    p_vals = 2 * (1 - stats.norm.cdf(np.abs(z_vals)))

    return {
        "coef": beta,
        "se": se_beta,
        "z": z_vals,
        "p": p_vals,
        "sigma": sigma,
        "ll": -float(res.fun),
        "n": int(len(y)),
        "n_cens": int((y <= lower).sum()),
        "n_uncens": int((y > lower).sum()),
        "converged": bool(res.success),
        "message": str(res.message),
    }


# ============================================================
# LOAD / PREPARE DATA
# ============================================================

def load_and_prepare(path: str, year_from: int) -> SampleBundle:
    df = pd.read_csv(path, low_memory=False)

    df["placement_dt"] = pd.to_datetime(df["placement_date"], errors="coerce")
    df["year"] = df["placement_dt"].dt.year
    sub = df[df["year"] >= year_from].copy()

    # -------------------------------
    # Sentiment variables
    # -------------------------------
    si_col = first_existing_col(sub, ["si_relevant", "si"])
    buzz_col = first_existing_col(sub, ["log_buzz_relevant", "log_buzz_all", "log_buzz"])
    share_neg_col = optional_existing_col(sub, ["share_negative_relevant", "share_negative"])
    share_pos_col = optional_existing_col(sub, ["share_positive_relevant", "share_positive"])

    si_raw = safe_numeric(sub, si_col)
    log_buzz_raw = safe_numeric(sub, buzz_col)
    share_neg_raw = safe_numeric(sub, share_neg_col, default=np.nan)
    share_pos_raw = safe_numeric(sub, share_pos_col, default=np.nan)

    sub["si"] = si_raw.fillna(0)
    sub["log_buzz"] = log_buzz_raw.fillna(0)
    sub["share_neg"] = share_neg_raw.fillna(0)
    sub["share_pos"] = share_pos_raw.fillna(0)

    # -------------------------------
    # Dependent variables
    # -------------------------------
    volume_ratio = pd.to_numeric(sub["volume_ratio"], errors="coerce")
    sub["underplacement"] = 1 - volume_ratio
    sub["coupon_reduction_bp"] = pd.to_numeric(sub["coupon_reduction_bp"], errors="coerce")

    # -------------------------------
    # Controls
    # -------------------------------
    sub["num_organizers"] = pd.to_numeric(sub["num_organizers"], errors="coerce")

    if "log_term" in sub.columns:
        sub["log_dur"] = pd.to_numeric(sub["log_term"], errors="coerce")
    elif "term" in sub.columns:
        term = pd.to_numeric(sub["term"], errors="coerce")
        sub["log_dur"] = np.log(term.clip(lower=1))
    else:
        raise KeyError("Нужна колонка log_term или term")

    sub["final_rating"] = normalize_rating(sub["final_rating"])
    sub["rating_num"] = sub["final_rating"].map(RATING_MAP)
    sub["ofz_yield"] = pd.to_numeric(sub["ofz_matched_yield"], errors="coerce")
    sub["rvi"] = pd.to_numeric(sub["rvi_close"], errors="coerce")

    if "is_floater" not in sub.columns:
        sub["is_floater"] = 0
    else:
        sub["is_floater"] = pd.to_numeric(sub["is_floater"], errors="coerce")

    # Ensure dummy/control numeric columns exist and are numeric.
    for col in ["is_debut", "hist_ever_reduced", "has_put", "hist_avg_volume_ratio"]:
        if col not in sub.columns:
            raise KeyError(f"Не найдена обязательная колонка: {col}")
        sub[col] = pd.to_numeric(sub[col], errors="coerce")

    sub = add_quality_flags(sub)
    sub["issuer_id"] = sub["issuer"].astype(str).astype("category").cat.codes

    # Drop missing values for common variables used in underplacement models.
    # Do not drop coupon_reduction_bp here, because coupon models use a conditional subsample.
    common_drop_vars = list(dict.fromkeys(
        ["underplacement", "issuer_id"] + SENT_VARS_SI + ["share_pos", "share_neg"]
    ))

    full_sub = sub.dropna(subset=common_drop_vars).copy()
    baseline_sub = apply_baseline_filters(full_sub)

    return SampleBundle(full=full_sub, baseline=baseline_sub)


# ============================================================
# ESTIMATION
# ============================================================

def estimate_models(sub: pd.DataFrame, sample_name: str = "Baseline", include_tobit: bool = True) -> Dict[str, Any]:
    """Estimate main models for one sample."""
    R: Dict[str, Any] = {"_sample_name": sample_name}

    # Underplacement models
    X_si = build_X(sub, SENT_VARS_SI, add_fe=True)
    X_sh = build_X(sub, SENT_VARS_SHARES, add_fe=True)
    y_up = sub["underplacement"].astype(float)

    R["underpl_ols_hc3"] = fit_ols(y_up, X_si, cov_type="HC3")
    R["underpl_ols_cluster"] = fit_ols(y_up, X_si, cov_type="cluster", groups=sub["issuer_id"])
    R["underpl_fraclogit_cluster"] = fit_frac_logit(y_up, X_si, groups=sub["issuer_id"])

    if include_tobit:
        R["underpl_tobit"] = tobit_fit(y_up.values, X_si.values, lower=0.0)
        R["underpl_tobit_shares"] = tobit_fit(y_up.values, X_sh.values, lower=0.0)

    # Coupon models: conditional on non-missing coupon_reduction_bp.
    sub_cr = sub.dropna(subset=["coupon_reduction_bp"]).copy()
    X_si_cr = build_X(sub_cr, SENT_VARS_SI, add_fe=True)
    X_sh_cr = build_X(sub_cr, SENT_VARS_SHARES, add_fe=True)
    y_cr = sub_cr["coupon_reduction_bp"].astype(float)

    R["coupon_ols_hc3"] = fit_ols(y_cr, X_si_cr, cov_type="HC3")
    R["coupon_ols_cluster"] = fit_ols(y_cr, X_si_cr, cov_type="cluster", groups=sub_cr["issuer_id"])
    R["coupon_ols_shares_hc3"] = fit_ols(y_cr, X_sh_cr, cov_type="HC3")

    sub_cr["coupon_reduction_bp_w"] = winsorize_series(
        sub_cr["coupon_reduction_bp"], lower=WINSOR_LOWER, upper=WINSOR_UPPER
    )
    y_cr_w = sub_cr["coupon_reduction_bp_w"].astype(float)
    R["coupon_ols_winsor_hc3"] = fit_ols(y_cr_w, X_si_cr, cov_type="HC3")
    R["coupon_ols_winsor_cluster"] = fit_ols(y_cr_w, X_si_cr, cov_type="cluster", groups=sub_cr["issuer_id"])

    # Metadata
    R["_n_underpl"] = int(len(sub))
    R["_n_coupon"] = int(len(sub_cr))
    R["_n_issuers"] = int(sub["issuer"].nunique())
    R["_cn_si"] = list(X_si.columns)
    R["_cn_sh"] = list(X_sh.columns)
    R["_cn_si_cr"] = list(X_si_cr.columns)
    R["_cn_sh_cr"] = list(X_sh_cr.columns)

    return R


def drop_top_cooks_d(sub: pd.DataFrame, q: float = 0.99) -> pd.DataFrame:
    """Drop top q-tail of Cook's D for underplacement OLS on baseline specification."""
    if len(sub) < 20:
        return sub.copy()

    X = build_X(sub, SENT_VARS_SI, add_fe=True)
    y = sub["underplacement"].astype(float)
    ols = sm.OLS(y, X).fit()
    cooks = ols.get_influence().cooks_distance[0]
    threshold = np.nanquantile(cooks, q)
    return sub.loc[cooks <= threshold].copy()


def make_robustness_samples(bundle: SampleBundle) -> Dict[str, pd.DataFrame]:
    full = bundle.full
    base = bundle.baseline

    samples = {
        "Baseline": base,
        "Full cleaned sample": full,
        "No ВЭБ.РФ only": full.loc[~full["is_special_issuer"]].copy(),
        "No zero organizers only": full.loc[~full["zero_organizers"]].copy(),
        f"Baseline without top {int((1 - COOK_DROP_Q) * 100)}% Cook D": drop_top_cooks_d(base, q=COOK_DROP_Q),
    }

    # Drop empty or too-small samples defensively.
    return {k: v for k, v in samples.items() if len(v) >= 30}


def estimate_all_samples(bundle: SampleBundle) -> Dict[str, Dict[str, Any]]:
    samples = make_robustness_samples(bundle)
    all_results: Dict[str, Dict[str, Any]] = {}

    for name, df_s in samples.items():
        include_tobit = name == "Baseline"
        all_results[name] = estimate_models(df_s, sample_name=name, include_tobit=include_tobit)

    return all_results


# ============================================================
# EXCEL OUTPUT
# ============================================================

def model_stat_value(res: Any, stat_name: str) -> str:
    if isinstance(res, dict):
        if stat_name == "N":
            return str(res.get("n", "—"))
        if stat_name == "Log-lik":
            return f"{res.get('ll', np.nan):.2f}"
        if stat_name == "Sigma":
            return f"{res.get('sigma', np.nan):.4f}"
        if stat_name == "Cens/Uncens":
            return f"{res.get('n_cens', '—')}/{res.get('n_uncens', '—')}"
        if stat_name == "Converged":
            return str(res.get("converged", "—"))
        return "—"

    if stat_name == "N":
        return str(int(res.nobs)) if hasattr(res, "nobs") else "—"
    if stat_name == "R²":
        return f"{res.rsquared:.4f}" if hasattr(res, "rsquared") else "—"
    if stat_name == "Adj. R²":
        return f"{res.rsquared_adj:.4f}" if hasattr(res, "rsquared_adj") else "—"
    if stat_name == "Log-lik":
        return f"{res.llf:.2f}" if hasattr(res, "llf") else "—"
    if stat_name == "AIC":
        return f"{res.aic:.2f}" if hasattr(res, "aic") else "—"
    return "—"


def write_model_table(
    ws,
    title: str,
    model_specs: List[Tuple[str, str, List[str]]],
    results: Dict[str, Any],
    start_row: int = 1,
) -> int:
    """
    Write a readable regression table.

    model_specs: list of (result_key, title, column_names_for_model)

    Difference from the earlier version:
    - one model = one Excel column;
    - stars are appended to the coefficient in the same cell;
    - standard errors are shown in the row below;
    - model headers are explicit and multi-line.
    """
    hf = Font(bold=True, size=10, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    sfill = PatternFill("solid", fgColor="E2EFDA")

    n_models = len(model_specs)

    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=13, name="Arial")
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=1 + n_models)

    row = start_row + 2
    ws.row_dimensions[row].height = 78

    ws.cell(row=row, column=1, value="Variable").font = hf
    ws.cell(row=row, column=1).fill = hfill
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="center")

    for j, (_, ttl, _) in enumerate(model_specs, start=2):
        cell = ws.cell(row=row, column=j, value=ttl)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    row += 1

    all_vars: List[str] = []
    for _, _, cn in model_specs:
        for v in cn:
            # Fixed effects are not printed coefficient-by-coefficient.
            if v.startswith("year_") or v.startswith("q_"):
                continue
            if v not in all_vars:
                all_vars.append(v)

    for var in all_vars:
        label = VAR_LABELS.get(var, var)
        ws.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10)
        ws.cell(row=row + 1, column=1, value="").font = Font(name="Arial", size=9, color="666666")

        for j, (key, _, cn) in enumerate(model_specs, start=2):
            res = results[key]
            coef, se, p = get_csp(res, var, cn)

            if coef is None:
                ws.cell(row=row, column=j, value="—").font = Font(name="Arial", size=10, color="999999")
                ws.cell(row=row + 1, column=j, value="").font = Font(name="Arial", size=9, color="999999")
                continue

            stars = sig_stars(p)
            coef_text = f"{coef:.4f}{stars}"
            se_text = f"({se:.4f})" if se is not None and pd.notna(se) else "(—)"

            coef_cell = ws.cell(row=row, column=j, value=coef_text)
            coef_cell.font = Font(name="Arial", size=10, bold=(p is not None and p < 0.10))
            coef_cell.alignment = Alignment(horizontal="right")
            if p is not None and p < 0.10:
                coef_cell.fill = sfill

            se_cell = ws.cell(row=row + 1, column=j, value=se_text)
            se_cell.font = Font(name="Arial", size=9, color="666666")
            se_cell.alignment = Alignment(horizontal="right")

        row += 2

    row += 1
    ws.cell(row=row, column=1, value="Model statistics").font = Font(bold=True, name="Arial")
    row += 1

    for stat_name in ["N", "R²", "Adj. R²", "Log-lik", "AIC", "Sigma", "Cens/Uncens", "Converged"]:
        ws.cell(row=row, column=1, value=stat_name).font = Font(name="Arial", size=10, italic=True)
        for j, (key, _, _) in enumerate(model_specs, start=2):
            ws.cell(row=row, column=j, value=model_stat_value(results[key], stat_name)).font = Font(name="Arial", size=10)
        row += 1

    row += 1
    notes = [
        "*** p<0.01, ** p<0.05, * p<0.1. Standard errors in parentheses.",
        "OLS HC3 = heteroskedasticity-robust SE; OLS cluster = SE clustered by issuer.",
        "Fractional logit is used for bounded underplacement; zeros are kept as true zeros.",
        "Underplacement = 1 − volume_ratio.",
        "Missing SI/buzz/shares values are set to 0.",
        "Baseline excludes ВЭБ.РФ/special technical placements, num_organizers=0, and book_date > placement_date.",
        "Coupon models are estimated on the subsample with non-missing coupon_reduction_bp.",
        "Year fixed effects included." if USE_YEAR_FE and not USE_QUARTER_FE else "Quarter fixed effects included." if USE_QUARTER_FE else "No time fixed effects included.",
    ]

    for note in notes:
        ws.cell(row=row, column=1, value=note).font = Font(name="Arial", size=9, italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1 + n_models)
        row += 1

    # Freeze panes and make the sheet readable immediately after opening.
    ws.freeze_panes = ws["B4"]
    ws.column_dimensions["A"].width = 36
    for i in range(2, 2 + n_models):
        ws.column_dimensions[get_column_letter(i)].width = 22

    return row


def write_robustness_snapshot(ws, all_results: Dict[str, Dict[str, Any]]) -> None:
    hf = Font(bold=True, size=11, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")

    ws["A1"] = "Robustness Snapshot: Sentiment Index Coefficient"
    ws["A1"].font = Font(bold=True, size=13, name="Arial")

    headers = [
        "Sample",
        "N underpl",
        "N coupon",
        "SI: underpl OLS HC3",
        "p",
        "SI: underpl cluster",
        "p",
        "SI: frac logit",
        "p",
        "SI: coupon OLS HC3",
        "p",
        "SI: coupon cluster",
        "p",
    ]

    row = 3
    for j, h in enumerate(headers, start=1):
        ws.cell(row=row, column=j, value=h).font = hf
        ws.cell(row=row, column=j).fill = hfill
        ws.cell(row=row, column=j).alignment = Alignment(wrap_text=True, horizontal="center")
    row += 1

    def pair(res, var, cn):
        coef, _, p = get_csp(res, var, cn)
        if coef is None:
            return "—", "—"
        return round(coef, 4), round(p, 4) if p is not None and pd.notna(p) else "—"

    for sample_name, R in all_results.items():
        cn_si = R["_cn_si"]
        vals = [
            sample_name,
            R["_n_underpl"],
            R["_n_coupon"],
        ]

        for key in [
            "underpl_ols_hc3",
            "underpl_ols_cluster",
            "underpl_fraclogit_cluster",
            "coupon_ols_hc3",
            "coupon_ols_cluster",
        ]:
            coef, p = pair(R[key], "si", cn_si)
            vals.extend([coef, p])

        for j, v in enumerate(vals, start=1):
            ws.cell(row=row, column=j, value=v).font = Font(name="Arial", size=10)
        row += 1

    ws.column_dimensions["A"].width = 32
    for i in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16


def write_descriptive_stats(ws, sub: pd.DataFrame, title: str) -> None:
    hf = Font(bold=True, size=11, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, name="Arial")

    headers = ["Variable", "N", "Mean", "Std", "Min", "p25", "Median", "p75", "Max"]
    row = 3
    for j, h in enumerate(headers, start=1):
        ws.cell(row=row, column=j, value=h).font = hf
        ws.cell(row=row, column=j).fill = hfill
    row += 1

    vars_to_show = [
        "coupon_reduction_bp",
        "underplacement",
        "volume_ratio",
        "si",
        "log_buzz",
        "share_pos",
        "share_neg",
        "num_organizers",
        "is_debut",
        "is_floater",
        "hist_ever_reduced",
        "hist_avg_volume_ratio",
        "rating_num",
        "ofz_yield",
        "rvi",
        "log_dur",
        "has_put",
    ]

    for var in vars_to_show:
        if var not in sub.columns:
            continue
        s = pd.to_numeric(sub[var], errors="coerce")
        if s.notna().sum() == 0:
            continue

        values = [
            int(s.notna().sum()),
            s.mean(),
            s.std(),
            s.min(),
            s.quantile(0.25),
            s.median(),
            s.quantile(0.75),
            s.max(),
        ]

        ws.cell(row=row, column=1, value=VAR_LABELS.get(var, var)).font = Font(name="Arial", size=10)
        for j, val in enumerate(values, start=2):
            ws.cell(row=row, column=j, value=round(float(val), 4) if pd.notna(val) else "—").number_format = "0.0000"
        row += 1

    ws.column_dimensions["A"].width = 36
    for i in range(2, 10):
        ws.column_dimensions[get_column_letter(i)].width = 12


def write_correlation_matrix(ws, sub: pd.DataFrame) -> None:
    hf = Font(bold=True, size=11, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")

    ws["A1"] = "Correlation Matrix — Baseline Sample"
    ws["A1"].font = Font(bold=True, size=13, name="Arial")

    corr_vars = [
        "coupon_reduction_bp",
        "underplacement",
        "si",
        "log_buzz",
        "num_organizers",
        "is_floater",
        "hist_ever_reduced",
        "hist_avg_volume_ratio",
        "rating_num",
        "ofz_yield",
        "rvi",
    ]
    corr_vars = [v for v in corr_vars if v in sub.columns]
    labels = [VAR_LABELS.get(v, v) for v in corr_vars]
    cm = sub[corr_vars].corr()

    row = 3
    for j, label in enumerate(labels, start=2):
        c = ws.cell(row=row, column=j, value=label)
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(text_rotation=45, wrap_text=True)
    row += 1

    for i, (var, label) in enumerate(zip(corr_vars, labels)):
        ws.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10)
        ws.cell(row=row, column=1).fill = hfill
        for j, var2 in enumerate(corr_vars, start=2):
            val = cm.loc[var, var2]
            cell = ws.cell(row=row, column=j, value=round(float(val), 3) if pd.notna(val) else "—")
            cell.number_format = "0.000"
            if i != (j - 2) and pd.notna(val) and abs(val) > 0.3:
                cell.font = Font(name="Arial", size=10, bold=True)
        row += 1

    ws.column_dimensions["A"].width = 32
    for i in range(2, len(corr_vars) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 12


def build_excel(all_results: Dict[str, Dict[str, Any]], bundle: SampleBundle, path: str) -> None:
    """Build Excel workbook with readable regression tables."""
    wb = Workbook()

    Rb = all_results["Baseline"]
    cn_si = Rb["_cn_si"]
    cn_sh = Rb["_cn_sh"]
    cn_si_cr = Rb["_cn_si_cr"]
    cn_sh_cr = Rb["_cn_sh_cr"]

    n_underpl = Rb["_n_underpl"]
    n_coupon = Rb["_n_coupon"]
    fe_label = "quarter" if USE_QUARTER_FE else "year" if USE_YEAR_FE else "none"

    # ------------------------------------------------------------
    # Main results: Tobit is intentionally not in the main table.
    # ------------------------------------------------------------
    main_models = [
        (
            "coupon_ols_hc3",
            f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: HC3 robust\nFE: {fe_label}\nN={n_coupon}",
            cn_si_cr,
        ),
        (
            "coupon_ols_cluster",
            f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={n_coupon}",
            cn_si_cr,
        ),
        (
            "underpl_ols_hc3",
            f"Y: underplacement\nDV: raw share\nModel: OLS\nSE: HC3 robust\nFE: {fe_label}\nN={n_underpl}",
            cn_si,
        ),
        (
            "underpl_ols_cluster",
            f"Y: underplacement\nDV: raw share\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={n_underpl}",
            cn_si,
        ),
        (
            "underpl_fraclogit_cluster",
            f"Y: underplacement\nDV: fractional share\nModel: fractional logit\nSE: cluster or HC1 fallback\nFE: {fe_label}\nN={n_underpl}",
            cn_si,
        ),
    ]

    ws1 = wb.active
    ws1.title = "Main Results"
    write_model_table(
        ws1,
        f"Table 1. Main Results — Baseline Sample ({YEAR_FROM}+)",
        main_models,
        Rb,
    )

    # ------------------------------------------------------------
    # Robustness snapshot across samples.
    # ------------------------------------------------------------
    ws2 = wb.create_sheet("Robustness Snapshot")
    write_robustness_snapshot(ws2, all_results)

    # ------------------------------------------------------------
    # Coupon robustness.
    # ------------------------------------------------------------
    coupon_robustness_models = [
        (
            "coupon_ols_hc3",
            f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: HC3 robust\nFE: {fe_label}\nN={n_coupon}",
            cn_si_cr,
        ),
        (
            "coupon_ols_cluster",
            f"Y: coupon reduction, bp\nDV: raw\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={n_coupon}",
            cn_si_cr,
        ),
        (
            "coupon_ols_winsor_hc3",
            f"Y: coupon reduction, bp\nDV: winsor 1/99%\nModel: OLS\nSE: HC3 robust\nFE: {fe_label}\nN={n_coupon}",
            cn_si_cr,
        ),
        (
            "coupon_ols_winsor_cluster",
            f"Y: coupon reduction, bp\nDV: winsor 1/99%\nModel: OLS\nSE: clustered by issuer\nFE: {fe_label}\nN={n_coupon}",
            cn_si_cr,
        ),
        (
            "coupon_ols_shares_hc3",
            f"Y: coupon reduction, bp\nX: pos/neg shares\nModel: OLS\nSE: HC3 robust\nFE: {fe_label}\nN={n_coupon}",
            cn_sh_cr,
        ),
    ]

    ws3 = wb.create_sheet("Coupon Robustness")
    write_model_table(
        ws3,
        "Table 2. Coupon Robustness — Baseline Sample",
        coupon_robustness_models,
        Rb,
    )

    # ------------------------------------------------------------
    # Tobit appendix.
    # ------------------------------------------------------------
    tobit_models = [
        (
            "underpl_tobit",
            f"Y: underplacement\nModel: Tobit\nX: SI + buzz\nSE: MLE Hessian\nFE: {fe_label}\nN={n_underpl}",
            cn_si,
        ),
        (
            "underpl_tobit_shares",
            f"Y: underplacement\nModel: Tobit\nX: pos/neg shares\nSE: MLE Hessian\nFE: {fe_label}\nN={n_underpl}",
            cn_sh,
        ),
    ]

    ws4 = wb.create_sheet("Tobit Appendix")
    write_model_table(
        ws4,
        "Appendix. Tobit Robustness — Baseline Sample",
        tobit_models,
        Rb,
    )

    # ------------------------------------------------------------
    # Descriptive statistics and correlations.
    # ------------------------------------------------------------
    ws5 = wb.create_sheet("Descriptive Baseline")
    write_descriptive_stats(
        ws5,
        bundle.baseline,
        f"Descriptive Statistics — Baseline Sample (N={len(bundle.baseline)})",
    )

    ws6 = wb.create_sheet("Descriptive Full")
    write_descriptive_stats(
        ws6,
        bundle.full,
        f"Descriptive Statistics — Full Cleaned Sample (N={len(bundle.full)})",
    )

    ws7 = wb.create_sheet("Correlations")
    write_correlation_matrix(ws7, bundle.baseline)

    wb.save(path)


# ============================================================
# CONSOLE OUTPUT
# ============================================================

def print_key_result(res: Any, var: str, cn: List[str]) -> str:
    coef, se, p = get_csp(res, var, cn)
    if coef is None:
        return "—"
    return f"{coef:.4f} ({se:.4f}), p={p:.4f} {sig_stars(p)}"


def print_sample_summary(bundle: SampleBundle) -> None:
    full = bundle.full
    base = bundle.baseline

    print("\n" + "=" * 70)
    print("SAMPLE / DATA QUALITY")
    print("=" * 70)
    print(f"Full cleaned sample       N={len(full)}, coupon_N={full['coupon_reduction_bp'].notna().sum()}, issuers={full['issuer'].nunique()}")
    print(f"Baseline sample           N={len(base)}, coupon_N={base['coupon_reduction_bp'].notna().sum()}, issuers={base['issuer'].nunique()}")
    print(f"book_date > placement     N={int(full['book_after_placement'].sum())}")
    print(f"Special issuer ВЭБ.РФ     N={int(full['is_special_issuer'].sum())}")
    print(f"num_organizers = 0        N={int(full['zero_organizers'].sum())}")
    print(f"underpl > 1% baseline     N={int((base['underplacement'] > 0.01).sum())}")
    print(f"coupon > 0.5bp baseline   N={int((base.dropna(subset=['coupon_reduction_bp'])['coupon_reduction_bp'] > 0.5).sum())}")


def print_results(all_results: Dict[str, Dict[str, Any]]) -> None:
    print("\n" + "=" * 70)
    print("REGRESSION RESULTS — SI SNAPSHOT")
    print("=" * 70)

    for sample_name, R in all_results.items():
        cn_si = R["_cn_si"]
        print(f"\n--- {sample_name} ---")
        print(f"N underpl={R['_n_underpl']}, N coupon={R['_n_coupon']}, issuers={R['_n_issuers']}")
        print(f"SI underpl OLS HC3:      {print_key_result(R['underpl_ols_hc3'], 'si', cn_si)}")
        print(f"SI underpl cluster:      {print_key_result(R['underpl_ols_cluster'], 'si', cn_si)}")
        print(f"SI underpl frac logit:   {print_key_result(R['underpl_fraclogit_cluster'], 'si', cn_si)}")
        print(f"SI coupon OLS HC3:       {print_key_result(R['coupon_ols_hc3'], 'si', cn_si)}")
        print(f"SI coupon cluster:       {print_key_result(R['coupon_ols_cluster'], 'si', cn_si)}")
        if "underpl_tobit" in R:
            tb = R["underpl_tobit"]
            print(f"Tobit convergence:       {tb['converged']} | {tb['message']}")


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    print("Загрузка и подготовка данных...")
    bundle = load_and_prepare(INPUT_FILE, YEAR_FROM)
    print_sample_summary(bundle)

    print("\nОценка моделей...")
    all_results = estimate_all_samples(bundle)
    print_results(all_results)

    print("\nСохранение Excel...")
    build_excel(all_results, bundle, OUTPUT_FILE)
    print(f"\n✅ Таблицы сохранены: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
