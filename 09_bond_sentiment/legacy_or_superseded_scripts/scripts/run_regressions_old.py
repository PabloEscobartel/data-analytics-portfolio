#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Регрессионный анализ успешности первичных размещений облигаций.
Модели: OLS (HC3, clustered SE), Tobit, Fractional Logit.
Выход: regression_results.xlsx с таблицами и описательной статистикой.

Требования: pip install pandas numpy statsmodels scipy openpyxl
Запуск: python run_regressions.py
"""

import pandas as pd
import numpy as np
import statsmodels.api as sm
from scipy import stats, optimize
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import warnings
warnings.filterwarnings('ignore')

# ============================================================
# НАСТРОЙКИ — МЕНЯТЬ ТУТ
# ============================================================

INPUT_FILE = "regression_dataset/regression_ready.csv"
OUTPUT_FILE = "regression_results.xlsx"
YEAR_FROM = 2022

RATING_MAP = {
    'AAA':15,'AA+':14,'AA':13,'AA-':12,'A+':11,'A':10,'A-':9,
    'BBB+':8,'BBB':7,'BBB-':6,'BB+':5,'BB':4,'BB-':3,
    'B+':2,'B':1,'B-':0,'C':-1
}

VAR_LABELS = {
    'const':'Constant', 'num_organizers':'Number of organizers',
    'is_debut':'Debut (dummy)', 'hist_ever_reduced':'Prior coupon tightening (dummy)',
    'hist_avg_volume_ratio':'Hist. avg volume ratio', 'rating_num':'Credit rating (ordinal)',
    'key_rate':'CBR key rate, %', 'log_dur':'ln(Duration, days)',
    'has_put':'Has put option (dummy)', 'si':'Sentiment Index (SI)',
    'log_buzz':'ln(Buzz)',
    'share_pos':'Share positive', 'share_neg':'Share negative',
    'ofz_yield': 'OFZ matched yield, %', 'rvi': 'RVI (volatility index)',
    'is_floater': 'Floating rate (dummy)',
}

BASE_VARS = ['num_organizers','is_debut','hist_ever_reduced','hist_avg_volume_ratio',
             'rating_num','ofz_yield', 'rvi','log_dur','has_put', 'is_floater']
SENT_VARS_SI = BASE_VARS + ['si','log_buzz']
SENT_VARS_SHARES = BASE_VARS + ['share_pos','share_neg','log_buzz']

# ============================================================
# TOBIT MLE
# ============================================================

def tobit_fit(y, X, lower=0):
    def neg_ll(params):
        beta, log_sigma = params[:-1], params[-1]
        sigma = np.exp(log_sigma)
        xb = X @ beta
        uncens = y > lower
        ll = 0.0
        if uncens.any():
            r = (y[uncens] - xb[uncens]) / sigma
            ll += np.sum(-0.5*np.log(2*np.pi) - np.log(sigma) - 0.5*r**2)
        cens = ~uncens
        if cens.any():
            z = (lower - xb[cens]) / sigma
            ll += np.sum(np.log(stats.norm.cdf(z) + 1e-12))
        return -ll

    ols = sm.OLS(y, X).fit()
    p0 = np.append(ols.params, np.log(np.std(ols.resid)))
    res = optimize.minimize(neg_ll, p0, method='Nelder-Mead',
                            options={'maxiter':80000,'xatol':1e-9})
    beta = res.x[:-1]
    sigma = np.exp(res.x[-1])

    eps = 1e-5
    n_p = len(res.x)
    H = np.zeros((n_p, n_p))
    for i in range(n_p):
        def fi(p, _i=i):
            return optimize.approx_fprime(p, neg_ll, eps)[_i]
        H[i] = optimize.approx_fprime(res.x, fi, eps)
    try:
        cov = np.linalg.inv(H)
        se = np.sqrt(np.abs(np.diag(cov)))
    except np.linalg.LinAlgError:
        se = np.full(n_p, np.nan)

    se_beta = se[:-1]
    z_vals = np.where(se_beta > 0, beta / se_beta, 0)
    p_vals = 2 * (1 - stats.norm.cdf(np.abs(z_vals)))
    return {'coef':beta,'se':se_beta,'z':z_vals,'p':p_vals,
            'sigma':sigma,'ll':-res.fun,'n':len(y),
            'n_cens':int((y<=lower).sum()),'n_uncens':int((y>lower).sum()),
            'converged':res.success}

# ============================================================
# ЗАГРУЗКА И ПОДГОТОВКА
# ============================================================

def first_existing_col(df, candidates):
    for col in candidates:
        if col in df.columns:
            return col
    raise KeyError(f"Не найдена ни одна из колонок: {', '.join(candidates)}")


def require_columns(df, columns):
    missing = [col for col in columns if col not in df.columns]
    if missing:
        raise KeyError(f"В {INPUT_FILE} нет нужных колонок: {', '.join(missing)}")


def load_and_prepare(path, year_from):
    df = pd.read_csv(path, low_memory=False)
    require_columns(df, [
        'placement_date', 'coupon_reduction_bp', 'volume_ratio',
        'num_organizers', 'is_debut', 'hist_ever_reduced',
        'hist_avg_volume_ratio', 'final_rating', 'ofz_matched_yield',
        'rvi_close', 'has_put', 'issuer',
    ])

    df['year'] = pd.to_datetime(df['placement_date']).dt.year
    sub = df[df['year'] >= year_from].copy()

    si_col = first_existing_col(sub, ['si_relevant', 'si'])
    buzz_col = first_existing_col(sub, ['log_buzz_relevant', 'log_buzz_all', 'log_buzz'])
    share_neg_col = first_existing_col(sub, ['share_negative_relevant', 'share_negative'])
    share_pos_col = first_existing_col(sub, ['share_positive_relevant', 'share_positive'])
    sub['si'] = sub[si_col].fillna(0)
    sub['log_buzz'] = sub[buzz_col].fillna(0)
    sub['share_neg'] = sub[share_neg_col].fillna(0)
    sub['share_pos'] = sub[share_pos_col].fillna(0)
    sub['underplacement'] = 1 - sub['volume_ratio']
    if 'duration_days' in sub.columns:
        sub['log_dur'] = np.log(pd.to_numeric(sub['duration_days'], errors='coerce').clip(lower=1))
    elif 'log_term' in sub.columns:
        sub['log_dur'] = pd.to_numeric(sub['log_term'], errors='coerce')
    else:
        term_col = first_existing_col(sub, ['term'])
        sub['log_dur'] = np.log(pd.to_numeric(sub[term_col], errors='coerce').clip(lower=1))
    sub['rating_num'] = sub['final_rating'].map(RATING_MAP)
    sub['ofz_yield'] = sub['ofz_matched_yield']
    sub['rvi'] = sub['rvi_close']
    sub['issuer_id'] = sub['issuer'].astype('category').cat.codes

    base_vars = ['underplacement', 'issuer_id'] + SENT_VARS_SI + SENT_VARS_SHARES
    base_vars = [v for v in base_vars if v != 'coupon_reduction_bp']
    return sub.dropna(subset=base_vars).copy()

# ============================================================
# ОЦЕНКА МОДЕЛЕЙ
# ============================================================

def estimate_all(sub):
    R = {}
    y_cr = sub['coupon_reduction_bp'].values
    y_up = sub['underplacement'].values
    X_si = sm.add_constant(sub[SENT_VARS_SI])
    X_sh = sm.add_constant(sub[SENT_VARS_SHARES])

    R['M1_OLS_coupon_HC3'] = sm.OLS(y_cr, X_si).fit(cov_type='HC3')
    R['M2_OLS_coupon_cluster'] = sm.OLS(y_cr, X_si).fit(
        cov_type='cluster', cov_kwds={'groups':sub['issuer_id']})
    R['M3_OLS_underpl_HC3'] = sm.OLS(y_up, X_si).fit(cov_type='HC3')
    R['M4_OLS_underpl_cluster'] = sm.OLS(y_up, X_si).fit(
        cov_type='cluster', cov_kwds={'groups':sub['issuer_id']})
    R['M5_Tobit_underpl'] = tobit_fit(y_up, X_si.values, lower=0)
    y_fl = sub['underplacement'].clip(0.001, 0.999)
    R['M6_FracLogit_underpl'] = sm.GLM(y_fl, X_si,
        family=sm.families.Binomial()).fit(cov_type='HC1')
    R['M7_OLS_coupon_shares'] = sm.OLS(y_cr, X_sh).fit(cov_type='HC3')
    R['M8_Tobit_underpl_shares'] = tobit_fit(y_up, X_sh.values, lower=0)
    return R

# ============================================================
# EXCEL
# ============================================================

def sig_stars(p):
    if np.isnan(p): return ''
    if p < 0.01: return '***'
    if p < 0.05: return '**'
    if p < 0.1: return '*'
    return ''

def get_csp(res, var, cn):
    if isinstance(res, dict):
        if var not in cn: return None,None,None
        i = cn.index(var)
        return res['coef'][i], res['se'][i], res['p'][i]
    else:
        if var in res.params.index:
            return res.params[var], res.bse[var], res.pvalues[var]
        return None,None,None

def write_model_table(ws, title, models, results, start_row=1):
    hf = Font(bold=True, size=11, name='Arial')
    hfill = PatternFill('solid', fgColor='D9E1F2')
    sfill = PatternFill('solid', fgColor='E2EFDA')

    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=13, name='Arial')
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2+len(models)*2)

    row = start_row + 2
    ws.cell(row=row, column=1, value='Variable').font = hf
    ws.cell(row=row, column=1).fill = hfill
    for j,(key,ttl,spec,cn) in enumerate(models):
        c = 2+j*2
        ws.cell(row=row, column=c, value=ttl).font = hf
        ws.cell(row=row, column=c).fill = hfill
        ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, horizontal='center')
        ws.merge_cells(start_row=row, start_column=c, end_row=row, end_column=c+1)
    row += 1

    all_v = []
    for _,_,_,cn in models:
        for v in cn:
            if v not in all_v: all_v.append(v)

    for var in all_v:
        lab = VAR_LABELS.get(var, var)
        ws.cell(row=row, column=1, value=lab).font = Font(name='Arial', size=10)
        for j,(key,_,_,cn) in enumerate(models):
            c = 2+j*2
            coef,se,p = get_csp(results[key], var, cn)
            if coef is None:
                ws.cell(row=row, column=c, value='—').font = Font(name='Arial',size=10,color='999999')
                continue
            st = sig_stars(p)
            cell = ws.cell(row=row, column=c, value=round(coef,4))
            cell.number_format = '0.0000'
            cell.font = Font(name='Arial',size=10,bold=(p<0.1))
            if p < 0.1: cell.fill = sfill
            ws.cell(row=row, column=c+1, value=st).font = Font(name='Arial',size=10)
            ws.cell(row=row+1, column=c, value=f'({se:.4f})').font = Font(name='Arial',size=9,color='666666')
        row += 2

    row += 1
    ws.cell(row=row, column=1, value='Model statistics').font = Font(bold=True, name='Arial')
    row += 1
    for sn, fn in [
        ('N', lambda r: int(r.nobs) if hasattr(r,'nobs') else r['n']),
        ('R²', lambda r: f"{r.rsquared:.4f}" if hasattr(r,'rsquared') else '—'),
        ('Adj. R²', lambda r: f"{r.rsquared_adj:.4f}" if hasattr(r,'rsquared_adj') else '—'),
        ('Log-lik', lambda r: f"{r['ll']:.2f}" if isinstance(r,dict) else '—'),
        ('Sigma', lambda r: f"{r['sigma']:.4f}" if isinstance(r,dict) else '—'),
        ('Cens/Uncens', lambda r: f"{r['n_cens']}/{r['n_uncens']}" if isinstance(r,dict) else '—'),
        ('SE type', lambda r: 'Tobit MLE' if isinstance(r,dict) else
                    ('Clustered' if 'cluster' in key else 'HC3' if 'HC3' in key else 'HC1')),
    ]:
        ws.cell(row=row, column=1, value=sn).font = Font(name='Arial',size=10,italic=True)
        for j,(key,_,_,_) in enumerate(models):
            ws.cell(row=row, column=2+j*2, value=str(fn(results[key]))).font = Font(name='Arial',size=10)
        row += 1

    row += 1
    for note in ['*** p<0.01, ** p<0.05, * p<0.1. SE in parentheses.',
                 'Underplacement = 1 − volume_ratio. Tobit left-censored at 0.',
                 'Sentiment=0 and buzz=0 for issues with no media coverage.']:
        ws.cell(row=row, column=1, value=note).font = Font(name='Arial',size=9,italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2+len(models)*2)
        row += 1

    ws.column_dimensions['A'].width = 34
    for i in range(2, 2+len(models)*2+1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    return row

def build_excel(R, sub, path):
    wb = Workbook()
    cn_si = ['const'] + SENT_VARS_SI
    cn_sh = ['const'] + SENT_VARS_SHARES

    # Sheet 1
    ws1 = wb.active; ws1.title = 'Main Results'
    models1 = [
        ('M1_OLS_coupon_HC3','OLS\ncoupon_reduction',SENT_VARS_SI,cn_si),
        ('M3_OLS_underpl_HC3','OLS\nunderplacement',SENT_VARS_SI,cn_si),
        ('M5_Tobit_underpl','Tobit\nunderplacement',SENT_VARS_SI,cn_si),
    ]
    write_model_table(ws1, f'Table 1. Determinants of Bond Placement Outcomes ({YEAR_FROM}+)', models1, R)

    # Sheet 2
    ws2 = wb.create_sheet('Robustness')
    models2 = [
        ('M2_OLS_coupon_cluster','OLS coupon\nclustered SE',SENT_VARS_SI,cn_si),
        ('M4_OLS_underpl_cluster','OLS underpl\nclustered SE',SENT_VARS_SI,cn_si),
        ('M6_FracLogit_underpl','Frac.Logit\nunderpl',SENT_VARS_SI,cn_si),
        ('M8_Tobit_underpl_shares','Tobit underpl\npos/neg',SENT_VARS_SHARES,cn_sh),
    ]
    write_model_table(ws2, 'Table 2. Robustness Checks', models2, R)

    # Sheet 3: Descriptive
    ws3 = wb.create_sheet('Descriptive Stats')
    hf = Font(bold=True,size=11,name='Arial')
    hfill = PatternFill('solid',fgColor='D9E1F2')
    ws3['A1'] = f'Table 3. Descriptive Statistics (N={len(sub)})'
    ws3['A1'].font = Font(bold=True,size=13,name='Arial')
    row = 3
    for j,h in enumerate(['Variable','N','Mean','Std','Min','p25','Median','p75','Max']):
        ws3.cell(row=row,column=j+1,value=h).font = hf
        ws3.cell(row=row,column=j+1).fill = hfill
    row += 1
    desc_vars = [
        'coupon_reduction_bp', 'underplacement', 'volume_ratio', 'si', 'log_buzz',
        'share_pos', 'share_neg', 'num_organizers', 'is_debut', 'hist_ever_reduced',
        'hist_avg_volume_ratio', 'rating_num', 'key_rate', 'ofz_yield', 'rvi',
        'log_dur', 'has_put',
    ]
    for var in [v for v in desc_vars if v in sub.columns]:
        s = sub[var]
        ws3.cell(row=row,column=1,value=VAR_LABELS.get(var,var)).font = Font(name='Arial',size=10)
        for j,v in enumerate([int(s.notna().sum()),s.mean(),s.std(),s.min(),
                              s.quantile(.25),s.median(),s.quantile(.75),s.max()]):
            ws3.cell(row=row,column=j+2,value=round(v,3)).number_format='0.000'
        row += 1
    ws3.column_dimensions['A'].width = 34
    for i in range(2,10): ws3.column_dimensions[get_column_letter(i)].width = 12

    # Sheet 4: Correlations
    ws4 = wb.create_sheet('Correlations')
    ws4['A1'] = 'Table 4. Correlation Matrix'
    ws4['A1'].font = Font(bold=True,size=13,name='Arial')
    cv = [
        'coupon_reduction_bp', 'underplacement', 'si', 'log_buzz', 'num_organizers',
        'hist_ever_reduced', 'hist_avg_volume_ratio', 'rating_num', 'key_rate',
        'ofz_yield', 'rvi',
    ]
    cv = [v for v in cv if v in sub.columns]
    cl = [VAR_LABELS.get(v,v) for v in cv]
    cm = sub[cv].corr()
    row = 3
    for j,l in enumerate(cl):
        c = ws4.cell(row=row,column=j+2,value=l)
        c.font = Font(bold=True,size=9,name='Arial'); c.fill = hfill
        c.alignment = Alignment(text_rotation=45,wrap_text=True)
    row += 1
    for i,(var,lab) in enumerate(zip(cv,cl)):
        ws4.cell(row=row,column=1,value=lab).font = Font(name='Arial',size=10)
        ws4.cell(row=row,column=1).fill = hfill
        for j,v2 in enumerate(cv):
            val = cm.loc[var,v2]
            c = ws4.cell(row=row,column=j+2,value=round(val,3))
            c.number_format='0.000'
            c.font = Font(name='Arial',size=10,bold=(abs(val)>0.3 and i!=j))
        row += 1
    ws4.column_dimensions['A'].width = 30
    for i in range(2,12): ws4.column_dimensions[get_column_letter(i)].width = 12

    wb.save(path)

# ============================================================
# КОНСОЛЬНЫЙ ВЫВОД
# ============================================================

def print_results(R, sub):
    cn = ['const'] + SENT_VARS_SI
    n = len(sub)
    print(f"\n{'='*70}")
    print(f"РЕЗУЛЬТАТЫ (N={n}, период {YEAR_FROM}+)")
    print(f"{'='*70}")
    print(f"coupon_reduction > 0: {(sub['coupon_reduction_bp']>0.5).sum()}")
    print(f"underplacement > 0: {(sub['underplacement']>0.01).sum()}")
    print(f"Эмитентов: {sub['issuer'].nunique()}")
    for name, key in [('Tobit underplacement','M5_Tobit_underpl'),
                      ('OLS coupon (HC3)','M1_OLS_coupon_HC3'),
                      ('OLS underplacement (HC3)','M3_OLS_underpl_HC3'),
                      ('OLS coupon clustered','M2_OLS_coupon_cluster'),
                      ('OLS underpl clustered','M4_OLS_underpl_cluster'),
                      ('Fractional Logit underpl','M6_FracLogit_underpl'),
                      ('OLS coupon shares','M7_OLS_coupon_shares'),
                      ('Tobit underpl shares','M8_Tobit_underpl_shares')]:
        res = R[key]
        cn_use = cn if 'shares' not in key else ['const'] + SENT_VARS_SHARES
        print(f"\n--- {name} ---")
        if isinstance(res, dict):
            print(f"LL={res['ll']:.2f}, sigma={res['sigma']:.4f}, cens={res['n_cens']}, uncens={res['n_uncens']}")
        else:
            if hasattr(res, 'rsquared'):
                print(f"R²={res.rsquared:.4f}, Adj.R²={res.rsquared_adj:.4f}")
            elif hasattr(res, 'llf'):
                print(f"LL={res.llf:.2f}, AIC={res.aic:.2f}")
        print(f"{'Variable':35s} {'Coef':>10s} {'SE':>10s} {'p':>8s}")
        print("-"*65)
        for i,var in enumerate(cn_use):
            coef,se,p = get_csp(res, var, cn_use)
            if coef is None: continue
            lab = VAR_LABELS.get(var,var)
            print(f"{lab:35s} {coef:10.4f} {se:10.4f} {p:8.4f} {sig_stars(p)}")

# ============================================================
# MAIN
# ============================================================

def main():
    print("Загрузка данных...")
    sub = load_and_prepare(INPUT_FILE, YEAR_FROM)
    print(f"N = {len(sub)}")

    print("Оценка моделей...")
    R = estimate_all(sub)

    print_results(R, sub)
    build_excel(R, sub, OUTPUT_FILE)
    print(f"\n✅ Таблицы сохранены: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
