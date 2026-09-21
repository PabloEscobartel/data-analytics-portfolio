#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Регрессионный анализ успешности первичных размещений облигаций.
Модели: OLS (HC3, clustered SE), Tobit, Fractional Logit.
Выход: regression_results.xlsx

Требования: pip install pandas numpy statsmodels scipy openpyxl
Запуск: python run_regressions.py
"""

import pandas as pd
import numpy as np
import statsmodels.api as sm
from scipy import stats, optimize
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import warnings
warnings.filterwarnings('ignore')

# ============================================================
# НАСТРОЙКИ
# ============================================================

INPUT_FILE = "regression_dataset/regression_ready.csv"
OUTPUT_FILE = "regression_results.xlsx"
YEAR_FROM = 2018

BASELINE_EXCLUDE_ISSUERS = {'ВЭБ.РФ'}
BASELINE_DROP_ZERO_ORGANIZERS = True
BASELINE_DROP_BOOK_AFTER_PLACEMENT = True

RATING_MAP = {
    'AAA':15,'AA+':14,'AA':13,'AA-':12,'A+':11,'A':10,'A-':9,
    'BBB+':8,'BBB':7,'BBB-':6,'BB+':5,'BB':4,'BB-':3,
    'B+':2,'B':1,'B-':0,'C':-1
}

VAR_LABELS = {
    'const':'Constant', 'num_organizers':'Number of organizers',
    'is_debut':'Debut (dummy)', 'hist_ever_reduced':'Prior tightening (dummy)',
    'hist_avg_volume_ratio':'Hist. avg volume ratio', 'rating_num':'Credit rating (ordinal)',
    'log_dur':'ln(Term to exit, days)', 'has_put':'Has put option (dummy)',
    'is_floater':'Floating rate (dummy)',
    'si':'Sentiment Index (SI)', 'log_buzz':'ln(Buzz)',
    'share_pos':'Share positive', 'share_neg':'Share negative',
    'ofz_yield':'OFZ matched yield, %', 'rvi':'RVI (volatility index)',
}

BASE_VARS = ['num_organizers', 'is_debut', 'hist_ever_reduced', 'hist_avg_volume_ratio',
             'rating_num', 'ofz_yield', 'rvi', 'log_dur', 'has_put', 'is_floater']
SENT_VARS_SI = BASE_VARS + ['si', 'log_buzz']
SENT_VARS_SHARES = BASE_VARS + ['share_pos', 'share_neg', 'log_buzz']

# ============================================================
# TOBIT MLE
# ============================================================

def tobit_fit(y, X, lower=0):
    def neg_ll(params):
        beta, log_sigma = params[:-1], params[-1]
        sigma = np.exp(log_sigma)
        xb = X @ beta
        uncens = y > lower
        ll = 0.0
        if uncens.any():
            r = (y[uncens] - xb[uncens]) / sigma
            ll += np.sum(-0.5*np.log(2*np.pi) - np.log(sigma) - 0.5*r**2)
        cens = ~uncens
        if cens.any():
            z = (lower - xb[cens]) / sigma
            ll += np.sum(np.log(stats.norm.cdf(z) + 1e-12))
        return -ll

    ols = sm.OLS(y, X).fit()
    p0 = np.append(ols.params, np.log(np.std(ols.resid)))
    res = optimize.minimize(neg_ll, p0, method='Nelder-Mead',
                            options={'maxiter':80000, 'xatol':1e-9})
    beta = res.x[:-1]
    sigma = np.exp(res.x[-1])
    eps = 1e-5
    n_p = len(res.x)
    H = np.zeros((n_p, n_p))
    for i in range(n_p):
        def fi(p, _i=i):
            return optimize.approx_fprime(p, neg_ll, eps)[_i]
        H[i] = optimize.approx_fprime(res.x, fi, eps)
    try:
        cov = np.linalg.inv(H)
        se = np.sqrt(np.abs(np.diag(cov)))
    except np.linalg.LinAlgError:
        se = np.full(n_p, np.nan)
    se_beta = se[:-1]
    z_vals = np.where(se_beta > 0, beta / se_beta, 0)
    p_vals = 2 * (1 - stats.norm.cdf(np.abs(z_vals)))
    return {'coef':beta, 'se':se_beta, 'z':z_vals, 'p':p_vals,
            'sigma':sigma, 'll':-res.fun, 'n':len(y),
            'n_cens':int((y<=lower).sum()), 'n_uncens':int((y>lower).sum()),
            'converged':res.success}

# ============================================================
# ЗАГРУЗКА
# ============================================================

def first_existing_col(df, candidates):
    for col in candidates:
        if col in df.columns:
            return col
    raise KeyError(f"Не найдена: {', '.join(candidates)}")

def normalize_rating(series):
    return (
        series
        .astype(str)
        .str.replace('–', '-', regex=False)
        .str.replace('—', '-', regex=False)
        .str.strip()
    )

def add_quality_flags(sub):
    sub['placement_dt'] = pd.to_datetime(sub['placement_date'], errors='coerce')
    sub['book_dt'] = pd.to_datetime(sub['book_date'], errors='coerce')
    sub['book_after_placement'] = (
        sub['book_dt'].notna()
        & sub['placement_dt'].notna()
        & (sub['book_dt'].dt.date > sub['placement_dt'].dt.date)
    )
    sub['is_special_issuer'] = sub['issuer'].astype(str).isin(BASELINE_EXCLUDE_ISSUERS)
    sub['zero_organizers'] = sub['num_organizers'].eq(0)
    return sub

def apply_baseline_filters(sub):
    mask = pd.Series(True, index=sub.index)
    if BASELINE_DROP_BOOK_AFTER_PLACEMENT:
        mask &= ~sub['book_after_placement']
    if BASELINE_EXCLUDE_ISSUERS:
        mask &= ~sub['is_special_issuer']
    if BASELINE_DROP_ZERO_ORGANIZERS:
        mask &= ~sub['zero_organizers']
    return sub.loc[mask].copy()

def load_and_prepare(path, year_from):
    df = pd.read_csv(path, low_memory=False)
    df['year'] = pd.to_datetime(df['placement_date'], errors='coerce').dt.year
    sub = df[df['year'] >= year_from].copy()

    # Сентимент: missing values are treated as neutral/zero signal.
    si_raw = pd.to_numeric(sub[first_existing_col(sub, ['si_relevant', 'si'])], errors='coerce')
    log_buzz_raw = pd.to_numeric(sub[first_existing_col(sub, ['log_buzz_relevant', 'log_buzz_all', 'log_buzz'])], errors='coerce')
    sub['si'] = si_raw.fillna(0)
    sub['log_buzz'] = log_buzz_raw.fillna(0)
    sub['share_neg'] = pd.to_numeric(sub[first_existing_col(sub, ['share_negative_relevant', 'share_negative'])], errors='coerce').fillna(0)
    sub['share_pos'] = pd.to_numeric(sub[first_existing_col(sub, ['share_positive_relevant', 'share_positive'])], errors='coerce').fillna(0)

    # Зависимые
    sub['underplacement'] = 1 - pd.to_numeric(sub['volume_ratio'], errors='coerce')
    sub['coupon_reduction_bp'] = pd.to_numeric(sub['coupon_reduction_bp'], errors='coerce')
    sub['num_organizers'] = pd.to_numeric(sub['num_organizers'], errors='coerce')

    # Контроли
    if 'log_term' in sub.columns:
        sub['log_dur'] = pd.to_numeric(sub['log_term'], errors='coerce')
    elif 'term' in sub.columns:
        sub['log_dur'] = np.log(pd.to_numeric(sub['term'], errors='coerce').clip(lower=1))
    sub['final_rating'] = normalize_rating(sub['final_rating'])
    sub['rating_num'] = sub['final_rating'].map(RATING_MAP)
    sub['ofz_yield'] = pd.to_numeric(sub['ofz_matched_yield'], errors='coerce')
    sub['rvi'] = pd.to_numeric(sub['rvi_close'], errors='coerce')
    if 'is_floater' not in sub.columns:
        sub['is_floater'] = 0
    sub = add_quality_flags(sub)
    sub['issuer_id'] = sub['issuer'].astype('category').cat.codes

    # Dropna по общим переменным (БЕЗ coupon_reduction_bp)
    drop_vars = list(set(['underplacement', 'issuer_id'] + SENT_VARS_SI + ['share_pos', 'share_neg']))
    full_sub = sub.dropna(subset=drop_vars).copy()
    return apply_baseline_filters(full_sub)

# ============================================================
# ОЦЕНКА
# ============================================================

def estimate_all(sub):
    R = {}
    y_up = sub['underplacement'].values
    X_si = sm.add_constant(sub[SENT_VARS_SI])
    X_sh = sm.add_constant(sub[SENT_VARS_SHARES])

    # Underplacement — вся выборка
    R['M3_OLS_underpl_HC3'] = sm.OLS(y_up, X_si).fit(cov_type='HC3')
    R['M4_OLS_underpl_cluster'] = sm.OLS(y_up, X_si).fit(
        cov_type='cluster', cov_kwds={'groups': sub['issuer_id']})
    R['M5_Tobit_underpl'] = tobit_fit(y_up, X_si.values, lower=0)
    R['M6_FracLogit_underpl'] = sm.GLM(
        sub['underplacement'].clip(0.001, 0.999), X_si,
        family=sm.families.Binomial()).fit(cov_type='HC1')
    R['M8_Tobit_underpl_shares'] = tobit_fit(y_up, X_sh.values, lower=0)

    # Coupon — подвыборка с ориентиром
    sub_cr = sub.dropna(subset=['coupon_reduction_bp']).copy()
    y_cr = sub_cr['coupon_reduction_bp'].values
    X_si_cr = sm.add_constant(sub_cr[SENT_VARS_SI])
    X_sh_cr = sm.add_constant(sub_cr[SENT_VARS_SHARES])

    R['M1_OLS_coupon_HC3'] = sm.OLS(y_cr, X_si_cr).fit(cov_type='HC3')
    R['M2_OLS_coupon_cluster'] = sm.OLS(y_cr, X_si_cr).fit(
        cov_type='cluster', cov_kwds={'groups': sub_cr['issuer_id']})
    R['M7_OLS_coupon_shares'] = sm.OLS(y_cr, X_sh_cr).fit(cov_type='HC3')

    R['_n_underpl'] = len(sub)
    R['_n_coupon'] = len(sub_cr)
    return R

# ============================================================
# EXCEL
# ============================================================

def sig_stars(p):
    if np.isnan(p): return ''
    if p < 0.01: return '***'
    if p < 0.05: return '**'
    if p < 0.1: return '*'
    return ''

def get_csp(res, var, cn):
    if isinstance(res, dict):
        if var not in cn: return None, None, None
        i = cn.index(var)
        return res['coef'][i], res['se'][i], res['p'][i]
    else:
        if var in res.params.index:
            return res.params[var], res.bse[var], res.pvalues[var]
        return None, None, None

def write_model_table(ws, title, models, results, start_row=1):
    hf = Font(bold=True, size=11, name='Arial')
    hfill = PatternFill('solid', fgColor='D9E1F2')
    sfill = PatternFill('solid', fgColor='E2EFDA')

    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=13, name='Arial')
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2+len(models)*2)

    row = start_row + 2
    ws.cell(row=row, column=1, value='Variable').font = hf
    ws.cell(row=row, column=1).fill = hfill
    for j, (key, ttl, spec, cn) in enumerate(models):
        c = 2 + j * 2
        ws.cell(row=row, column=c, value=ttl).font = hf
        ws.cell(row=row, column=c).fill = hfill
        ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, horizontal='center')
        ws.merge_cells(start_row=row, start_column=c, end_row=row, end_column=c+1)
    row += 1

    all_v = []
    for _, _, _, cn in models:
        for v in cn:
            if v not in all_v: all_v.append(v)

    for var in all_v:
        lab = VAR_LABELS.get(var, var)
        ws.cell(row=row, column=1, value=lab).font = Font(name='Arial', size=10)
        for j, (key, _, _, cn) in enumerate(models):
            c = 2 + j * 2
            coef, se, p = get_csp(results[key], var, cn)
            if coef is None:
                ws.cell(row=row, column=c, value='—').font = Font(name='Arial', size=10, color='999999')
                continue
            st = sig_stars(p)
            cell = ws.cell(row=row, column=c, value=round(coef, 4))
            cell.number_format = '0.0000'
            cell.font = Font(name='Arial', size=10, bold=(p < 0.1))
            if p < 0.1: cell.fill = sfill
            ws.cell(row=row, column=c+1, value=st).font = Font(name='Arial', size=10)
            ws.cell(row=row+1, column=c, value=f'({se:.4f})').font = Font(name='Arial', size=9, color='666666')
        row += 2

    row += 1
    ws.cell(row=row, column=1, value='Model statistics').font = Font(bold=True, name='Arial')
    row += 1
    for sn, fn in [
        ('N', lambda r: int(r.nobs) if hasattr(r, 'nobs') else r['n']),
        ('R²', lambda r: f"{r.rsquared:.4f}" if hasattr(r, 'rsquared') else '—'),
        ('Adj. R²', lambda r: f"{r.rsquared_adj:.4f}" if hasattr(r, 'rsquared_adj') else '—'),
        ('Log-lik', lambda r: f"{r['ll']:.2f}" if isinstance(r, dict) else
                    (f"{r.llf:.2f}" if hasattr(r, 'llf') and not hasattr(r, 'rsquared') else '—')),
        ('Sigma', lambda r: f"{r['sigma']:.4f}" if isinstance(r, dict) else '—'),
        ('Cens/Uncens', lambda r: f"{r['n_cens']}/{r['n_uncens']}" if isinstance(r, dict) else '—'),
    ]:
        ws.cell(row=row, column=1, value=sn).font = Font(name='Arial', size=10, italic=True)
        for j, (key, _, _, _) in enumerate(models):
            ws.cell(row=row, column=2+j*2, value=str(fn(results[key]))).font = Font(name='Arial', size=10)
        row += 1

    row += 1
    for note in ['*** p<0.01, ** p<0.05, * p<0.1. SE in parentheses.',
                 'Underplacement = 1 − volume_ratio. Tobit left-censored at 0.',
                 'Missing SI/buzz values are set to 0.',
                 'Baseline excludes ВЭБ.РФ/special technical placements, num_organizers=0, and book_date > placement_date.',
                 'Coupon models: subset with guidance (orient_type ≠ none).']:
        ws.cell(row=row, column=1, value=note).font = Font(name='Arial', size=9, italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2+len(models)*2)
        row += 1
    ws.column_dimensions['A'].width = 34
    for i in range(2, 2+len(models)*2+1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    return row

def build_excel(R, sub, path):
    wb = Workbook()
    cn_si = ['const'] + SENT_VARS_SI
    cn_sh = ['const'] + SENT_VARS_SHARES
    hf = Font(bold=True, size=11, name='Arial')
    hfill = PatternFill('solid', fgColor='D9E1F2')

    n_up = R['_n_underpl']
    n_cr = R['_n_coupon']

    ws1 = wb.active; ws1.title = 'Main Results'
    write_model_table(ws1, f'Table 1. Placement Outcomes ({YEAR_FROM}–2025)', [
        ('M1_OLS_coupon_HC3', f'OLS coupon\n(N={n_cr})', SENT_VARS_SI, cn_si),
        ('M3_OLS_underpl_HC3', f'OLS underpl\n(N={n_up})', SENT_VARS_SI, cn_si),
        ('M5_Tobit_underpl', f'Tobit underpl\n(N={n_up})', SENT_VARS_SI, cn_si),
    ], R)

    ws2 = wb.create_sheet('Robustness')
    write_model_table(ws2, 'Table 2. Robustness Checks', [
        ('M2_OLS_coupon_cluster', f'OLS coupon\ncluster\n(N={n_cr})', SENT_VARS_SI, cn_si),
        ('M4_OLS_underpl_cluster', f'OLS underpl\ncluster\n(N={n_up})', SENT_VARS_SI, cn_si),
        ('M6_FracLogit_underpl', f'FracLogit\nunderpl\n(N={n_up})', SENT_VARS_SI, cn_si),
        ('M8_Tobit_underpl_shares', f'Tobit underpl\nshares\n(N={n_up})', SENT_VARS_SHARES, cn_sh),
    ], R)

    # Descriptive
    ws3 = wb.create_sheet('Descriptive Stats')
    ws3['A1'] = f'Table 3. Descriptive Statistics (N_underpl={n_up}, N_coupon={n_cr})'
    ws3['A1'].font = Font(bold=True, size=13, name='Arial')
    row = 3
    for j, h in enumerate(['Variable', 'N', 'Mean', 'Std', 'Min', 'p25', 'Median', 'p75', 'Max']):
        ws3.cell(row=row, column=j+1, value=h).font = hf; ws3.cell(row=row, column=j+1).fill = hfill
    row += 1
    for var in ['coupon_reduction_bp', 'underplacement', 'volume_ratio', 'si', 'log_buzz',
                'share_pos', 'share_neg', 'num_organizers', 'is_debut', 'is_floater',
                'hist_ever_reduced', 'hist_avg_volume_ratio', 'rating_num', 'ofz_yield', 'rvi', 'log_dur', 'has_put']:
        if var not in sub.columns: continue
        s = sub[var]
        ws3.cell(row=row, column=1, value=VAR_LABELS.get(var, var)).font = Font(name='Arial', size=10)
        for j, v in enumerate([int(s.notna().sum()), s.mean(), s.std(), s.min(),
                              s.quantile(.25), s.median(), s.quantile(.75), s.max()]):
            ws3.cell(row=row, column=j+2, value=round(v, 3)).number_format = '0.000'
        row += 1
    ws3.column_dimensions['A'].width = 34
    for i in range(2, 10): ws3.column_dimensions[get_column_letter(i)].width = 12

    # Correlations
    ws4 = wb.create_sheet('Correlations')
    ws4['A1'] = 'Table 4. Correlation Matrix'
    ws4['A1'].font = Font(bold=True, size=13, name='Arial')
    cv = [v for v in ['coupon_reduction_bp', 'underplacement', 'si', 'log_buzz', 'num_organizers',
          'is_floater', 'hist_ever_reduced', 'hist_avg_volume_ratio', 'rating_num', 'ofz_yield', 'rvi']
          if v in sub.columns]
    cl = [VAR_LABELS.get(v, v) for v in cv]
    cm = sub[cv].corr()
    row = 3
    for j, l in enumerate(cl):
        c = ws4.cell(row=row, column=j+2, value=l)
        c.font = Font(bold=True, size=9, name='Arial'); c.fill = hfill
        c.alignment = Alignment(text_rotation=45, wrap_text=True)
    row += 1
    for i, (var, lab) in enumerate(zip(cv, cl)):
        ws4.cell(row=row, column=1, value=lab).font = Font(name='Arial', size=10)
        ws4.cell(row=row, column=1).fill = hfill
        for j, v2 in enumerate(cv):
            val = cm.loc[var, v2]
            c = ws4.cell(row=row, column=j+2, value=round(val, 3))
            c.number_format = '0.000'
            c.font = Font(name='Arial', size=10, bold=(abs(val) > 0.3 and i != j))
        row += 1
    ws4.column_dimensions['A'].width = 30
    for i in range(2, 14): ws4.column_dimensions[get_column_letter(i)].width = 12

    wb.save(path)

# ============================================================
# КОНСОЛЬ
# ============================================================

def print_results(R, sub):
    cn = ['const'] + SENT_VARS_SI
    sub_cr = sub.dropna(subset=['coupon_reduction_bp'])
    print(f"\n{'='*70}")
    print(f"РЕЗУЛЬТАТЫ (период {YEAR_FROM}+)")
    print(f"{'='*70}")
    print(f"N underplacement: {R['_n_underpl']}")
    print(f"N coupon: {R['_n_coupon']}")
    print(f"coupon > 0: {(sub_cr['coupon_reduction_bp'] > 0.5).sum()}")
    print(f"coupon < 0: {(sub_cr['coupon_reduction_bp'] < -0.5).sum()}")
    print(f"underpl > 0: {(sub['underplacement'] > 0.01).sum()}")
    print(f"Эмитентов: {sub['issuer'].nunique()}")
    print(f"Флоутеров: {int(sub['is_floater'].sum())}")

    for name, key in [('Tobit underplacement', 'M5_Tobit_underpl'),
                      ('OLS coupon (HC3)', 'M1_OLS_coupon_HC3'),
                      ('OLS underplacement (HC3)', 'M3_OLS_underpl_HC3'),
                      ('OLS coupon clustered', 'M2_OLS_coupon_cluster'),
                      ('OLS underpl clustered', 'M4_OLS_underpl_cluster'),
                      ('Fractional Logit underpl', 'M6_FracLogit_underpl'),
                      ('OLS coupon shares', 'M7_OLS_coupon_shares'),
                      ('Tobit underpl shares', 'M8_Tobit_underpl_shares')]:
        res = R[key]
        cn_use = cn if 'shares' not in key else ['const'] + SENT_VARS_SHARES
        print(f"\n--- {name} ---")
        if isinstance(res, dict):
            print(f"LL={res['ll']:.2f}, sigma={res['sigma']:.4f}, cens={res['n_cens']}, uncens={res['n_uncens']}")
        elif hasattr(res, 'rsquared'):
            print(f"R²={res.rsquared:.4f}, Adj.R²={res.rsquared_adj:.4f}, N={int(res.nobs)}")
        elif hasattr(res, 'llf'):
            print(f"LL={res.llf:.2f}, AIC={res.aic:.2f}, N={int(res.nobs)}")
        print(f"{'Variable':35s} {'Coef':>10s} {'SE':>10s} {'p':>8s}")
        print("-" * 65)
        for var in cn_use:
            coef, se, p = get_csp(res, var, cn_use)
            if coef is None: continue
            print(f"{VAR_LABELS.get(var, var):35s} {coef:10.4f} {se:10.4f} {p:8.4f} {sig_stars(p)}")

def main():
    print("Загрузка данных...")
    sub = load_and_prepare(INPUT_FILE, YEAR_FROM)
    print(f"N (underplacement) = {len(sub)}")
    print(f"N (coupon) = {sub['coupon_reduction_bp'].notna().sum()}")
    print("Оценка моделей...")
    R = estimate_all(sub)
    print_results(R, sub)
    build_excel(R, sub, OUTPUT_FILE)
    print(f"\n✅ Таблицы сохранены: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
